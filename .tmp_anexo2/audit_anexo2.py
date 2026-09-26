from __future__ import annotations

import csv
import re
import sys
import zipfile
from collections import Counter, defaultdict
from datetime import date, datetime
from pathlib import Path
from xml.etree import ElementTree as ET

from docx import Document
from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from infrastructure.readers.excel_reader import ExcelReader
from infrastructure.repositories.catalog_repository import ExcelCatalogRepository


ROOT = Path(r"C:\taully_demand_forecast")
DOCX = ROOT / "outputs" / "01a0b651-7620-73f0-8f6a-8b68b14f7a17" / "Anexo_2_Instrumentos_de_recoleccion_de_datos.docx"
HISTORY = ROOT / "data" / "historial_demanda.csv"
INVENTORY = ROOT / "data" / "inventario_producto_Taully_2026-04-01_a_2026-09-15.xlsx"
THESIS = Path(r"C:\Users\Acer\OneDrive\Tesis Sr\LIMA-NORTE_PI_GURRERRO.docx")

EXPECTED_INSTRUMENTS = [
    "Instrumento 1 Cantidad demandada mediante inventario",
    "Instrumento 2 Índice de salida de inventario",
    "Instrumento 3 Tasa de quiebre de stock",
    "Instrumento 4 Estacionalidad diaria",
    "Instrumento 5 Error de pronóstico",
    "Instrumento 6 Volumen de demanda mediante ventas",
    "Instrumento 7 Patrón de demanda estacional",
    "Instrumento 8 Precisión del pronóstico",
]
EXPECTED_CATEGORIES = ["ABARROTES", "BEBIDAS", "GOLOSINAS", "HELADOS", "LIMPIEZA"]


def fail(message: str) -> None:
    print(f"ERROR: {message}")
    raise SystemExit(1)


def clean(value: object) -> str:
    return re.sub(r"\s+", " ", str(value or "")).strip()


def parse_history() -> list[dict[str, object]]:
    rows: list[dict[str, object]] = []
    with HISTORY.open("r", encoding="utf-8-sig", newline="") as f:
        for row in csv.DictReader(f):
            rows.append({
                "date": datetime.strptime(row["date"], "%Y-%m-%d").date(),
                "category": clean(row["category"]),
                "quantity": int(float(row["quantity"])),
            })
    return rows


def date_value(value: object) -> date:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, str):
        for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y"):
            try:
                return datetime.strptime(value.strip(), fmt).date()
            except ValueError:
                pass
    raise ValueError(f"No se reconoce fecha: {value!r}")


def find_header(sheet, expected: str) -> int:
    for row_no in range(1, min(sheet.max_row, 15) + 1):
        values = [clean(cell.value).upper() for cell in sheet[row_no]]
        if expected.upper() in values:
            return row_no
    raise ValueError(f"No se halló encabezado {expected!r} en {sheet.title}")


def inventory_sales_by_category() -> tuple[Counter, dict[str, int]]:
    wb = load_workbook(INVENTORY, read_only=True, data_only=True)
    ws = wb["Datos inventario"]
    header_row = find_header(ws, "Cantidad_Vendida")
    header = [clean(cell.value) for cell in ws[header_row]]
    index = {value: pos for pos, value in enumerate(header)}
    required = {"Fecha", "Categoría", "Cantidad_Vendida", "Stock_Inicial", "Entradas", "Stock_Disponible", "Stock_Final", "Dias_Sin_Stock"}
    missing = required - set(index)
    if missing:
        fail("Faltan columnas de inventario: " + ", ".join(sorted(missing)))
    totals: Counter = Counter()
    counts = {"rows": 0, "formula_errors": 0, "dqs_errors": 0}
    for row in ws.iter_rows(min_row=header_row + 1, values_only=True):
        if not any(value is not None for value in row):
            continue
        counts["rows"] += 1
        category = clean(row[index["Categoría"]])
        sold = int(row[index["Cantidad_Vendida"]] or 0)
        initial = int(row[index["Stock_Inicial"]] or 0)
        entry = int(row[index["Entradas"]] or 0)
        available = int(row[index["Stock_Disponible"]] or 0)
        final = int(row[index["Stock_Final"]] or 0)
        dqs = int(row[index["Dias_Sin_Stock"]] or 0)
        totals[category] += sold
        if available != initial + entry or final != max(0, available - sold):
            counts["formula_errors"] += 1
        if dqs not in (0, 1):
            counts["dqs_errors"] += 1
    return totals, counts


def raw_report_checks(history: list[dict[str, object]]) -> list[str]:
    reports = sorted((ROOT / "data").glob("reporte_ventas_*.xlsx"))
    if len(reports) != 168:
        fail(f"Se hallaron {len(reports)} reportes diarios; se esperaban 168.")
    reader = ExcelReader()
    catalog = ExcelCatalogRepository(ROOT / "data" / "catalogo_maestro.xlsx")
    totals: Counter = Counter()
    observed_dates: set[date] = set()
    unmatched: list[str] = []
    for report in reports:
        sales = reader.read_sales(str(report))
        for sale in sales:
            report_date = sale.date.date()
            observed_dates.add(report_date)
            product = catalog.get_product(sale.product_name)
            if product is None:
                unmatched.append(f"{report.name}: {sale.product_name}")
                continue
            totals[(report_date, product.category)] += int(sale.quantity)
    if unmatched:
        fail("Hay productos de ventas sin categoría en los reportes: " + "; ".join(unmatched[:5]))
    history_totals = {(r["date"], r["category"]): r["quantity"] for r in history}
    differences = []
    for key in sorted(set(totals) | set(history_totals)):
        if totals.get(key, 0) != history_totals.get(key, 0):
            differences.append(f"{key}: ventas={totals.get(key, 0)}, historial={history_totals.get(key, 0)}")
    if differences:
        fail("El historial no coincide con los reportes diarios: " + "; ".join(differences[:5]))
    if observed_dates != set(r["date"] for r in history):
        fail("Las fechas de los reportes diarios no coinciden con las fechas del historial.")
    return ["Los 168 reportes de ventas, clasificados con el catálogo maestro, coinciden registro por registro con el historial consolidado."]


def current_document_checks(history: list[dict[str, object]]) -> list[str]:
    doc = Document(DOCX)
    headings = [clean(p.text) for p in doc.paragraphs if clean(p.text).startswith("Instrumento ")]
    if headings != EXPECTED_INSTRUMENTS:
        fail("La secuencia de instrumentos del Word no coincide con las ocho fichas aprobadas.")
    if len(doc.tables) != 17:
        fail(f"El Word tiene {len(doc.tables)} tablas; se esperaban 17.")
    if any(not clean(cell.text) for table in doc.tables for row in table.rows for cell in row.cells):
        fail("Hay una celda vacía en una tabla del Word.")

    # The first table is the relation. Each instrument follows with one metadata table and one result table.
    data_tables = [doc.tables[2 + i * 2] for i in range(8)]
    category_totals = Counter()
    for row in history:
        category_totals[row["category"]] += row["quantity"]
    category_rows = [[clean(c.text) for c in r.cells] for r in data_tables[0].rows[1:]]
    for row in category_rows:
        if int(row[5].replace(",", "")) != category_totals[row[1]]:
            fail(f"Cantidad demandada incorrecta para {row[1]}.")

    sep_1_15 = [r for r in history if date(2026, 9, 1) <= r["date"] <= date(2026, 9, 15)]
    sep15 = {r["category"]: r["quantity"] for r in history if r["date"] == date(2026, 9, 15)}
    sep_avg = Counter()
    for row in sep_1_15:
        sep_avg[row["category"]] += row["quantity"]
    sep_avg = {category: qty / 15 for category, qty in sep_avg.items()}
    season_rows = [[clean(c.text) for c in r.cells] for r in data_tables[3].rows[1:]]
    for row in season_rows:
        category, vd, avg, ie = row[1], int(row[2]), float(row[3]), float(row[4])
        if vd != sep15[category] or round(avg, 2) != round(sep_avg[category], 2) or round(ie, 2) != round(vd / sep_avg[category], 2):
            fail(f"Estacionalidad diaria incorrecta para {category}.")

    history_days = len(set(row["date"] for row in history))
    reference_avg = {category: category_totals[category] / history_days for category in EXPECTED_CATEGORIES}
    pattern_rows = [[clean(c.text) for c in r.cells] for r in data_tables[6].rows[1:]]
    for row in pattern_rows:
        category, vp, vpr, ie = row[1], float(row[2]), float(row[3]), float(row[4])
        if round(vp, 2) != round(sep_avg[category], 2) or round(vpr, 2) != round(reference_avg[category], 2) or round(ie, 2) != round(vp / vpr, 2):
            fail(f"Patrón estacional incorrecto para {category}.")

    # Values for SI, EN, SF, ISI, DQS, TQS, DP and MAPE must stay explicitly unavailable.
    unavailable_tables = (data_tables[0], data_tables[1], data_tables[2], data_tables[4], data_tables[7])
    if not all("N.A." in " ".join(clean(c.text) for r in table.rows for c in r.cells) for table in unavailable_tables):
        fail("Un indicador sin fuente real de inventario o pronóstico fue consignado como si tuviera resultado.")

    # Form 4 has a dimensional notation correction: the average is daily, not monthly.
    meta_4 = " ".join(clean(c.text) for r in doc.tables[7].rows for c in r.cells)
    if "VPD" not in meta_4:
        fail("El instrumento 4 todavía usa VPM aunque el promedio empleado es diario.")

    return [
        f"8 instrumentos y 17 tablas presentes; sin celdas vacías.",
        f"Cifras de ventas: {len(history)} registros, {history_days} días, {sum(category_totals.values()):,} unidades.",
        "Las tablas de demanda, estacionalidad diaria y patrón estacional coinciden con historial_demanda.csv.",
        "Los indicadores que requieren kardex o pares pronóstico-real permanecen marcados como N.A. de forma explícita.",
    ]


def thesis_checks() -> list[str]:
    doc = Document(THESIS)
    text = "\n".join(clean(p.text) for p in doc.paragraphs)
    if "Anexo 2" not in text:
        fail("No se halló Anexo 2 en la tesis fuente.")
    start = text.find("Anexo 2")
    end = text.find("Anexo 7", start)
    annex = text[start:end if end != -1 else None]
    required_terms = ["Cantidad demandada", "Índice de Salida de Inventario", "Tasa de quiebre de stock", "Estacionalidad", "Error de pronóstico", "Volumen de demanda", "Precisión del Pronóstico"]
    missing = [term for term in required_terms if term.lower() not in annex.lower()]
    if missing:
        fail("La tesis fuente no contiene los instrumentos esperados: " + ", ".join(missing))
    return ["La tesis fuente contiene las ocho fichas de Anexo 2 que el Word desarrolla como instrumentos 1 a 8."]


def xml_checks() -> list[str]:
    ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
    with zipfile.ZipFile(DOCX) as archive:
        root = ET.fromstring(archive.read("word/document.xml"))
    page_breaks = root.findall(".//w:br[@w:type='page']", ns)
    if len(page_breaks) != 8:
        fail(f"El Word contiene {len(page_breaks)} saltos de página; se esperaban 8.")
    return ["El archivo DOCX abre como paquete válido y contiene 8 saltos de página para separar las fichas."]


def main() -> None:
    if not DOCX.exists() or not HISTORY.exists() or not INVENTORY.exists() or not THESIS.exists():
        fail("Falta uno de los archivos requeridos para la auditoría.")
    history = parse_history()
    if len(history) != 840:
        fail(f"historial_demanda.csv contiene {len(history)} registros, no 840.")
    if sorted(set(r["category"] for r in history)) != EXPECTED_CATEGORIES:
        fail("Las cinco categorías del historial no coinciden con las del Anexo.")
    if min(r["date"] for r in history) != date(2026, 4, 1) or max(r["date"] for r in history) != date(2026, 9, 15):
        fail("El periodo del historial no es del 1 de abril al 15 de septiembre de 2026.")
    inventory_totals, inventory_counts = inventory_sales_by_category()
    history_totals = Counter()
    for row in history:
        history_totals[row["category"]] += row["quantity"]
    if inventory_totals != history_totals:
        fail("Las ventas agrupadas del inventario no coinciden con historial_demanda.csv.")
    if inventory_counts["formula_errors"] or inventory_counts["dqs_errors"]:
        fail("La hoja de inventario contiene inconsistencias aritméticas internas.")

    results = []
    results.extend(xml_checks())
    results.extend(thesis_checks())
    results.extend(raw_report_checks(history))
    results.extend(current_document_checks(history))
    results.append(f"El archivo de inventario tiene {inventory_counts['rows']:,} filas y sus ventas por categoría coinciden exactamente con el historial.")
    results.append("Las columnas de stock del inventario se mantienen fuera de los resultados reales: su procedencia no es un kardex validado por la empresa.")
    print("AUDITORÍA APROBADA")
    for item in results:
        print(f"OK: {item}")


if __name__ == "__main__":
    main()
