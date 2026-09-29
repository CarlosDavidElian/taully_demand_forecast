"""Auditoría de contenido y forma de los Anexos 2 finales."""

from pathlib import Path

from docx import Document


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "outputs" / "01a0b651-7620-73f0-8f6a-8b68b14f7a17"

EXPECTED_ROWS = [6, 6, 6, 6, 6, 6, 6, 6, 6, 6, 6, 16, 6, 6, 6, 6, 18]
META_TABLES = [0, 2, 4, 6, 8, 10, 12, 14]
META_LABELS = ["INDICADOR:", "INVESTIGADOR:", "LUGAR DE ESTUDIO:", "FÓRMULA:", "DONDE:", "Momento:"]
INSTRUMENTS = [
    "Cantidad demandada",
    "Índice de salida de inventario",
    "Tasa de quiebre de stock",
    "Estacionalidad",
    "Error de pronóstico",
    "Volumen de demanda",
    "Patrón de demanda",
    "Precisión del pronóstico",
]


def text_of(document):
    return "\n".join(
        [paragraph.text for paragraph in document.paragraphs]
        + [cell.text for table in document.tables for row in table.rows for cell in row.cells]
    )


def verify(label, filename, metric_tokens):
    document = Document(OUTPUT / filename)
    assert len(document.tables) == 17, f"{label}: se esperaban 17 tablas y se obtuvieron {len(document.tables)}"
    actual_rows = [len(table.rows) for table in document.tables]
    assert actual_rows == EXPECTED_ROWS, f"{label}: estructura de filas inesperada {actual_rows}"
    for index in META_TABLES:
        labels = [row.cells[0].text.strip() for row in document.tables[index].rows]
        assert labels == META_LABELS, f"{label}: metadatos de tabla {index} no coinciden con Anexo 2: {labels}"
    body = text_of(document)
    for instrument in INSTRUMENTS:
        assert instrument.lower() in body.lower(), f"{label}: falta instrumento {instrument}"
    assert body.lower().count("postest" if label == "Postest" else "pretest") >= 9, f"{label}: rótulos de momento incompletos"
    assert "595 observaciones" in body, f"{label}: falta cantidad de observaciones"
    assert "01/04 al 15/09/2026" in body, f"{label}: falta fuente temporal completa"
    assert "11 748 registros producto-fecha" in body, f"{label}: falta base de inventario"
    assert "no una certificación de kardex" in body, f"{label}: falta alcance de la base de inventario"
    assert "30 116" in body and "4 528" in body, f"{label}: faltan cifras de SI o EN de inventario"
    assert "24,75%" in body and "19,64%" in body, f"{label}: faltan resultados ISI o TQS"
    assert "No se calcula ISI" not in body and "No se calcula TQS" not in body, f"{label}: conserva texto de tabla no calculable"
    assert "15 productos reales" in body, f"{label}: falta aclaración de extracto de productos"
    assert "N.A." not in body and "N/D" not in body, f"{label}: hay marcadores ambiguos"
    assert "Fuente de datos" not in body and "Periodo" not in body, f"{label}: conserva filas ajenas al Anexo 2 original"
    for token in metric_tokens:
        assert token in body, f"{label}: falta resultado {token}"
    assert all(cell.text.strip() for table in document.tables for row in table.rows for cell in row.cells), f"{label}: hay celdas vacías"
    print(f"{label}: Anexo 2, 8 instrumentos, 17 cortes y cifras verificados.")


verify(
    "Pretest",
    "Anexo_2_Pretest_validacion_historica_rodante_hasta_2026-09-15.docx",
    ["MAE: 10,42", "RMSE: 14,26", "MAPE: 26,09%", "WAPE: 19,96%"],
)
verify(
    "Postest",
    "Anexo_2_Postest_validacion_historica_rodante_hasta_2026-09-15.docx",
    ["MAE: 10,02", "RMSE: 13,79", "MAPE: 24,67%", "WAPE: 19,19%", "1,42 puntos porcentuales", "11 de las 17 semanas"],
)
print("AUDITORÍA ANEXO 2 APROBADA")
