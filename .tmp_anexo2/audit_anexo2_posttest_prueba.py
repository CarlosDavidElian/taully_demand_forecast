from __future__ import annotations

import sys
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

from docx import Document

ROOT = Path(r"C:\taully_demand_forecast")
DOCX = ROOT / "outputs" / "01a0b651-7620-73f0-8f6a-8b68b14f7a17" / "Anexo_2_Postest_evento_de_prueba_simulado.docx"
sys.path.insert(0, str(ROOT / ".tmp_anexo2"))

from build_anexo2_posttest_prueba import generate_test_event


def clean(value: object) -> str:
    return " ".join(str(value or "").split())


def number(value: object) -> float:
    return float(str(value).replace(",", "").replace("%", ""))


def rows(table):
    return [[clean(cell.text) for cell in row.cells] for row in table.rows][1:]


def fail(message: str):
    raise AssertionError(message)


def main():
    expected_categories, expected_products = generate_test_event()
    document = Document(DOCX)
    headings = [clean(paragraph.text) for paragraph in document.paragraphs if clean(paragraph.text).startswith("Instrumento ")]
    if len(headings) != 8 or len(document.tables) != 17:
        fail("El postest no tiene las ocho fichas y sus 17 tablas.")
    text = " ".join(clean(paragraph.text) for paragraph in document.paragraphs)
    if "evento de prueba" not in text.lower() or "no representa ventas" not in text.lower():
        fail("El documento no identifica que los datos son simulados de prueba.")
    with zipfile.ZipFile(DOCX) as archive:
        xml = ET.fromstring(archive.read("word/document.xml"))
    ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
    if len(xml.findall(".//w:br[@w:type='page']", ns)) != 8:
        fail("El postest no conserva los ocho saltos de página.")

    expected_by_product = {item["product"]: item for item in expected_products}
    if len(expected_by_product) != 15:
        fail("El evento no contiene 15 productos distintos.")
    balances = {}
    for row in rows(document.tables[2]):
        product, si, en, sf, cd = row[1], number(row[3]), number(row[4]), number(row[5]), number(row[6])
        expected = expected_by_product.get(product)
        if expected is None or cd != si + en - sf:
            fail(f"CD no cuadra en {product}.")
        if (si, en, sf, cd) != (expected["si"], expected["en"], expected["sf"], expected["cd"]):
            fail(f"El balance no coincide con la fuente en {product}.")
        balances[product] = (si, en, sf, cd)
    if len(balances) != 15:
        fail("No se muestran 15 productos en el instrumento 1.")

    for row in rows(document.tables[4]):
        product, cd, si, en, isi = row[1], number(row[3]), number(row[4]), number(row[5]), number(row[6])
        expected = expected_by_product.get(product)
        if expected is None or (si, en, cd) != (balances[product][0], balances[product][1], balances[product][3]):
            fail(f"ISI no usa el mismo balance en {product}.")
        if round(cd / (si + en) * 100, 2) != round(isi, 2):
            fail(f"ISI incorrecto en {product}.")

    for row in rows(document.tables[6]):
        product, dqs, dd, tqs = row[1], number(row[3]), number(row[4]), number(row[5])
        expected = expected_by_product.get(product)
        if expected is None or dd != 7 or round(dqs / dd * 100, 2) != round(tqs, 2) or int(dqs) != expected["dqs"]:
            fail(f"TQS incorrecta en {product}.")

    for row in rows(document.tables[12]):
        product, actual, cd = row[1], number(row[3]), number(row[4])
        expected = expected_by_product.get(product)
        if expected is None or actual != cd or actual != expected["cd"]:
            fail(f"Volumen de demanda incorrecto en {product}.")

    for table_index in (8, 10, 14, 16):
        for row in rows(document.tables[table_index]):
            category = row[1]
            if category not in expected_categories:
                fail(f"Categoría inesperada: {category}.")
    for row in rows(document.tables[8]):
        category, vd, vpd, index = row[1], number(row[2]), number(row[3]), number(row[4])
        expected = expected_categories[category]
        if round(vd, 2) != round(expected["actual_last"], 2) or round(vpd, 2) != round(expected["actual_average"], 2) or round(index, 2) != round(expected["daily_index"], 2):
            fail(f"Estacionalidad diaria incorrecta en {category}.")
    for table_index in (10, 16):
        for row in rows(document.tables[table_index]):
            category, actual, forecast, mape = row[1], number(row[2]), number(row[3]), number(row[4])
            expected = expected_categories[category]
            if round(actual, 2) != round(expected["actual_total"], 2) or round(forecast, 2) != round(expected["forecast_total"], 2) or round(mape, 2) != round(expected["mape"], 2):
                fail(f"MAPE incorrecto en {category}.")
    for row in rows(document.tables[14]):
        category, vp, vpr, index = row[1], number(row[2]), number(row[3]), number(row[4])
        expected = expected_categories[category]
        if round(vp, 2) != round(expected["actual_average"], 2) or round(vpr, 2) != round(expected["reference_average"], 2) or round(index, 2) != round(expected["pattern_index"], 2):
            fail(f"Patrón de demanda incorrecto en {category}.")
    print("AUDITORÍA APROBADA")
    print("OK: 8 instrumentos, 17 tablas y 15 productos de prueba.")
    print("OK: CD, ISI, TQS, estacionalidad y MAPE cuadran con el evento simulado reproducible.")
    print("OK: el documento deja claro que no representa ventas ni kardex reales.")


if __name__ == "__main__":
    main()
