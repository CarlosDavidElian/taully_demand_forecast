"""Verifica estructura y cifras de los documentos de Anexo 2 generados."""

from pathlib import Path

from docx import Document


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "outputs" / "01a0b651-7620-73f0-8f6a-8b68b14f7a17"

FILES = {
    "Pretest": OUTPUT / "Anexo_2_Pretest_instrumentos_7_dias_validacion_historica.docx",
    "Postest": OUTPUT / "Anexo_2_Postest_instrumentos_7_dias_validacion_historica.docx",
}

EXPECTED_HEADINGS = [
    "Cantidad demandada",
    "Índice de salida de inventario",
    "Tasa de quiebre de stock",
    "Estacionalidad",
    "Error de pronóstico",
    "Volumen de demanda",
    "Patrón de demanda",
    "Precisión del pronóstico",
]
EXPECTED_ROWS = [7, 2, 7, 2, 7, 2, 7, 6, 7, 6, 7, 6, 7, 6, 7, 6]


def table_text(table):
    return "\n".join(cell.text.strip() for row in table.rows for cell in row.cells)


def verify(label, path):
    doc = Document(path)
    assert len(doc.tables) == 16, f"{label}: se esperaban 16 tablas y hay {len(doc.tables)}"
    assert [len(t.rows) for t in doc.tables] == EXPECTED_ROWS, f"{label}: filas de tablas inesperadas"

    body = "\n".join(p.text for p in doc.paragraphs) + "\n" + "\n".join(table_text(t) for t in doc.tables)
    for heading in EXPECTED_HEADINGS:
        assert heading in body, f"{label}: falta el instrumento {heading}"
    for forbidden in ("N.A.", "N/D", "simulad", "16 al 22", "16–22"):
        assert forbidden.lower() not in body.lower(), f"{label}: contiene texto no permitido: {forbidden}"
    assert "19 al 25 de mayo de 2026" in body, f"{label}: falta el período común de 7 días"
    assert "18 de mayo de 2026" in body, f"{label}: falta el corte de entrenamiento"
    assert all(cell.text.strip() for table in doc.tables for row in table.rows for cell in row.cells), f"{label}: hay celdas vacías"
    assert "No se calcula ISI" in body, f"{label}: falta la advertencia de ISI no verificable"
    assert "No se calcula TQS" in body, f"{label}: falta la advertencia de TQS no verificable"

    if label == "Pretest":
        for metric in ("MAE 11,47", "RMSE 15,66", "MAPE 31,90%", "WAPE 24,17%"):
            assert metric in body, f"{label}: falta {metric}"
    else:
        for metric in ("MAE 10,29", "RMSE 14,34", "MAPE 29,67%", "WAPE 21,68%", "2,23 puntos porcentuales"):
            assert metric in body, f"{label}: falta {metric}"
    print(f"{label}: estructura, período, instrumentos y cifras verificados.")


for name, document_path in FILES.items():
    verify(name, document_path)

print("AUDITORÍA APROBADA")
