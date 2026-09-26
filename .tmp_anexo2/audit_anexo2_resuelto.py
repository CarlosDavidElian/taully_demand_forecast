from __future__ import annotations

import csv
import json
import sys
from collections import Counter
from datetime import date, datetime
from pathlib import Path
from xml.etree import ElementTree as ET
import zipfile

import pandas as pd
from docx import Document

ROOT = Path(r"C:\taully_demand_forecast")
DOCX = ROOT / "outputs" / "01a0b651-7620-73f0-8f6a-8b68b14f7a17" / "Anexo_2_vista_previa_inventario_calculado.docx"
HISTORY = ROOT / "data" / "historial_demanda.csv"

sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT / ".tmp_inventory_demo"))
sys.path.insert(0, str(ROOT / ".venv" / "Lib" / "site-packages"))

from sklearn.model_selection import train_test_split

from application.services.train_service import TrainService
from config.settings import FORECAST_FEATURES
from infrastructure.ml.model_trainer import ModelTrainer
from audit_inventory_test import read_exported, reconstruct, source_sales_match_raw_reports


def fail(message: str) -> None:
    raise AssertionError(message)


def clean(value: object) -> str:
    return " ".join(str(value or "").split())


def number(value: object) -> float:
    return float(str(value).replace(",", "").replace("%", ""))


def table_rows(table):
    return [[clean(cell.text) for cell in row.cells] for row in table.rows]


def history_rows():
    with HISTORY.open(encoding="utf-8-sig", newline="") as stream:
        rows = [
            {
                "date": datetime.strptime(row["date"], "%Y-%m-%d").date(),
                "category": clean(row["category"]),
                "quantity": int(float(row["quantity"])),
            }
            for row in csv.DictReader(stream)
        ]
    if len(rows) != 840:
        fail(f"El historial tiene {len(rows)} registros en vez de 840.")
    if min(row["date"] for row in rows) != date(2026, 4, 1) or max(row["date"] for row in rows) != date(2026, 9, 15):
        fail("El periodo del historial no coincide con el Anexo.")
    return rows


def expected_mape(history: pd.DataFrame):
    result = {}
    for category in sorted(history["category"].unique()):
        frame = TrainService._build_training_frame(history[history["category"] == category].copy())
        trainer = ModelTrainer()
        metrics = trainer.train(
            frame[FORECAST_FEATURES].values,
            frame["quantity"].values,
            seasonal_feature_index=FORECAST_FEATURES.index("seasonal_mean_4"),
            model_name=f"model_{category}",
        )
        _, x_test, _, y_test = train_test_split(
            frame[FORECAST_FEATURES].values,
            frame["quantity"].values,
            test_size=0.2,
            shuffle=False,
        )
        result[category] = {
            "dr": float(y_test.sum()),
            "dp": float(trainer.best_model.predict(x_test).sum()),
            "mape": float(metrics["mape"]),
        }
    return result


def main() -> None:
    if not DOCX.exists():
        fail("No se encontró el Anexo resuelto.")
    rows = history_rows()
    history = pd.read_csv(HISTORY)
    history["date"] = pd.to_datetime(history["date"])

    # Verificación independiente de reportes diarios -> detalle fuente -> Excel de inventario.
    source_sales_match_raw_reports()
    exported, inventory_issues = read_exported()
    generated, inventory_categories = reconstruct()
    if set(exported) != set(generated) or inventory_issues["duplicates"] or inventory_issues["balance_errors"] or inventory_issues["dqs_errors"]:
        fail("El inventario de prueba no coincide con la lógica de origen.")
    for key, expected in generated.items():
        actual = exported[key]
        observed = {
            "category": clean(actual["Categoría"]),
            "demand": int(actual["Cantidad_Vendida"]),
            "si": int(actual["Stock_Inicial"]),
            "entries": int(actual["Entradas"]),
            "available": int(actual["Stock_Disponible"]),
            "sf": int(actual["Stock_Final"]),
            "dqs": int(actual["Dias_Sin_Stock"]),
        }
        if observed != expected:
            fail(f"La fila de inventario {key} no coincide con la lógica de origen.")

    document = Document(DOCX)
    headings = [clean(p.text) for p in document.paragraphs if clean(p.text).startswith("Instrumento ")]
    if len(headings) != 8 or len(document.tables) != 17:
        fail("El Anexo no contiene las ocho fichas y sus tablas esperadas.")
    all_text = " ".join(clean(cell.text) for table in document.tables for row in table.rows for cell in row.cells)
    all_paragraphs = " ".join(clean(p.text) for p in document.paragraphs)
    if "N.A." in all_text:
        fail("Aún existe un N.A. en la versión resuelta.")
    if "Vista previa" not in all_paragraphs or "160 unidades" not in all_paragraphs:
        fail("El Anexo no comunica que el stock mostrado es calculado de prueba.")

    ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
    with zipfile.ZipFile(DOCX) as archive:
        xml = ET.fromstring(archive.read("word/document.xml"))
    if len(xml.findall(".//w:br[@w:type='page']", ns)) != 8:
        fail("El Anexo no contiene los ocho saltos de página esperados.")

    # Instrumentos 1, 2 y 3: comprobar las fórmulas de inventario en cada categoría.
    quantities = {}
    for row in table_rows(document.tables[2])[1:]:
        category, si, en, sf, cd = row[1], number(row[2]), number(row[3]), number(row[4]), number(row[5])
        if cd != si + en - sf:
            fail(f"CD no cuadra en {category}.")
        quantities[category] = (si, en, sf, cd)
    for row in table_rows(document.tables[4])[1:]:
        category, cd, si, en, isi = row[1], number(row[2]), number(row[3]), number(row[4]), number(row[5])
        if (si, en, cd) != (quantities[category][0], quantities[category][1], quantities[category][3]):
            fail(f"ISI no usa los mismos datos de inventario en {category}.")
        if round(cd / (si + en) * 100, 2) != round(isi, 2):
            fail(f"ISI incorrecto en {category}.")
    for row in table_rows(document.tables[6])[1:]:
        category, dqs, dd, tqs = row[1], number(row[2]), number(row[3]), number(row[4])
        if dd != 168 or round(dqs / dd * 100, 2) != round(tqs, 2):
            fail(f"TQS incorrecto en {category}.")
        if int(dqs) != len(inventory_categories[category]["dqs_category_days"]):
            fail(f"DQS no coincide con el inventario de prueba en {category}.")

    category_totals = Counter()
    for row in rows:
        category_totals[row["category"]] += row["quantity"]
    if sum(category_totals.values()) != 43105:
        fail("El total de ventas no coincide con 43,105.")

    # Instrumentos 4 y 7: revisar promedios, índices y referencia histórica.
    sep = [row for row in rows if date(2026, 9, 1) <= row["date"] <= date(2026, 9, 15)]
    sep_totals = Counter(row["category"] for row in sep)
    sep_quantities = Counter()
    sep15 = {}
    for row in sep:
        sep_quantities[row["category"]] += row["quantity"]
        if row["date"] == date(2026, 9, 15):
            sep15[row["category"]] = row["quantity"]
    reference_days = len({row["date"] for row in rows})
    for row in table_rows(document.tables[8])[1:]:
        category, vd, vpd, ie = row[1], number(row[2]), number(row[3]), number(row[4])
        expected_vpd = sep_quantities[category] / 15
        if vd != sep15[category] or round(vpd, 2) != round(expected_vpd, 2) or round(ie, 2) != round(vd / expected_vpd, 2):
            fail(f"Estacionalidad diaria incorrecta en {category}.")
    for row in table_rows(document.tables[14])[1:]:
        category, vp, vpr, ie = row[1], number(row[2]), number(row[3]), number(row[4])
        expected_vp = sep_quantities[category] / 15
        expected_vpr = category_totals[category] / reference_days
        if round(vp, 2) != round(expected_vp, 2) or round(vpr, 2) != round(expected_vpr, 2) or round(ie, 2) != round(vp / vpr, 2):
            fail(f"Patrón de demanda incorrecto en {category}.")

    # Instrumentos 5 y 8: comprobar MAPE, DR y DP con validación histórica independiente.
    validation = expected_mape(history)
    for table_index in (10, 16):
        for row in table_rows(document.tables[table_index])[1:]:
            category, dr, dp, mape = row[1], number(row[2]), number(row[3]), number(row[4])
            expected = validation[category]
            if round(dr, 2) != round(expected["dr"], 2) or round(dp, 2) != round(expected["dp"], 2) or round(mape, 2) != round(expected["mape"], 2):
                fail(f"MAPE, DR o DP no coincide en {category}.")

    print("AUDITORÍA APROBADA")
    print("OK: los 168 reportes diarios coinciden producto por producto y fecha por fecha con el detalle que alimenta el inventario.")
    print("OK: las 11,748 filas del inventario de prueba coinciden con la lógica de cálculo y cada saldo por fila es correcto.")
    print("OK: las ocho fichas, sus 17 tablas, fórmulas de inventario, estacionalidad y MAPE validado no presentan discrepancias.")
    print("OK: no hay N.A. en el Anexo resuelto y se conserva la advertencia de que el inventario es calculado de prueba.")
    print("LÍMITE: los resultados de inventario no son kardex real de la empresa; el propio documento lo identifica como vista previa.")


if __name__ == "__main__":
    main()
