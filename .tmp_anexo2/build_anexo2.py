from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT = Path(r"C:\taully_demand_forecast\outputs\01a0b651-7620-73f0-8f6a-8b68b14f7a17\Anexo_2_Instrumentos_de_recoleccion_de_datos.docx")

NAVY = "17365D"
BLUE = "DCE6F1"
PALE_BLUE = "F3F7FB"
GRAY = "D9D9D9"
TEXT = "1F1F1F"


def set_run_font(run, size=None, bold=None, color=None, italic=None):
    run.font.name = "Arial"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if italic is not None:
        run.font.italic = italic
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)
    shd.set(qn("w:val"), "clear")


def set_cell_border(cell, color=GRAY, size="8"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right"):
        tag = qn(f"w:{edge}")
        element = borders.find(tag)
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_cell_margin(cell, top=90, start=110, bottom=90, end=110):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        element = tc_mar.find(qn(f"w:{side}"))
        if element is None:
            element = OxmlElement(f"w:{side}")
            tc_mar.append(element)
        element.set(qn("w:w"), str(value))
        element.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    node = OxmlElement("w:tblHeader")
    node.set(qn("w:val"), "true")
    tr_pr.append(node)


def cell_text(cell, text, size=9.5, bold=False, color=TEXT, align=None):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.line_spacing = 1.05
    if align is not None:
        p.alignment = align
    run = p.add_run(str(text))
    set_run_font(run, size=size, bold=bold, color=color)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    set_cell_margin(cell)
    set_cell_border(cell)


def style_table(table, widths, header=True):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for row_idx, row in enumerate(table.rows):
        for col_idx, cell in enumerate(row.cells):
            cell.width = Inches(widths[col_idx])
            if header and row_idx == 0:
                set_cell_shading(cell, NAVY)
                for p in cell.paragraphs:
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    for run in p.runs:
                        set_run_font(run, size=9, bold=True, color="FFFFFF")
            elif row_idx % 2 == 0:
                set_cell_shading(cell, PALE_BLUE)
            else:
                set_cell_shading(cell, "FFFFFF")
    if header:
        set_repeat_table_header(table.rows[0])


def add_paragraph(doc, text="", size=11, bold=False, italic=False, align=None, space_before=0, space_after=6):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.15
    if align is not None:
        p.alignment = align
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold, italic=italic, color="000000")
    return p


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.space_before = Pt(4 if level == 1 else 2)
    p.paragraph_format.space_after = Pt(7)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    set_run_font(run, size=13 if level == 1 else 11, bold=True, color="000000")
    return p


def add_meta_table(doc, indicator, formula, definitions, data_source, period):
    rows = [
        ("Indicador", indicator),
        ("Investigador", "Guerrero Mora, Carlos David Elian"),
        ("Lugar de estudio", "Minimarket Taully, Comas"),
        ("Formula", formula),
        ("Definicion operativa", definitions),
        ("Fuente de datos", data_source),
        ("Periodo", period),
    ]
    table = doc.add_table(rows=len(rows), cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for i, (label, value) in enumerate(rows):
        left, right = table.rows[i].cells
        left.width = Inches(1.55)
        right.width = Inches(5.35)
        cell_text(left, label, size=9.2, bold=True, color="FFFFFF")
        cell_text(right, value, size=9.2)
        set_cell_shading(left, NAVY)
        set_cell_shading(right, "FFFFFF" if i % 2 else PALE_BLUE)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def add_data_table(doc, headers, rows, widths, numeric_cols=None):
    numeric_cols = set(numeric_cols or [])
    table = doc.add_table(rows=1, cols=len(headers))
    for ci, h in enumerate(headers):
        cell_text(table.rows[0].cells[ci], h, size=9, bold=True, color="FFFFFF", align=WD_ALIGN_PARAGRAPH.CENTER)
    for row in rows:
        cells = table.add_row().cells
        for ci, value in enumerate(row):
            align = WD_ALIGN_PARAGRAPH.RIGHT if ci in numeric_cols else WD_ALIGN_PARAGRAPH.LEFT
            cell_text(cells[ci], value, size=9.2, align=align)
    style_table(table, widths, header=True)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_note(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.05
    lead = p.add_run("Nota. ")
    set_run_font(lead, size=9, bold=True, italic=True)
    body = p.add_run(text)
    set_run_font(body, size=9, italic=True)
    return p


def add_page_break(doc):
    doc.add_page_break()


def build_document():
    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)

    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor(0, 0, 0)

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(8)
    run = title.add_run("Anexo 2 Instrumentos de recoleccion de datos")
    set_run_font(run, size=16, bold=True, color="000000")

    add_paragraph(
        doc,
        "Las fichas presentan los indicadores empleados para registrar y analizar la demanda diaria por categoria de producto. Los resultados disponibles se obtuvieron del historial procesado entre el 1 de abril y el 15 de septiembre de 2026.",
        size=10.5,
        align=WD_ALIGN_PARAGRAPH.JUSTIFY,
        space_after=8,
    )
    add_heading(doc, "Fuente y alcance de los datos", level=1)
    add_data_table(
        doc,
        ["Elemento", "Registro verificado"],
        [
            ("Reportes diarios de ventas", "168 reportes comprendidos entre el 1 de abril y el 15 de septiembre de 2026."),
            ("Registros consolidados", "840 observaciones diarias por categoria."),
            ("Categorias analizadas", "Abarrotes, Bebidas, Golosinas, Helados y Limpieza."),
            ("Unidades vendidas", "43,105 unidades en el periodo evaluado."),
            ("Productos y marcas", "200 productos vendidos y 70 marcas asociadas al catalogo."),
        ],
        [2.0, 4.9],
    )
    add_note(
        doc,
        "Las ventas permiten calcular cantidad demandada y estacionalidad. El indice de salida de inventario y la tasa de quiebre de stock requieren un kardex real con stock inicial, entradas, stock final y dias sin stock. Esos campos no se obtienen de los reportes de venta.",
    )

    add_page_break(doc)
    add_heading(doc, "Ficha 1 Cantidad demandada", level=1)
    add_meta_table(
        doc,
        "Cantidad demandada",
        "CD = sumatoria de cantidades vendidas",
        "CD es el total de unidades vendidas por categoria durante el periodo evaluado.",
        "Reportes diarios de ventas consolidados por fecha y categoria.",
        "Del 1 de abril al 15 de septiembre de 2026.",
    )
    add_data_table(
        doc,
        ["N", "Categoria", "Cantidad demandada CD"],
        [
            ("1", "ABARROTES", "8,577"),
            ("2", "BEBIDAS", "15,046"),
            ("3", "GOLOSINAS", "12,498"),
            ("4", "HELADOS", "4,200"),
            ("5", "LIMPIEZA", "2,784"),
            ("", "TOTAL", "43,105"),
        ],
        [0.5, 3.0, 3.4],
        numeric_cols=[0, 2],
    )
    add_note(doc, "La cantidad demandada se registra directamente desde las ventas reales consolidadas. No equivale a dinero ni al numero de productos distintos.")

    add_page_break(doc)
    add_heading(doc, "Ficha 2 Indice de salida de inventario", level=1)
    add_meta_table(
        doc,
        "Indice de salida de inventario",
        "ISI = CD / (SI + EN)",
        "ISI expresa la proporcion de la mercaderia disponible que se vendio en el periodo. CD es cantidad demandada, SI es stock inicial y EN son entradas.",
        "Kardex o reporte de inventario validado por producto y fecha.",
        "Del 1 de abril al 15 de septiembre de 2026.",
    )
    add_data_table(
        doc,
        ["N", "Categoria", "CD", "SI", "EN", "ISI", "Estado"],
        [
            ("1", "ABARROTES", "8,577", "N.A.", "N.A.", "N.A.", "Requiere kardex real"),
            ("2", "BEBIDAS", "15,046", "N.A.", "N.A.", "N.A.", "Requiere kardex real"),
            ("3", "GOLOSINAS", "12,498", "N.A.", "N.A.", "N.A.", "Requiere kardex real"),
            ("4", "HELADOS", "4,200", "N.A.", "N.A.", "N.A.", "Requiere kardex real"),
            ("5", "LIMPIEZA", "2,784", "N.A.", "N.A.", "N.A.", "Requiere kardex real"),
        ],
        [0.35, 1.35, 0.75, 0.65, 0.65, 0.65, 2.1],
        numeric_cols=[0, 2, 3, 4, 5],
    )
    add_note(doc, "N.A. significa no aplicable con los reportes actuales. No se debe reemplazar SI o EN con valores estimados si el indicador se presentara como resultado real de la empresa.")

    add_page_break(doc)
    add_heading(doc, "Ficha 3 Tasa de quiebre de stock", level=1)
    add_meta_table(
        doc,
        "Tasa de quiebre de stock",
        "TQS = DQS / DD x 100",
        "TQS mide el porcentaje de dias con quiebre de stock. DQS es el numero de dias sin disponibilidad y DD el numero total de dias evaluados.",
        "Kardex, control de inventario o registro diario de productos agotados.",
        "Del 1 de abril al 15 de septiembre de 2026.",
    )
    add_data_table(
        doc,
        ["N", "Categoria", "DQS", "DD", "TQS", "Estado"],
        [
            ("1", "ABARROTES", "N.A.", "168", "N.A.", "Requiere registro de agotados"),
            ("2", "BEBIDAS", "N.A.", "168", "N.A.", "Requiere registro de agotados"),
            ("3", "GOLOSINAS", "N.A.", "168", "N.A.", "Requiere registro de agotados"),
            ("4", "HELADOS", "N.A.", "168", "N.A.", "Requiere registro de agotados"),
            ("5", "LIMPIEZA", "N.A.", "168", "N.A.", "Requiere registro de agotados"),
        ],
        [0.35, 1.5, 0.75, 0.65, 0.75, 2.5],
        numeric_cols=[0, 2, 3, 4],
    )
    add_note(doc, "El periodo tiene 168 dias de ventas. La ausencia de una venta no demuestra por si sola que hubo quiebre de stock; se necesita el registro de disponibilidad de inventario.")

    add_page_break(doc)
    add_heading(doc, "Ficha 4 Estacionalidad", level=1)
    add_meta_table(
        doc,
        "Estacionalidad",
        "IE = VPD del periodo / VPD de referencia",
        "IE es el indice estacional. VPD del periodo es la venta promedio diaria entre el 1 y el 15 de septiembre. VPD de referencia es la venta promedio diaria entre el 1 de abril y el 15 de septiembre.",
        "Historial de demanda consolidado por fecha y categoria.",
        "Periodo analizado: 1 al 15 de septiembre de 2026. Referencia: 1 de abril al 15 de septiembre de 2026.",
    )
    add_data_table(
        doc,
        ["N", "Categoria", "VPD periodo", "VPD referencia", "IE"],
        [
            ("1", "ABARROTES", "51.00", "51.05", "1.00"),
            ("2", "BEBIDAS", "88.87", "89.56", "0.99"),
            ("3", "GOLOSINAS", "69.27", "74.39", "0.93"),
            ("4", "HELADOS", "23.73", "25.00", "0.95"),
            ("5", "LIMPIEZA", "14.67", "16.57", "0.89"),
        ],
        [0.45, 2.0, 1.45, 1.55, 1.45],
        numeric_cols=[0, 2, 3, 4],
    )
    add_note(doc, "Un indice cercano a 1 indica un comportamiento similar a la referencia. Un valor menor que 1 muestra una demanda promedio diaria inferior en el periodo analizado.")

    add_page_break(doc)
    add_heading(doc, "Ficha 5 Error de pronostico", level=1)
    add_meta_table(
        doc,
        "Error de pronostico",
        "MAPE = (100 / n) x sumatoria de |(DR - DP) / DR|",
        "DR es demanda real, DP es demanda pronosticada y n es el numero de observaciones comparadas. El sistema tambien muestra MAE, RMSE y WAPE al finalizar el entrenamiento.",
        "Salida de entrenamiento y pronosticos generados por el sistema.",
        "Validacion historica definida durante el entrenamiento.",
    )
    add_data_table(
        doc,
        ["Dato requerido", "Disponibilidad actual", "Uso en la ficha"],
        [
            ("Demanda real DR", "Disponible en 840 registros diarios por categoria.", "Permite evaluar el pronostico cuando se compara con la misma fecha."),
            ("Demanda pronosticada DP", "Se genera al ejecutar el pronostico.", "Debe conservarse junto con la fecha pronosticada."),
            ("MAPE", "No se conserva como registro descargable en la interfaz actual.", "Se calcula al disponer de pares DR y DP del mismo periodo."),
            ("MAE, RMSE y WAPE", "El sistema los muestra despues de entrenar.", "Se registran en el reporte de metricas de cada entrenamiento."),
        ],
        [1.65, 2.15, 3.1],
    )
    add_note(doc, "Para completar el MAPE final se debe guardar la demanda pronosticada y, al concluir el periodo, compararla con la demanda real registrada. El pronostico futuro por si solo no permite calcular un error.")

    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build_document()
