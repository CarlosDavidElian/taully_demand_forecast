from pathlib import Path
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

out_dir = Path('outputs')
out_dir.mkdir(exist_ok=True)
out_file = out_dir / 'anexo_2_base.docx'

doc = Document()
section = doc.sections[0]
section.left_margin = 720
section.right_margin = 720
section.top_margin = 720
section.bottom_margin = 720

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = p.add_run('Anexo 2')
run.bold = True
run.font.size = Pt(18)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = p.add_run('Instrumentos de recolección de datos')
run.bold = True
run.font.size = Pt(16)

doc.add_paragraph('')
items = [
    '1. Cantidad demandada',
    '2. Índice de salida de inventario',
    '3. Tasa de quiebre de stock',
    '4. Estacionalidad de la demanda',
    '5. Error de pronóstico',
    '6. Volumen de demanda',
    '7. Patrón de demanda',
    '8. Precisión del pronóstico',
]
for item in items:
    p = doc.add_paragraph()
    p.add_run(f'{item}')
    p.runs[0].bold = True
    p.runs[0].font.size = Pt(12)
    for label in ['Indicador:', 'Investigador:', 'Lugar de estudio:', 'Fórmula:', 'Dónde:', 'Momento:']:
        doc.add_paragraph(label)
    doc.add_paragraph('')

p = doc.add_paragraph()
p.add_run('Nota: este documento es una versión base para Anexo 2. Puede ajustarse con los datos reales del estudio y la estructura final de la tesis.')
p.runs[0].italic = True

doc.save(out_file)
print(f'Archivo generado: {out_file.resolve()}')
