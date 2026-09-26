from pathlib import Path
import sys

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from build_anexo2 import (
    add_data_table,
    add_heading,
    add_meta_table,
    add_note,
    add_page_break,
    add_paragraph,
    set_run_font,
)


PREVIEW = "--vista-previa" in sys.argv
OFFICIAL_OUT = Path(r"C:\taully_demand_forecast\outputs\01a0b651-7620-73f0-8f6a-8b68b14f7a17\Anexo_2_Instrumentos_de_recoleccion_de_datos.docx")
PREVIEW_OUT = Path(r"C:\taully_demand_forecast\outputs\01a0b651-7620-73f0-8f6a-8b68b14f7a17\Anexo_2_vista_previa_inventario_calculado.docx")
OUT = PREVIEW_OUT if PREVIEW else OFFICIAL_OUT

CATEGORIES = [
    ("1", "ABARROTES", "8,577"),
    ("2", "BEBIDAS", "15,046"),
    ("3", "GOLOSINAS", "12,498"),
    ("4", "HELADOS", "4,200"),
    ("5", "LIMPIEZA", "2,784"),
]

INVENTORY_PREVIEW = [
    ("1", "ABARROTES", "565", "8,371", "362", "8,574", "95.95%", "3", "1.79%"),
    ("2", "BEBIDAS", "1,077", "14,523", "629", "14,971", "95.97%", "33", "19.64%"),
    ("3", "GOLOSINAS", "959", "12,115", "646", "12,428", "95.06%", "44", "26.19%"),
    ("4", "HELADOS", "311", "4,107", "226", "4,192", "94.88%", "6", "3.57%"),
    ("5", "LIMPIEZA", "245", "2,697", "162", "2,780", "94.49%", "3", "1.79%"),
]

MAPE_PREVIEW = [
    ("1", "ABARROTES", "1,558", "1,544.00", "23.48%"),
    ("2", "BEBIDAS", "2,638", "2,609.59", "15.47%"),
    ("3", "GOLOSINAS", "1,985", "2,114.58", "23.27%"),
    ("4", "HELADOS", "716", "733.75", "26.51%"),
    ("5", "LIMPIEZA", "432", "513.25", "39.38%"),
]


def setup_document():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)

    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    return doc


def title_page(doc):
    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(8)
    run = title.add_run(
        "Anexo 2 Vista previa de instrumentos con inventario calculado"
        if PREVIEW else "Anexo 2 Instrumentos de recolección de datos"
    )
    set_run_font(run, size=16, bold=True, color="000000")

    intro = (
        "Esta vista previa permite revisar las tablas de inventario con valores calculados a partir de los reportes de venta del 1 de abril al 15 de septiembre de 2026. No reemplaza el kardex real de la empresa ni debe utilizarse como resultado definitivo de tesis."
        if PREVIEW else
        "Este anexo presenta los ocho instrumentos definidos en la tesis para el pretest. Las cifras de ventas y estacionalidad corresponden al historial procesado entre el 1 de abril y el 15 de septiembre de 2026."
    )
    add_paragraph(doc, intro, size=10.5, align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after=8)
    add_heading(doc, "Relación de instrumentos", level=1)
    add_data_table(
        doc,
        ["N", "Instrumento", "Momento", "Estado de datos"],
        [
            ("1", "Cantidad demandada mediante inventario", "Pretest", "Vista previa calculada. Kardex real pendiente.") if PREVIEW else ("1", "Cantidad demandada mediante inventario", "Pretest", "Cantidad demandada disponible. Stock real pendiente."),
            ("2", "Índice de salida de inventario", "Pretest", "Vista previa calculada. Kardex real pendiente.") if PREVIEW else ("2", "Índice de salida de inventario", "Pretest", "Pendiente de kardex real."),
            ("3", "Tasa de quiebre de stock", "Pretest", "Vista previa calculada. Registro real pendiente.") if PREVIEW else ("3", "Tasa de quiebre de stock", "Pretest", "Pendiente de registro real de agotados."),
            ("4", "Estacionalidad diaria", "Pretest", "Disponible desde ventas reales."),
            ("5", "Error de pronóstico", "Pretest", "MAPE histórico calculado.") if PREVIEW else ("5", "Error de pronóstico", "Pretest", "Requiere pares de demanda real y pronosticada."),
            ("6", "Volumen de demanda mediante ventas", "Pretest", "Disponible desde ventas reales."),
            ("7", "Patrón de demanda estacional", "Pretest", "Disponible desde ventas reales."),
            ("8", "Precisión del pronóstico", "Pretest", "MAPE histórico calculado.") if PREVIEW else ("8", "Precisión del pronóstico", "Pretest", "Requiere pares de demanda real y pronosticada."),
        ],
        [0.4, 2.7, 0.8, 3.0],
    )
    note = (
        "La vista previa muestra stock calculado, no registrado por la tienda. La diferencia entre la demanda de los reportes (43,105) y la cantidad atendida por el stock calculado (42,945) es de 160 unidades."
        if PREVIEW else
        "Los 168 reportes diarios generan 840 registros consolidados en cinco categorías, con 43,105 unidades vendidas. Los campos de stock solo pueden usarse como resultados reales cuando provienen del kardex de la empresa."
    )
    add_note(doc, note)


def instrument_1(doc):
    add_heading(doc, "Instrumento 1 Cantidad demandada mediante inventario", level=1)
    add_meta_table(
        doc,
        "Cantidad demandada",
        "CD = (SI + EN) - SF",
        "CD es la cantidad demandada. SI es stock inicial, EN son entradas y SF es stock final.",
        "Kardex real y reportes de venta consolidados.",
        "Pretest. Del 1 de abril al 15 de septiembre de 2026.",
    )
    rows = [(n, cat, si, en, sf, cd) for n, cat, si, en, sf, cd, _, _, _ in INVENTORY_PREVIEW] if PREVIEW else [(n, cat, "N.A.", "N.A.", "N.A.", cd) for n, cat, cd in CATEGORIES]
    add_data_table(doc, ["N", "Categoría", "SI", "EN", "SF", "CD"], rows, [0.4, 2.0, 1.0, 1.0, 1.0, 1.2], numeric_cols=[0, 2, 3, 4, 5])
    add_note(doc, "Vista previa: CD se calcula como SI + EN - SF. Estos valores son de un inventario calculado desde ventas y no sustituyen el kardex real." if PREVIEW else "La CD proviene de ventas reales. SI, EN y SF no se consignan porque no existe kardex real para aplicar esta fórmula de inventario.")


def instrument_2(doc):
    add_heading(doc, "Instrumento 2 Índice de salida de inventario", level=1)
    add_meta_table(
        doc,
        "Índice de salida de inventario",
        "ISI = CD / (SI + EN)",
        "ISI mide la proporción de mercadería disponible que se vendió durante el periodo.",
        "Kardex real con stock inicial y entradas por producto o categoría.",
        "Pretest. Del 1 de abril al 15 de septiembre de 2026.",
    )
    rows = [(n, cat, cd, si, en, isi) for n, cat, si, en, _, cd, isi, _, _ in INVENTORY_PREVIEW] if PREVIEW else [(n, cat, cd, "N.A.", "N.A.", "N.A.") for n, cat, cd in CATEGORIES]
    add_data_table(doc, ["N", "Categoría", "CD", "SI", "EN", "ISI"], rows, [0.4, 2.0, 1.2, 1.0, 1.0, 1.2], numeric_cols=[0, 2, 3, 4, 5])
    add_note(doc, "Vista previa: ISI se calcula con las cifras del inventario calculado. No representa el índice real de salida de la tienda." if PREVIEW else "No se calcula ISI porque faltarían SI y EN reales. Un valor automático derivado de las ventas no representa la salida real de inventario de la tienda.")


def instrument_3(doc):
    add_heading(doc, "Instrumento 3 Tasa de quiebre de stock", level=1)
    add_meta_table(
        doc,
        "Tasa de quiebre de stock",
        "TQS = (DQS / DD) x 100",
        "TQS mide el porcentaje de días con quiebre de stock. DQS son días sin disponibilidad y DD son días evaluados.",
        "Kardex o control diario de productos agotados.",
        "Pretest. Del 1 de abril al 15 de septiembre de 2026.",
    )
    rows = [(n, cat, dqs, "168", tqs) for n, cat, _, _, _, _, _, dqs, tqs in INVENTORY_PREVIEW] if PREVIEW else [(n, cat, "N.A.", "168", "N.A.") for n, cat, _ in CATEGORIES]
    add_data_table(doc, ["N", "Categoría", "DQS", "DD", "TQS"], rows, [0.4, 2.3, 1.1, 1.1, 1.2], numeric_cols=[0, 2, 3, 4])
    add_note(doc, "Vista previa: DQS cuenta los días con al menos un producto de la categoría sin stock calculado suficiente. No equivale a un registro real de agotados." if PREVIEW else "Los 168 días corresponden al periodo de ventas. Un día sin venta no demuestra por sí solo que hubo quiebre; por ello DQS queda pendiente de un registro real de agotados.")


def instrument_4(doc):
    add_heading(doc, "Instrumento 4 Estacionalidad diaria", level=1)
    add_meta_table(
        doc,
        "Estacionalidad",
        "IE = VD / VPD",
        "IE es índice estacional. VD es venta diaria del 15 de septiembre y VPD es venta promedio diaria del 1 al 15 de septiembre.",
        "Historial de demanda consolidado por fecha y categoría.",
        "Pretest. Referencia: del 1 al 15 de septiembre de 2026.",
    )
    rows = [
        ("1", "ABARROTES", "53", "51.00", "1.04"),
        ("2", "BEBIDAS", "94", "88.87", "1.06"),
        ("3", "GOLOSINAS", "67", "69.27", "0.97"),
        ("4", "HELADOS", "25", "23.73", "1.05"),
        ("5", "LIMPIEZA", "16", "14.67", "1.09"),
    ]
    add_data_table(doc, ["N", "Categoría", "VD", "VPD", "IE"], rows, [0.4, 2.1, 1.15, 1.35, 1.1], numeric_cols=[0, 2, 3, 4])
    add_note(doc, "Un IE mayor a 1 muestra que la venta del día fue superior a su promedio diario del periodo. Un IE menor a 1 indica una venta inferior al promedio.")


def error_table(doc):
    rows = MAPE_PREVIEW if PREVIEW else [
        ("1", "ABARROTES", "53", "N.A.", "N.A."),
        ("2", "BEBIDAS", "94", "N.A.", "N.A."),
        ("3", "GOLOSINAS", "67", "N.A.", "N.A."),
        ("4", "HELADOS", "25", "N.A.", "N.A."),
        ("5", "LIMPIEZA", "16", "N.A.", "N.A."),
    ]
    return add_data_table(doc, ["N", "Categoría", "DR", "DP", "MAPE"], rows, [0.4, 2.1, 1.0, 1.35, 1.25], numeric_cols=[0, 2, 3, 4])


def instrument_5(doc):
    add_heading(doc, "Instrumento 5 Error de pronóstico", level=1)
    add_meta_table(
        doc,
        "Error de pronóstico",
        "MAPE = (100 / n) x sumatoria de |(DR - DP) / DR|",
        "DR es demanda real, DP es demanda pronosticada y n es el número de observaciones comparadas.",
        "Pronóstico generado por el sistema y reporte de ventas de la misma fecha.",
        "Pretest. Comparación histórica definida al entrenar el modelo.",
    )
    error_table(doc)
    add_note(doc, "Vista previa: DR y DP son los totales de las 28 observaciones diarias validadas del 19 de agosto al 15 de septiembre de 2026. MAPE se calcula observación por observación antes de promediar." if PREVIEW else "DR está disponible desde las ventas del 15 de septiembre. DP y MAPE quedan pendientes hasta conservar el pronóstico efectuado para esa misma fecha y contrastarlo con la venta real.")


def instrument_6(doc):
    add_heading(doc, "Instrumento 6 Volumen de demanda mediante ventas", level=1)
    add_meta_table(
        doc,
        "Cantidad demandada",
        "CD = sumatoria de ventas realizadas",
        "CD es la cantidad demandada obtenida directamente de las unidades vendidas.",
        "Reportes diarios de ventas procesados por el sistema.",
        "Pretest. Del 1 de abril al 15 de septiembre de 2026.",
    )
    rows = [
        ("1", "PACK MINI COCA COLA + INCA KOLA + FANTA + SPRITE 300 ML", "714", "714"),
        ("2", "SPORADE APPLE ICE SIN AZÚCAR 500 ML", "673", "673"),
        ("3", "GASEOSA COCA COLA PLUS LATA 320 ML", "608", "608"),
        ("4", "GASEOSA COCA COLA ORIGINAL 600 ML", "591", "591"),
        ("5", "SPORADE UVA 500 ML", "577", "577"),
    ]
    add_data_table(doc, ["N", "Producto", "Cantidad vendida", "CD"], rows, [0.4, 4.1, 1.1, 0.9], numeric_cols=[0, 2, 3])
    add_note(doc, "Se presentan los cinco productos con mayor cantidad vendida en el periodo. La ficha completa se alimenta con los 200 productos registrados en los reportes.")


def instrument_7(doc):
    add_heading(doc, "Instrumento 7 Patrón de demanda estacional", level=1)
    add_meta_table(
        doc,
        "Estacionalidad",
        "IE = VP / VPR",
        "VP es la venta promedio diaria del periodo analizado y VPR es la venta promedio diaria de referencia.",
        "Historial de demanda consolidado por fecha y categoría.",
        "Pretest. Periodo: 1 al 15 de septiembre. Referencia: 1 de abril al 15 de septiembre de 2026.",
    )
    rows = [
        ("1", "ABARROTES", "51.00", "51.05", "1.00"),
        ("2", "BEBIDAS", "88.87", "89.56", "0.99"),
        ("3", "GOLOSINAS", "69.27", "74.39", "0.93"),
        ("4", "HELADOS", "23.73", "25.00", "0.95"),
        ("5", "LIMPIEZA", "14.67", "16.57", "0.89"),
    ]
    add_data_table(doc, ["N", "Categoría", "VP", "VPR", "IE"], rows, [0.4, 2.1, 1.15, 1.35, 1.1], numeric_cols=[0, 2, 3, 4])
    add_note(doc, "Esta ficha muestra el comportamiento promedio del periodo frente a su referencia histórica. Un IE cercano a 1 indica estabilidad respecto a la referencia.")


def instrument_8(doc):
    add_heading(doc, "Instrumento 8 Precisión del pronóstico", level=1)
    add_meta_table(
        doc,
        "Error de pronóstico",
        "MAPE = (100 / n) x sumatoria de |(DR - DP) / DR|",
        "La precisión se determina al comparar demanda real y pronóstico para el mismo día y categoría.",
        "Reporte de métricas del entrenamiento, pronóstico descargado y ventas reales posteriores.",
        "Pretest. Comparación histórica definida al entrenar el modelo.",
    )
    error_table(doc)
    add_note(doc, "Vista previa: los pares DR y DP pertenecen a la validación histórica de 28 días, del 19 de agosto al 15 de septiembre de 2026. El sistema también muestra MAE, RMSE y WAPE luego de entrenar." if PREVIEW else "El sistema muestra MAE, RMSE y WAPE luego de entrenar. Para registrar el MAPE en esta ficha se deben conservar los pares DR y DP de cada fecha evaluada.")


def build_document():
    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc = setup_document()
    title_page(doc)
    for instrument in (instrument_1, instrument_2, instrument_3, instrument_4, instrument_5, instrument_6, instrument_7, instrument_8):
        add_page_break(doc)
        instrument(doc)
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build_document()
