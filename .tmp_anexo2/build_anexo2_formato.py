from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from build_anexo2 import add_data_table, add_heading, add_meta_table, add_page_break, set_run_font


OUT = Path(r"C:\taully_demand_forecast\outputs\01a0b651-7620-73f0-8f6a-8b68b14f7a17\Anexo_2_Instrumentos_de_recoleccion_de_datos_formato.docx")
ROWS = ["1", "2", "3", "4", "5", ".", ".", "X"]


def empty_rows(columns: int):
    return [(number, *("" for _ in range(columns - 1))) for number in ROWS]


def setup_document() -> Document:
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    normal = document.styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    return document


def title(document: Document) -> None:
    paragraph = document.add_paragraph(style="Title")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_after = Pt(8)
    run = paragraph.add_run("Anexo 2 Instrumentos de recolección de datos")
    set_run_font(run, size=16, bold=True, color="000000")


def form(document: Document, heading: str, indicator: str, formula: str, definition: str, source: str, headers, widths) -> None:
    add_heading(document, heading, level=1)
    add_meta_table(
        document,
        indicator,
        formula,
        definition,
        source,
        "Pretest",
    )
    add_data_table(document, headers, empty_rows(len(headers)), widths)


def build() -> None:
    OUT.parent.mkdir(parents=True, exist_ok=True)
    document = setup_document()
    title(document)

    form(
        document,
        "Instrumento de recolección de datos del indicador para el pretest Cantidad demandada",
        "Cantidad demandada",
        "CD = (SI + EN) - SF",
        "CD es cantidad demandada. SI es stock inicial. EN son entradas. SF es stock final.",
        "Kardex y registros de venta de la empresa.",
        ["N°", "SI", "EN", "SF", "CD"],
        [0.7, 1.5, 1.5, 1.5, 1.5],
    )
    add_page_break(document)
    form(
        document,
        "Instrumento de recolección de datos del indicador para el pretest Índice de salida de inventario",
        "Índice de salida de inventario",
        "ISI = CD / (SI + EN)",
        "ISI es índice de salida de inventario. CD es cantidad demandada. SI es stock inicial. EN son entradas.",
        "Kardex y registros de venta de la empresa.",
        ["N°", "CD", "Stock inicial SI", "Entradas EN", "Índice de salida de inventario ISI"],
        [0.55, 1.1, 1.4, 1.35, 2.3],
    )
    add_page_break(document)
    form(
        document,
        "Instrumento de recolección de datos del indicador para el pretest Tasa de quiebre de stock",
        "Tasa de quiebre de stock",
        "TQS = (DQS / DD) x 100",
        "TQS es tasa de quiebre de stock. DQS son días con quiebre de stock. DD son días evaluados.",
        "Kardex o registro diario de productos agotados.",
        ["N°", "DQS", "DD", "TQS"],
        [0.7, 1.8, 1.8, 1.8],
    )
    add_page_break(document)
    form(
        document,
        "Instrumento de recolección de datos del indicador para el pretest Estacionalidad",
        "Estacionalidad",
        "IE = VD / VPM",
        "IE es índice estacional. VD es venta diaria. VPM es venta promedio mensual.",
        "Reportes de venta de la empresa.",
        ["N°", "IE", "VD", "VPM"],
        [0.7, 1.8, 1.8, 1.8],
    )
    add_page_break(document)
    form(
        document,
        "Instrumento de recolección de datos del indicador para el pretest Error de pronóstico",
        "Error de pronóstico",
        "MAPE = (100 / n) x Σ |(DR - DP) / DR|",
        "MAPE es error porcentual absoluto medio. DR es demanda real. DP es demanda pronosticada. n es número de observaciones.",
        "Pronósticos generados y reportes de venta de la empresa.",
        ["N°", "Demanda real DR", "Demanda pronosticada DP", "MAPE"],
        [0.7, 1.7, 2.4, 1.5],
    )
    add_page_break(document)
    form(
        document,
        "Instrumento de recolección de datos de la dimensión para el pretest Volumen de demanda",
        "Cantidad demandada",
        "CD = Σ ventas realizadas",
        "CD es cantidad demandada. Σ es la sumatoria de ventas realizadas.",
        "Reportes de venta de la empresa.",
        ["N°", "Producto", "Cantidad vendida", "Cantidad demandada CD"],
        [0.7, 2.8, 1.6, 1.8],
    )
    add_page_break(document)
    form(
        document,
        "Instrumento de recolección de datos de la dimensión para el pretest Patrón de demanda",
        "Estacionalidad",
        "IE = VP / VPR",
        "IE es índice estacional. VP es venta del período. VPR es venta promedio de referencia.",
        "Reportes de venta de la empresa.",
        ["N°", "Venta del período VP", "Venta promedio VPR", "Índice estacional IE"],
        [0.7, 2.0, 2.0, 1.8],
    )
    add_page_break(document)
    form(
        document,
        "Instrumento de recolección de datos de la dimensión para el pretest Precisión del pronóstico",
        "Error de pronóstico",
        "MAPE = (100 / n) x Σ |(DR - DP) / DR|",
        "MAPE es error porcentual absoluto medio. DR es demanda real. DP es demanda pronosticada. n es número de observaciones.",
        "Pronósticos generados y reportes de venta de la empresa.",
        ["N°", "Demanda real DR", "Demanda pronosticada DP", "MAPE"],
        [0.7, 1.7, 2.4, 1.5],
    )
    document.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
