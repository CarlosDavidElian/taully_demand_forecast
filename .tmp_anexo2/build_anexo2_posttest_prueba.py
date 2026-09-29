"""Construye un par pretest/postest para una prueba controlada de siete días.

Ambos escenarios usan las mismas fechas, productos y demanda de referencia
simulada. Solo cambia la estimación que se compara con esa referencia. Esto
permite comprobar las fórmulas del sistema, pero no representa ventas ni
kardex reales ni demuestra una mejora real de la tienda.
"""

from __future__ import annotations

import csv
import math
import sys
from collections import defaultdict
from datetime import timedelta
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

VENV_PACKAGES = Path(r"C:\taully_demand_forecast\.venv\Lib\site-packages")
if VENV_PACKAGES.exists():
    sys.path.insert(0, str(VENV_PACKAGES))
WORKSPACE = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(WORKSPACE))

from application.services.catalog_service import CatalogService
from application.services.predict_service import PredictService
from config.settings import CATALOG_FILE, DATA_DIR
from infrastructure.readers.excel_reader import ExcelReader
from infrastructure.repositories.catalog_repository import ExcelCatalogRepository
from infrastructure.repositories.csv_repository import CSVDemandRepository
from interfaces.web import _build_seasonal_product_forecasts, _load_product_sale_history

from build_anexo2 import add_data_table, add_heading, add_meta_table, add_note, add_page_break, add_paragraph, set_run_font


ROOT = Path(r"C:\taully_demand_forecast")
OUTPUT_DIR = ROOT / "outputs" / "01a0b651-7620-73f0-8f6a-8b68b14f7a17"
HISTORY = DATA_DIR / "historial_demanda.csv"
TEST_DAYS = 7


def fmt(value: float | int, decimals: int = 0) -> str:
    return f"{value:,.{decimals}f}"


def load_history_totals() -> tuple[dict[str, float], int]:
    totals: defaultdict[str, float] = defaultdict(float)
    dates = set()
    with HISTORY.open(encoding="utf-8-sig", newline="") as stream:
        for row in csv.DictReader(stream):
            totals[str(row["category"]).strip()] += float(row["quantity"])
            dates.add(str(row["date"]))
    return dict(totals), len(dates)


def scenario_label(scenario: str) -> str:
    if scenario == "pretest":
        return "Pretest"
    if scenario == "posttest":
        return "Postest"
    raise ValueError("El escenario debe ser 'pretest' o 'posttest'.")


def output_path(scenario: str) -> Path:
    return OUTPUT_DIR / f"Anexo_2_{scenario_label(scenario)}_evento_prueba_7_dias.docx"


def observed_variation(category: str, day_index: int) -> float:
    """Demanda de referencia común para ambos escenarios simulados."""
    base = (0.04, -0.02, 0.06, -0.04, 0.01, 0.03, -0.03)[day_index]
    scale = {
        "ABARROTES": 0.85,
        "BEBIDAS": 1.00,
        "GOLOSINAS": 0.75,
        "HELADOS": 0.90,
        "LIMPIEZA": 1.10,
    }.get(category, 1.0)
    return base * scale


def baseline_adjustment(category: str, day_index: int) -> float:
    """Desviación controlada de una estimación base usada solo en el pretest."""
    base = (-0.11, 0.09, -0.13, 0.10, -0.08, 0.12, -0.10)[day_index]
    scale = {
        "ABARROTES": 0.85,
        "BEBIDAS": 1.00,
        "GOLOSINAS": 0.75,
        "HELADOS": 0.90,
        "LIMPIEZA": 1.10,
    }.get(category, 1.0)
    return base * scale


def forecast_adjustment(scenario: str, category: str, day_index: int) -> float:
    return baseline_adjustment(category, day_index) if scenario == "pretest" else 0.0


def generate_test_event(scenario: str = "posttest"):
    scenario_label(scenario)
    service = PredictService(CSVDemandRepository())
    predictions = service.predict_future(TEST_DAYS)
    serialized = {
        category: [
            {"date": demand.date.date().isoformat(), "quantity": float(demand.quantity)}
            for demand in demands
        ]
        for category, demands in predictions.items()
    }
    first_date = min(item["date"] for demands in serialized.values() for item in demands)
    catalog = CatalogService(ExcelCatalogRepository(CATALOG_FILE))
    product_history, source_dates = _load_product_sale_history(
        catalog,
        __import__("datetime").date.fromisoformat(first_date) - timedelta(days=1),
    )
    product_predictions = _build_seasonal_product_forecasts(
        serialized,
        product_history,
        source_dates,
        {},
    )

    categories = {}
    product_totals: defaultdict[str, defaultdict[str, dict[str, float]]] = defaultdict(
        lambda: defaultdict(lambda: {"forecast": 0.0, "actual": 0.0, "selection": 0.0})
    )
    for category in sorted(serialized):
        daily = serialized[category]
        forecast_values = [
            round(row["quantity"] * (1 + forecast_adjustment(scenario, category, index)), 2)
            for index, row in enumerate(daily)
        ]
        forecast_total = round(sum(forecast_values), 2)
        actual_values = [
            round(row["quantity"] * (1 + observed_variation(category, index)), 2)
            for index, row in enumerate(daily)
        ]
        actual_total = round(sum(actual_values), 2)
        mape = sum(
            abs(actual - forecast) / actual
            for actual, forecast in zip(actual_values, forecast_values)
            if actual > 0
        ) / len(actual_values) * 100
        wape = sum(abs(actual - forecast) for actual, forecast in zip(actual_values, forecast_values)) / sum(actual_values) * 100
        categories[category] = {
            "forecast_total": forecast_total,
            "actual_total": actual_total,
            "forecast_last": daily[-1]["quantity"],
            "actual_last": actual_values[-1],
            "actual_average": actual_total / TEST_DAYS,
            "mape": mape,
            "wape": wape,
            "first_date": daily[0]["date"],
            "last_date": daily[-1]["date"],
        }

        for day_index, daily_product in enumerate(product_predictions[category]):
            for product in daily_product["products"]:
                name = str(product["name"])
                product_totals[category][name]["forecast"] += round(
                    float(product["quantity"]) * (1 + forecast_adjustment(scenario, category, day_index)), 2
                )
                product_totals[category][name]["selection"] += float(product["quantity"])
                product_totals[category][name]["actual"] += round(
                    float(product["quantity"]) * (1 + observed_variation(category, day_index)), 2
                )

    selected_products = []
    for category in sorted(product_totals):
        ranking = sorted(
            product_totals[category].items(),
            key=lambda item: (-item[1]["selection"], item[0]),
        )[:3]
        for product, totals in ranking:
            cd = max(1, math.ceil(totals["actual"]))
            si = max(3, math.ceil(cd * 0.25))
            available = max(si, math.ceil(cd * 1.10))
            en = available - si
            sf = available - cd
            selected_products.append(
                {
                    "product": product,
                    "category": category,
                    "forecast": round(totals["forecast"], 2),
                    "actual": round(totals["actual"], 2),
                    "cd": cd,
                    "si": si,
                    "en": en,
                    "sf": sf,
                    "isi": cd / available,
                    "dqs": 0,
                    "dd": TEST_DAYS,
                    "tqs": 0.0,
                }
            )

    history_totals, history_days = load_history_totals()
    for category, item in categories.items():
        item["reference_average"] = history_totals[category] / history_days
        item["daily_index"] = item["actual_last"] / item["actual_average"] if item["actual_average"] else 0
        item["pattern_index"] = item["actual_average"] / item["reference_average"] if item["reference_average"] else 0
    return categories, selected_products


def setup_document() -> Document:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    return doc


def test_source(scenario: str) -> str:
    if scenario == "pretest":
        return "Evento de prueba simulado con una estimación base y una demanda de referencia simulada."
    return "Evento de prueba simulado con el pronóstico del modelo activo y una demanda de referencia simulada."


def test_period(categories: dict[str, dict[str, float]], scenario: str) -> str:
    sample = next(iter(categories.values()))
    return f"{scenario_label(scenario)} simulado. Del {sample['first_date'][8:10]}/{sample['first_date'][5:7]}/{sample['first_date'][:4]} al {sample['last_date'][8:10]}/{sample['last_date'][5:7]}/{sample['last_date'][:4]}."


def build_document(scenario: str = "posttest"):
    event_label = scenario_label(scenario)
    categories, products = generate_test_event(scenario)
    period = test_period(categories, scenario)
    out = output_path(scenario)
    out.parent.mkdir(parents=True, exist_ok=True)
    doc = setup_document()

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(8)
    run = title.add_run(f"Anexo 2 {event_label} simulado de siete días para prueba controlada")
    set_run_font(run, size=16, bold=True, color="000000")
    add_paragraph(
        doc,
        f"Este anexo registra el {event_label.lower()} de una prueba controlada de siete días. El pretest y el postest usan exactamente el mismo período, las mismas cinco categorías, los mismos 15 productos y 105 observaciones producto-día. La demanda de referencia y los datos de inventario se simularon para comprobar las fórmulas, las descargas y la presentación de resultados. Solo cambia la estimación comparada. No representa ventas ni kardex reales de la empresa ni demuestra una mejora real de la tienda.",
        size=10.5,
        align=WD_ALIGN_PARAGRAPH.JUSTIFY,
        space_after=8,
    )
    add_heading(doc, "Relación de instrumentos", level=1)
    add_data_table(
        doc,
        ["N", "Instrumento", "Momento", "Fuente"],
        [
            ("1", "Cantidad demandada mediante inventario", f"{event_label} simulado", "Evento de prueba"),
            ("2", "Índice de salida de inventario", f"{event_label} simulado", "Evento de prueba"),
            ("3", "Tasa de quiebre de stock", f"{event_label} simulado", "Evento de prueba"),
            ("4", "Estacionalidad diaria", f"{event_label} simulado", "Evento de prueba"),
            ("5", "Error de pronóstico", f"{event_label} simulado", "Estimación frente a referencia simulada"),
            ("6", "Volumen de demanda mediante ventas", f"{event_label} simulado", "Evento de prueba"),
            ("7", "Patrón de demanda estacional", f"{event_label} simulado", "Evento de prueba e historial"),
            ("8", "Precisión del pronóstico", f"{event_label} simulado", "Estimación frente a referencia simulada"),
        ],
        [0.4, 2.8, 1.1, 2.6],
    )
    add_note(doc, "Los 15 productos corresponden a tres productos con mayor pronóstico en cada una de las cinco categorías. Se conservan sin cambios en ambos escenarios. Los valores de SI, EN, SF y DQS son únicamente datos de prueba.")

    add_page_break(doc)
    add_heading(doc, "Instrumento 1 Cantidad demandada mediante inventario", level=1)
    add_meta_table(doc, "Cantidad demandada", "CD = (SI + EN) - SF", "CD es la cantidad atendida en el evento de prueba. SI es stock inicial, EN son entradas y SF es stock final.", test_source(scenario), period)
    rows = [(str(i), item["product"], item["category"], fmt(item["si"]), fmt(item["en"]), fmt(item["sf"]), fmt(item["cd"])) for i, item in enumerate(products, 1)]
    add_data_table(doc, ["N", "Producto", "Categoría", "SI", "EN", "SF", "CD"], rows, [0.35, 2.55, 0.8, 0.55, 0.65, 0.55, 0.65], numeric_cols=[0, 3, 4, 5, 6])
    add_note(doc, "En cada producto, CD cuadra con SI + EN - SF. Los saldos se generaron solo para ejecutar el evento de prueba.")

    add_page_break(doc)
    add_heading(doc, "Instrumento 2 Índice de salida de inventario", level=1)
    add_meta_table(doc, "Índice de salida de inventario", "ISI = CD / (SI + EN)", "ISI mide la proporción de mercadería disponible que salió durante el evento de prueba.", test_source(scenario), period)
    rows = [(str(i), item["product"], item["category"], fmt(item["cd"]), fmt(item["si"]), fmt(item["en"]), f"{item['isi'] * 100:.2f}%") for i, item in enumerate(products, 1)]
    add_data_table(doc, ["N", "Producto", "Categoría", "CD", "SI", "EN", "ISI"], rows, [0.35, 2.45, 0.8, 0.6, 0.55, 0.65, 0.6], numeric_cols=[0, 3, 4, 5, 6])
    add_note(doc, "ISI se calcula con los datos de inventario simulados de esta prueba. No representa el indicador real de la tienda.")

    add_page_break(doc)
    add_heading(doc, "Instrumento 3 Tasa de quiebre de stock", level=1)
    add_meta_table(doc, "Tasa de quiebre de stock", "TQS = (DQS / DD) x 100", "TQS mide el porcentaje de días con quiebre. DQS son los días sin disponibilidad y DD los días evaluados.", test_source(scenario), period)
    rows = [(str(i), item["product"], item["category"], str(item["dqs"]), str(item["dd"]), f"{item['tqs'] * 100:.2f}%") for i, item in enumerate(products, 1)]
    add_data_table(doc, ["N", "Producto", "Categoría", "DQS", "DD", "TQS"], rows, [0.35, 2.65, 0.9, 0.55, 0.65, 0.6], numeric_cols=[0, 3, 4, 5])
    add_note(doc, "La regla de reposición del evento de prueba mantiene DQS en 0. No debe interpretarse como ausencia de quiebres reales en la tienda.")

    add_page_break(doc)
    add_heading(doc, "Instrumento 4 Estacionalidad diaria", level=1)
    add_meta_table(doc, "Estacionalidad", "IE = VD / VPD", "VD es la demanda simulada del último día del evento y VPD es el promedio diario simulado del periodo.", test_source(scenario), period)
    rows = [(str(i), cat, fmt(data["actual_last"], 2), fmt(data["actual_average"], 2), fmt(data["daily_index"], 2)) for i, (cat, data) in enumerate(sorted(categories.items()), 1)]
    add_data_table(doc, ["N", "Categoría", "VD", "VPD", "IE"], rows, [0.4, 2.1, 1.15, 1.35, 1.1], numeric_cols=[0, 2, 3, 4])
    add_note(doc, "El índice estacional compara el último día simulado con el promedio de los siete días de prueba.")

    error_rows = [(str(i), cat, fmt(data["actual_total"], 2), fmt(data["forecast_total"], 2), f"{data['mape']:.2f}%") for i, (cat, data) in enumerate(sorted(categories.items()), 1)]
    add_page_break(doc)
    add_heading(doc, "Instrumento 5 Error de pronóstico", level=1)
    add_meta_table(doc, "Error de pronóstico", "MAPE = (100 / n) x sumatoria de |(DR - DP) / DR|", "DR es demanda simulada observada, DP es la estimación comparada y n son siete días de prueba.", test_source(scenario), period)
    add_data_table(doc, ["N", "Categoría", "DR", "DP", "MAPE"], error_rows, [0.4, 2.1, 1.0, 1.35, 1.25], numeric_cols=[0, 2, 3, 4])
    add_note(doc, "MAPE se calculó día por día contra las observaciones simuladas y luego se promedió. DR y DP de la tabla son los totales del evento.")

    add_page_break(doc)
    add_heading(doc, "Instrumento 6 Volumen de demanda mediante ventas", level=1)
    add_meta_table(doc, "Cantidad demandada", "CD = sumatoria de cantidades atendidas", "CD es la cantidad simulada atendida por producto durante el evento de prueba.", test_source(scenario), period)
    rows = [(str(i), item["product"], item["category"], fmt(item["cd"]), fmt(item["cd"])) for i, item in enumerate(products, 1)]
    add_data_table(doc, ["N", "Producto", "Categoría", "Cantidad atendida", "CD"], rows, [0.35, 3.4, 0.9, 0.9, 0.65], numeric_cols=[0, 3, 4])
    add_note(doc, "La cantidad atendida es simulada para la prueba. Coincide con CD de cada producto del instrumento 1.")

    add_page_break(doc)
    add_heading(doc, "Instrumento 7 Patrón de demanda estacional", level=1)
    add_meta_table(doc, "Estacionalidad", "IE = VP / VPR", "VP es el promedio diario simulado del evento y VPR es el promedio diario histórico del 1 de abril al 15 de septiembre de 2026.", "Evento de prueba simulado e historial de demanda real consolidado.", period)
    rows = [(str(i), cat, fmt(data["actual_average"], 2), fmt(data["reference_average"], 2), fmt(data["pattern_index"], 2)) for i, (cat, data) in enumerate(sorted(categories.items()), 1)]
    add_data_table(doc, ["N", "Categoría", "VP", "VPR", "IE"], rows, [0.4, 2.1, 1.15, 1.35, 1.1], numeric_cols=[0, 2, 3, 4])
    add_note(doc, "El patrón compara la demanda media simulada de la prueba con la referencia histórica real disponible.")

    add_page_break(doc)
    add_heading(doc, "Instrumento 8 Precisión del pronóstico", level=1)
    add_meta_table(doc, "Precisión del pronóstico", "MAPE = (100 / n) x sumatoria de |(DR - DP) / DR|", "La precisión del evento se revisa al comparar la demanda simulada observada y la estimación para cada día y categoría.", test_source(scenario), period)
    add_data_table(doc, ["N", "Categoría", "DR", "DP", "MAPE"], error_rows, [0.4, 2.1, 1.0, 1.35, 1.25], numeric_cols=[0, 2, 3, 4])
    add_note(doc, "Este resultado confirma el funcionamiento del cálculo en un escenario de prueba. Como DR fue simulado, la diferencia entre pretest y postest no constituye evidencia de una mejora real. Para una evaluación real, DR debe proceder de ventas posteriores a la implementación.")

    doc.save(out)
    print(out)


if __name__ == "__main__":
    build_document("pretest" if "--pretest" in sys.argv else "posttest")
