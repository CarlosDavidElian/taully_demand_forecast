"""Crea Pretest y Postest siguiendo las ocho fichas de Anexo 2.

Los resultados predictivos provienen de una validación histórica rodante de
17 semanas reproducida con los servicios de entrenamiento y predicción de la
aplicación. Las fichas de inventario se calculan sobre el archivo disponible;
el documento no lo presenta como un kardex certificado por la tienda.
"""

from __future__ import annotations

import json
from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT_DIR = ROOT / "outputs" / "01a0b651-7620-73f0-8f6a-8b68b14f7a17"
TEMPLATE = OUTPUT_DIR / "Anexo_2_Instrumentos_de_recoleccion_de_datos_formato.docx"
RESULTS = ROOT / ".tmp_anexo2" / "validacion_rodante_2026-05-20_a_2026-09-15.json"
HISTORY = ROOT / "data" / "historial_demanda.csv"
PRODUCT_MIX = ROOT / "data" / "productos_por_categoria.csv"
INVENTORY = ROOT / "data" / "inventario_producto_Taully_2026-04-01_a_2026-09-15.xlsx"

NAVY = "1F4E78"
PALE_BLUE = "EDF3F8"
BORDER = "D9D9D9"


def set_run_font(run, *, size=8.7, bold=False, italic=False, color="000000"):
    run.font.name = "Arial"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)


def write_paragraph(paragraph, text, *, size=9.3, bold=False, italic=False, align=None, after=3):
    paragraph.clear()
    if align is not None:
        paragraph.alignment = align
    paragraph.paragraph_format.space_after = Pt(after)
    run = paragraph.add_run(str(text))
    set_run_font(run, size=size, bold=bold, italic=italic)


def write_cell(cell, text, *, size=8.4, bold=False, align=None, color="000000"):
    paragraph = cell.paragraphs[0]
    paragraph.clear()
    paragraph.paragraph_format.space_after = Pt(0)
    if align is not None:
        paragraph.alignment = align
    run = paragraph.add_run(str(text))
    set_run_font(run, size=size, bold=bold, color=color)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def delete_until(table, rows):
    while len(table.rows) > rows:
        table._tbl.remove(table.rows[-1]._tr)


def fill_meta(table, indicator, formula, where, moment):
    """Restaura los seis campos exactos de cada ficha del Anexo 2 original."""
    delete_until(table, 6)
    labels = [
        "INDICADOR:",
        "INVESTIGADOR:",
        "LUGAR DE ESTUDIO:",
        "FÓRMULA:",
        "DONDE:",
        "Momento:",
    ]
    values = [
        indicator,
        "Guerrero Mora, Carlos David Elian",
        "Minimarket Taully, Comas",
        formula,
        where,
        moment,
    ]
    for row, label, value in zip(table.rows, labels, values):
        write_cell(row.cells[0], label, size=8.25, bold=True, color="FFFFFF")
        write_cell(row.cells[1], value, size=8.25)


def fill_data_table(table, headers, rows, numeric_columns=()):
    delete_until(table, len(rows) + 1)
    while len(table.rows) < len(rows) + 1:
        table.add_row()
    for index, header in enumerate(headers):
        write_cell(
            table.rows[0].cells[index],
            header,
            size=7.8,
            bold=True,
            align=WD_ALIGN_PARAGRAPH.CENTER,
            color="FFFFFF",
        )
    for row_number, values in enumerate(rows, start=1):
        for index, value in enumerate(values):
            write_cell(
                table.rows[row_number].cells[index],
                value,
                size=8.05,
                align=WD_ALIGN_PARAGRAPH.RIGHT if index in numeric_columns else WD_ALIGN_PARAGRAPH.LEFT,
            )


def fill_unavailable_table(table, headers, message):
    delete_until(table, 2)
    for index, header in enumerate(headers):
        write_cell(
            table.rows[0].cells[index],
            header,
            size=7.8,
            bold=True,
            align=WD_ALIGN_PARAGRAPH.CENTER,
            color="FFFFFF",
        )
    merged = table.rows[1].cells[0]
    for cell in table.rows[1].cells[1:]:
        merged = merged.merge(cell)
    write_cell(merged, message, size=8.0, align=WD_ALIGN_PARAGRAPH.LEFT)


def fmt_number(value, decimals=0):
    return f"{float(value):,.{decimals}f}".replace(",", " ").replace(".", ",")


def fmt_pct(value):
    return f"{float(value):.2f}%".replace(".", ",")


def shading(cell, color):
    tc_pr = cell._tc.get_or_add_tcPr()
    node = tc_pr.find(qn("w:shd"))
    if node is None:
        node = OxmlElement("w:shd")
        tc_pr.append(node)
    node.set(qn("w:fill"), color)


def borders(cell, color=BORDER):
    tc_pr = cell._tc.get_or_add_tcPr()
    node = tc_pr.first_child_found_in("w:tcBorders")
    if node is None:
        node = OxmlElement("w:tcBorders")
        tc_pr.append(node)
    for edge in ("top", "left", "bottom", "right"):
        tag = qn(f"w:{edge}")
        element = node.find(tag)
        if element is None:
            element = OxmlElement(f"w:{edge}")
            node.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "4")
        element.set(qn("w:color"), color)


def add_reproduction_table(document, rows):
    table = document.add_table(rows=1, cols=7)
    table.style = "Table Grid"
    headers = ["N°", "Corte", "Semana evaluada", "MAE", "RMSE", "MAPE", "WAPE"]
    for index, label in enumerate(headers):
        cell = table.rows[0].cells[index]
        shading(cell, NAVY)
        borders(cell)
        write_cell(cell, label, size=7.2, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, color="FFFFFF")
    for row_index, values in enumerate(rows):
        cells = table.add_row().cells
        for index, value in enumerate(values):
            if row_index % 2:
                shading(cells[index], PALE_BLUE)
            borders(cells[index])
            write_cell(
                cells[index],
                value,
                size=7.1,
                align=WD_ALIGN_PARAGRAPH.RIGHT if index >= 3 else WD_ALIGN_PARAGRAPH.CENTER,
            )
    return table


def historical_pattern_rows():
    history = pd.read_csv(HISTORY)
    history["date"] = pd.to_datetime(history["date"])
    final_day = history[history["date"] == "2026-09-15"].set_index("category")["quantity"]
    september_average = history[(history["date"] >= "2026-09-01") & (history["date"] <= "2026-09-15")].groupby("category")["quantity"].mean()
    evaluation_average = history[(history["date"] >= "2026-05-20") & (history["date"] <= "2026-09-15")].groupby("category")["quantity"].mean()
    initial_average = history[(history["date"] >= "2026-04-01") & (history["date"] <= "2026-05-19")].groupby("category")["quantity"].mean()
    daily = [
        (f"{index + 1}\n{category}", fmt_number(final_day[category] / september_average[category], 2), fmt_number(final_day[category]), fmt_number(september_average[category], 2))
        for index, category in enumerate(sorted(final_day.index))
    ]
    period = [
        (f"{index + 1}\n{category}", fmt_number(evaluation_average[category], 2), fmt_number(initial_average[category], 2), fmt_number(evaluation_average[category] / initial_average[category], 2))
        for index, category in enumerate(sorted(evaluation_average.index))
    ]
    return daily, period


def forecast_rows(results, method):
    return [
        (
            f"{index + 1}\n{category}",
            fmt_number(detail["actual_total"]),
            fmt_number(detail[f"{method}_total"], 2),
            fmt_pct(detail[method]["mape"]),
        )
        for index, (category, detail) in enumerate(sorted(results["by_category"].items()))
    ]


def demand_rows():
    """Muestra 15 productos reales: los tres con mayor venta por categoría.

    El Anexo 2 solicita producto, no categoría. La tabla es un extracto
    declarado de 15 productos de los 200 del catálogo histórico completo.
    """
    products = pd.read_csv(PRODUCT_MIX)
    selected = (
        products.sort_values(["category", "historical_quantity", "product"], ascending=[True, False, True])
        .groupby("category", sort=True)
        .head(3)
        .sort_values(["category", "historical_quantity", "product"], ascending=[True, False, True])
        .reset_index(drop=True)
    )
    return [
        (
            str(index + 1),
            f"{row.product}\n{row.category}",
            fmt_number(row.historical_quantity),
            fmt_number(row.historical_quantity),
        )
        for index, row in enumerate(selected.itertuples(index=False))
    ]


def inventory_rows():
    """Construye las tres tablas de inventario desde todos los registros disponibles.

    SI, EN y SF se acumulan por categoría en los registros producto-fecha.
    DQS cuenta fechas de una categoría que tienen al menos una alerta de
    ``Dias_Sin_Stock``. Esta definición evita contar varias veces la misma
    fecha cuando se registran varios productos de una categoría.
    """
    inventory = pd.read_excel(INVENTORY)
    inventory["Fecha"] = pd.to_datetime(inventory["Fecha"])
    evaluated_days = int(inventory["Fecha"].nunique())
    if evaluated_days != 168:
        raise ValueError(f"Se esperaban 168 días en inventario y se encontraron {evaluated_days}.")

    demand = []
    index = []
    stockouts = []
    for number, (category, group) in enumerate(inventory.groupby("Categoría", sort=True), start=1):
        stock_initial = int(group["Stock_Inicial"].sum())
        entries = int(group["Entradas"].sum())
        stock_final = int(group["Stock_Final"].sum())
        demand_formula = stock_initial + entries - stock_final
        isi = 100 * demand_formula / (stock_initial + entries)
        stockout_days = int(group.loc[group["Dias_Sin_Stock"] > 0, "Fecha"].nunique())
        tqs = 100 * stockout_days / evaluated_days
        label = f"{number}\n{category}"
        demand.append(
            (label, fmt_number(stock_initial), fmt_number(entries), fmt_number(stock_final), fmt_number(demand_formula))
        )
        index.append(
            (label, fmt_number(demand_formula), fmt_number(stock_initial), fmt_number(entries), fmt_pct(isi))
        )
        stockouts.append((label, fmt_number(stockout_days), fmt_number(evaluated_days), fmt_pct(tqs)))
    return demand, index, stockouts


def reproduction_rows(results, method):
    return [
        (
            fold["number"],
            pd.Timestamp(fold["cutoff_date"]).strftime("%d/%m/%Y"),
            f"{pd.Timestamp(fold['test_start']).strftime('%d/%m')}–{pd.Timestamp(fold['test_end']).strftime('%d/%m/%Y')}",
            fmt_number(fold["global_metrics"][method]["mae"], 2),
            fmt_number(fold["global_metrics"][method]["rmse"], 2),
            fmt_pct(fold["global_metrics"][method]["mape"]),
            fmt_pct(fold["global_metrics"][method]["wape"]),
        )
        for fold in results["folds"]
    ]


def set_heading_paragraph(paragraph, text):
    write_paragraph(paragraph, text, size=10.2, bold=True, after=7)


def build(label, results):
    is_pretest = label == "Pretest"
    method_key = "pms" if is_pretest else "model"
    method_name = "Promedio móvil simple recursivo de siete días (PMS-7)" if is_pretest else "Modelo de pronóstico del programa"
    output = OUTPUT_DIR / f"Anexo_2_{label}_validacion_historica_rodante_hasta_2026-09-15.docx"
    document = Document(TEMPLATE)

    write_paragraph(
        document.paragraphs[0],
        f"Anexo 2 Instrumentos de recolección de datos {label}",
        size=15,
        bold=True,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        after=8,
    )
    method_note = (
        "Pretest: promedio móvil simple recursivo de siete días (PMS-7), cálculo externo reproducible."
        if is_pretest
        else "Postest: pronósticos obtenidos con los servicios de entrenamiento y predicción del programa."
    )
    note = document.paragraphs[1].insert_paragraph_before(
        "Nota metodológica. Datos de ventas: 168 reportes consolidados en 840 registros, cinco categorías, del 01/04 al 15/09/2026. "
        "Se comparan 17 cortes históricos de siete días, con pruebas del 20/05 al 15/09/2026. "
        "Los Instrumentos 1 a 3 se calcularon con los 11 748 registros producto-fecha del archivo de inventario disponible; "
        "son cálculos sobre esa base y no una certificación de kardex emitida por la tienda. "
        f"{method_note}"
    )
    write_paragraph(note, note.text, size=8.25, italic=True, align=WD_ALIGN_PARAGRAPH.JUSTIFY, after=5)
    for paragraph in document.paragraphs[1:]:
        if "pretest" in paragraph.text.lower():
            text = paragraph.text.replace("pretest", label.lower()).replace("Pretest", label)
            set_heading_paragraph(paragraph, text)

    daily_pattern, period_pattern = historical_pattern_rows()
    forecast = forecast_rows(results, method_key)
    demand = demand_rows()
    inventory_demand, inventory_index, inventory_stockouts = inventory_rows()

    fill_meta(
        document.tables[0],
        "Cantidad Demandada",
        "CD = (SI + EN) - SF",
        "CD = Cantidad Demandada calculada\nSI = Stock Inicial acumulado\nEN = Entradas acumuladas\nSF = Stock Final acumulado",
        label,
    )
    fill_data_table(
        document.tables[1],
        ["N° / categoría", "SI", "EN", "SF", "CD"],
        inventory_demand,
        numeric_columns=(1, 2, 3, 4),
    )

    fill_meta(
        document.tables[2],
        "Índice de salida de inventario",
        "ISI = [CD / (SI + EN)] x 100",
        "ISI = Índice de Salida de Inventario\nCD = Cantidad Demandada calculada\nSI = Stock Inicial acumulado\nEN = Entradas acumuladas",
        label,
    )
    fill_data_table(
        document.tables[3],
        ["N° / categoría", "CD", "Stock inicial SI", "Entradas EN", "ISI"],
        inventory_index,
        numeric_columns=(1, 2, 3, 4),
    )

    fill_meta(
        document.tables[4],
        "Tasa de quiebre de stock",
        "TQS = (DQS / DD) x 100",
        "TQS = Tasa de Quiebre de Stock\nDQS = Fechas con al menos una alerta de stock\nDD = Días Evaluados",
        label,
    )
    fill_data_table(
        document.tables[5],
        ["N° / categoría", "DQS", "DD", "TQS"],
        inventory_stockouts,
        numeric_columns=(1, 2, 3),
    )

    fill_meta(
        document.tables[6],
        "Estacionalidad",
        "IE = VD / VPM",
        "IE = Índice Estacional\nVD = Venta Diaria\nVPM = Venta Promedio Mensual",
        label,
    )
    fill_data_table(
        document.tables[7],
        ["N° / categoría", "IE", "VD", "VPM"],
        daily_pattern,
        numeric_columns=(1, 2, 3),
    )

    fill_meta(
        document.tables[8],
        "Error de pronóstico",
        "MAPE = (100 / n) x sumatoria de |(DR - DP) / DR|",
        "MAPE = Error Porcentual Absoluto Medio\nDR = Demanda Real\nDP = Demanda Pronosticada\nn = Número de observaciones",
        label,
    )
    fill_data_table(
        document.tables[9],
        ["N° / categoría", "Demanda Real (DR)", "Demanda Pronosticada (DP)", "MAPE"],
        forecast,
        numeric_columns=(1, 2, 3),
    )

    fill_meta(
        document.tables[10],
        "Cantidad Demandada",
        "CD = Σ Ventas realizadas",
        "CD = Cantidad Demandada\nΣ = Sumatoria de ventas realizadas",
        label,
    )
    fill_data_table(
        document.tables[11],
        ["N°", "Producto", "Cantidad vendida", "Cantidad demandada (CD)"],
        demand,
        numeric_columns=(2, 3),
    )

    fill_meta(
        document.tables[12],
        "Estacionalidad",
        "IE = VP / VPR",
        "IE = Índice Estacional\nVP = Venta del Período\nVPR = Venta Promedio de Referencia",
        label,
    )
    fill_data_table(
        document.tables[13],
        ["N° / categoría", "Venta del Período (VP)", "Venta Promedio (VPR)", "Índice Estacional (IE)"],
        period_pattern,
        numeric_columns=(1, 2, 3),
    )

    fill_meta(
        document.tables[14],
        "Precisión del pronóstico",
        "MAPE = (100 / n) x sumatoria de |(DR - DP) / DR|",
        "MAPE = Error Porcentual Absoluto Medio\nDR = Demanda Real\nDP = Demanda Pronosticada\nn = Número de observaciones\nΣ = Sumatoria de errores porcentuales absolutos",
        label,
    )
    fill_data_table(
        document.tables[15],
        ["N° / categoría", "Demanda Real (DR)", "Demanda Pronosticada (DP)", "MAPE"],
        forecast,
        numeric_columns=(1, 2, 3),
    )

    metrics = results["overall"][method_key]
    document.add_page_break()
    heading = document.add_paragraph(style="Heading 1")
    write_paragraph(heading, "Resultado global de la validación histórica", size=11, bold=True, after=4)
    summary = document.add_paragraph()
    write_paragraph(
        summary,
        f"{label}: {method_name}. Se evaluaron 595 observaciones (5 categorías por 7 días por 17 cortes), con una demanda real acumulada de {fmt_number(results['overall']['actual_total'])} unidades. MAE: {fmt_number(metrics['mae'], 2)}; RMSE: {fmt_number(metrics['rmse'], 2)}; MAPE: {fmt_pct(metrics['mape'])}; WAPE: {fmt_pct(metrics['wape'])}.",
        size=9.2,
        after=4,
    )
    if not is_pretest:
        pms_metrics = results["overall"]["pms"]
        mape_difference = pms_metrics["mape"] - metrics["mape"]
        wape_difference = pms_metrics["wape"] - metrics["wape"]
        weekly_wins = sum(
            fold["global_metrics"]["model"]["mape"] < fold["global_metrics"]["pms"]["mape"]
            for fold in results["folds"]
        )
        note = document.add_paragraph()
        write_paragraph(
            note,
            f"Comparado con el pretest, el MAPE global es {fmt_number(mape_difference, 2)} puntos porcentuales menor y el WAPE global es {fmt_number(wape_difference, 2)} puntos menor. El modelo obtuvo menor MAPE en {weekly_wins} de las 17 semanas; no fue mejor en todas las semanas, incluido el último corte.",
            size=9.2,
            after=4,
        )

    heading = document.add_paragraph(style="Heading 1")
    write_paragraph(heading, "Registro de reproducción en el programa", size=11, bold=True, after=3)
    instructions = document.add_paragraph()
    write_paragraph(
        instructions,
        "Para reproducir cada fila del postest: seleccione horizonte de 7 días, ingrese la fecha de Corte de prueba, pulse Entrenar modelo y luego Generar pronóstico. Las métricas de esta tabla se calcularon después al contrastar esos pronósticos con las ventas reales; la interfaz muestra por separado su validación interna de entrenamiento.",
        size=8.8,
        after=4,
    )
    add_reproduction_table(document, reproduction_rows(results, method_key))
    source_note = document.add_paragraph()
    write_paragraph(
        source_note,
        "Nota: pretest y postest usan las mismas semanas, categorías y ventas reales para comparar métodos sin mezclar cambios de demanda entre períodos. Por esa razón, los instrumentos de inventario muestran las mismas cifras en ambos momentos: usan la misma base de 11 748 registros producto-fecha. En los Instrumentos 1 y 2, CD se obtiene con SI + EN - SF; en el Instrumento 6, CD suma ventas registradas, por lo que pueden diferir si el saldo final de un registro quedó limitado a cero. El Instrumento 6 muestra 15 productos reales, los tres con mayor venta de cada categoría, de un catálogo histórico de 200 productos; sus cantidades corresponden al historial completo del 01/04 al 15/09/2026 y no representan la totalidad del catálogo. Esta validación histórica mide precisión predictiva; no demuestra por sí sola un cambio causal de inventario de la tienda.",
        size=8.5,
        italic=True,
        after=2,
    )

    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    document.save(output)
    return output


def main():
    results = json.loads(RESULTS.read_text(encoding="utf-8"))
    print(build("Pretest", results))
    print(build("Postest", results))


if __name__ == "__main__":
    main()
