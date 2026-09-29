"""Completa las ocho fichas de Anexo 2 con la validación histórica real.

El documento conserva la estructura de los instrumentos de la tesis. Las
fichas de pretest y postest usan las mismas ventas reales de siete días; solo
cambia el método de pronóstico comparado.
"""

from __future__ import annotations

import json
from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
TEMPLATE = ROOT / "outputs" / "01a0b651-7620-73f0-8f6a-8b68b14f7a17" / "Anexo_2_Instrumentos_de_recoleccion_de_datos_formato.docx"
RESULTS = ROOT / ".tmp_anexo2" / "backtest_real_2026-05-19_a_2026-05-25.json"
HISTORY = ROOT / "data" / "historial_demanda.csv"
OUTPUT_DIR = ROOT / "outputs" / "01a0b651-7620-73f0-8f6a-8b68b14f7a17"


def set_run_font(run, size=9.0, bold=False, italic=False):
    run.font.name = "Arial"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = RGBColor(0, 0, 0)


def write_cell(cell, text, *, bold=False, align=None, size=8.6):
    paragraph = cell.paragraphs[0]
    paragraph.clear()
    paragraph.paragraph_format.space_after = Pt(0)
    if align is not None:
        paragraph.alignment = align
    run = paragraph.add_run(str(text))
    set_run_font(run, size=size, bold=bold)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def write_paragraph(paragraph, text, *, bold=False, italic=False, align=None, size=10.0):
    paragraph.clear()
    paragraph.paragraph_format.space_after = Pt(4)
    if align is not None:
        paragraph.alignment = align
    run = paragraph.add_run(text)
    set_run_font(run, size=size, bold=bold, italic=italic)


def delete_until(table, total_rows):
    while len(table.rows) > total_rows:
        table._tbl.remove(table.rows[-1]._tr)


def fill_data_table(table, headers, rows):
    delete_until(table, len(rows) + 1)
    for index, header in enumerate(headers):
        write_cell(table.rows[0].cells[index], header, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, size=8.0)
    for row_index, values in enumerate(rows, start=1):
        for index, value in enumerate(values):
            alignment = WD_ALIGN_PARAGRAPH.RIGHT if index > 0 else WD_ALIGN_PARAGRAPH.LEFT
            write_cell(table.rows[row_index].cells[index], value, align=alignment, size=8.2)


def fill_unavailable_table(table, headers, message):
    delete_until(table, 2)
    for index, header in enumerate(headers):
        write_cell(table.rows[0].cells[index], header, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, size=8.0)
    merged = table.rows[1].cells[0]
    for cell in table.rows[1].cells[1:]:
        merged = merged.merge(cell)
    write_cell(merged, message, align=WD_ALIGN_PARAGRAPH.LEFT, size=8.1)


def fill_meta(table, indicator, formula, definition, source, period):
    values = [
        indicator,
        "Guerrero Mora, Carlos David Elian",
        "Minimarket Taully, Comas",
        formula,
        definition,
        source,
        period,
    ]
    for row, value in zip(table.rows, values):
        write_cell(row.cells[1], value, size=8.7)


def num(value, decimals=0):
    return f"{float(value):,.{decimals}f}".replace(",", " ").replace(".", ",")


def percentage(value):
    return f"{float(value):.2f}%".replace(".", ",")


def seasonal_data():
    history = pd.read_csv(HISTORY)
    history["date"] = pd.to_datetime(history["date"])
    april = history[(history["date"] >= "2026-04-01") & (history["date"] <= "2026-04-30")].groupby("category")["quantity"].mean()
    final_day = history[history["date"] == "2026-05-25"].set_index("category")["quantity"]
    period = history[(history["date"] >= "2026-05-19") & (history["date"] <= "2026-05-25")].groupby("category")["quantity"].mean()
    reference = history[(history["date"] >= "2026-04-01") & (history["date"] <= "2026-05-18")].groupby("category")["quantity"].mean()
    daily = [(category, num(final_day[category]), num(april[category], 2), num(final_day[category] / april[category], 2)) for category in sorted(april.index)]
    pattern = [(category, num(period[category], 2), num(reference[category], 2), num(period[category] / reference[category], 2)) for category in sorted(period.index)]
    return daily, pattern


def metric_rows(data, forecast_key, metric_key):
    return [
        (
            category,
            num(sum(detail["actual"])),
            num(sum(detail[forecast_key]), 2),
            percentage(detail[metric_key]["mape"]),
        )
        for category, detail in sorted(data["categories"].items())
    ]


def demand_rows(data):
    return [
        (category, num(sum(detail["actual"])), num(sum(detail["actual"])))
        for category, detail in sorted(data["categories"].items())
    ]


def build(kind, data):
    is_pretest = kind == "pretest"
    label = "Pretest" if is_pretest else "Postest"
    forecast_key = "baseline_pms_7" if is_pretest else "model_forecast"
    metric_key = "baseline_metrics" if is_pretest else "model_metrics"
    metric = data["all_category_day_observations"][metric_key]
    output = OUTPUT_DIR / f"Anexo_2_{label}_instrumentos_7_dias_validacion_historica.docx"
    document = Document(TEMPLATE)

    write_paragraph(document.paragraphs[0], f"Anexo 2 Instrumentos de recolección de datos {label}", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, size=15)
    introduction = document.paragraphs[1].insert_paragraph_before(
        "Las ocho fichas se aplican a una validación histórica de siete días. El pretest usa el promedio móvil simple de siete días y el postest usa el modelo seleccionado por el programa. Ambos métodos se contrastan con las mismas ventas reales del 19 al 25 de mayo de 2026, usando datos de entrenamiento solo hasta el 18 de mayo de 2026."
    )
    introduction.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    introduction.paragraph_format.space_after = Pt(8)
    for run in introduction.runs:
        set_run_font(run, size=9.5)

    for paragraph in document.paragraphs:
        if "para el pretest" in paragraph.text.lower():
            replacement = paragraph.text.replace("pretest", label.lower()).replace("Pretest", label)
            write_paragraph(paragraph, replacement, bold=True, size=10.5)

    period = f"{label}. Validación histórica del 19 al 25 de mayo de 2026. Corte de entrenamiento: 18 de mayo de 2026."
    method_source = (
        "Cálculo externo transparente del promedio móvil simple de siete días y reportes diarios de ventas."
        if is_pretest
        else "Pronósticos obtenidos con el mismo servicio de entrenamiento y predicción del programa, y reportes diarios de ventas."
    )
    daily_rows, pattern_rows = seasonal_data()
    errors = metric_rows(data, forecast_key, metric_key)
    demand = demand_rows(data)

    # Instrumento 1: no se imputan saldos de kardex inexistentes.
    fill_meta(
        document.tables[0],
        "Cantidad demandada",
        "CD = (SI + EN) - SF",
        "CD se obtiene del movimiento real de inventario cuando existe un kardex completo con stock inicial, entradas y stock final.",
        "Kardex real diario requerido. El archivo disponible no contiene una serie continua verificable.",
        period,
    )
    fill_unavailable_table(
        document.tables[1],
        ["N°", "SI", "EN", "SF", "CD"],
        "No se registra este cálculo: no se cuenta con kardex real continuo para los cinco grupos y los siete días. La cantidad demandada comprobable por ventas se presenta en el Instrumento 6.",
    )

    # Instrumento 2.
    fill_meta(
        document.tables[2],
        "Índice de salida de inventario",
        "ISI = CD / (SI + EN)",
        "ISI expresa la proporción de la mercadería disponible que salió en el período.",
        "Kardex real diario requerido. No se infiere ISI a partir de ventas aisladas.",
        period,
    )
    fill_unavailable_table(
        document.tables[3],
        ["N°", "CD", "Stock inicial SI", "Entradas EN", "Índice de salida de inventario ISI"],
        "No se calcula ISI con el inventario disponible, porque sus saldos no forman un kardex continuo por producto y fecha. Reportar un porcentaje en estas condiciones sería no verificable.",
    )

    # Instrumento 3.
    fill_meta(
        document.tables[4],
        "Tasa de quiebre de stock",
        "TQS = (DQS / DD) x 100",
        "TQS mide los días con quiebre de stock dentro de los días evaluados.",
        "Registro diario completo de productos agotados requerido. No se puede sustituir por ausencia de venta.",
        period,
    )
    fill_unavailable_table(
        document.tables[5],
        ["N°", "DQS", "DD", "TQS"],
        "No se calcula TQS: el archivo no registra la disponibilidad de todo el catálogo en cada día. Un producto sin venta no demuestra que se encontraba disponible ni agotado.",
    )

    # Instrumento 4.
    fill_meta(
        document.tables[6],
        "Estacionalidad",
        "IE = VD / VPM",
        "IE compara la venta diaria final del período con la venta promedio diaria del mes completo anterior.",
        "Reportes diarios de venta consolidados. VPM corresponde al promedio diario de abril de 2026.",
        period,
    )
    fill_data_table(document.tables[7], ["Categoría", "VD", "VPM", "IE"], daily_rows)

    # Instrumento 5.
    fill_meta(
        document.tables[8],
        "Error de pronóstico",
        "MAPE = (100 / n) x sumatoria de |(DR - DP) / DR|",
        "MAPE mide el error porcentual medio al comparar la demanda real y la demanda pronosticada para cada categoría y día.",
        method_source,
        period,
    )
    fill_data_table(document.tables[9], ["Categoría", "Demanda real DR", "Demanda pronosticada DP", "MAPE"], errors)

    # Instrumento 6.
    fill_meta(
        document.tables[10],
        "Cantidad demandada",
        "CD = sumatoria de ventas realizadas",
        "CD corresponde a las unidades vendidas consolidadas durante los siete días evaluados.",
        "Reportes diarios de venta consolidados por categoría.",
        period,
    )
    fill_data_table(document.tables[11], ["Categoría", "Cantidad vendida", "Cantidad demandada CD", "Período"], [(cat, qty, cd, "7 días") for cat, qty, cd in demand])

    # Instrumento 7.
    fill_meta(
        document.tables[12],
        "Estacionalidad",
        "IE = VP / VPR",
        "IE compara el promedio diario vendido en el período con el promedio diario de referencia anterior a la fecha de corte.",
        "Reportes diarios de venta consolidados. VPR usa del 01 de abril al 18 de mayo de 2026.",
        period,
    )
    fill_data_table(document.tables[13], ["Categoría", "Venta del período VP", "Venta promedio VPR", "Índice estacional IE"], pattern_rows)

    # Instrumento 8.
    fill_meta(
        document.tables[14],
        "Precisión del pronóstico",
        "MAPE = (100 / n) x sumatoria de |(DR - DP) / DR|",
        "La precisión se determina comparando el pronóstico y la demanda real en las 35 observaciones de categoría y día.",
        method_source,
        period,
    )
    fill_data_table(document.tables[15], ["Categoría", "Demanda real DR", "Demanda pronosticada DP", "MAPE"], errors)

    summary = document.add_paragraph()
    summary.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    summary.paragraph_format.space_before = Pt(6)
    summary.paragraph_format.space_after = Pt(3)
    summary_text = (
        f"Resultado global del {label.lower()}: MAE {num(metric['mae'], 2)}, RMSE {num(metric['rmse'], 2)}, MAPE {percentage(metric['mape'])} y WAPE {percentage(metric['wape'])}, calculados sobre 35 observaciones."
    )
    run = summary.add_run(summary_text)
    set_run_font(run, size=9.5, bold=True)
    if not is_pretest:
        improvement = data["all_category_day_observations"]["baseline_metrics"]["mape"] - metric["mape"]
        note = document.add_paragraph()
        note.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        note.paragraph_format.space_after = Pt(4)
        run = note.add_run(f"Comparado con el pretest, el MAPE disminuye {num(improvement, 2)} puntos porcentuales. Esta diferencia corresponde a una validación histórica de métodos y no a una medición causal de inventario.")
        set_run_font(run, size=9.5)

    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    document.save(output)
    return output


def main():
    data = json.loads(RESULTS.read_text(encoding="utf-8"))
    print(build("pretest", data))
    print(build("postest", data))


if __name__ == "__main__":
    main()
