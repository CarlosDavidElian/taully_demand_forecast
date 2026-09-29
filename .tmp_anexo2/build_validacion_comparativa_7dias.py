"""Genera dos fichas de validación histórica con datos reales de siete días.

Las fichas comparan el promedio móvil simple con el modelo de machine learning
en la misma ventana histórica. No modifican archivos de la aplicación.
"""

from __future__ import annotations

import json
from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
RESULTS = ROOT / ".tmp_anexo2" / "backtest_real_2026-05-19_a_2026-05-25.json"
HISTORY = ROOT / "data" / "historial_demanda.csv"
OUTPUT_DIR = ROOT / "outputs" / "01a0b651-7620-73f0-8f6a-8b68b14f7a17"

NAVY = "1F4E78"
PALE_BLUE = "EEF4F8"
LIGHT_GRAY = "F6F8FA"
BORDER = "D9D9D9"
TEXT = "000000"
MONTHS_ES = {
    1: "enero", 2: "febrero", 3: "marzo", 4: "abril", 5: "mayo", 6: "junio",
    7: "julio", 8: "agosto", 9: "septiembre", 10: "octubre", 11: "noviembre", 12: "diciembre",
}


def set_font(run, size=10.0, bold=False, color=TEXT, italic=False):
    run.font.name = "Arial"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)


def set_cell_shading(cell, color):
    props = cell._tc.get_or_add_tcPr()
    shading = props.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        props.append(shading)
    shading.set(qn("w:fill"), color)


def set_cell_border(cell, color=BORDER):
    props = cell._tc.get_or_add_tcPr()
    borders = props.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        props.append(borders)
    for edge in ("top", "left", "bottom", "right"):
        tag = qn(f"w:{edge}")
        element = borders.find(tag)
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "4")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_cell_margin(cell, top=90, start=90, bottom=90, end=90):
    props = cell._tc.get_or_add_tcPr()
    margins = props.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        props.append(margins)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    props = row._tr.get_or_add_trPr()
    element = OxmlElement("w:tblHeader")
    element.set(qn("w:val"), "true")
    props.append(element)


def fmt_number(value, decimals=0):
    value = float(value)
    text = f"{value:,.{decimals}f}"
    return text.replace(",", " ").replace(".", ",")


def fmt_percent(value):
    return f"{float(value):.2f}".replace(".", ",") + "%"


def fmt_long_date(value):
    stamp = pd.Timestamp(value)
    return f"{stamp.day:02d} de {MONTHS_ES[stamp.month]} de {stamp.year}"


def set_width(cell, inches):
    cell.width = Inches(inches)
    tc_width = cell._tc.get_or_add_tcPr().first_child_found_in("w:tcW")
    if tc_width is None:
        tc_width = OxmlElement("w:tcW")
        cell._tc.get_or_add_tcPr().append(tc_width)
    tc_width.set(qn("w:w"), str(int(inches * 1440)))
    tc_width.set(qn("w:type"), "dxa")


def add_table(doc, headers, rows, widths, numeric_columns=(), font_size=8.5):
    table = doc.add_table(rows=1, cols=len(headers))
    table.autofit = False
    table.style = "Table Grid"
    header = table.rows[0]
    set_repeat_table_header(header)
    for index, text in enumerate(headers):
        cell = header.cells[index]
        set_width(cell, widths[index])
        set_cell_shading(cell, NAVY)
        set_cell_border(cell)
        set_cell_margin(cell)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_after = Pt(0)
        run = paragraph.add_run(str(text))
        set_font(run, size=font_size, bold=True, color="FFFFFF")
    for row_index, row_values in enumerate(rows):
        cells = table.add_row().cells
        for index, value in enumerate(row_values):
            cell = cells[index]
            set_width(cell, widths[index])
            if row_index % 2:
                set_cell_shading(cell, PALE_BLUE)
            set_cell_border(cell)
            set_cell_margin(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            paragraph = cell.paragraphs[0]
            paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT if index in numeric_columns else WD_ALIGN_PARAGRAPH.LEFT
            paragraph.paragraph_format.space_after = Pt(0)
            run = paragraph.add_run(str(value))
            set_font(run, size=font_size)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def add_heading(doc, text, level=1):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(8 if level == 1 else 5)
    paragraph.paragraph_format.space_after = Pt(4)
    run = paragraph.add_run(text)
    set_font(run, size=12 if level == 1 else 10.5, bold=True)
    return paragraph


def add_body(doc, text, bold_prefix=None):
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    paragraph.paragraph_format.space_after = Pt(5)
    paragraph.paragraph_format.line_spacing = 1.08
    if bold_prefix and text.startswith(bold_prefix):
        run = paragraph.add_run(bold_prefix)
        set_font(run, size=9.8, bold=True)
        run = paragraph.add_run(text[len(bold_prefix):])
        set_font(run, size=9.8)
    else:
        run = paragraph.add_run(text)
        set_font(run, size=9.8)
    return paragraph


def configure_document():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.7)
    section.right_margin = Inches(0.7)
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    normal.font.size = Pt(9.8)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    return doc


def controlled_period_rows(result):
    cutoff = pd.Timestamp(result["cutoff_date"])
    start = pd.Timestamp(result["period_start"])
    end = pd.Timestamp(result["period_end"])
    training_start = pd.Timestamp(result["training_start"])
    return [
        ("Corte de entrenamiento", fmt_long_date(cutoff)),
        ("Período de prueba", f"{start.day:02d} al {fmt_long_date(end)}"),
        ("Entrenamiento disponible", f"{fmt_long_date(training_start)} al {fmt_long_date(cutoff)}"),
        ("Observaciones evaluadas", "35 observaciones: 5 categorías por 7 días"),
        ("Demanda real del período", f"{fmt_number(sum(result['overall']['actual']))} unidades"),
        ("Fuente", "Reportes diarios de ventas consolidados por el programa"),
    ]


def pattern_rows(result):
    history = pd.read_csv(HISTORY)
    history["date"] = pd.to_datetime(history["date"])
    start = pd.Timestamp(result["period_start"])
    end = pd.Timestamp(result["period_end"])
    cutoff = pd.Timestamp(result["cutoff_date"])
    event = history[(history["date"] >= start) & (history["date"] <= end)]
    reference = history[(history["date"] >= pd.Timestamp(result["training_start"])) & (history["date"] <= cutoff)]
    averages = event.groupby("category")["quantity"].mean()
    references = reference.groupby("category")["quantity"].mean()
    return [
        (category, fmt_number(averages[category], 2), fmt_number(references[category], 2), fmt_number(averages[category] / references[category], 2))
        for category in sorted(averages.index)
    ]


def category_rows(result, forecast_key, metrics_key):
    values = []
    for category, detail in sorted(result["categories"].items()):
        item_metrics = detail[metrics_key]
        values.append(
            (
                category,
                fmt_number(sum(detail["actual"])),
                fmt_number(sum(detail[forecast_key]), 2),
                fmt_number(item_metrics["mae"], 2),
                fmt_number(item_metrics["rmse"], 2),
                fmt_percent(item_metrics["mape"]),
            )
        )
    return values


def daily_rows(result, forecast_key):
    return [
        (
            pd.Timestamp(date_value).strftime("%d/%m/%Y"),
            fmt_number(actual),
            fmt_number(prediction, 2),
            fmt_number(abs(actual - prediction), 2),
        )
        for date_value, actual, prediction in zip(result["dates"], result["overall"]["actual"], result["overall"][forecast_key])
    ]


def write_document(kind, result):
    if kind == "pms":
        title = "Anexo 2 Ficha A Modelo de referencia PMS"
        output = OUTPUT_DIR / "Anexo_2_Ficha_A_PMS_validacion_historica_7_dias.docx"
        method = "Promedio móvil simple recursivo de siete días"
        forecast_key = "baseline_pms_7"
        metrics_key = "baseline_metrics"
        metric_summary = result["all_category_day_observations"]["baseline_metrics"]
        conclusion = (
            "La ficha A constituye la línea base. Sus métricas se calculan sobre las 35 observaciones de categoría y día. "
            "Se usa para contrastar el modelo de machine learning contra el mismo conjunto de ventas reales."
        )
    else:
        title = "Anexo 2 Ficha B Modelo de machine learning"
        output = OUTPUT_DIR / "Anexo_2_Ficha_B_modelo_ML_validacion_historica_7_dias.docx"
        method = "Modelo seleccionado por el programa, entrenado únicamente hasta la fecha de corte"
        forecast_key = "model_forecast"
        metrics_key = "model_metrics"
        metric_summary = result["all_category_day_observations"]["model_metrics"]
        difference = metric_summary["mape"] - result["all_category_day_observations"]["baseline_metrics"]["mape"]
        direction = "menor" if difference < 0 else "mayor"
        assessment = "sí evidencia una mejora del modelo frente al método base en este período." if difference < 0 else "no permite afirmar una mejora del modelo frente al método base en este período."
        conclusion = (
            f"En esta ventana el MAPE del modelo es {fmt_percent(metric_summary['mape'])}, {fmt_percent(abs(difference))} puntos porcentuales {direction} que el PMS. "
            f"Por tanto, esta evidencia {assessment}"
        )

    doc = configure_document()
    heading = doc.add_paragraph(style="Title")
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    heading.paragraph_format.space_after = Pt(3)
    run = heading.add_run(title)
    set_font(run, size=16, bold=True)
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(9)
    run = subtitle.add_run("Evaluación comparativa mediante validación histórica de siete días")
    set_font(run, size=10.5, italic=True)

    add_body(
        doc,
        "Las dos fichas comparan métodos de pronóstico usando exactamente la misma ventana de ventas reales. "
        "La comparación es un backtesting histórico con orden cronológico, no una medición causal antes y después de implementar el sistema.",
    )

    add_heading(doc, "Control del período evaluado")
    add_table(doc, ["Elemento", "Registro"], controlled_period_rows(result), [2.0, 4.85], font_size=8.8)

    add_heading(doc, "Método aplicado")
    method_rows = [("Método", method)]
    if kind == "ml":
        method_rows.extend(
            (category, f"{detail['trained_model']} | WAPE de validación interna: {fmt_percent(detail['training_wape'])}")
            for category, detail in sorted(result["categories"].items())
        )
    else:
        method_rows.append(("Regla de proyección", "Cada día proyectado usa el promedio de las siete observaciones disponibles más recientes; los días futuros se proyectan de forma recursiva."))
    add_table(doc, ["Elemento", "Aplicación"], method_rows, [2.0, 4.85], font_size=8.5)

    add_heading(doc, "Resultado general")
    general_rows = [
        ("MAE", fmt_number(metric_summary["mae"], 2)),
        ("RMSE", fmt_number(metric_summary["rmse"], 2)),
        ("MAPE", fmt_percent(metric_summary["mape"])),
        ("WAPE", fmt_percent(metric_summary["wape"])),
    ]
    add_table(doc, ["Métrica", "Resultado en 35 observaciones"], general_rows, [2.0, 4.85], font_size=9)
    add_body(doc, conclusion)

    doc.add_page_break()
    add_heading(doc, "Resultados por categoría")
    add_body(
        doc,
        "DR corresponde a la suma de la demanda real de los siete días. DP corresponde a la suma de las proyecciones para esos mismos siete días. MAE, RMSE y MAPE se calculan día por día dentro de cada categoría.",
    )
    add_table(
        doc,
        ["Categoría", "DR", "DP", "MAE", "RMSE", "MAPE"],
        category_rows(result, forecast_key, metrics_key),
        [1.35, 0.78, 0.95, 0.87, 0.93, 0.95],
        numeric_columns=(1, 2, 3, 4, 5),
        font_size=8.3,
    )

    add_heading(doc, "Control diario del período")
    add_table(
        doc,
        ["Fecha", "Demanda real", "Pronóstico", "Error absoluto"],
        daily_rows(result, forecast_key),
        [1.55, 1.65, 1.65, 1.65],
        numeric_columns=(1, 2, 3),
        font_size=8.5,
    )

    add_heading(doc, "Patrón de demanda observable")
    add_body(
        doc,
        "El índice de patrón se presenta como IE = VP / VPR, donde VP es el promedio diario del período evaluado y VPR es el promedio diario de referencia anterior al corte. Un valor menor que 1 indica una demanda media menor que la referencia histórica.",
    )
    add_table(
        doc,
        ["Categoría", "VP", "VPR", "IE"],
        pattern_rows(result),
        [2.0, 1.35, 1.7, 1.2],
        numeric_columns=(1, 2, 3),
        font_size=8.5,
    )

    add_heading(doc, "Alcance de los indicadores de inventario")
    add_body(
        doc,
        "La cantidad demandada se respalda en las ventas reales consolidadas: 1 671 unidades en el período. No se reportan el índice de salida de inventario ni la tasa de quiebre de stock como resultados reales, porque el archivo de inventario disponible no contiene un kardex continuo para todos los productos y los siete días. Calcularlos con registros incompletos produciría un resultado no verificable.",
    )
    add_body(
        doc,
        "Para un postest causal de implementación se requieren ventas e inventario registrados después de utilizar el modelo. Esta ficha compara métodos contra la misma demanda real y no debe interpretarse como evidencia de un cambio operativo de la tienda.",
    )

    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    doc.save(output)
    return output


def main():
    result = json.loads(RESULTS.read_text(encoding="utf-8"))
    for kind in ("pms", "ml"):
        print(write_document(kind, result))


if __name__ == "__main__":
    main()
