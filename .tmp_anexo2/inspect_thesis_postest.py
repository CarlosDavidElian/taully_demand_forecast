"""Extrae pasajes pertinentes de la tesis para una revisión de alcance."""

from pathlib import Path
import sys

from docx import Document


SOURCE = Path(r"C:\Users\Acer\OneDrive\Tesis Sr\LIMA-NORTE_PI_GURRERRO.docx")
TERMS = (
    "postest",
    "pretest",
    "inventario",
    "índice de salida",
    "quiebre",
    "mape",
    "pronóstico",
    "pronostico",
    "sistema",
    "pms",
    "export",
)


def compact(text: str) -> str:
    return " ".join(text.split())


def main() -> None:
    sys.stdout.reconfigure(encoding="utf-8")
    document = Document(SOURCE)
    blocks: list[tuple[str, str]] = []
    blocks.extend((f"P{index + 1}", compact(paragraph.text)) for index, paragraph in enumerate(document.paragraphs))
    for table_index, table in enumerate(document.tables, start=1):
        for row_index, row in enumerate(table.rows, start=1):
            text = compact(" | ".join(cell.text for cell in row.cells))
            blocks.append((f"T{table_index}R{row_index}", text))

    print(f"Párrafos={len(document.paragraphs)}; tablas={len(document.tables)}")
    if "--core" in sys.argv:
        print("\n=== PÁRRAFOS CENTRALES (problema, objetivos, hipótesis e indicadores) ===")
        for index in [70, 72, 73, 83, 84, 85, 89, 90, 91, 150, 165, 177, 193, 208, 218, 232, 248]:
            if index <= len(document.paragraphs):
                print(f"P{index}: {compact(document.paragraphs[index - 1].text)}")
        print("\n=== TABLA DE CONSISTENCIA ===")
        table = document.tables[21]
        for row_index, row in enumerate(table.rows, start=1):
            print(f"T22R{row_index}: {compact(' | '.join(cell.text for cell in row.cells))}")
        return
    for term in TERMS:
        matches = [(label, text) for label, text in blocks if term in text.lower()]
        print(f"\n=== {term.upper()} ({len(matches)}) ===")
        for label, text in matches[:80]:
            print(f"{label}: {text}")


if __name__ == "__main__":
    main()
