from __future__ import annotations

import json
import sys
from collections import Counter, defaultdict
from datetime import date, datetime, timedelta
from pathlib import Path

from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from infrastructure.readers.excel_reader import ExcelReader
from infrastructure.repositories.catalog_repository import ExcelCatalogRepository


ROOT = Path(r"C:\taully_demand_forecast")
SOURCE = ROOT / ".tmp_inventory_demo" / "product_sales_from_reports.json"
WORKBOOK = ROOT / "data" / "inventario_producto_Taully_2026-04-01_a_2026-09-15.xlsx"
SHEET = "Datos inventario"


def fail(message: str) -> None:
    raise RuntimeError(message)


def date_only(value: object) -> date:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    return datetime.strptime(str(value), "%Y-%m-%d").date()


def integer(value: object) -> int:
    return int(float(value or 0))


def read_exported() -> tuple[dict[tuple[date, str], dict], dict[str, int]]:
    wb = load_workbook(WORKBOOK, data_only=True, read_only=True)
    ws = wb[SHEET]
    headers = [str(cell.value).strip() if cell.value is not None else "" for cell in ws[1]]
    needed = ["Fecha", "Producto", "Categoría", "Cantidad_Vendida", "Stock_Inicial", "Entradas", "Stock_Disponible", "Stock_Final", "Dias_Sin_Stock"]
    if any(col not in headers for col in needed):
        fail("Faltan columnas requeridas en el inventario.")
    index = {value: idx for idx, value in enumerate(headers)}
    rows: dict[tuple[date, str], dict] = {}
    issues = {"rows": 0, "duplicates": 0, "balance_errors": 0, "dqs_errors": 0}
    for values in ws.iter_rows(min_row=2, values_only=True):
        if not any(value is not None for value in values):
            continue
        issues["rows"] += 1
        record = {header: values[idx] for header, idx in index.items()}
        key = (date_only(record["Fecha"]), str(record["Producto"]).strip())
        if key in rows:
            issues["duplicates"] += 1
        rows[key] = record
        si = integer(record["Stock_Inicial"])
        en = integer(record["Entradas"])
        available = integer(record["Stock_Disponible"])
        sf = integer(record["Stock_Final"])
        demand = integer(record["Cantidad_Vendida"])
        dqs = integer(record["Dias_Sin_Stock"])
        if available != si + en or sf != max(0, available - demand):
            issues["balance_errors"] += 1
        if dqs not in (0, 1):
            issues["dqs_errors"] += 1
    return rows, issues


def reconstruct() -> tuple[dict[tuple[date, str], dict], dict[str, dict]]:
    data = json.loads(SOURCE.read_text(encoding="utf-8"))
    metadata, sales = data["metadata"], data["records"]
    start = datetime.strptime(metadata["start_date"], "%Y-%m-%d").date()
    dates = [start + timedelta(days=index) for index in range(metadata["date_count"])]
    product_info: dict[str, dict] = {}
    sales_by_key: dict[tuple[date, str], int] = {}
    for sale in sales:
        product_info[sale["product"]] = sale
        sales_by_key[(datetime.strptime(sale["date"], "%Y-%m-%d").date(), sale["product"])] = int(sale["quantity"])

    by_sale_key: dict[tuple[date, str], dict] = {}
    category = defaultdict(lambda: {
        "si": 0, "entries_all": 0, "entries_exported": 0, "sf": 0,
        "demand": 0, "fulfilled": 0, "shortage": 0,
        "dqs_events": 0, "dqs_category_days": set(), "zero_sale_entry_events": 0,
        "products_without_sale_last_day": set(),
    })
    for product, source in product_info.items():
        daily = [sales_by_key.get((current_date, product), 0) for current_date in dates]
        average = sum(daily) / len(dates)
        highest_sale = max(daily)
        target = max(int(-(-(average * 10) // 1)), highest_sale * 2, 1)
        # Equivalent to Math.ceil for this nonnegative value.
        reorder = max(int(-(-average * 3 // 1)), 1)
        category_metrics = category[source["category"]]
        stock_initial = target
        category_metrics["si"] += stock_initial
        product_last_sale = None
        for current_date, demand in zip(dates, daily):
            entries = target - stock_initial if stock_initial <= reorder else 0
            available = stock_initial + entries
            shortage = max(0, demand - available)
            fulfilled = min(demand, available)
            stock_final = max(0, available - demand)
            dqs = int(shortage > 0)
            category_metrics["entries_all"] += entries
            category_metrics["demand"] += demand
            category_metrics["fulfilled"] += fulfilled
            category_metrics["shortage"] += shortage
            category_metrics["dqs_events"] += dqs
            if dqs:
                category_metrics["dqs_category_days"].add(current_date)
            if demand > 0:
                product_last_sale = current_date
                category_metrics["entries_exported"] += entries
                by_sale_key[(current_date, product)] = {
                    "category": source["category"], "demand": demand, "si": stock_initial,
                    "entries": entries, "available": available, "sf": stock_final, "dqs": dqs,
                }
            elif entries > 0:
                category_metrics["zero_sale_entry_events"] += 1
            stock_initial = stock_final
        category_metrics["sf"] += stock_initial
        if product_last_sale != dates[-1]:
            category_metrics["products_without_sale_last_day"].add(product)
    return by_sale_key, category


def source_sales_match_raw_reports() -> None:
    """Confirma que el detalle que alimentó el inventario coincide con los Excel diarios."""
    data = json.loads(SOURCE.read_text(encoding="utf-8"))
    expected = Counter()
    for item in data["records"]:
        key = (datetime.strptime(item["date"], "%Y-%m-%d").date(), item["product"])
        expected[key] += int(item["quantity"])
    reports = sorted((ROOT / "data").glob("reporte_ventas_*.xlsx"))
    if len(reports) != 168:
        fail(f"Se esperaban 168 reportes de venta y se encontraron {len(reports)}.")
    reader = ExcelReader()
    catalog = ExcelCatalogRepository(ROOT / "data" / "catalogo_maestro.xlsx")
    observed = Counter()
    unclassified = []
    for report in reports:
        for sale in reader.read_sales(str(report)):
            product = catalog.get_product(sale.product_name)
            if product is None:
                unclassified.append(f"{report.name}: {sale.product_name}")
                continue
            observed[(sale.date.date(), product.product_name)] += int(sale.quantity)
    if unclassified:
        fail("Hay ventas sin producto clasificado: " + "; ".join(unclassified[:5]))
    mismatches = [
        f"{key}: fuente={expected.get(key, 0)}, reporte={observed.get(key, 0)}"
        for key in sorted(set(expected) | set(observed))
        if expected.get(key, 0) != observed.get(key, 0)
    ]
    if mismatches:
        fail("El detalle de ventas no coincide con los reportes: " + "; ".join(mismatches[:5]))


def main() -> None:
    exported, export_issues = read_exported()
    expected, categories = reconstruct()
    source_sales_match_raw_reports()
    if len(exported) != len(expected):
        fail(f"El archivo tiene {len(exported)} filas y el detalle de ventas tiene {len(expected)}.")
    mismatches = []
    for key, expected_record in expected.items():
        actual = exported.get(key)
        if actual is None:
            mismatches.append(f"Falta {key}")
            continue
        tested = {
            "category": str(actual["Categoría"]).strip(),
            "demand": integer(actual["Cantidad_Vendida"]),
            "si": integer(actual["Stock_Inicial"]),
            "entries": integer(actual["Entradas"]),
            "available": integer(actual["Stock_Disponible"]),
            "sf": integer(actual["Stock_Final"]),
            "dqs": integer(actual["Dias_Sin_Stock"]),
        }
        if tested != expected_record:
            mismatches.append(f"{key}: archivo={tested}, esperado={expected_record}")
    extras = set(exported) - set(expected)
    if mismatches or extras or export_issues["duplicates"] or export_issues["balance_errors"] or export_issues["dqs_errors"]:
        fail("El archivo no coincide con la lógica de origen. " + (mismatches + [f"Extras: {len(extras)}"])[0])

    print("AUDITORÍA DEL INVENTARIO DE PRUEBA")
    print(f"Filas exportadas: {export_issues['rows']:,}; todas coinciden con el generador.")
    print("Las cantidades de venta del generador coinciden con los 168 reportes originales, producto por producto y fecha por fecha.")
    print("Saldo por fila: Stock_Disponible = Stock_Inicial + Entradas y Stock_Final = max(0, disponible - cantidad): correcto en todas las filas.")
    print("\nRESUMEN POR CATEGORÍA")
    print("Categoría | SI inicial | Entradas (periodo) | SF al 15/09 | Demanda registrada | Venta atendida | Demanda no atendida | DQS eventos | DQS días categoría")
    total = Counter()
    for name in sorted(categories):
        m = categories[name]
        for field in ("si", "entries_all", "sf", "demand", "fulfilled", "shortage", "dqs_events"):
            total[field] += m[field]
        print(
            f"{name} | {m['si']:,} | {m['entries_all']:,} | {m['sf']:,} | {m['demand']:,} | "
            f"{m['fulfilled']:,} | {m['shortage']:,} | {m['dqs_events']:,} | {len(m['dqs_category_days']):,}"
        )
    print(
        f"TOTAL | {total['si']:,} | {total['entries_all']:,} | {total['sf']:,} | {total['demand']:,} | "
        f"{total['fulfilled']:,} | {total['shortage']:,} | {total['dqs_events']:,} | -"
    )
    missing_entry_events = sum(m["zero_sale_entry_events"] for m in categories.values())
    products_without_final_row = sum(len(m["products_without_sale_last_day"]) for m in categories.values())
    print("\nHALLAZGOS")
    print(f"Reposiciones que ocurrieron en días sin venta y no quedaron en el Excel: {missing_entry_events:,}.")
    print(f"Productos sin fila de cierre al 15/09 porque no vendieron ese día: {products_without_final_row:,}.")
    print("Conclusión: el archivo reproduce una simulación de stock a partir de ventas. No constituye un kardex real ni permite reemplazar N/D como evidencia de tesis.")


if __name__ == "__main__":
    main()
