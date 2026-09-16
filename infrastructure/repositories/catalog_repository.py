"""Repositorio del catálogo maestro almacenado en Excel."""

from __future__ import annotations

from datetime import date, datetime, time
from pathlib import Path
import re
from typing import Iterable, List, Optional
import unicodedata
from zipfile import BadZipFile

from openpyxl import load_workbook
from openpyxl.utils.datetime import to_excel
from openpyxl.utils.exceptions import InvalidFileException

from config.settings import CATALOG_FILE
from domain.entities.product import Product
from domain.interfaces.repositories import ProductCatalogRepository


class ExcelCatalogRepository(ProductCatalogRepository):
    """Carga el catálogo desde la hoja ``PRODUCTOS`` del archivo maestro.

    La hoja puede contener pestañas auxiliares (por ejemplo, fuentes y
    marcas). Al elegir explícitamente ``PRODUCTOS`` evitamos que el orden de
    las hojas cambie el catálogo que usa la aplicación.
    """

    PRODUCT_SHEET = "PRODUCTOS"
    REQUIRED_COLUMNS = ("PRODUCTO", "FAMILIA", "CATEGORIA")

    def __init__(self, catalog_path: Path | str | None = None):
        self.catalog_path = Path(catalog_path) if catalog_path is not None else CATALOG_FILE
        self._products: list[Product] = []
        self._products_by_key: dict[str, Product] = {}
        self._load_catalog()

    def _load_catalog(self) -> None:
        if not self.catalog_path.exists():
            raise FileNotFoundError(f"Catálogo no encontrado en {self.catalog_path}")

        workbook = None
        try:
            workbook = load_workbook(self.catalog_path, read_only=True, data_only=True)
            sheet = self._get_product_sheet(workbook.sheetnames, workbook)
            rows = sheet.iter_rows(values_only=False)
            header_cells = next(rows, None)
            if header_cells is None:
                raise ValueError("El catálogo no contiene encabezados.")

            columns = self._get_columns(header_cells)
            missing_columns = [column for column in self.REQUIRED_COLUMNS if column not in columns]
            if missing_columns:
                missing = ", ".join(missing_columns)
                raise ValueError(f"El catálogo debe tener las columnas: {missing}.")

            products: list[Product] = []
            products_by_key: dict[str, Product] = {}
            for row_number, cells in enumerate(rows, start=2):
                values = {column: cells[index].value for column, index in columns.items() if index < len(cells)}
                if not any(value is not None and str(value).strip() for value in values.values()):
                    continue

                product_name = self._required_text(values.get("PRODUCTO"), "PRODUCTO", row_number)
                family = self._required_text(values.get("FAMILIA"), "FAMILIA", row_number)
                category = self._required_text(values.get("CATEGORIA"), "CATEGORIA", row_number)
                brand = self._optional_text(values.get("MARCA"))
                cost = self._parse_cost(values.get("COSTO"), row_number, workbook.epoch)

                product = Product(
                    product_name=product_name,
                    family=family,
                    category=category,
                    brand=brand,
                    cost=cost,
                )
                key = self._normalise_key(product_name)
                if key in products_by_key:
                    duplicate = products_by_key[key].product_name
                    raise ValueError(
                        f"El catálogo repite el producto '{product_name}' en la fila {row_number} "
                        f"(ya existe como '{duplicate}')."
                    )
                products.append(product)
                products_by_key[key] = product

            if not products:
                raise ValueError("El catálogo no contiene productos válidos en la hoja PRODUCTOS.")

            self._products = products
            self._products_by_key = products_by_key
        except (OSError, BadZipFile, InvalidFileException, ValueError) as exc:
            if isinstance(exc, ValueError):
                raise
            raise ValueError(f"No se pudo leer el catálogo: {exc}") from exc
        finally:
            if workbook is not None:
                workbook.close()

    @classmethod
    def _get_product_sheet(cls, sheet_names: Iterable[str], workbook):
        for sheet_name in sheet_names:
            if sheet_name.strip().upper() == cls.PRODUCT_SHEET:
                return workbook[sheet_name]
        raise ValueError("El catálogo debe incluir una hoja llamada PRODUCTOS.")

    @staticmethod
    def _get_columns(header_cells) -> dict[str, int]:
        columns: dict[str, int] = {}
        for index, cell in enumerate(header_cells):
            if cell.value is None:
                continue
            header = str(cell.value).strip().upper()
            if header:
                columns.setdefault(header, index)
        return columns

    @staticmethod
    def _required_text(value: object, column: str, row_number: int) -> str:
        text = ExcelCatalogRepository._optional_text(value)
        if not text:
            raise ValueError(f"La columna {column} está vacía en la fila {row_number} del catálogo.")
        return text

    @staticmethod
    def _optional_text(value: object) -> str:
        return str(value).strip() if value is not None else ""

    @staticmethod
    def _parse_cost(value: object, row_number: int, epoch) -> float:
        """Convierte el costo y corrige seriales Excel leídos como fecha/hora.

        Algunos archivos de Excel guardan correctamente el formato monetario,
        pero ``openpyxl`` entrega seriales pequeños como ``datetime`` o
        ``time``. Convertirlos otra vez a serial preserva, por ejemplo, 1.20
        en vez de convertirlo erróneamente en una fecha de 1900.
        """
        if value is None or (isinstance(value, str) and not value.strip()):
            return 0.0
        if isinstance(value, datetime):
            cost = float(to_excel(value, epoch))
        elif isinstance(value, date):
            cost = float(to_excel(value, epoch))
        elif isinstance(value, time):
            cost = (
                value.hour * 3600 + value.minute * 60 + value.second + value.microsecond / 1_000_000
            ) / 86400
        elif isinstance(value, (int, float)):
            cost = float(value)
        else:
            text = str(value).strip().replace("S/", "").replace("s/", "")
            text = re.sub(r"\s+", "", text)
            if text.count(",") == 1 and text.count(".") == 0:
                text = text.replace(",", ".")
            elif text.count(",") == 1 and text.count(".") >= 1:
                if text.rfind(",") > text.rfind("."):
                    text = text.replace(".", "").replace(",", ".")
                else:
                    text = text.replace(",", "")
            try:
                cost = float(text)
            except (TypeError, ValueError):
                # El costo es un dato complementario. Marcas como "POR
                # CONFIRMAR" se conservan como costo pendiente en vez de
                # impedir que el resto del catálogo opere.
                return 0.0

        if cost < 0:
            raise ValueError(f"El costo de la fila {row_number} no puede ser negativo.")
        return cost

    @staticmethod
    def _normalise_key(product_name: str) -> str:
        decomposed = unicodedata.normalize("NFKD", product_name)
        without_accents = "".join(char for char in decomposed if not unicodedata.combining(char))
        return re.sub(r"\s+", " ", without_accents).strip().upper()

    def get_product(self, product_name: str) -> Optional[Product]:
        return self._products_by_key.get(self._normalise_key(product_name))

    def get_all_products(self) -> List[Product]:
        return list(self._products)
