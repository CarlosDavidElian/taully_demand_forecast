from __future__ import annotations

from datetime import date
from pathlib import Path
from tempfile import TemporaryDirectory
import unittest

from openpyxl import Workbook, load_workbook

from application.services.inventory_service import InventoryService


class InventoryServiceTests(unittest.TestCase):
    HEADERS = [
        "Fecha",
        "Producto",
        "Categoría",
        "Cantidad_Vendida",
        "Stock_Inicial",
        "Entradas",
        "Stock_Disponible",
        "Stock_Final",
        "Dias_Sin_Stock",
    ]

    def _write_inventory(self, path: Path, include_stock_final: bool = True) -> None:
        headers = list(self.HEADERS)
        if not include_stock_final:
            headers.remove("Stock_Final")
        rows = [
            [date(2026, 4, 1), "Producto A", "ABARROTES", 5, 10, 2, 12, 7, 0],
            [date(2026, 4, 1), "Producto B", "ABARROTES", 4, 8, 0, 8, 4, 1],
            [date(2026, 4, 2), "Producto A", "ABARROTES", 3, 7, 5, 12, 9, 0],
            [date(2026, 4, 1), "Producto C", "BEBIDAS", 2, 4, 1, 5, 3, 0],
            [date(2026, 4, 2), "Producto C", "BEBIDAS", 1, 3, 0, 3, 0, 1],
        ]
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Datos inventario"
        sheet.append(headers)
        for row in rows:
            if not include_stock_final:
                row = [value for index, value in enumerate(row) if self.HEADERS[index] != "Stock_Final"]
            sheet.append(row)
        workbook.save(path)

    def test_calculates_category_metrics_with_unique_stockout_dates(self):
        with TemporaryDirectory() as folder:
            inventory_path = Path(folder) / "inventario.xlsx"
            self._write_inventory(inventory_path)

            metrics = InventoryService().calculate_category_metrics(inventory_path)

        self.assertEqual([metric.category for metric in metrics], ["ABARROTES", "BEBIDAS"])
        grocery, drinks = metrics
        self.assertEqual(grocery.records, 3)
        # SI/SF son saldos de período: el primero y último por producto, no
        # la suma de saldos de cada día.
        self.assertEqual(grocery.stock_initial, 18)
        self.assertEqual(grocery.entries, 7)
        self.assertEqual(grocery.stock_final, 13)
        self.assertEqual(grocery.demand_quantity, 12)
        self.assertAlmostEqual(grocery.inventory_output_index, 48.0)
        self.assertEqual(grocery.stockout_days, 1)
        self.assertEqual(grocery.evaluated_days, 2)
        self.assertEqual(grocery.registered_days, 2)
        self.assertAlmostEqual(grocery.coverage_rate, 100.0)
        self.assertAlmostEqual(grocery.stockout_rate, 50.0)

        self.assertEqual(drinks.demand_quantity, 5)
        self.assertAlmostEqual(drinks.inventory_output_index, 100.0)
        self.assertEqual(drinks.stockout_days, 1)
        self.assertEqual(drinks.evaluated_days, 2)
        self.assertEqual(drinks.registered_days, 2)
        self.assertAlmostEqual(drinks.stockout_rate, 50.0)

    def test_filters_the_evaluation_period_inclusively(self):
        with TemporaryDirectory() as folder:
            inventory_path = Path(folder) / "inventario.xlsx"
            self._write_inventory(inventory_path)

            metrics = InventoryService().calculate_category_metrics(
                inventory_path,
                start_date="2026-04-02",
                end_date="2026-04-02",
            )

        grocery, drinks = metrics
        self.assertEqual((grocery.stock_initial, grocery.entries, grocery.stock_final, grocery.demand_quantity), (7, 5, 9, 3))
        self.assertAlmostEqual(grocery.inventory_output_index, 25.0)
        self.assertEqual((grocery.stockout_days, grocery.evaluated_days, grocery.stockout_rate), (0, 1, 0.0))
        self.assertEqual((drinks.stock_initial, drinks.entries, drinks.stock_final, drinks.demand_quantity), (3, 0, 0, 3))
        self.assertAlmostEqual(drinks.inventory_output_index, 100.0)
        self.assertEqual((drinks.stockout_days, drinks.evaluated_days, drinks.stockout_rate), (1, 1, 100.0))

    def test_uses_calendar_days_and_reports_missing_coverage(self):
        with TemporaryDirectory() as folder:
            inventory_path = Path(folder) / "inventario.xlsx"
            self._write_inventory(inventory_path)

            metrics = InventoryService().calculate_category_metrics(
                inventory_path,
                start_date="2026-04-01",
                end_date="2026-04-03",
            )

        grocery = next(metric for metric in metrics if metric.category == "ABARROTES")
        self.assertEqual(grocery.evaluated_days, 3)
        self.assertEqual(grocery.registered_days, 2)
        self.assertAlmostEqual(grocery.coverage_rate, 66.6666666667)
        self.assertAlmostEqual(grocery.stockout_rate, 100 / 3)

    def test_reports_nonblocking_sales_and_stock_reconciliation(self):
        with TemporaryDirectory() as folder:
            inventory_path = Path(folder) / "inventario.xlsx"
            self._write_inventory(inventory_path)

            reconciliation = InventoryService().reconciliation_summary(inventory_path)

        self.assertEqual(reconciliation["rows_with_difference"], 1)
        self.assertFalse(reconciliation["reconciled"])
        self.assertEqual(reconciliation["sales_total"], 15.0)
        self.assertEqual(reconciliation["balance_demand"], 17.0)
        self.assertEqual(reconciliation["difference"], -2.0)

    def test_import_creates_content_addressed_copy_without_overwriting_previous_copy(self):
        with TemporaryDirectory() as folder:
            root = Path(folder)
            inventory_path = root / "inventario.xlsx"
            self._write_inventory(inventory_path)
            service = InventoryService(storage_directory=root / "inventory_uploads")

            first = service.import_inventory(inventory_path)
            second = service.import_inventory(inventory_path)
            original_copy = first.stored_path.read_bytes()

            workbook = load_workbook(inventory_path)
            workbook.active.append([date(2026, 4, 3), "Producto D", "BEBIDAS", 1, 2, 0, 2, 1, 0])
            workbook.save(inventory_path)
            third = service.import_inventory(inventory_path)

            self.assertIsNotNone(first.inventory_id)
            self.assertEqual(first.inventory_id, second.inventory_id)
            self.assertEqual(first.stored_path, second.stored_path)
            self.assertTrue(first.stored_path.exists())
            self.assertNotEqual(first.inventory_id, third.inventory_id)
            self.assertEqual(first.stored_path.read_bytes(), original_copy)
            self.assertEqual(len(list((root / "inventory_uploads").glob("*.xlsx"))), 2)
            self.assertEqual((first.records, first.categories), (5, 2))

    def test_rejects_missing_required_column(self):
        with TemporaryDirectory() as folder:
            inventory_path = Path(folder) / "inventario.xlsx"
            self._write_inventory(inventory_path, include_stock_final=False)

            with self.assertRaisesRegex(ValueError, "Stock_Final"):
                InventoryService().calculate_category_metrics(inventory_path)

    def test_rejects_inconsistent_stock_available_and_final_stock(self):
        with TemporaryDirectory() as folder:
            inventory_path = Path(folder) / "inventario.xlsx"
            self._write_inventory(inventory_path)
            workbook = load_workbook(inventory_path)
            workbook.active["G2"] = 11  # SI 10 + EN 2 = 12, no 11.
            workbook.save(inventory_path)

            with self.assertRaisesRegex(ValueError, "Stock_Disponible"):
                InventoryService().calculate_category_metrics(inventory_path)

            self._write_inventory(inventory_path)
            workbook = load_workbook(inventory_path)
            workbook.active["H2"] = 13  # No puede superar el disponible (12).
            workbook.save(inventory_path)

            with self.assertRaisesRegex(ValueError, "Stock_Final"):
                InventoryService().calculate_category_metrics(inventory_path)

    def test_reads_the_inventory_file_available_in_the_project(self):
        inventory_path = Path(__file__).resolve().parents[1] / "data" / "inventario_producto_Taully_2026-04-01_a_2026-09-15.xlsx"

        metrics = InventoryService().calculate_category_metrics(inventory_path)

        self.assertEqual([metric.category for metric in metrics], ["ABARROTES", "BEBIDAS", "GOLOSINAS", "HELADOS", "LIMPIEZA"])
        self.assertTrue(all(metric.evaluated_days > 0 for metric in metrics))
        self.assertTrue(all(metric.stockout_days <= metric.evaluated_days for metric in metrics))
