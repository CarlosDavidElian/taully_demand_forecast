from pathlib import Path
from docx import Document
from docx.shared import Cm,Pt
from lxml import etree
R=Path('C:/taully_demand_forecast'); O=R/'outputs/anexo2_30dias_hasta_15septiembre'; C=R/'tmp/guia_capturas'
d=Document(O/'Guia_detallada_verificacion_para_jefatura.docx')
def capture(title,filename,notes,arrows):
 d.add_page_break();d.add_heading(title,1)
 for n in notes:d.add_paragraph(n)
 p=d.add_paragraph();p.paragraph_format.space_after=Pt(0)
 p.add_run().add_picture(str(C/filename),width=Cm(17))
 # Native Word vector arrows over an unchanged screenshot.
 for x,y,xx,yy in arrows:
  scale=481.89/990
  xml=f'''<w:pict xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" xmlns:v="urn:schemas-microsoft-com:vml"><v:line style="position:absolute;margin-left:0;margin-top:0;z-index:251659264;mso-position-horizontal-relative:text;mso-position-vertical-relative:line" from="{x*scale}pt,{y*scale}pt" to="{xx*scale}pt,{yy*scale}pt" strokecolor="#D32020" strokeweight="3pt"><v:stroke endarrow="block"/></v:line></w:pict>'''
  p.add_run()._r.append(etree.fromstring(xml))
capture('Captura 1 Elegir horizonte y fecha','config.png',[
 '1. En Horizonte elija 30 días. 2. En Corte de prueba ingrese 16/08/2026 para el postest. Las flechas rojas señalan ambos campos.',
 'Para el pretest use estos mismos controles con corte 17/07/2026 y horizonte de 30 días. Después baje hasta Postest e indicadores de inventario.'],[(610,395,805,395),(610,437,805,437)])
if (C/'ejecutar.png').exists():capture('Captura 2 Ejecutar la evaluación','ejecutar.png',[
 '3. Pulse Ejecutar Postest en el bloque de la derecha. Espere a que termine; durante el cálculo el botón muestra Evaluando.',
 'El inventario activo de la izquierda es demostrativo. Su tabla general abarca todo el historial y no debe confundirse con la ventana de 30 días.'],[(710,390,710,460)])
if (C/'resultado.png').exists():capture('Captura 3 Revisar y descargar','resultado.png',[
 '4. Compruebe la tarjeta Modelo predictivo para el postest: MAPE 26,21 %. 5. Pulse Descargar Postest y siga los cálculos de Excel de esta guía.',
 'Para el pretest, ejecute el corte 17/07/2026 y lea la tarjeta PMS-7: MAPE 23,76 %. Guarde cada exportación por separado.'],[(700,180,790,180),(195,245,195,293)])
d.save(O/'Guia_verificacion_con_capturas.docx')
print(O/'Guia_verificacion_con_capturas.docx')


