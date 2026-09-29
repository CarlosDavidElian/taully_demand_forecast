import ast
import csv
import json
import math
from pathlib import Path
from statistics import mean
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path('C:/taully_demand_forecast')
OUT=ROOT/'outputs/postest_visualizado'
doc=Document(OUT/'Guia_verificacion_inventario_postest.docx')
doc.styles['Subtitle'].font.color.rgb=RGBColor(0,0,0)
# Reuse only the formatting helpers, without executing the earlier builder.
tree=ast.parse((ROOT/'tmp/build_guia_verificacion.py').read_text(encoding='utf-8'))
for node in tree.body:
    if isinstance(node,ast.FunctionDef) and node.name in {'p','table'}:
        exec(compile(ast.Module(body=[node],type_ignores=[]),'<helpers>','exec'))
doc.paragraphs[0].text='Guía de verificación del pretest y postest'
doc.paragraphs[0].style='Title'
doc.paragraphs[1].text='Minimarket Taully  |  Inventario y fichas de las tres dimensiones'
doc.paragraphs[2].text=('Esta guía explica cómo revisar los indicadores de inventario del postest y las tres dimensiones del anexo: volumen de demanda, patrón de demanda y precisión del pronóstico. Incluye seis fichas, una por dimensión y momento, con datos disponibles del proyecto y pasos para reproducir sus cálculos.')
hist=list(csv.DictReader((ROOT/'data/historial_demanda.csv').open(encoding='utf-8-sig')))
cats=sorted({r['category'] for r in hist})
lookup={(r['date'],r['category']):float(r['quantity']) for r in hist}
assert len(lookup)==len(hist), 'Hay registros repetidos por fecha y categoría'
payload=json.loads((ROOT/'data/posttests/postest_20260927t191247z_20260908_7d_70578bbb.json').read_text(encoding='utf-8'))
result=payload['result']
ref={c:mean(float(r['quantity']) for r in hist if r['category']==c and '2026-08-01'<=r['date']<='2026-08-31') for c in cats}
assert all(sum(r['category']==c and '2026-08-01'<=r['date']<='2026-08-31' for r in hist)==31 for c in cats)
def fmt(x,n=2):return f'{x:.{n}f}'.replace('.',',')
def stats(obs):
    errs=[a-b for _,_,a,b in obs]
    return dict(dr=sum(a for _,_,a,b in obs),dp=sum(b for _,_,a,b in obs),mape=mean(abs((a-b)/a) for _,_,a,b in obs if a!=0)*100,mae=mean(abs(e) for e in errs),rmse=math.sqrt(mean(e*e for e in errs)))
allmetrics={}
for moment,start,end,cutoff in [('Pretest','2026-09-02','2026-09-08','2026-09-01'),('Postest','2026-09-09','2026-09-15','2026-09-08')]:
    obs=[]
    if moment=='Pretest':
        for c in cats:
            train=sorted((r['date'],float(r['quantity'])) for r in hist if r['category']==c and r['date']<=cutoff)
            vals=[v for _,v in train[-7:]]
            actual=sorted((d,v) for (d,cat),v in lookup.items() if cat==c and start<=d<=end)
            assert len(actual)==7 and len(vals)==7
            for date,a in actual:
                pred=mean(vals[-7:]); vals.append(pred); obs.append((date,c,a,pred))
    else:
        for r in result['observations']:
            assert lookup[(r['date'],r['category'])]==r['actual_demand']
            obs.append((r['date'],r['category'],r['actual_demand'],r['model_prediction']))
    assert len(obs)==35
    per={c:stats([r for r in obs if r[1]==c]) for c in cats}
    total=stats(obs); allmetrics[moment]=total
    period=f'{start[8:10]}/09/2026 al {end[8:10]}/09/2026'
    for dim,indicator,formula,definition in [
        ('Volumen de demanda','Cantidad demandada','CD = Σ ventas realizadas','CD se expresa en unidades vendidas. Σ representa la suma de las cantidades vendidas de los siete días.'),
        ('Patrón de demanda','Estacionalidad','IE = VP / VPR','VP = promedio diario vendido en los siete días evaluados. VPR = promedio diario de agosto de 2026. Ambos se expresan en unidades por día; IE es un índice sin unidad.'),
        ('Precisión del pronóstico','Error de pronóstico','MAPE = (100/n) × Σ |(DR − DP) / DR|','DR = venta observada; DP = pronóstico; n = observaciones con DR distinta de cero. El valor absoluto evita que los errores positivos y negativos se compensen.')]:
        doc.add_page_break()
        p(f'{moment}  {dim}','Heading 1')
        p('Instrumento de recolección de datos de la dimensión','Subtitle')
        table(['Campo','Registro'],[
            ('Indicador',indicator),('Investigador','Guerrero Mora, Carlos David Elian'),('Lugar de estudio','Minimarket Taully, Comas'),('Momento y período',f'{moment}. {period}. Siete días.'),('Fórmula',formula),('Donde',definition)
        ],[1.3,5.5])
        if dim=='Volumen de demanda':
            table(['N°','Categoría','Cantidad vendida','CD en unidades'],[[i,c,fmt(per[c]['dr'],0),fmt(per[c]['dr'],0)] for i,c in enumerate(cats,1)]+[['','TOTAL',fmt(total['dr'],0),fmt(total['dr'],0)]],[.4,2.1,2.15,2.15])
            p('Cómo verificar','Heading 2')
            p(f'En historial_demanda.csv, filtra las fechas {period}. Agrupa category y suma quantity. Si usas una tabla dinámica, coloca category en Filas y quantity en Valores con la operación Suma. ABARROTES debe sumar {fmt(per["ABARROTES"]["dr"],0)} unidades y el total {fmt(total["dr"],0)}.')
            p('Contrasta esa suma con las cantidades de los reportes POS de los mismos días, usando la misma clasificación de productos. La ficha original dice Producto; esta versión resume por categoría, que es la unidad de análisis del modelo. Si se exige detalle por producto, debe conservarse como respaldo desde los reportes.')
            p('La CD de esta ficha se obtiene de ventas, no de SI + EN − SF. El inventario y las ventas requieren conciliación; no sustituyas un resultado por el otro. Ventas observadas tampoco permiten medir por sí solas la demanda no atendida por falta de stock.')
        elif dim=='Patrón de demanda':
            table(['N°','Categoría','VP diario','VPR diario','IE'],[[i,c,fmt(per[c]['dr']/7),fmt(ref[c]),fmt((per[c]['dr']/7)/ref[c])] for i,c in enumerate(cats,1)]+[['','TOTAL',fmt(total['dr']/7),fmt(sum(ref.values())),fmt((total['dr']/7)/sum(ref.values()))]],[.4,2.1,1.45,1.45,1.4])
            p('Cómo verificar','Heading 2')
            p('Suma las unidades de cada categoría en la semana y divide entre 7 para obtener VP. Suma sus unidades del 1 al 31 de agosto de 2026 y divide entre 31 para obtener VPR. Divide VP entre VPR conservando toda la precisión; redondea solo al mostrar el resultado. Se usa agosto como referencia común anterior a los dos períodos.')
            c='ABARROTES'; p(f'Ejemplo ABARROTES: VP = {fmt(per[c]["dr"],0)} ÷ 7 = {fmt(per[c]["dr"]/7)} unidades/día; VPR = {fmt(ref[c]*31,0)} ÷ 31 = {fmt(ref[c])}; IE = {fmt((per[c]["dr"]/7)/ref[c])}.')
            p('En una hoja de Excel, con VP en C2 y VPR en D2, escribe =C2/D2 en E2. No apliques formato Porcentaje. IE mayor que 1 indica venta por encima de la referencia; menor que 1, por debajo. Este cociente describe variación frente a agosto y no demuestra por sí solo un patrón estacional recurrente ni una mejora causada por el modelo.')
        else:
            table(['N°','Categoría','DR semanal','DP semanal','MAPE diario'],[[i,c,fmt(per[c]['dr'],0),fmt(per[c]['dp']),fmt(per[c]['mape'])+' %'] for i,c in enumerate(cats,1)]+[['','GLOBAL',fmt(total['dr'],0),fmt(total['dp']),fmt(total['mape'])+' %']],[.4,1.75,1.45,1.5,1.7])
            method='PMS-7 recursivo calculado con información hasta el 01/09/2026' if moment=='Pretest' else 'modelo del programa con corte 08/09/2026, ejecución guardada del 27/09/2026'
            p(f'Fuente de DP: {method}. Fuente de DR: historial_demanda.csv. DR y DP de la tabla son sumas semanales; el MAPE se obtiene de las siete observaciones diarias por categoría, no de esos dos totales.')
            p('Cómo verificar','Heading 2')
            p('En una hoja auxiliar coloca Fecha, Categoría, DR y DP en A:D. En E2 escribe =SI(C2=0;"";ABS((C2-D2)/C2)) y copia para las 35 observaciones. Con formato Porcentaje, =PROMEDIO(E2:E36) devuelve el MAPE global. Para cada categoría, promedia solo sus siete errores. Aquí las 35 demandas son positivas.')
            if moment=='Pretest':
                p('Reconstrucción de DP: para cada categoría, promedia las ventas del 26/08 al 01/09 para pronosticar el 02/09. Añade ese pronóstico al final de la serie, retira el dato más antiguo y promedia los últimos siete valores para el siguiente día. Continúa hasta el 08/09 sin introducir ventas posteriores al corte. PMS-7 es una referencia calculada; no se ha acreditado que fuera el método usado por la tienda.')
            else:
                p('En el programa, ejecuta el postest con corte 08/09 y horizonte 7 días. Descarga el Excel y usa las filas diarias de demanda observada y predicción del modelo para esta comprobación. El PMS-7 que aparece junto al modelo evalúa la misma semana del 09 al 15/09: su MAPE de 31,54 % no es el MAPE del pretest del 02 al 08/09.')
            p(f'Control adicional: MAE = {fmt(total["mae"])} y RMSE = {fmt(total["rmse"])} unidades, calculados sobre los 35 errores diarios. El MAPE global es {fmt(total["mape"])} %. Son evaluaciones históricas: comparar semanas distintas no aísla el efecto del modelo ni acredita una intervención en la tienda.')

for style in doc.styles:
    for border in list(style.element.iter(qn('w:pBdr'))):border.getparent().remove(border)
doc.core_properties.title='Guía de verificación del pretest y postest con fichas de dimensiones'
doc.save(OUT/'Guia_verificacion_pretest_postest_ampliada.docx')
print(json.dumps(allmetrics,ensure_ascii=False))
print(OUT/'Guia_verificacion_pretest_postest_ampliada.docx')
