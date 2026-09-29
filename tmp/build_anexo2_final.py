from __future__ import annotations

import json
from datetime import datetime
from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "outputs" / "01a0b651-7620-73f0-8f6a-8b68b14f7a17" / "Anexo_2_Instrumentos_de_recoleccion_de_datos_validado.docx"
HISTORY = ROOT / "data" / "historial_demanda.csv"
POSTTESTS = ROOT / "data" / "posttests"

NAVY = "17365D"
BLUE = "DDEBF7"
PALE = "F4F7FA"
LIGHT = "D9E2F3"
GRID = "B8C4CE"
TEXT = "1F2933"
RED = "9C1C1C"


def set_run_font(run, size=10.5, bold=None, italic=None, color=TEXT):
    run.font.name = "Arial"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_cell_fill(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color=GRID, size="6"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right"):
        tag = qn(f"w:{edge}")
        border = borders.find(tag)
        if border is None:
            border = OxmlElement(f"w:{edge}")
            borders.append(border)
        border.set(qn("w:val"), "single")
        border.set(qn("w:sz"), size)
        border.set(qn("w:space"), "0")
        border.set(qn("w:color"), color)


def set_cell_margins(cell, top=85, start=105, bottom=85, end=105):
    tc_pr = cell._tc.get_or_add_tcPr()
    mar = tc_pr.first_child_found_in("w:tcMar")
    if mar is None:
        mar = OxmlElement("w:tcMar")
        tc_pr.append(mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def clear_paragraph_borders(paragraph):
    p_pr = paragraph._p.get_or_add_pPr()
    borders = p_pr.find(qn("w:pBdr"))
    if borders is not None:
        p_pr.remove(borders)


def clear_style_borders(style):
    p_pr = style.element.get_or_add_pPr()
    borders = p_pr.find(qn("w:pBdr"))
    if borders is not None:
        p_pr.remove(borders)


def prevent_row_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    if tr_pr.find(qn("w:cantSplit")) is None:
        tr_pr.append(OxmlElement("w:cantSplit"))


def set_cell_text(cell, value, *, size=9.2, bold=False, color=TEXT, align=WD_ALIGN_PARAGRAPH.LEFT):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.alignment = align
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.0
    run = paragraph.add_run(str(value))
    set_run_font(run, size=size, bold=bold, color=color)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    set_cell_margins(cell)
    set_cell_border(cell)


def add_table(doc, headers, rows, widths, *, number_columns=(), font_size=8.9):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    header = table.rows[0]
    set_repeat_table_header(header)
    prevent_row_split(header)
    for col, (cell, value, width) in enumerate(zip(header.cells, headers, widths)):
        cell.width = Inches(width)
        set_cell_fill(cell, NAVY)
        set_cell_text(cell, value, size=font_size, bold=True, color="FFFFFF", align=WD_ALIGN_PARAGRAPH.CENTER)
    for row_index, values in enumerate(rows):
        cells = table.add_row().cells
        prevent_row_split(table.rows[-1])
        for col, (cell, value, width) in enumerate(zip(cells, values, widths)):
            cell.width = Inches(width)
            if row_index % 2 == 1:
                set_cell_fill(cell, PALE)
            align = WD_ALIGN_PARAGRAPH.RIGHT if col in number_columns else WD_ALIGN_PARAGRAPH.LEFT
            if col == 0 and len(headers) > 2:
                align = WD_ALIGN_PARAGRAPH.CENTER
            set_cell_text(cell, value, size=font_size, align=align)
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            cell.width = Inches(width)
    return table


def add_metadata_table(doc, rows):
    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    widths = (1.76, 5.13)
    header = table.rows[0]
    set_repeat_table_header(header)
    prevent_row_split(header)
    for cell, value, width in zip(header.cells, ("Campo", "Registro"), widths):
        cell.width = Inches(width)
        set_cell_fill(cell, NAVY)
        set_cell_text(cell, value, size=8.8, bold=True, color="FFFFFF", align=WD_ALIGN_PARAGRAPH.CENTER)
    for index, (label, value) in enumerate(rows):
        cells = table.add_row().cells
        prevent_row_split(table.rows[-1])
        cells[0].width = Inches(widths[0])
        cells[1].width = Inches(widths[1])
        set_cell_fill(cells[0], NAVY)
        set_cell_text(cells[0], label, size=8.8, bold=True, color="FFFFFF")
        if index % 2:
            set_cell_fill(cells[1], PALE)
        set_cell_text(cells[1], value, size=8.9)
    return table


def add_paragraph(doc, text="", *, size=10.3, bold=False, italic=False, align=None, before=0, after=5, color=TEXT):
    paragraph = doc.add_paragraph()
    if align is not None:
        paragraph.alignment = align
    paragraph.paragraph_format.space_before = Pt(before)
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.line_spacing = 1.12
    run = paragraph.add_run(text)
    set_run_font(run, size=size, bold=bold, italic=italic, color=color)
    return paragraph


def add_heading(doc, text):
    paragraph = doc.add_paragraph(style="Heading 1")
    paragraph.paragraph_format.keep_with_next = True
    paragraph.paragraph_format.space_before = Pt(10)
    paragraph.paragraph_format.space_after = Pt(4)
    run = paragraph.add_run(text)
    set_run_font(run, size=11.5, bold=True, color="000000")
    return paragraph


def fmt_number(value, decimals=2):
    if abs(float(value) - round(float(value))) < 1e-9 and decimals == 0:
        return f"{int(round(float(value))):,}"
    return f"{float(value):,.{decimals}f}"


def fmt_percent(value):
    return f"{float(value):.2f}%"


def add_spacer(doc, points=4):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(points)
    p.paragraph_format.space_before = Pt(0)
    return p


def main():
    latest = max(POSTTESTS.glob("*.json"), key=lambda item: item.stat().st_mtime)
    payload = json.loads(latest.read_text(encoding="utf-8"))
    result = payload["result"]
    cutoff = result["cutoff_date"]
    period = result["period"]
    days = int(result["horizon_days"])

    history = pd.read_csv(HISTORY, parse_dates=["date"])
    start = pd.Timestamp(period["start_date"])
    end = pd.Timestamp(period["end_date"])
    cutoff_timestamp = pd.Timestamp(cutoff)
    period_data = history[(history["date"] >= start) & (history["date"] <= end)].copy()
    prior_start = cutoff_timestamp - pd.Timedelta(days=27)
    prior_data = history[(history["date"] >= prior_start) & (history["date"] <= cutoff_timestamp)].copy()
    actual_by_category = period_data.groupby("category", sort=True)["quantity"].sum()
    reference_weekly = prior_data.groupby("category", sort=True)["quantity"].sum() / 4
    categories = list(result["categories_evaluated"])

    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.58)
    section.bottom_margin = Inches(0.58)
    section.left_margin = Inches(0.78)
    section.right_margin = Inches(0.78)
    section.header_distance = Inches(0.25)
    section.footer_distance = Inches(0.25)

    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    normal.font.size = Pt(10.3)
    normal.font.color.rgb = RGBColor.from_string(TEXT)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.12

    title = doc.styles["Title"]
    title.font.name = "Arial"
    title._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    title._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    title.font.size = Pt(14)
    title.font.bold = True
    title.font.color.rgb = RGBColor(0, 0, 0)
    title.paragraph_format.space_before = Pt(0)
    title.paragraph_format.space_after = Pt(6)
    title.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    clear_style_borders(title)

    heading = doc.styles["Heading 1"]
    heading.font.name = "Arial"
    heading._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    heading._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    heading.font.size = Pt(11.5)
    heading.font.bold = True
    heading.font.color.rgb = RGBColor(0, 0, 0)

    title_paragraph = doc.add_paragraph(style="Title")
    title_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    clear_paragraph_borders(title_paragraph)
    title_run = title_paragraph.add_run("Anexo 2 Instrumentos de recolección de datos")
    set_run_font(title_run, size=14, bold=True, color="000000")

    add_paragraph(
        doc,
        "Este anexo reúne las cinco fichas de registro definidas para evaluar la demanda diaria por categoría de producto. "
        "Como referencia comparativa, PMS-7 se presenta como Pretest y el modelo predictivo como Postest. "
        "Ambos se contrastan con las mismas ventas reales registradas durante siete días.",
        after=4,
    )
    add_paragraph(
        doc,
        "La comparación es una validación histórica y respeta el orden temporal: el entrenamiento llega hasta la fecha de corte y las ventas posteriores se usan solo para evaluar el resultado. "
        "Los indicadores de inventario se dejan como ficha de registro hasta contar con el kardex oficial de la tienda.",
        after=7,
    )

    add_heading(doc, "Información de aplicación")
    add_metadata_table(
        doc,
        [
            ("Investigador", "Guerrero Mora, Carlos David Elian"),
            ("Lugar de estudio", "Minimarket Taully, Comas"),
            ("Técnica", "Análisis documental de ventas consolidadas y kardex oficial cuando corresponda"),
            ("Fecha de corte", "8 de septiembre de 2026"),
            ("Período evaluado", "Del 9 al 15 de septiembre de 2026"),
            ("Observaciones", f"{result['observation_count']} observaciones diarias: 5 categorías durante {days} días"),
            ("Fuente de ventas", f"Historial procesado: {len(history):,} registros, {history['category'].nunique()} categorías y {history['quantity'].sum():,.0f} unidades acumuladas"),
        ],
    )

    add_heading(doc, "Instrumento de cantidad demandada para Pretest y Postest")
    add_metadata_table(
        doc,
        [
            ("Indicador", "Cantidad demandada"),
            ("Fórmula operativa", "CD = suma de las ventas reales por categoría durante el período evaluado"),
            ("Fuente de datos", "Historial de demanda consolidado por fecha y categoría"),
            ("Aplicación", "La misma demanda real se utiliza para contrastar el Pretest y el Postest"),
        ],
    )
    add_spacer(doc, 3)
    demand_rows = []
    for index, category in enumerate(categories, start=1):
        demand_rows.append((index, category, fmt_number(actual_by_category[category], 0)))
    demand_rows.append(("", "Total", fmt_number(actual_by_category.sum(), 0)))
    demand_table = add_table(
        doc,
        ["N", "Categoría", "Cantidad demandada CD"],
        demand_rows,
        [0.48, 3.22, 2.14],
        number_columns=(2,),
        font_size=9.0,
    )
    for cell in demand_table.rows[-1].cells:
        set_cell_fill(cell, BLUE)
        for run in cell.paragraphs[0].runs:
            run.bold = True
    add_paragraph(doc, "Nota. La demanda real total del período fue de 1,671 unidades.", size=8.8, italic=True, after=4)

    add_heading(doc, "Instrumento de índice de salida de inventario")
    add_metadata_table(
        doc,
        [
            ("Indicador", "Índice de salida de inventario"),
            ("Fórmula operativa", "ISI = CD / (SI + EN) x 100"),
            ("Fuente de datos", "Kardex oficial con stock inicial, entradas y stock final por producto o categoría"),
            ("Criterio de cálculo", "SI es el primer saldo del período, EN suma las entradas y SF es el último saldo del período. Para ISI, CD se obtiene del balance SI + EN - SF y debe conciliarse con las ventas registradas."),
            ("Estado de aplicación", "El sistema está preparado para calcularlo al cargar el kardex oficial. No se consignan cifras de inventario de demostración como evidencia."),
        ],
    )
    add_spacer(doc, 3)
    add_table(
        doc,
        ["Campo de registro", "Dato que debe registrar el kardex oficial", "Uso en el cálculo"],
        [
            ("Fecha", "Fecha de cada movimiento o cierre", "Delimita el período evaluado"),
            ("Producto y categoría", "Identificador y categoría comercial", "Permite agrupar los registros"),
            ("SI y EN", "Stock inicial y entradas", "Determina el stock disponible"),
            ("SF", "Stock final del período", "Permite obtener CD por balance"),
            ("ISI", "Resultado calculado por el sistema", "Mide la proporción de salida del stock disponible"),
        ],
        [1.52, 3.16, 2.18],
        font_size=8.7,
    )

    add_heading(doc, "Instrumento de tasa de quiebre de stock")
    add_metadata_table(
        doc,
        [
            ("Indicador", "Tasa de quiebre de stock"),
            ("Fórmula operativa", "TQS = DQS / DD x 100"),
            ("Fuente de datos", "Kardex oficial con fecha, producto, categoría y días sin stock"),
            ("Criterio de cálculo", "DQS cuenta los días con quiebre. DD representa los días calendario del período evaluado."),
            ("Estado de aplicación", "El sistema está preparado para calcularlo al cargar el kardex oficial. No se consignan cifras de inventario de demostración como evidencia."),
        ],
    )
    add_spacer(doc, 3)
    add_table(
        doc,
        ["Campo de registro", "Dato que debe registrar el kardex oficial", "Uso en el cálculo"],
        [
            ("Fecha", "Fecha de control por producto", "Determina los días evaluados"),
            ("Producto y categoría", "Identificador y categoría comercial", "Permite identificar el quiebre por categoría"),
            ("Días sin stock", "Marca o número de días con existencia agotada", "Permite obtener DQS"),
            ("DD", "Cantidad de días del período", "Es el denominador de la tasa"),
            ("TQS", "Resultado calculado por el sistema", "Expresa la frecuencia de desabastecimiento"),
        ],
        [1.52, 3.16, 2.18],
        font_size=8.7,
    )

    add_heading(doc, "Instrumento de estacionalidad para Pretest y Postest")
    add_metadata_table(
        doc,
        [
            ("Indicador", "Índice estacional"),
            ("Fórmula operativa", "IE = VP / VPR"),
            ("Definición de VP", "Ventas reales acumuladas en el período evaluado"),
            ("Definición de VPR", "Promedio semanal de las cuatro semanas completas anteriores al corte, del 12 de agosto al 8 de septiembre de 2026"),
            ("Fuente de datos", "Historial de demanda consolidado por fecha y categoría"),
        ],
    )
    add_spacer(doc, 3)
    seasonal_rows = []
    for index, category in enumerate(categories, start=1):
        vp = actual_by_category[category]
        vpr = reference_weekly[category]
        seasonal_rows.append((index, category, fmt_number(vp, 0), fmt_number(vpr, 2), f"{vp / vpr:.2f}"))
    total_vp = actual_by_category.sum()
    total_vpr = reference_weekly.sum()
    seasonal_rows.append(("", "Total", fmt_number(total_vp, 0), fmt_number(total_vpr, 2), f"{total_vp / total_vpr:.2f}"))
    seasonal_table = add_table(
        doc,
        ["N", "Categoría", "VP", "VPR", "IE"],
        seasonal_rows,
        [0.48, 2.45, 1.15, 1.38, 0.85],
        number_columns=(2, 3, 4),
        font_size=8.9,
    )
    for cell in seasonal_table.rows[-1].cells:
        set_cell_fill(cell, BLUE)
        for run in cell.paragraphs[0].runs:
            run.bold = True
    add_paragraph(doc, "Nota. Un IE menor que 1 indica ventas del período por debajo del promedio de referencia.", size=8.8, italic=True, after=4)

    baseline = result["baseline"]
    model = result["model"]

    add_heading(doc, "Instrumento de error de pronóstico")
    add_metadata_table(
        doc,
        [
            ("Indicador", "Error de pronóstico"),
            ("Referencia Pretest", "PMS-7, promedio móvil simple de siete días calculado de forma recursiva"),
            ("Resultado Postest", "Modelo predictivo del sistema, entrenado solo con datos disponibles hasta el 8 de septiembre de 2026"),
            ("Fuente de datos", "Pronósticos de ambos métodos y ventas reales posteriores a la fecha de corte"),
            ("Métricas", "MAPE es el indicador operacionalizado. WAPE, MAE y RMSE se reportan como métricas complementarias del sistema. Un valor menor representa menor error."),
            ("Período", "Del 9 al 15 de septiembre de 2026"),
        ],
    )
    add_spacer(doc, 3)
    error_rows = []
    for index, category in enumerate(categories, start=1):
        pre_metric = baseline["by_category"][category]
        post_metric = model["by_category"][category]
        error_rows.append((
            index,
            category,
            fmt_number(pre_metric["actual_total"], 0),
            fmt_number(pre_metric["prediction_total"], 2),
            fmt_percent(pre_metric["mape"]),
            fmt_number(post_metric["prediction_total"], 2),
            fmt_percent(post_metric["mape"]),
        ))
    base_metrics = baseline["metrics"]
    model_metrics = model["metrics"]
    error_rows.append((
        "",
        "Total",
        fmt_number(base_metrics["actual_total"], 0),
        fmt_number(base_metrics["prediction_total"], 2),
        fmt_percent(base_metrics["mape"]),
        fmt_number(model_metrics["prediction_total"], 2),
        fmt_percent(model_metrics["mape"]),
    ))
    error_table = add_table(
        doc,
        ["N", "Categoría", "DR", "DP PMS-7", "MAPE PMS-7", "DP modelo", "MAPE modelo"],
        error_rows,
        [0.36, 1.28, 0.65, 0.90, 0.96, 0.90, 0.96],
        number_columns=(2, 3, 4, 5, 6),
        font_size=8.25,
    )
    for cell in error_table.rows[-1].cells:
        set_cell_fill(cell, BLUE)
        for run in cell.paragraphs[0].runs:
            run.bold = True
    add_paragraph(doc, "Nota. DR: demanda real. DP: demanda pronosticada. MAPE se calcula sobre las observaciones diarias de cada categoría.", size=8.6, italic=True, after=4)

    add_heading(doc, "Resumen de la validación histórica")
    comparison_rows = []
    labels = (("WAPE", "wape", "%"), ("MAPE", "mape", "%"), ("MAE", "mae", "unidades"), ("RMSE", "rmse", "unidades"))
    for label, key, unit in labels:
        base_value = base_metrics[key]
        post_value = model_metrics[key]
        difference = post_value - base_value
        formatted_base = fmt_percent(base_value) if unit == "%" else fmt_number(base_value)
        formatted_post = fmt_percent(post_value) if unit == "%" else fmt_number(post_value)
        formatted_diff = ("+" if difference >= 0 else "") + (fmt_percent(difference) if unit == "%" else fmt_number(difference))
        result_text = "PMS-7 tuvo menor error" if base_value < post_value else "Modelo predictivo tuvo menor error"
        comparison_rows.append((label, formatted_base, formatted_post, formatted_diff, result_text))
    add_table(
        doc,
        ["Métrica", "Referencia PMS-7", "Modelo predictivo", "Diferencia", "Resultado"],
        comparison_rows,
        [1.08, 1.23, 1.25, 1.05, 2.28],
        number_columns=(1, 2, 3),
        font_size=8.6,
    )
    conclusion = add_paragraph(
        doc,
        "Interpretación. En el período evaluado, PMS-7 registró menor WAPE, MAPE, MAE y RMSE que el modelo predictivo. "
        "Por ello, este resultado histórico de siete días no demuestra una mejora del modelo frente a la línea base. "
        "Para una conclusión de tesis más sólida se deben ejecutar y reportar varias fechas de corte, manteniendo el mismo procedimiento temporal.",
        size=9.2,
        after=3,
        color=RED,
    )

    footer = section.footer
    footer_paragraph = footer.paragraphs[0]
    footer_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_paragraph.paragraph_format.space_before = Pt(0)
    footer_paragraph.paragraph_format.space_after = Pt(0)
    footer_run = footer_paragraph.add_run("Anexo 2 Instrumentos de recolección de datos")
    set_run_font(footer_run, size=8, color="5B6670")

    doc.core_properties.title = "Anexo 2 Instrumentos de recolección de datos"
    doc.core_properties.author = "Guerrero Mora, Carlos David Elian"
    doc.core_properties.subject = "Instrumentos de recolección de datos para la evaluación histórica de demanda"
    doc.core_properties.comments = ""

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
