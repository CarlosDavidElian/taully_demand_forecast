from __future__ import annotations

import json
from pathlib import Path

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = (
    ROOT
    / "outputs"
    / "01a0b651-7620-73f0-8f6a-8b68b14f7a17"
    / "Anexo_2_Instrumentos_Pretest_y_Postest_corregido.docx"
)
POSTTEST = ROOT / "data" / "posttests" / "postest_20260927t191247z_20260908_7d_70578bbb.json"

NAVY = "17365D"
PALE_BLUE = "EDF3F8"
PALE_GRAY = "F6F8FA"
GRID = "D1DAE3"
TEXT = "1F2933"
MUTED = "56616F"


def set_run_font(run, size=10.0, bold=None, italic=None, color=TEXT):
    run.font.name = "Arial"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def clear_style_borders(style):
    p_pr = style.element.get_or_add_pPr()
    borders = p_pr.find(qn("w:pBdr"))
    if borders is not None:
        p_pr.remove(borders)


def clear_paragraph_borders(paragraph):
    p_pr = paragraph._p.get_or_add_pPr()
    borders = p_pr.find(qn("w:pBdr"))
    if borders is not None:
        p_pr.remove(borders)


def set_cell_fill(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color=GRID, size="6"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right"):
        tag = qn(f"w:{edge}")
        border = borders.find(tag)
        if border is None:
            border = OxmlElement(f"w:{edge}")
            borders.append(border)
        border.set(qn("w:val"), "single")
        border.set(qn("w:sz"), size)
        border.set(qn("w:space"), "0")
        border.set(qn("w:color"), color)


def set_cell_margins(cell, top=92, start=112, bottom=92, end=112):
    tc_pr = cell._tc.get_or_add_tcPr()
    margins = tc_pr.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        tc_pr.append(margins)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def prevent_row_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    if tr_pr.find(qn("w:cantSplit")) is None:
        tr_pr.append(OxmlElement("w:cantSplit"))


def set_cell_text(cell, value, size=9.0, bold=False, color=TEXT, align=WD_ALIGN_PARAGRAPH.LEFT):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.alignment = align
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.03
    run = paragraph.add_run(str(value))
    set_run_font(run, size=size, bold=bold, color=color)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    set_cell_margins(cell)
    set_cell_border(cell)


def add_table(doc, headers, rows, widths, *, number_columns=(), font_size=8.9):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    header = table.rows[0]
    set_repeat_table_header(header)
    prevent_row_split(header)
    for index, (cell, value, width) in enumerate(zip(header.cells, headers, widths)):
        cell.width = Inches(width)
        set_cell_fill(cell, NAVY)
        set_cell_text(
            cell,
            value,
            size=font_size,
            bold=True,
            color="FFFFFF",
            align=WD_ALIGN_PARAGRAPH.CENTER,
        )
    for row_index, values in enumerate(rows):
        cells = table.add_row().cells
        prevent_row_split(table.rows[-1])
        for column, (cell, value, width) in enumerate(zip(cells, values, widths)):
            cell.width = Inches(width)
            if row_index % 2:
                set_cell_fill(cell, PALE_GRAY)
            alignment = WD_ALIGN_PARAGRAPH.RIGHT if column in number_columns else WD_ALIGN_PARAGRAPH.LEFT
            if column == 0 and len(headers) > 2:
                alignment = WD_ALIGN_PARAGRAPH.CENTER
            set_cell_text(cell, value, size=font_size, align=alignment)
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            cell.width = Inches(width)
    return table


def add_metadata_table(doc, rows):
    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    widths = (1.72, 5.18)
    header = table.rows[0]
    set_repeat_table_header(header)
    prevent_row_split(header)
    for cell, value, width in zip(header.cells, ("Campo", "Registro"), widths):
        cell.width = Inches(width)
        set_cell_fill(cell, NAVY)
        set_cell_text(cell, value, size=8.8, bold=True, color="FFFFFF", align=WD_ALIGN_PARAGRAPH.CENTER)
    for index, (label, value) in enumerate(rows):
        cells = table.add_row().cells
        prevent_row_split(table.rows[-1])
        cells[0].width = Inches(widths[0])
        cells[1].width = Inches(widths[1])
        set_cell_fill(cells[0], NAVY)
        set_cell_text(cells[0], label, size=8.75, bold=True, color="FFFFFF")
        if index % 2:
            set_cell_fill(cells[1], PALE_GRAY)
        set_cell_text(cells[1], value, size=8.9)
    return table


def add_paragraph(doc, text="", *, size=10.0, bold=False, italic=False, align=None, before=0, after=5, color=TEXT):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(before)
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.line_spacing = 1.12
    if align is not None:
        paragraph.alignment = align
    run = paragraph.add_run(text)
    set_run_font(run, size=size, bold=bold, italic=italic, color=color)
    return paragraph


def add_note(doc, label, text):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(4)
    paragraph.paragraph_format.space_after = Pt(5)
    paragraph.paragraph_format.line_spacing = 1.1
    label_run = paragraph.add_run(label)
    set_run_font(label_run, size=8.85, bold=True, italic=True, color=MUTED)
    text_run = paragraph.add_run(text)
    set_run_font(text_run, size=8.85, italic=True, color=MUTED)
    return paragraph


def add_heading(doc, text, level=1):
    paragraph = doc.add_paragraph(style=f"Heading {level}")
    paragraph.paragraph_format.keep_with_next = True
    paragraph.paragraph_format.space_before = Pt(10 if level == 1 else 6)
    paragraph.paragraph_format.space_after = Pt(4)
    run = paragraph.add_run(text)
    set_run_font(run, size=11.3 if level == 1 else 10.2, bold=True, color="000000")
    return paragraph


def set_keep_with_next(paragraph):
    paragraph.paragraph_format.keep_with_next = True
    return paragraph


def add_page_break(doc):
    doc.add_page_break()


def fmt(value, decimals=2):
    return f"{float(value):,.{decimals}f}"


def pct(value):
    return f"{float(value):.2f}%"


def add_instrument_header(doc, number, name):
    add_heading(doc, f"Instrumento {number} {name}")
    add_paragraph(
        doc,
        "Ficha para registrar el indicador con el mismo criterio en O1 Pretest y O2 Postest. "
        "Las dos mediciones deben usar igual unidad de análisis, categorías comparables y una ventana de tiempo equivalente.",
        size=9.1,
        after=5,
        color=MUTED,
    )


def add_moment_table(doc, pretest, posttest):
    return add_table(
        doc,
        ["Momento", "Propósito de la medición", "Fuente que se conserva"],
        [
            ("O1 Pretest", pretest, "Registro original y evidencia de la fuente"),
            ("O2 Postest", posttest, "Registro posterior y evidencia de la fuente"),
        ],
        [1.12, 3.52, 2.26],
        font_size=8.7,
    )


def add_field_table(doc, rows):
    return add_table(
        doc,
        ["Campo", "Qué debe registrarse", "Uso en el indicador"],
        rows,
        [1.48, 3.34, 2.08],
        font_size=8.7,
    )


def main():
    result = json.loads(POSTTEST.read_text(encoding="utf-8"))["result"]
    model = result["model"]["metrics"]
    baseline = result["baseline"]["metrics"]

    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.60)
    section.bottom_margin = Inches(0.62)
    section.left_margin = Inches(0.80)
    section.right_margin = Inches(0.80)
    section.header_distance = Inches(0.25)
    section.footer_distance = Inches(0.28)

    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    normal.font.size = Pt(10.0)
    normal.font.color.rgb = RGBColor.from_string(TEXT)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.12

    title_style = doc.styles["Title"]
    title_style.font.name = "Arial"
    title_style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    title_style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    title_style.font.size = Pt(14)
    title_style.font.bold = True
    title_style.font.color.rgb = RGBColor(0, 0, 0)
    title_style.paragraph_format.space_before = Pt(0)
    title_style.paragraph_format.space_after = Pt(7)
    title_style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    clear_style_borders(title_style)

    for style_name, size in (("Heading 1", 11.3), ("Heading 2", 10.2)):
        style = doc.styles[style_name]
        style.font.name = "Arial"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        clear_style_borders(style)

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    clear_paragraph_borders(title)
    set_run_font(title.add_run("Anexo 2 Instrumentos de recolección de datos"), size=14, bold=True, color="000000")

    add_paragraph(
        doc,
        "El presente anexo reúne los instrumentos para medir los indicadores de demanda, inventario, estacionalidad y error de pronóstico. "
        "Cada ficha se aplica en dos momentos comparables: O1 Pretest, antes de la implementación del modelo predictivo, y O2 Postest, después de su aplicación.",
        after=5,
    )
    add_paragraph(
        doc,
        "La comparación requiere la misma unidad de análisis, las mismas categorías comerciales, registros trazables y períodos de igual duración. "
        "Los datos de ventas se obtienen del historial consolidado; los indicadores de inventario se calculan exclusivamente con el kardex oficial de la tienda.",
        after=7,
    )

    add_table(
        doc,
        ["Secuencia", "Descripción", "Evidencia que se conserva"],
        [
            ("O1 Pretest", "Medición inicial de los cinco indicadores antes de aplicar el modelo.", "Fichas de registro, reporte de ventas y kardex del período inicial."),
            ("X Intervención", "Implementación y uso del modelo predictivo para pronosticar la demanda.", "Reporte procesado, fecha de corte, horizonte y archivo de salida."),
            ("O2 Postest", "Medición posterior de los mismos indicadores con el mismo criterio de registro.", "Fichas de registro, reporte de ventas y kardex del período posterior."),
        ],
        [1.22, 3.62, 2.06],
        font_size=8.8,
    )
    add_note(
        doc,
        "Nota metodológica. ",
        "PMS-7 es un método de referencia para comparar el error de pronóstico. No sustituye la medición O1 Pretest ni se presenta como un resultado experimental previo.",
    )

    add_heading(doc, "Criterios de aplicación")
    add_metadata_table(
        doc,
        [
            ("Investigador", "Guerrero Mora, Carlos David Elian"),
            ("Lugar de estudio", "Minimarket Taully, Comas"),
            ("Técnica", "Análisis documental de reportes de ventas, reportes de inventario y salidas del sistema"),
            ("Unidad de análisis", "Registro diario por producto o categoría comercial, según el indicador"),
            ("Criterio comparativo", "Mismo número de días, categorías equivalentes y fuente verificable en O1 y O2"),
            ("Regla de trazabilidad", "Conservar el archivo fuente, fecha de extracción, período evaluado y cálculo aplicado"),
        ],
    )
    add_note(
        doc,
        "Regla de calidad. ",
        "No se reemplazan SI, EN, SF, DQS o DD con estimaciones. Para los indicadores de inventario se requiere el kardex oficial con sus movimientos y cierres.",
    )

    # Instrumento 1
    add_page_break(doc)
    add_instrument_header(doc, "1", "Cantidad demandada")
    add_metadata_table(
        doc,
        [
            ("Indicador", "Cantidad demandada"),
            ("Fórmula", "CD = (SI + EN) - SF"),
            ("Definición", "Unidades demandadas o vendidas durante el período evaluado."),
            ("Unidad de medida", "Unidades"),
            ("Fuente principal", "Kardex oficial y reporte de ventas para la conciliación"),
        ],
    )
    add_paragraph(
        doc,
        "La cantidad demandada se calcula con el balance de inventario. Cuando exista un reporte de ventas, la suma de ventas se usa para conciliar el resultado del balance, no para reemplazar los campos de inventario.",
        size=9.0,
        after=5,
    )
    add_field_table(
        doc,
        [
            ("Período", "Fecha de inicio y fecha de cierre de la medición", "Delimita el cálculo"),
            ("Producto o categoría", "Código o nombre y categoría comercial", "Permite consolidar los registros"),
            ("SI", "Stock existente al inicio del período", "Componente del stock disponible"),
            ("EN", "Unidades ingresadas durante el período", "Componente del stock disponible"),
            ("SF", "Stock existente al cierre del período", "Permite obtener CD"),
            ("CD", "Resultado de la fórmula por producto o categoría", "Indicador de demanda"),
        ],
    )
    add_paragraph(doc, "Aplicación por momento", size=9.5, bold=True, before=7, after=3)
    add_moment_table(
        doc,
        "Registrar CD del período anterior a la implementación del modelo.",
        "Registrar CD de un período posterior de igual duración para comparar el cambio.",
    )
    add_note(
        doc,
        "Fuente. ",
        "Kardex oficial del minimarket y reporte de ventas consolidado del mismo período.",
    )

    # Instrumento 2
    add_page_break(doc)
    add_instrument_header(doc, "2", "Índice de salida de inventario")
    add_metadata_table(
        doc,
        [
            ("Indicador", "Índice de salida de inventario"),
            ("Fórmula", "ISI = CD / (SI + EN) x 100"),
            ("Definición", "Proporción del stock disponible que salió durante el período evaluado."),
            ("Unidad de medida", "Porcentaje"),
            ("Fuente principal", "Kardex oficial con stock inicial, entradas y stock final"),
        ],
    )
    add_paragraph(
        doc,
        "El índice muestra qué parte de la mercadería disponible se convirtió en salida. CD debe provenir del mismo balance de inventario y del mismo período registrado en esta ficha.",
        size=9.0,
        after=5,
    )
    add_field_table(
        doc,
        [
            ("Período", "Fecha de inicio y fecha de cierre de la medición", "Asegura comparabilidad"),
            ("Producto o categoría", "Código o nombre y categoría comercial", "Define el nivel de consolidación"),
            ("SI", "Stock inicial", "Parte del denominador"),
            ("EN", "Entradas de mercadería", "Parte del denominador"),
            ("CD", "Cantidad demandada calculada", "Numerador"),
            ("ISI", "Resultado expresado en porcentaje", "Mide la salida del inventario"),
        ],
    )
    add_paragraph(doc, "Aplicación por momento", size=9.5, bold=True, before=7, after=3)
    add_moment_table(
        doc,
        "Calcular ISI con el kardex del período previo a la intervención.",
        "Calcular ISI con el kardex del período posterior, usando la misma duración.",
    )
    add_note(
        doc,
        "Control de cálculo. ",
        "El porcentaje debe calcularse por producto o categoría antes de elaborar el consolidado del período.",
    )

    # Instrumento 3
    add_page_break(doc)
    add_instrument_header(doc, "3", "Tasa de quiebre de stock")
    add_metadata_table(
        doc,
        [
            ("Indicador", "Tasa de quiebre de stock"),
            ("Fórmula", "TQS = (DQS / DD) x 100"),
            ("Definición", "Porcentaje de días en que el producto o categoría no tuvo stock disponible."),
            ("Unidad de medida", "Porcentaje"),
            ("Fuente principal", "Kardex oficial o control diario de existencias"),
        ],
    )
    add_paragraph(
        doc,
        "La tasa permite identificar la frecuencia de desabastecimiento. Un día cuenta como quiebre cuando el registro de inventario evidencia stock agotado para el producto o categoría evaluada.",
        size=9.0,
        after=5,
    )
    add_field_table(
        doc,
        [
            ("Período", "Fecha de inicio y fecha de cierre", "Define los días evaluados"),
            ("Producto o categoría", "Código o nombre y categoría comercial", "Identifica el objeto medido"),
            ("DQS", "Número de días con stock agotado", "Numerador"),
            ("DD", "Número total de días observados", "Denominador"),
            ("TQS", "Resultado expresado en porcentaje", "Mide el desabastecimiento"),
        ],
    )
    add_paragraph(doc, "Aplicación por momento", size=9.5, bold=True, before=7, after=3)
    add_moment_table(
        doc,
        "Contar los días sin stock en el período anterior a la intervención.",
        "Contar los días sin stock en un período posterior de igual duración.",
    )
    add_note(
        doc,
        "Control de fuente. ",
        "Las ventas no permiten inferir por sí solas un quiebre de stock; se requiere el registro de existencia o el kardex.",
    )

    # Instrumento 4
    add_page_break(doc)
    add_instrument_header(doc, "4", "Estacionalidad")
    add_metadata_table(
        doc,
        [
            ("Indicador", "Índice estacional"),
            ("Fórmula", "IE = VD / VPM"),
            ("Definición", "Relación entre las ventas del período observado y el promedio de ventas del período de referencia."),
            ("Unidad de medida", "Índice"),
            ("Fuente principal", "Historial de demanda consolidado por fecha y categoría"),
        ],
    )
    add_paragraph(
        doc,
        "El índice estacional permite identificar si la demanda observada está por encima o por debajo de su comportamiento promedio. La definición del período de referencia debe conservarse igual en O1 y O2.",
        size=9.0,
        after=5,
    )
    add_field_table(
        doc,
        [
            ("Período observado", "Fechas de la ventana que se evalúa", "Define VD"),
            ("Período de referencia", "Fechas usadas para calcular el promedio comparable", "Define VPM"),
            ("Producto o categoría", "Nombre o categoría comercial", "Permite consolidar ventas"),
            ("VD", "Ventas del período observado", "Numerador"),
            ("VPM", "Promedio de ventas del período de referencia", "Denominador"),
            ("IE", "Resultado del cociente VD / VPM", "Mide el comportamiento estacional"),
        ],
    )
    add_paragraph(doc, "Aplicación por momento", size=9.5, bold=True, before=7, after=3)
    add_moment_table(
        doc,
        "Calcular IE sobre un período anterior a la intervención con su referencia definida.",
        "Calcular IE sobre un período posterior de igual duración y criterio de referencia.",
    )
    add_note(
        doc,
        "Interpretación. ",
        "Un índice mayor que 1 indica ventas por encima del promedio de referencia; un índice menor que 1 indica ventas por debajo de ese promedio.",
    )

    # Instrumento 5
    add_page_break(doc)
    add_instrument_header(doc, "5", "Error de pronóstico")
    add_metadata_table(
        doc,
        [
            ("Indicador", "Error de pronóstico"),
            ("Fórmula principal", "MAPE = (100 / n) x suma de |(DR - DP) / DR|"),
            ("Definición", "Error porcentual medio entre la demanda real y la demanda pronosticada."),
            ("Unidad de medida", "Porcentaje"),
            ("Fuente principal", "Exportación del Postest del programa y ventas reales posteriores al corte"),
            ("Métricas complementarias", "WAPE, MAE y RMSE; un valor menor representa menor error"),
        ],
    )
    add_paragraph(
        doc,
        "Para el Postest, el programa entrena solo con datos hasta la fecha de corte, pronostica los días posteriores y compara el pronóstico con las ventas reales de esos mismos días. "
        "Las métricas internas del entrenamiento no deben registrarse como Postest.",
        size=9.0,
        after=5,
    )
    add_field_table(
        doc,
        [
            ("Fecha de corte", "Última fecha disponible para entrenar", "Separa entrenamiento y evaluación"),
            ("Horizonte", "Número de días que se pronostican", "Define el período evaluado"),
            ("DR", "Demanda real posterior al corte", "Valor de contraste"),
            ("DP", "Demanda pronosticada por el modelo", "Valor evaluado"),
            ("n", "Cantidad de observaciones comparadas", "Base de cálculo del MAPE"),
            ("MAPE, WAPE, MAE, RMSE", "Métricas exportadas por el programa", "Evalúan el error de pronóstico"),
        ],
    )
    add_paragraph(doc, "Aplicación por momento", size=9.5, bold=True, before=7, after=3)
    add_moment_table(
        doc,
        "Registrar el error del método definido antes de la intervención, con sus ventas reales y período equivalente.",
        "Registrar el error del modelo aplicado, con ventas reales posteriores al corte y el mismo horizonte.",
    )
    add_note(
        doc,
        "Referencia PMS-7. ",
        "El programa calcula PMS-7 como comparación de desempeño. Su resultado se registra como línea base metodológica, no como O1 Pretest.",
    )

    # Verified program record
    add_page_break(doc)
    add_heading(doc, "Registro de aplicación del Postest verificable en el programa")
    add_paragraph(
        doc,
        "La siguiente aplicación corresponde a una validación hacia adelante registrada por el sistema. El resultado se conserva en la descarga de Postest y puede repetirse con el mismo corte y horizonte.",
        after=5,
    )
    add_metadata_table(
        doc,
        [
            ("Fecha de corte", "8 de septiembre de 2026"),
            ("Horizonte", "7 días"),
            ("Período evaluado", "Del 9 al 15 de septiembre de 2026"),
            ("Observaciones", "35 observaciones: 5 categorías durante 7 días"),
            ("Demanda real total", "1,671 unidades"),
            ("Procedimiento", "Entrenar hasta el corte, pronosticar el horizonte y comparar con las ventas reales posteriores"),
        ],
    )
    add_paragraph(doc, "Resultados del error de pronóstico", size=9.7, bold=True, before=7, after=3)
    table = add_table(
        doc,
        ["Método", "WAPE", "MAPE", "MAE", "RMSE", "Demanda pronosticada"],
        [
            ("Modelo predictivo", pct(model["wape"]), pct(model["mape"]), fmt(model["mae"]), fmt(model["rmse"]), fmt(model["prediction_total"])),
            ("PMS-7 referencia", pct(baseline["wape"]), pct(baseline["mape"]), fmt(baseline["mae"]), fmt(baseline["rmse"]), fmt(baseline["prediction_total"])),
        ],
        [1.56, 0.86, 0.86, 0.86, 0.86, 1.90],
        number_columns=(1, 2, 3, 4, 5),
        font_size=8.65,
    )
    for cell in table.rows[1].cells:
        set_cell_fill(cell, PALE_BLUE)
    add_note(
        doc,
        "Resultado de esta ejecución. ",
        "PMS-7 tuvo menor WAPE, MAPE, MAE y RMSE en este período. Por ello, esta única corrida no demuestra una mejora del modelo frente a la referencia.",
    )
    add_heading(doc, "Procedimiento de verificación")
    add_table(
        doc,
        ["Paso", "Acción en el programa", "Comprobación"],
        [
            ("1", "Seleccionar horizonte de 7 días.", "El selector muestra 7 días."),
            ("2", "Ingresar 08/09/2026 en Corte de prueba.", "La fecha corresponde al último día de entrenamiento."),
            ("3", "Ir a Postest e indicadores de inventario y seleccionar Ejecutar Postest.", "Se genera la evaluación histórica hacia adelante."),
            ("4", "Comprobar período, registros y métricas.", "Debe mostrar 09/09/2026 a 15/09/2026, 35 observaciones y 5 categorías."),
            ("5", "Seleccionar Descargar Postest.", "El Excel debe incluir Resumen postest, Métricas categoría y Detalle diario."),
        ],
        [0.48, 3.57, 2.45],
        font_size=8.65,
    )
    add_note(
        doc,
        "Evidencia. ",
        "La descarga de Postest conserva la fecha de corte, el período evaluado, las ventas reales, los pronósticos y las métricas. Para ISI y TQS, adjuntar además el kardex oficial del período correspondiente.",
    )

    footer = section.footer
    footer_p = footer.paragraphs[0]
    footer_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_p.paragraph_format.space_before = Pt(0)
    footer_p.paragraph_format.space_after = Pt(0)
    set_run_font(footer_p.add_run("Anexo 2 Instrumentos de recolección de datos"), size=8, color=MUTED)

    doc.core_properties.title = "Anexo 2 Instrumentos de recolección de datos"
    doc.core_properties.author = "Guerrero Mora, Carlos David Elian"
    doc.core_properties.subject = "Instrumentos de Pretest y Postest para evaluación de demanda e inventario"
    doc.core_properties.comments = ""

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
