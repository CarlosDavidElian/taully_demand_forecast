from pathlib import Path

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = (
    ROOT
    / "outputs"
    / "01a0b651-7620-73f0-8f6a-8b68b14f7a17"
    / "Anexo_2_Instrumentos_Completos_Pretest_Postest.docx"
)

NAVY = "17365D"
PALE_BLUE = "EDF3F8"
PALE_GRAY = "F6F8FA"
GRID = "D1DAE3"
TEXT = "1F2933"
MUTED = "56616F"

CATEGORIES = ("ABARROTES", "BEBIDAS", "GOLOSINAS", "HELADOS", "LIMPIEZA")


def set_run_font(run, size=10.0, bold=None, italic=None, color=TEXT):
    run.font.name = "Arial"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def clear_style_borders(style):
    p_pr = style.element.get_or_add_pPr()
    borders = p_pr.find(qn("w:pBdr"))
    if borders is not None:
        p_pr.remove(borders)


def clear_paragraph_borders(paragraph):
    p_pr = paragraph._p.get_or_add_pPr()
    borders = p_pr.find(qn("w:pBdr"))
    if borders is not None:
        p_pr.remove(borders)


def set_cell_fill(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color=GRID, size="6"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right"):
        tag = qn(f"w:{edge}")
        border = borders.find(tag)
        if border is None:
            border = OxmlElement(f"w:{edge}")
            borders.append(border)
        border.set(qn("w:val"), "single")
        border.set(qn("w:sz"), size)
        border.set(qn("w:space"), "0")
        border.set(qn("w:color"), color)


def set_cell_margins(cell, top=88, start=110, bottom=88, end=110):
    tc_pr = cell._tc.get_or_add_tcPr()
    margins = tc_pr.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        tc_pr.append(margins)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def prevent_row_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    if tr_pr.find(qn("w:cantSplit")) is None:
        tr_pr.append(OxmlElement("w:cantSplit"))


def set_cell_text(cell, value, *, size=9.0, bold=False, italic=False, color=TEXT, align=WD_ALIGN_PARAGRAPH.LEFT):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.alignment = align
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.03
    run = paragraph.add_run(str(value))
    set_run_font(run, size=size, bold=bold, italic=italic, color=color)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    set_cell_margins(cell)
    set_cell_border(cell)


def add_table(doc, headers, rows, widths, *, number_columns=(), font_size=8.8):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    header = table.rows[0]
    set_repeat_table_header(header)
    prevent_row_split(header)
    for column, (cell, value, width) in enumerate(zip(header.cells, headers, widths)):
        cell.width = Inches(width)
        set_cell_fill(cell, NAVY)
        set_cell_text(
            cell,
            value,
            size=font_size,
            bold=True,
            color="FFFFFF",
            align=WD_ALIGN_PARAGRAPH.CENTER,
        )
    for row_index, values in enumerate(rows):
        cells = table.add_row().cells
        prevent_row_split(table.rows[-1])
        for column, (cell, value, width) in enumerate(zip(cells, values, widths)):
            cell.width = Inches(width)
            if row_index % 2:
                set_cell_fill(cell, PALE_GRAY)
            align = WD_ALIGN_PARAGRAPH.RIGHT if column in number_columns else WD_ALIGN_PARAGRAPH.LEFT
            if column == 0 and len(headers) > 2:
                align = WD_ALIGN_PARAGRAPH.CENTER
            set_cell_text(cell, value, size=font_size, align=align)
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            cell.width = Inches(width)
    return table


def add_metadata_table(doc, rows):
    return add_table(doc, ("Campo", "Registro"), rows, (1.72, 5.18), font_size=8.8)


def add_paragraph(doc, text="", *, size=10.0, bold=False, italic=False, before=0, after=5, color=TEXT, align=None):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(before)
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.line_spacing = 1.12
    if align is not None:
        paragraph.alignment = align
    set_run_font(paragraph.add_run(text), size=size, bold=bold, italic=italic, color=color)
    return paragraph


def add_heading(doc, text, level=1):
    paragraph = doc.add_paragraph(style=f"Heading {level}")
    paragraph.paragraph_format.keep_with_next = True
    paragraph.paragraph_format.space_before = Pt(10 if level == 1 else 6)
    paragraph.paragraph_format.space_after = Pt(4)
    set_run_font(paragraph.add_run(text), size=11.4 if level == 1 else 10.2, bold=True, color="000000")
    return paragraph


def add_note(doc, label, text):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(4)
    paragraph.paragraph_format.space_after = Pt(4)
    paragraph.paragraph_format.line_spacing = 1.08
    set_run_font(paragraph.add_run(label), size=8.7, bold=True, italic=True, color=MUTED)
    set_run_font(paragraph.add_run(text), size=8.7, italic=True, color=MUTED)
    return paragraph


def add_page_break(doc):
    doc.add_page_break()


def blank(value=""):
    return value


def category_rows(headers):
    rows = []
    for index, category in enumerate(CATEGORIES, start=1):
        values = [index, category] + [blank() for _ in headers[2:]]
        rows.append(values)
    rows.append(["", "TOTAL"] + [blank() for _ in headers[2:]])
    return rows


def make_form(doc, *, moment, number, indicator, formula, definition, unit, source, fields, note):
    add_heading(doc, f"{moment} Instrumento {number} {indicator}")
    add_paragraph(
        doc,
        "Ficha de registro para aplicar el mismo indicador por categoría comercial durante un período definido.",
        size=9.1,
        color=MUTED,
        after=5,
    )
    add_metadata_table(
        doc,
        [
            ("Indicador", indicator),
            ("Investigador", "Guerrero Mora, Carlos David Elian"),
            ("Lugar de estudio", "Minimarket Taully, Comas"),
            ("Momento", moment),
            ("Fórmula", formula),
            ("Definición operativa", definition),
            ("Unidad de medida", unit),
            ("Fuente requerida", source),
            ("Período", "Del ____ / ____ / ______ al ____ / ____ / ______"),
            ("Código de evidencia", "____________________________________________"),
        ],
    )
    add_paragraph(doc, "Registro por categoría", size=9.6, bold=True, before=7, after=3)
    widths = {
        5: (0.40, 1.52, 1.10, 1.10, 2.78),
        6: (0.40, 1.25, 0.95, 0.95, 0.95, 2.40),
        7: (0.36, 1.16, 0.82, 0.82, 0.82, 0.82, 1.85),
    }[len(fields)]
    table = add_table(doc, fields, category_rows(fields), widths, font_size=8.5)
    for cell in table.rows[-1].cells:
        set_cell_fill(cell, PALE_BLUE)
        for run in cell.paragraphs[0].runs:
            run.bold = True
    add_note(doc, "Criterio de registro. ", note)
    add_paragraph(doc, "Responsable del registro: ____________________________________________", size=9.0, before=6, after=0)


def main():
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.58)
    section.bottom_margin = Inches(0.60)
    section.left_margin = Inches(0.80)
    section.right_margin = Inches(0.80)
    section.header_distance = Inches(0.25)
    section.footer_distance = Inches(0.28)

    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    normal.font.size = Pt(10.0)
    normal.font.color.rgb = RGBColor.from_string(TEXT)
    normal.paragraph_format.line_spacing = 1.12
    normal.paragraph_format.space_after = Pt(5)

    title_style = doc.styles["Title"]
    title_style.font.name = "Arial"
    title_style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    title_style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    title_style.font.size = Pt(14)
    title_style.font.bold = True
    title_style.font.color.rgb = RGBColor(0, 0, 0)
    title_style.paragraph_format.space_before = Pt(0)
    title_style.paragraph_format.space_after = Pt(7)
    title_style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    clear_style_borders(title_style)
    for style_name, size in (("Heading 1", 11.4), ("Heading 2", 10.2)):
        style = doc.styles[style_name]
        style.font.name = "Arial"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        clear_style_borders(style)

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    clear_paragraph_borders(title)
    set_run_font(title.add_run("Anexo 2 Instrumentos de recolección de datos"), size=14, bold=True, color="000000")
    add_paragraph(
        doc,
        "Este anexo contiene las fichas para registrar los cinco indicadores del estudio en los dos momentos de medición. "
        "Cada ficha debe completarse con evidencia primaria de la tienda y conservar su archivo de respaldo.",
        after=5,
    )
    add_paragraph(
        doc,
        "O1 Pretest corresponde a la medición previa a la implementación del modelo. X corresponde a la implementación y uso del modelo predictivo. "
        "O2 Postest corresponde a la medición posterior, usando igual período de observación, categorías y unidad de medida.",
        after=7,
    )
    add_table(
        doc,
        ("Secuencia", "Aplicación", "Evidencia mínima"),
        [
            ("O1 Pretest", "Registrar los cinco indicadores antes de aplicar el modelo.", "Reporte de ventas y kardex del período inicial."),
            ("X Intervención", "Procesar ventas, entrenar el modelo y generar pronóstico.", "Reporte procesado, fecha de corte y archivo de salida."),
            ("O2 Postest", "Registrar los mismos cinco indicadores después de aplicar el modelo.", "Reporte de ventas, kardex y descarga del Postest."),
        ],
        (1.18, 3.64, 2.08),
        font_size=8.7,
    )
    add_note(
        doc,
        "Comparación de pronóstico. ",
        "PMS-7 es una línea de referencia para contrastar el error del modelo. No sustituye la ficha O1 Pretest.",
    )
    add_heading(doc, "Instrumentos incluidos")
    add_table(
        doc,
        ("N", "Indicador", "Fórmula"),
        [
            ("1", "Cantidad demandada", "CD = (SI + EN) - SF"),
            ("2", "Índice de salida de inventario", "ISI = CD / (SI + EN) x 100"),
            ("3", "Tasa de quiebre de stock", "TQS = (DQS / DD) x 100"),
            ("4", "Estacionalidad", "IE = VD / VPM"),
            ("5", "Error de pronóstico", "MAPE = (100 / n) x suma de |(DR - DP) / DR|"),
        ],
        (0.50, 3.32, 3.08),
        font_size=8.7,
    )
    add_note(
        doc,
        "Regla de fuente. ",
        "Para CD, ISI y TQS se requiere un kardex oficial con stock inicial, entradas, stock final y días sin stock. Para IE se requiere ventas reales consolidadas. Para el error de pronóstico se requiere la exportación del Postest y las ventas reales posteriores al corte.",
    )

    # O1 Pretest forms
    add_page_break(doc)
    make_form(
        doc,
        moment="O1 Pretest",
        number="1",
        indicator="Cantidad demandada",
        formula="CD = (SI + EN) - SF",
        definition="Unidades demandadas durante el período evaluado.",
        unit="Unidades",
        source="Kardex oficial y reporte de ventas del mismo período.",
        fields=("N", "Categoría", "SI", "EN", "SF", "CD"),
        note="SI es el stock al inicio; EN son las entradas; SF es el stock al cierre. CD se obtiene por balance y se concilia con ventas.",
    )
    add_page_break(doc)
    make_form(
        doc,
        moment="O1 Pretest",
        number="2",
        indicator="Índice de salida de inventario",
        formula="ISI = CD / (SI + EN) x 100",
        definition="Proporción del stock disponible que salió durante el período evaluado.",
        unit="Porcentaje",
        source="Kardex oficial con stock inicial, entradas y stock final.",
        fields=("N", "Categoría", "CD", "SI", "EN", "ISI"),
        note="Registrar los valores de CD, SI y EN del mismo período. El resultado se expresa en porcentaje.",
    )
    add_page_break(doc)
    make_form(
        doc,
        moment="O1 Pretest",
        number="3",
        indicator="Tasa de quiebre de stock",
        formula="TQS = (DQS / DD) x 100",
        definition="Porcentaje de días en que la categoría tuvo stock agotado.",
        unit="Porcentaje",
        source="Kardex oficial o control diario de existencias.",
        fields=("N", "Categoría", "DQS", "DD", "TQS"),
        note="DQS es el número de días sin stock. DD es el número total de días del período observado.",
    )
    add_page_break(doc)
    make_form(
        doc,
        moment="O1 Pretest",
        number="4",
        indicator="Estacionalidad",
        formula="IE = VD / VPM",
        definition="Relación entre las ventas del período observado y el promedio de ventas del período de referencia.",
        unit="Índice",
        source="Reporte de ventas reales consolidado por fecha y categoría.",
        fields=("N", "Categoría", "VD", "VPM", "IE"),
        note="VD son las ventas del período observado. VPM es el promedio de ventas del período de referencia definido por el investigador.",
    )
    add_page_break(doc)
    make_form(
        doc,
        moment="O1 Pretest",
        number="5",
        indicator="Error de pronóstico",
        formula="MAPE = (100 / n) x suma de |(DR - DP) / DR|",
        definition="Error porcentual medio entre la demanda real y la demanda pronosticada.",
        unit="Porcentaje y unidades",
        source="Método de referencia definido antes de la intervención y ventas reales del mismo período.",
        fields=("N", "Categoría", "DR", "DP referencia", "MAPE", "MAE", "RMSE"),
        note="DR es la demanda real. DP referencia es el pronóstico del método previo. Registrar MAPE, MAE y RMSE usando las mismas observaciones.",
    )

    # O2 Postest forms
    add_page_break(doc)
    make_form(
        doc,
        moment="O2 Postest",
        number="1",
        indicator="Cantidad demandada",
        formula="CD = (SI + EN) - SF",
        definition="Unidades demandadas durante el período evaluado después de aplicar el modelo.",
        unit="Unidades",
        source="Kardex oficial y reporte de ventas del mismo período.",
        fields=("N", "Categoría", "SI", "EN", "SF", "CD"),
        note="Mantener la misma duración y categorías usadas en O1. CD se obtiene por balance y se concilia con ventas.",
    )
    add_page_break(doc)
    make_form(
        doc,
        moment="O2 Postest",
        number="2",
        indicator="Índice de salida de inventario",
        formula="ISI = CD / (SI + EN) x 100",
        definition="Proporción del stock disponible que salió durante el período posterior a la implementación.",
        unit="Porcentaje",
        source="Kardex oficial con stock inicial, entradas y stock final.",
        fields=("N", "Categoría", "CD", "SI", "EN", "ISI"),
        note="Usar el mismo procedimiento de cálculo de O1 y registrar la fuente del kardex.",
    )
    add_page_break(doc)
    make_form(
        doc,
        moment="O2 Postest",
        number="3",
        indicator="Tasa de quiebre de stock",
        formula="TQS = (DQS / DD) x 100",
        definition="Porcentaje de días en que la categoría tuvo stock agotado después de aplicar el modelo.",
        unit="Porcentaje",
        source="Kardex oficial o control diario de existencias.",
        fields=("N", "Categoría", "DQS", "DD", "TQS"),
        note="Mantener igual número de días que en O1. Cada día sin stock debe estar respaldado por el control de existencias.",
    )
    add_page_break(doc)
    make_form(
        doc,
        moment="O2 Postest",
        number="4",
        indicator="Estacionalidad",
        formula="IE = VD / VPM",
        definition="Relación entre las ventas posteriores y el promedio de ventas del período de referencia.",
        unit="Índice",
        source="Reporte de ventas reales consolidado por fecha y categoría.",
        fields=("N", "Categoría", "VD", "VPM", "IE"),
        note="Mantener el mismo criterio de período de referencia empleado en O1 para que la comparación sea válida.",
    )
    add_page_break(doc)
    make_form(
        doc,
        moment="O2 Postest",
        number="5",
        indicator="Error de pronóstico",
        formula="MAPE = (100 / n) x suma de |(DR - DP) / DR|",
        definition="Error porcentual medio entre la demanda real posterior al corte y el pronóstico del modelo.",
        unit="Porcentaje y unidades",
        source="Descarga de Postest del programa y ventas reales posteriores a la fecha de corte.",
        fields=("N", "Categoría", "DR", "DP modelo", "MAPE", "MAE", "RMSE"),
        note="Seleccionar fecha de corte y horizonte. El Postest debe comparar únicamente ventas reales posteriores al corte con el pronóstico del modelo.",
    )
    add_paragraph(doc, "Datos de ejecución del Postest", size=9.6, bold=True, before=7, after=3)
    add_table(
        doc,
        ("Campo", "Dato que debe conservarse"),
        [
            ("Fecha de corte", "Última fecha usada para entrenar el modelo."),
            ("Horizonte", "Número de días pronosticados."),
            ("Período evaluado", "Fechas posteriores al corte comparadas con ventas reales."),
            ("Archivo de evidencia", "Excel descargado desde el botón Descargar Postest."),
            ("Comparación", "MAPE, MAE y RMSE del modelo y de PMS-7 como referencia."),
        ],
        (1.72, 5.18),
        font_size=8.6,
    )

    # Traceability page
    add_page_break(doc)
    add_heading(doc, "Hoja de trazabilidad de los instrumentos")
    add_paragraph(
        doc,
        "Completar esta hoja al aplicar las fichas. Permite identificar la fuente, el período y el responsable de cada medición.",
        after=5,
    )
    add_table(
        doc,
        ("Momento", "Instrumento", "Período", "Código del reporte o kardex", "Responsable"),
        [
            ("O1 Pretest", "1 Cantidad demandada", "", "", ""),
            ("O1 Pretest", "2 Índice de salida", "", "", ""),
            ("O1 Pretest", "3 Quiebre de stock", "", "", ""),
            ("O1 Pretest", "4 Estacionalidad", "", "", ""),
            ("O1 Pretest", "5 Error de pronóstico", "", "", ""),
            ("O2 Postest", "1 Cantidad demandada", "", "", ""),
            ("O2 Postest", "2 Índice de salida", "", "", ""),
            ("O2 Postest", "3 Quiebre de stock", "", "", ""),
            ("O2 Postest", "4 Estacionalidad", "", "", ""),
            ("O2 Postest", "5 Error de pronóstico", "", "", ""),
        ],
        (1.15, 1.65, 1.10, 2.02, 0.98),
        font_size=8.2,
    )
    add_heading(doc, "Documentos que se deben conservar")
    add_table(
        doc,
        ("Documento", "Uso"),
        [
            ("Reporte de ventas real", "Sustenta VD, demanda real y comparación de pronósticos."),
            ("Kardex oficial", "Sustenta SI, EN, SF y días sin stock para CD, ISI y TQS."),
            ("Descarga de Postest", "Sustenta fecha de corte, horizonte, demanda real, pronóstico y métricas."),
            ("Ficha de trazabilidad", "Vincula cada resultado con su fuente y período."),
        ],
        (2.02, 4.88),
        font_size=8.8,
    )
    add_note(
        doc,
        "Cierre. ",
        "Registrar resultados únicamente cuando la fuente corresponde a la evidencia primaria de la tienda y puede ser verificada.",
    )

    footer = section.footer
    footer_paragraph = footer.paragraphs[0]
    footer_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_paragraph.paragraph_format.space_before = Pt(0)
    footer_paragraph.paragraph_format.space_after = Pt(0)
    set_run_font(footer_paragraph.add_run("Anexo 2 Instrumentos de recolección de datos"), size=8, color=MUTED)

    doc.core_properties.title = "Anexo 2 Instrumentos de recolección de datos"
    doc.core_properties.author = "Guerrero Mora, Carlos David Elian"
    doc.core_properties.subject = "Fichas de Pretest y Postest para el estudio de pronóstico de demanda"
    doc.core_properties.comments = ""

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
