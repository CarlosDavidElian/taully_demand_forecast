from pathlib import Path
from collections import defaultdict
from datetime import datetime
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import openpyxl

ROOT = Path('C:/taully_demand_forecast')
OUT = ROOT / 'outputs/postest_visualizado'
OUT.mkdir(exist_ok=True, parents=True)
source = ROOT / 'data/inventario_producto_Taully_2026-04-01_a_2026-09-15.xlsx'
book = openpyxl.load_workbook(source, read_only=True, data_only=True)
rows = list(book['Datos inventario'].values)[1:]
groups = defaultdict(list)
for row in rows:
    if datetime(2026, 9, 9) <= row[0] <= datetime(2026, 9, 15):
        groups[(row[4], row[1])].append(row)
metrics = defaultdict(lambda: [0, 0, 0, 0])
for (category, product), records in groups.items():
    records.sort(key=lambda r: r[0])
    vals = metrics[category]
    vals[0] += records[0][6]
    vals[1] += sum(r[7] for r in records)
    vals[2] += records[-1][9]
    vals[3] += sum(r[5] for r in records)
assert sum(v[0] + v[1] - v[2] for v in metrics.values()) == 1311
assert sum(v[3] for v in metrics.values()) == 1671
book.close()

doc = Document()
s = doc.sections[0]
s.page_width, s.page_height = Inches(8.27), Inches(11.69)
s.top_margin = s.bottom_margin = Inches(.70)
s.left_margin = s.right_margin = Inches(.72)
for name in ['Normal', 'Title', 'Heading 1', 'Heading 2']:
    st = doc.styles[name]
    st.font.name = 'Arial'
    st.font.color.rgb = RGBColor(0,0,0)
doc.styles['Normal'].font.size = Pt(10.5)
doc.styles['Normal'].paragraph_format.space_after = Pt(7)
doc.styles['Normal'].paragraph_format.line_spacing = 1.10
doc.styles['Title'].font.size = Pt(21)
doc.styles['Heading 1'].font.size = Pt(14)
doc.styles['Heading 2'].font.size = Pt(11.5)

def p(text, style=None):
    return doc.add_paragraph(text, style)

def table(headers, values, widths):
    t = doc.add_table(rows=1, cols=len(headers))
    t.autofit = False
    for col, width in zip(t.columns, widths): col.width = Inches(width)
    for cell, text in zip(t.rows[0].cells, headers): cell.text = text
    for vals in values:
        for cell, val in zip(t.add_row().cells, vals): cell.text = str(val)
    for i, row in enumerate(t.rows):
        trpr = row._tr.get_or_add_trPr()
        trpr.append(OxmlElement('w:cantSplit'))
        if i == 0: trpr.append(OxmlElement('w:tblHeader'))
        for j, cell in enumerate(row.cells):
            cell.width = Inches(widths[j])
            props = cell._tc.get_or_add_tcPr()
            sh = OxmlElement('w:shd'); sh.set(qn('w:fill'), '17365D' if i == 0 else ('F0F4F7' if i%2 == 0 else 'FFFFFF')); props.append(sh)
            borders = OxmlElement('w:tcBorders')
            for edge in ['top','left','bottom','right']:
                el = OxmlElement('w:'+edge); el.set(qn('w:val'),'single'); el.set(qn('w:sz'),'4'); el.set(qn('w:color'),'D9D9D9'); borders.append(el)
            props.append(borders)
            mar=OxmlElement('w:tcMar')
            for edge in ['top','bottom','left','right']:
                el=OxmlElement('w:'+edge); el.set(qn('w:w'),'90'); el.set(qn('w:type'),'dxa'); mar.append(el)
            props.append(mar)
            for para in cell.paragraphs:
                para.paragraph_format.space_after = Pt(0)
                for run in para.runs:
                    run.font.size = Pt(9.5)
                    if i == 0: run.font.bold = True; run.font.color.rgb = RGBColor(255,255,255)
    p('') .paragraph_format.space_after = Pt(0)
    return t

p('Guía para verificar los resultados de inventario del postest', 'Title')
p('Minimarket Taully  |  Período del 9 al 15 de septiembre de 2026')
p('Esta guía permite reconstruir los valores de stock inicial, entradas y stock final desde el inventario, y comprobar la cantidad calculada por balance y el índice de salida. Una coincidencia matemática confirma el cálculo; la validez de los datos debe comprobarse con los registros del negocio.')
p('1 Localizar los resultados', 'Heading 1')
for text in [
    'Abre el programa en http://127.0.0.1:5000. En Corte de prueba selecciona 08/09/2026 y en Horizonte elige 7 días.',
    'En Postest e indicadores de inventario, pulsa Ejecutar Postest. Comprueba que el resultado indique 09/09/2026 al 15/09/2026.',
    'Pulsa Descargar Postest y abre la hoja Inventario del Excel. Revisa el nombre del archivo de origen que aparece sobre la tabla.',
    'Abre también el archivo original inventario_producto_Taully_2026-04-01_a_2026-09-15.xlsx, ubicado en la carpeta data del proyecto. Su hoja es Datos inventario. Trabaja sobre copias para conservar los originales.'
]: p(text, 'List Number')
p('2 Reconocer los datos', 'Heading 1')
table(['Campo', 'Qué representa en el cálculo del programa'], [
    ('SI', 'Suma del primer stock inicial disponible de cada producto dentro del período.'),
    ('EN', 'Suma de todas las entradas de los productos durante el período.'),
    ('SF', 'Suma del último stock final disponible de cada producto dentro del período.'),
    ('CD', 'Salida calculada por balance: SI + EN − SF.'),
    ('ISI', 'Porcentaje de salida: CD ÷ (SI + EN) × 100.')
], [0.65, 6.15])
p('Origen de los valores', 'Heading 2')
p('SI, EN y SF proceden del archivo de inventario. CD e ISI se calculan con esos datos; no son predicciones de machine learning. La etiqueta CD del programa no prueba por sí sola la demanda real: también deben conciliarse ventas, mermas, devoluciones y ajustes.')

doc.add_page_break()
p('3 Reconstruir las cifras desde el archivo original', 'Heading 1')
p('En Datos inventario, las columnas relevantes son A Fecha, B Producto, E Categoría, F Cantidad_Vendida, G Stock_Inicial, H Entradas y J Stock_Final.')
for text in [
    'Activa los filtros de Excel. Filtra Fecha desde el 09/09/2026 hasta el 15/09/2026, incluyendo ambos días, y Categoría por ABARROTES.',
    'Ordena por Producto y después por Fecha, de la más antigua a la más reciente. Verifica que no haya registros duplicados del mismo producto y día.',
    'En una hoja auxiliar, crea una sola fila por producto. Anota el Stock_Inicial de su primera fila visible, suma sus Entradas de los siete días y anota el Stock_Final de su última fila visible. Conserva las fechas de esos dos saldos para poder revisarlas.',
    'Suma las tres columnas de la hoja auxiliar para obtener SI, EN y SF de ABARROTES. Repite el procedimiento para BEBIDAS, GOLOSINAS, HELADOS y LIMPIEZA.'
]: p(text, 'List Number')
p('No sumes los stocks de todos los días: repetirías las mismas existencias. Las entradas sí se suman porque son movimientos. Si faltan registros del primer o último día, el programa usa el primer o último registro disponible del producto; anota esa limitación y solicita los saldos de las fechas correctas para validar el período completo.')
p('Ejemplo de ABARROTES', 'Heading 2')
p('SI = 383 unidades; EN = 175 unidades; SF = 276 unidades.\nCD = 383 + 175 − 276 = 282 unidades.\nISI = 282 ÷ (383 + 175) × 100 = 50,54 %.')
p('Valores de referencia del archivo utilizado', 'Heading 2')
ref=[]
for cat,v in sorted(metrics.items()):
    si,en,sf,_=v; cd=si+en-sf
    ref.append([cat, f'{si:.0f}', f'{en:.0f}', f'{sf:.0f}', f'{cd:.0f}', f'{cd/(si+en)*100:.2f} %'.replace('.',',')])
si,en,sf=[sum(v[i] for v in metrics.values()) for i in range(3)]
ref.append(['TOTAL',str(int(si)),str(int(en)),str(int(sf)),str(int(si+en-sf)),f'{(si+en-sf)/(si+en)*100:.2f} %'.replace('.',',')])
table(['Categoría','SI','EN','SF','CD','ISI'],ref,[1.7,.85,.85,.85,.85,1.7])
p('El ISI total se calcula con los totales: 1311 ÷ (1965 + 786) × 100 = 47,66 %. No se obtiene promediando los cinco porcentajes. La fila TOTAL se añade para esta comprobación; la exportación del programa presenta las categorías.')

doc.add_page_break()
p('4 Comprobar las fórmulas en el Excel exportado', 'Heading 1')
p('En la hoja Inventario, la fila 6 contiene los títulos: A Categoría, B SI, C EN, D SF, E CD y F ISI. ABARROTES está en la fila 7 y LIMPIEZA en la fila 11. Conserva esos resultados y realiza la revisión en una hoja nueva.')
p('Copia A6:F11 a una hoja llamada Verificación, manteniendo las mismas posiciones. En esa copia, añade las siguientes columnas. Usa Excel en español y aplica las fórmulas desde la fila 7 hasta la 11.')
table(['Celda', 'Fórmula o título', 'Qué comprueba'],[
    ('G6 / G7', 'CD recalculada\n=B7+C7-D7', 'Debe coincidir con E7.'),
    ('H6 / H7', 'ISI recalculado\n=G7/(B7+C7)', 'Aplica formato Porcentaje con dos decimales. Debe verse igual que F7.'),
    ('I6 / I7', 'Diferencia de CD\n=G7-E7', 'El resultado esperado es 0.')
],[.8,2.6,3.4])
p('El exportador guarda ISI como un número, por ejemplo 50,54, y le añade un signo % visual. La fórmula de H7 devuelve 0,5054 y el formato Porcentaje muestra 50,54 %. Por eso, para una comparación numérica usa H7*100 frente a F7; no compares H7 directamente con F7.')
p('Para el total, escribe =SUMA(B7:B11) en B12 y repite para C12, D12 y E12. En G12 escribe =B12+C12-D12 y en H12 escribe =G12/(B12+C12), con formato Porcentaje. Si SI + EN es cero, registra ISI como no calculable y revisa el caso antes de interpretarlo.')
p('5 Contrastar con ventas y registrar la revisión', 'Heading 1')
p('Filtra Cantidad_Vendida del inventario por el mismo período y categorías, y contrástala con los reportes POS. En el archivo utilizado, las ventas registradas suman 1671 unidades y la salida por balance es 1311: existe una diferencia de 360 unidades. No deben presentarse ambas cifras como equivalentes sin conciliarlas.')
p('Revisa fechas faltantes, duplicados, continuidad del saldo final de un día con el inicial del siguiente, entradas, devoluciones, mermas y ajustes. Conserva el inventario original, el Excel exportado y la hoja auxiliar por producto. Registra fecha de revisión, responsable y diferencias encontradas.')
p('Alcance de esta comprobación', 'Heading 2')
p('El inventario usado por el programa está identificado como demostración. Estos valores sirven para reproducir el cálculo técnico; para sustentar resultados reales de la tesis se requiere el kardex del negocio, sus reportes de ventas y la conciliación de los movimientos. Repetir el cálculo con éxito no convierte los datos demostrativos en evidencia real.')
doc.core_properties.title='Guía para verificar los resultados de inventario del postest'
doc.core_properties.subject='Verificación del período 09/09/2026 al 15/09/2026'
for style in doc.styles:
    for border in list(style.element.iter(qn('w:pBdr'))):
        border.getparent().remove(border)
for paragraph in doc.paragraphs:
    for border in list(paragraph._p.iter(qn('w:pBdr'))):
        border.getparent().remove(border)
doc.save(OUT / 'Guia_verificacion_inventario_postest.docx')
print(OUT / 'Guia_verificacion_inventario_postest.docx')
print(ref)
