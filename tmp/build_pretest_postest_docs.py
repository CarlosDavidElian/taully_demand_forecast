from __future__ import annotations

import json
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from build_anexo2_final import (
    NAVY,
    TEXT,
    add_heading,
    add_metadata_table,
    add_paragraph,
    add_spacer,
    add_table,
    clear_paragraph_borders,
    clear_style_borders,
    fmt_number,
    fmt_percent,
    set_run_font,
)


ROOT = Path(__file__).resolve().parents[1]
OUTPUT_DIR = ROOT / "outputs" / "01a0b651-7620-73f0-8f6a-8b68b14f7a17"
POSTTEST_DIR = ROOT / "data" / "posttests"


def spanish_date(value: str) -> str:
    months = (
        "enero", "febrero", "marzo", "abril", "mayo", "junio",
        "julio", "agosto", "septiembre", "octubre", "noviembre", "diciembre",
    )
    parsed = date.fromisoformat(value)
    return f"{parsed.day} de {months[parsed.month - 1]} de {parsed.year}"


def configure_document(doc: Document, title_text: str, footer_text: str) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.60)
    section.left_margin = Inches(0.78)
    section.right_margin = Inches(0.78)
    section.header_distance = Inches(0.25)
    section.footer_distance = Inches(0.25)

    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    normal.font.size = Pt(10.4)
    normal.font.color.rgb = RGBColor.from_string(TEXT)
    normal.paragraph_format.line_spacing = 1.12
    normal.paragraph_format.space_after = Pt(5)

    title = doc.styles["Title"]
    title.font.name = "Arial"
    title._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    title._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    title.font.size = Pt(14)
    title.font.bold = True
    title.font.color.rgb = RGBColor(0, 0, 0)
    title.paragraph_format.space_before = Pt(0)
    title.paragraph_format.space_after = Pt(6)
    title.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    clear_style_borders(title)

    heading = doc.styles["Heading 1"]
    heading.font.name = "Arial"
    heading._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    heading._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    heading.font.size = Pt(11.5)
    heading.font.bold = True
    heading.font.color.rgb = RGBColor(0, 0, 0)

    paragraph = doc.add_paragraph(style="Title")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    clear_paragraph_borders(paragraph)
    run = paragraph.add_run(title_text)
    set_run_font(run, size=14, bold=True, color="000000")

    footer = section.footer
    footer_p = footer.paragraphs[0]
    footer_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_p.paragraph_format.space_before = Pt(0)
    footer_p.paragraph_format.space_after = Pt(0)
    footer_run = footer_p.add_run(footer_text)
    set_run_font(footer_run, size=8, color="5B6670")


def add_result_table(doc: Document, categories: list[str], by_category: dict, *, prediction_label: str):
    rows = []
    for number, category in enumerate(categories, start=1):
        metric = by_category[category]
        rows.append((
            number,
            category,
            fmt_number(metric["actual_total"], 0),
            fmt_number(metric["prediction_total"], 2),
            fmt_percent(metric["wape"]),
            fmt_percent(metric["mape"]),
            fmt_number(metric["mae"]),
            fmt_number(metric["rmse"]),
        ))
    return rows


def add_metrics_table(doc: Document, metrics: dict, method: str):
    add_table(
        doc,
        ["Métrica", "Resultado", "Lectura"],
        [
            ("WAPE", fmt_percent(metrics["wape"]), "Error ponderado respecto de la demanda total"),
            ("MAPE", fmt_percent(metrics["mape"]), "Error porcentual absoluto medio"),
            ("MAE", fmt_number(metrics["mae"]), "Error absoluto medio en unidades"),
            ("RMSE", fmt_number(metrics["rmse"]), "Penaliza los errores de mayor magnitud"),
            ("Método", method, "Un valor menor de error representa mejor ajuste"),
        ],
        [1.15, 1.32, 4.42],
        number_columns=(1,),
        font_size=8.9,
    )


def build_pretest(result: dict, run_id: str) -> Path:
    doc = Document()
    configure_document(doc, "Pretest Validación histórica PMS 7", "Pretest Validación histórica PMS 7")

    period = result["period"]
    baseline = result["baseline"]
    categories = list(result["categories_evaluated"])
    metrics = baseline["metrics"]

    add_paragraph(
        doc,
        "Esta ficha registra la línea base PMS-7 para comparar la exactitud del modelo predictivo. "
        "PMS-7 utiliza un promedio móvil simple de siete días y se contrasta con las mismas ventas reales que se emplean en el Postest.",
        after=6,
    )
    add_heading(doc, "Datos de aplicación")
    add_metadata_table(
        doc,
        [
            ("Investigador", "Guerrero Mora, Carlos David Elian"),
            ("Lugar de estudio", "Minimarket Taully, Comas"),
            ("Fecha de corte", spanish_date(result["cutoff_date"])),
            ("Período evaluado", f"Del {spanish_date(period['start_date'])} al {spanish_date(period['end_date'])}"),
            ("Método", "PMS-7, promedio móvil simple de siete días calculado de forma recursiva"),
            ("Observaciones", f"{result['observation_count']} observaciones diarias de {len(categories)} categorías"),
            ("Fuente", "Ventas reales consolidadas y pronóstico base generado antes del período evaluado"),
        ],
    )
    add_heading(doc, "Resultados por categoría")
    rows = add_result_table(doc, categories, baseline["by_category"], prediction_label="DP PMS-7")
    rows.append((
        "", "Total", fmt_number(metrics["actual_total"], 0), fmt_number(metrics["prediction_total"], 2),
        fmt_percent(metrics["wape"]), fmt_percent(metrics["mape"]), fmt_number(metrics["mae"]), fmt_number(metrics["rmse"]),
    ))
    table = add_table(
        doc,
        ["N", "Categoría", "DR", "DP PMS-7", "WAPE", "MAPE", "MAE", "RMSE"],
        rows,
        [0.36, 1.32, 0.66, 0.88, 0.74, 0.74, 0.61, 0.70],
        number_columns=(2, 3, 4, 5, 6, 7),
        font_size=8.25,
    )
    for cell in table.rows[-1].cells:
        for run in cell.paragraphs[0].runs:
            run.bold = True
    add_paragraph(doc, "Nota. DR: demanda real. DP: demanda pronosticada. MAPE se calcula a partir de las observaciones diarias de cada categoría.", size=8.6, italic=True, after=5)

    add_heading(doc, "Resumen de métricas")
    add_metrics_table(doc, metrics, "PMS-7")
    add_paragraph(
        doc,
        "Alcance. Esta es una validación histórica de siete días. La ficha representa la línea base de comparación y no mide una intervención realizada en la tienda.",
        size=9.2,
        after=2,
    )

    doc.core_properties.title = "Pretest Validación histórica PMS 7"
    doc.core_properties.author = "Guerrero Mora, Carlos David Elian"
    doc.core_properties.subject = f"Pretest PMS-7. Ejecución {run_id}."
    output = OUTPUT_DIR / "Pretest_PMS_7_validacion_historica_2026_09_09_a_2026_09_15.docx"
    doc.save(output)
    return output


def build_postest(result: dict, run_id: str) -> Path:
    doc = Document()
    configure_document(doc, "Postest Validación histórica Modelo predictivo", "Postest Validación histórica Modelo predictivo")
    # El Postest contiene una tabla adicional de comparación; se ajustan sus márgenes
    # para conservar la ficha completa en una sola página sin reducir la legibilidad.
    doc.sections[0].top_margin = Inches(0.52)
    doc.sections[0].bottom_margin = Inches(0.46)

    period = result["period"]
    model = result["model"]
    baseline = result["baseline"]
    categories = list(result["categories_evaluated"])
    metrics = model["metrics"]

    add_paragraph(
        doc,
        "Esta ficha registra la evaluación del modelo predictivo con las mismas ventas reales usadas para la referencia PMS-7. "
        "El modelo se entrenó únicamente con información anterior a la fecha de corte.",
        after=4,
    )
    add_heading(doc, "Datos de aplicación")
    add_metadata_table(
        doc,
        [
            ("Investigador", "Guerrero Mora, Carlos David Elian"),
            ("Lugar de estudio", "Minimarket Taully, Comas"),
            ("Fecha de corte", spanish_date(result["cutoff_date"])),
            ("Período evaluado", f"Del {spanish_date(period['start_date'])} al {spanish_date(period['end_date'])}"),
            ("Método", "Modelo predictivo del sistema entrenado con información anterior al período evaluado"),
            ("Observaciones", f"{result['observation_count']} observaciones diarias de {len(categories)} categorías"),
            ("Fuente", "Ventas reales consolidadas y pronóstico del modelo generado antes del período evaluado"),
        ],
    )
    add_heading(doc, "Resultados por categoría")
    rows = add_result_table(doc, categories, model["by_category"], prediction_label="DP modelo")
    rows.append((
        "", "Total", fmt_number(metrics["actual_total"], 0), fmt_number(metrics["prediction_total"], 2),
        fmt_percent(metrics["wape"]), fmt_percent(metrics["mape"]), fmt_number(metrics["mae"]), fmt_number(metrics["rmse"]),
    ))
    table = add_table(
        doc,
        ["N", "Categoría", "DR", "DP modelo", "WAPE", "MAPE", "MAE", "RMSE"],
        rows,
        [0.36, 1.32, 0.66, 0.88, 0.74, 0.74, 0.61, 0.70],
        number_columns=(2, 3, 4, 5, 6, 7),
        font_size=8.25,
    )
    for cell in table.rows[-1].cells:
        for run in cell.paragraphs[0].runs:
            run.bold = True
    add_paragraph(doc, "Nota. DR: demanda real. DP: demanda pronosticada. MAPE se calcula a partir de las observaciones diarias de cada categoría.", size=8.6, italic=True, after=5)

    add_heading(doc, "Resumen de métricas")
    add_metrics_table(doc, metrics, "Modelo predictivo")

    add_heading(doc, "Comparación con la referencia")
    comparison_rows = []
    for label, key, unit in (("WAPE", "wape", "%"), ("MAPE", "mape", "%"), ("MAE", "mae", "unidades"), ("RMSE", "rmse", "unidades")):
        base_value = baseline["metrics"][key]
        model_value = metrics[key]
        if unit == "%":
            base_text, model_text, difference = fmt_percent(base_value), fmt_percent(model_value), fmt_percent(model_value - base_value)
        else:
            base_text, model_text, difference = fmt_number(base_value), fmt_number(model_value), fmt_number(model_value - base_value)
        comparison_rows.append((label, base_text, model_text, "+" + difference if not difference.startswith("-") else difference, "PMS-7 tuvo menor error" if base_value < model_value else "Modelo predictivo tuvo menor error"))
    add_table(
        doc,
        ["Métrica", "PMS-7", "Modelo", "Diferencia", "Resultado"],
        comparison_rows,
        [1.08, 1.15, 1.15, 1.05, 2.38],
        number_columns=(1, 2, 3),
        font_size=8.6,
    )
    add_paragraph(
        doc,
        "Interpretación. Para esta fecha de corte, PMS-7 obtuvo menor WAPE, MAPE, MAE y RMSE. "
        "Es una validación histórica y no evidencia una mejora del modelo frente a la línea base.",
        size=8.8,
        after=2,
        color="9C1C1C",
    )

    doc.core_properties.title = "Postest Validación histórica Modelo predictivo"
    doc.core_properties.author = "Guerrero Mora, Carlos David Elian"
    doc.core_properties.subject = f"Postest del modelo predictivo. Ejecución {run_id}."
    output = OUTPUT_DIR / "Postest_modelo_predictivo_validacion_historica_2026_09_09_a_2026_09_15.docx"
    doc.save(output)
    return output


def main() -> None:
    latest = max(POSTTEST_DIR.glob("*.json"), key=lambda item: item.stat().st_mtime)
    payload = json.loads(latest.read_text(encoding="utf-8"))
    result = payload["result"]
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    pretest = build_pretest(result, payload["run_id"])
    postest = build_postest(result, payload["run_id"])
    print(pretest)
    print(postest)


if __name__ == "__main__":
    main()
