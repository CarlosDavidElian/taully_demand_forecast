from pathlib import Path
from copy import deepcopy
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml.ns import qn

root = Path('C:/taully_demand_forecast')
source = next((root/'outputs').rglob('Anexo_2_Instrumentos_Resultados_Tecnicos_Pretest_Postest_2026-09-02_a_2026-09-15.docx'))
doc = Document(source)
body = doc._element.body
selected = []
active = False
for el in list(body):
    text = ''.join(el.itertext()) if False else ''.join(el.xpath('.//w:t/text()'))
    if el.tag == qn('w:p') and text in ['O1 Pretest Instrumento 5 Error de pronóstico', 'O2 Postest Instrumento 5 Error de pronóstico']:
        active = True
    if active:
        selected.append(deepcopy(el))
    if active and text.startswith('Responsable del registro:'):
        active = False
for el in list(body):
    if el.tag != qn('w:sectPr'):
        body.remove(el)
count = 0
for el in selected:
    text = ''.join(el.xpath('.//w:t/text()'))
    if text.startswith('O2 Postest Instrumento 5'):
        doc.add_page_break()
    body.insert(len(body)-1, el)

for table in doc.tables:
    heads = [c.text for c in table.rows[0].cells]
    if 'RMSE' in heads:
        for row in table.rows:
            for index in [6,5]:
                row._tr.remove(row._tr.tc_lst[index])
        for index in [6,5]:
            table._tbl.tblGrid.remove(table._tbl.tblGrid.gridCol_lst[index])
        widths = [0.4,1.6,1.1,1.35,1.2]
        for col,width in zip(table.columns,widths):
            col.width=Inches(width)
        for row in table.rows:
            for cell,width in zip(row.cells,widths):
                cell.width=Inches(width)
        table.rows[0].cells[2].text='DR semanal'
        table.rows[0].cells[3].text='DP semanal'
        for cell in table.rows[0].cells:
            for run in cell.paragraphs[0].runs:
                run.font.color.rgb=RGBColor(255,255,255)
                run.font.bold=True
                run.font.size=Pt(10)
    else:
        for row in table.rows:
            if row.cells[0].text=='Unidad de medida':
                row.cells[1].text='Porcentaje (%)'

for p in doc.paragraphs:
    if p.text.startswith('Criterio de cálculo.'):
        if 'PMS-7' in p.text:
            p.text='Cálculo del MAPE: se comparan DR y DP de cada día; n = 7 por categoría y n = 35 en el total. DR y DP de la tabla son sumas semanales y no se usan como una sola observación. PMS-7 es una referencia histórica reconstruida, no un método previo de la tienda acreditado.'
        else:
            p.text='Cálculo del MAPE: se comparan DR y DP de cada día; n = 7 por categoría y n = 35 en el total. DR y DP de la tabla son sumas semanales y no se usan como una sola observación. La evaluación histórica no acredita por sí sola una intervención real ni una mejora.'
        for run in p.runs:
            run.font.size=Pt(10)
    if p.text.startswith('Ficha de registro'):
        p.text='Precisión del pronóstico. Indicador del instrumento: error de pronóstico medido mediante MAPE.'

out=root/'outputs/postest_visualizado/Fichas_error_pronostico_pretest_postest_MAPE.docx'
doc.save(out)
check=Document(out)
alltext=' '.join(p.text for p in check.paragraphs)+' '.join(c.text for t in check.tables for r in t.rows for c in r.cells)
assert not any(x in alltext for x in ['RMSE','MAE','WAPE'])
assert len(check.tables)==4
print(out)
