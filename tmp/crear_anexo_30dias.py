from pathlib import Path
from copy import deepcopy
from datetime import date
import json,csv,zipfile,hashlib
from lxml import etree
from docx import Document
from docx.table import Table
from docx.shared import Pt,Cm,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
R=Path('C:/taully_demand_forecast');O=R/'outputs/anexo2_30dias_hasta_15septiembre'
S=Path('C:/Users/Acer/Downloads/LIMA-NORTE_PI_GURRERRO_removed_removed.docx');source=Document(S)
A=json.loads((O/'calculos_30dias_verificados.json').read_text(encoding='utf-8'));cats=A['categories']
D=list(csv.DictReader((O/'detalle_30dias_verificado.csv').open(encoding='utf-8-sig')))
def sha(b):return hashlib.sha256(b).hexdigest()
with zipfile.ZipFile(S) as z: parts={n:sha(z.read(n)) for n in z.namelist()}
(R/'tmp/anexo30_reference_manifest.json').write_text(json.dumps({'source':str(S),'sha256':sha(S.read_bytes()),'parts':parts},indent=2),encoding='utf-8')
def f(n):return f'{n:,.2f}'.replace(',',' ').replace('.',',')
def it(n):return f'{n:,.0f}'.replace(',',' ')
def dt(s):return date.fromisoformat(s).strftime('%d/%m/%Y')
titles=['Cantidad demandada','Índice de Salida de Inventario','Tasa de quiebre de stock','Estacionalidad','Error de pronóstico','Volumen de demanda','PATRÓN DE DEMANDA','Precisión del Pronóstico']
expected={}
def fill(cell,text,bold=False,size=12):
 para=cell.paragraphs[0];rpr=deepcopy(para.runs[0]._r.rPr) if para.runs and para.runs[0]._r.rPr is not None else None
 cell.text=str(text)
 pp=cell.paragraphs[0];pp.alignment=WD_ALIGN_PARAGRAPH.CENTER;pp.paragraph_format.space_before=Pt(2);pp.paragraph_format.space_after=Pt(2);pp.paragraph_format.line_spacing=1
 rr=pp.runs[0];rr.font.name='Arial';rr.font.size=Pt(size);rr.bold=bold;rr.font.color.rgb=RGBColor(0,0,0)
 return pp
def text(doc,s,bold=False,size=11,center=False,style=None):
 pp=doc.add_paragraph(style=style);pp.paragraph_format.space_after=Pt(8);pp.paragraph_format.space_before=Pt(0);pp.paragraph_format.line_spacing=1.05
 pp.alignment=WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.LEFT
 rr=pp.add_run(s);rr.bold=bold;rr.font.name='Arial';rr.font.size=Pt(size);rr.font.color.rgb=RGBColor(0,0,0)
 return pp
def attach(doc,el):doc._element.body.insert(len(doc._element.body)-1,el);return Table(el,doc._body)
def clean_rows(t):
 for row in t.rows:
  pr=row._tr.get_or_add_trPr()
  for old in list(pr.findall(qn('w:trHeight'))):pr.remove(old)
  pr.append(OxmlElement('w:cantSplit'))
  for cell in row.cells:
   for p in cell.paragraphs:
    p.paragraph_format.keep_with_next=False;p.paragraph_format.keep_together=False
    p.paragraph_format.space_before=Pt(0);p.paragraph_format.space_after=Pt(0)
def data_table(doc,idx,rows,vol=False):
 el=deepcopy(source.tables[1+2*idx]._tbl);t=attach(doc,el);model=deepcopy(t.rows[1]._tr)
 for row in list(t.rows)[1:]:t._tbl.remove(row._tr)
 for values in rows:
  elr=deepcopy(model);t._tbl.append(elr)
  for ci,(cell,s) in enumerate(zip(t.rows[-1].cells,values)):
   p=fill(cell,s,size=10.5 if vol else 12)
   if vol and ci==1:p.alignment=WD_ALIGN_PARAGRAPH.LEFT
 clean_rows(t);t.alignment=WD_TABLE_ALIGNMENT.CENTER
 t.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
 if vol:
  t.autofit=False
  for c,w in zip(t.columns,[1.0,7.5,3.0,4.0]):c.width=Cm(w)
  for row in t.rows:
   for c,w in zip(row.cells,[1.0,7.5,3.0,4.0]):c.width=Cm(w)
 return t
def formulas(doc):
 out=[]
 for t in doc.tables:
  for r in t.rows:
   if r.cells[0].text.strip().startswith('FÓRMULA'):
    out.append(''.join(r.cells[1]._tc.xpath('.//w:t/text()|.//m:t/text()')).replace(' ',''))
 return out
for label in ['pre','post']:
 v=A['periods'][label];moment='Pretest' if label=='pre' else 'Postest';g=v['global'];inv=v['inventory'];tot=v['inventory_total']
 doc=Document(S);body=doc._element.body;sect=deepcopy(source.sections[-1]._sectPr)
 for tag in ['headerReference','footerReference','pgNumType']:
  for el in list(sect.findall(qn('w:'+tag))):sect.remove(el)
 for el in list(body):body.remove(el)
 body.append(sect)
 def p(s,bold=False,size=11,center=False):return text(doc,s,bold,size,center)
 def newpage():doc.add_page_break()
 p('Anexo 2',True,18,True);p(moment+' de 30 días',True,18,True)
 p(dt(v['start'])+' al '+dt(v['end']),True,14,True)
 p('Evaluación retrospectiva para revisión metodológica',True,12,True)
 p('Se conservan los ocho instrumentos, sus fórmulas y los encabezados del Anexo 2 de la tesis. Se completan con los datos del período indicado y se amplían las filas cuando el detalle lo requiere.')
 p('El postest de 30 días termina el 15 de septiembre de 2026. El pretest utiliza los 30 días inmediatamente anteriores. No se ha confirmado la fecha de uso real del sistema ni el procedimiento previo del negocio; estos nombres identifican las ventanas de comparación, sin acreditar una intervención.')
 p(('El pretest usa el PMS-7 reconstruido desde el corte del 17/07/2026. Puede reproducirse con la columna PMS-7 de la evaluación del programa.' if label=='pre' else 'El postest utiliza los pronósticos del programa entrenado con datos hasta el 16/08/2026 y evaluado sobre las ventas del 17/08 al 15/09/2026.'))
 p('Resultado de ventas: '+it(g['DR'])+' unidades. Pronóstico acumulado: '+f(g['DP'])+' unidades. MAPE: '+f(g['MAPE'])+' %. Se utilizan 30 pares diarios por categoría y 150 pares para el global.')
 p('Las fichas de inventario proceden de un archivo demostrativo. Sus cifras no deben presentarse como mediciones reales de la tienda. Es necesario reemplazarlo por el kardex real y conciliar sus movimientos.')
 p('Lectura de las fichas',True,12)
 p('En las fichas 1, 2, 3, 5, 7 y 8: N° 1 = ABARROTES; 2 = BEBIDAS; 3 = GOLOSINAS; 4 = HELADOS; 5 = LIMPIEZA. TOTAL o GLOBAL resume las cinco categorías.')
 p('La ficha 4 registra 30 días del negocio, sumando las cinco categorías por fecha. La ficha 6 conserva el campo Producto con el detalle individual de los reportes de ventas.')
 p('Referencia común para IE: junio de 2026, anterior a ambos períodos. VPM se interpreta como promedio diario de ese mes y VPR como su equivalente de 30 días; esta definición debe validarse con el docente.')
 for idx,title in enumerate(titles):
  if idx==0:newpage()
  else:
   sp=deepcopy(source.sections[idx-1]._sectPr)
   for tag in ['headerReference','footerReference','pgNumType']:
    for ee in list(sp.findall(qn('w:'+tag))):sp.remove(ee)
   bp=doc.add_paragraph();bp._p.get_or_add_pPr().append(sp)
  p('Instrumento de recolección de datos '+('del indicador' if idx<5 else 'de la dimensión')+' para el '+moment.lower(),True,12,True)
  p(title,True,12,True)
  meta=attach(doc,deepcopy(source.tables[2*idx]._tbl));clean_rows(meta)
  fill(meta.rows[-1].cells[1],moment+' retrospectivo\n'+dt(v['start'])+' al '+dt(v['end']),False,11)
  p('')
  if idx in [0,1,2,4,6,7]:p('N° 1 Abarrotes · 2 Bebidas · 3 Golosinas · 4 Helados · 5 Limpieza',size=9)
  if idx==0:
   rows=[[n+1]+[it(inv[c][k]) for k in ['SI','EN','SF','CD']] for n,c in enumerate(cats)]+[['TOTAL']+[it(tot[k]) for k in ['SI','EN','SF','CD']]]
   data_table(doc,idx,rows)
   p('Fuente: inventario demostrativo del período. El balance CD suma '+it(tot['CD'])+' y las ventas suman '+it(g['DR'])+' unidades. Diferencia pendiente de conciliación: '+it(tot['sales_difference'])+' unidades. Los saldos corresponden a la primera y última fila disponibles por producto, que pueden no ser las fechas límite. No es evidencia del kardex real.',size=10)
  elif idx==1:
   rows=[[n+1]+[it(inv[c][k]) for k in ['CD','SI','EN']]+[f(inv[c]['ISI'])+' %'] for n,c in enumerate(cats)]+[['TOTAL']+[it(tot[k]) for k in ['CD','SI','EN']]+[f(tot['ISI'])+' %']]
   data_table(doc,idx,rows)
   p('Fuente: inventario demostrativo. Se conserva el cociente original y se expresa como porcentaje. El total se calcula dividiendo la suma de CD entre la suma de SI y EN; no es el promedio simple de los porcentajes. Requiere kardex real conciliado.',size=10)
  elif idx==2:
   rows=[[n+1,inv[c]['DQS'],30,f(inv[c]['TQS'])+' %'] for n,c in enumerate(cats)]+[['TOTAL',tot['DQS'],150,f(tot['TQS'])+' %']]
   data_table(doc,idx,rows)
   p('Fuente: inventario demostrativo. DQS cuenta fechas distintas con quiebre por categoría. DD = 30 por categoría. El total evalúa 150 categoría-días, no 150 días distintos del negocio. No acredita quiebres reales de la tienda.',size=10)
  elif idx==3:
   p('Cada fila representa un día del negocio. VPM es el promedio diario de junio de 2026: '+f(sum(A['reference_daily'].values()))+' unidades. N° 1 corresponde al '+dt(v['start'])+' y la numeración avanza un día por fila.',size=10)
   rows=[]
   for n,day in enumerate(v['days']):
    vd=sum(float(x['actual']) for x in D if x['period']==label and x['date']==day);vpm=sum(A['reference_daily'].values())
    rows.append([n+1,f(vd/vpm),it(vd),f(vpm)])
   data_table(doc,idx,rows[:13]);newpage();p(moment+' Estacionalidad continuación',True,12,True);p('Período '+dt(v['start'])+' al '+dt(v['end'])+'. Las filas conservan la numeración diaria de la página anterior.',size=10)
   data_table(doc,idx,rows[13:])
   p('Fuente: historial y reportes diarios de ventas. Se utiliza el VPM sin redondear. La definición operativa de VPM como promedio diario del mes de referencia debe confirmarse con el docente. Los índices no prueban por sí solos una estacionalidad recurrente.',size=10)
  elif idx in [4,7]:
   rows=[[n+1,it(v['by_category'][c]['DR']),f(v['by_category'][c]['DP']),f(v['by_category'][c]['MAPE'])+' %'] for n,c in enumerate(cats)]+[['GLOBAL',it(g['DR']),f(g['DP']),f(g['MAPE'])+' %']]
   data_table(doc,idx,rows)
   p('DR y DP son acumulados de 30 días. El MAPE se calcula con los pares diarios: n = 30 por categoría y n = 150 para el global. No hay DR igual a cero. No se calcula el MAPE dividiendo los acumulados de esta tabla.',size=10)
   p(('Fuente: PMS-7 reconstruido con corte '+dt(v['cutoff'])+'. En el programa y Excel use PMS-7 MAPE para reproducir este pretest; no la columna del modelo. No se ha acreditado que PMS-7 fuese el método previo de la tienda.' if label=='pre' else 'Fuente: ejecución guardada del programa con corte '+dt(v['cutoff'])+' y horizonte de 30 días. En el Excel use MODELO MAPE. El resultado procede de una evaluación retrospectiva.'),size=10)
   if idx==7:p('Esta ficha usa las mismas observaciones que Error de pronóstico y no constituye una muestra adicional.',size=10)
  elif idx==5:
   prods=sorted(A['products'][label].items());rows=[[n+1,k,it(q),it(q)] for n,(k,q) in enumerate(prods)]+[['','TOTAL',it(g['DR']),it(g['DR'])]]
   p('Fuente: suma de unidades vendidas por producto en los 30 reportes diarios. El listado contiene '+str(len(prods))+' productos. Se conserva el encabezado Producto del instrumento.',size=10)
   data_table(doc,idx,rows[:8],True)
   for offset in range(8,len(rows),22):
    newpage();p(moment+' Volumen de demanda continuación',True,12,True);p(dt(v['start'])+' al '+dt(v['end'])+' · Productos '+str(offset+1)+' a '+str(min(offset+22,len(prods))),size=10)
    data_table(doc,idx,rows[offset:offset+22],True)
   p('CD en esta ficha es la suma de las ventas realizadas. Total: '+it(g['DR'])+' unidades. Se diferencia del CD calculado por balance del inventario demostrativo. Las ventas observadas no incluyen demanda no atendida.',size=10)
  elif idx==6:
   rows=[[n+1,it(v['by_category'][c]['VP']),f(v['by_category'][c]['VPR']),f(v['by_category'][c]['IE'])] for n,c in enumerate(cats)]+[['TOTAL',it(g['VP']),f(g['VPR']),f(g['IE'])]]
   data_table(doc,idx,rows)
   p('Fuente: historial y reportes diarios. VP es la venta acumulada del período. VPR es el equivalente de 30 días según el promedio diario de junio de 2026, una referencia común previa a ambas ventanas. Como junio tiene 30 días, ese equivalente coincide con las unidades vendidas en junio. No se redondean los valores antes de dividir.',size=10)
  expected.setdefault(label,[]).append({'index':idx,'rows':rows})
 out=O/(moment+'_Anexo2_30dias.docx')
 xml=etree.tostring(doc._element,encoding='UTF-8',xml_declaration=True,standalone=True)
 with zipfile.ZipFile(S) as src,zipfile.ZipFile(out,'w',zipfile.ZIP_DEFLATED) as dest:
  for item in src.infolist():dest.writestr(item,xml if item.filename=='word/document.xml' else src.read(item.filename))
 with zipfile.ZipFile(out) as dest:
  assert all(sha(dest.read(n))==h for n,h in parts.items() if n!='word/document.xml')
 reread=Document(out);assert formulas(reread)==formulas(source)[:8]
 assert len(reread.sections)==len(source.sections)==8
 print(out,'products',len(A['products'][label]))
(O/'tablas_esperadas.json').write_text(json.dumps(expected,ensure_ascii=False,indent=2),encoding='utf-8')
assert sha(S.read_bytes())==json.loads((R/'tmp/anexo30_reference_manifest.json').read_text())['sha256']
