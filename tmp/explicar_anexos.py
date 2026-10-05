from pathlib import Path
import json,csv,math
from datetime import date,timedelta
from statistics import mean
from docx import Document
from docx.shared import Cm,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
R=Path('C:/taully_demand_forecast')
inputs=json.loads((R/'tmp/explicacion_anexos.json').read_text(encoding='utf-8'))
audit=json.loads((R/'outputs/anexo2_30dias_hasta_15septiembre/calculos_30dias_verificados.json').read_text(encoding='utf-8'))
detail=list(csv.DictReader((R/'outputs/anexo2_30dias_hasta_15septiembre/detalle_30dias_verificado.csv').open(encoding='utf-8-sig')))
cats=['ABARROTES','BEBIDAS','GOLOSINAS','HELADOS','LIMPIEZA']
labels=['Pretest','Postest']; keys=['pre','post']; checks=0
def num(t):return float(t.replace(' ','').replace('%','').replace(',','.'))
def fmt(x,n=0):return f'{x:,.{n}f}'.replace(',',' ').replace('.',',')
def eq(a,b,tol=.00501):
 global checks
 assert abs(a-b)<tol,(a,b)
 checks+=1
for idx,key in enumerate(keys):
 ts=inputs[idx]['tables']; q=audit['periods'][key]
 for i,c in enumerate(cats):
  inv=q['inventory'][c]
  for j,k in enumerate(['SI','EN','SF','CD'],1):eq(num(ts[1][i+1][j]),inv[k])
  eq(inv['SI']+inv['EN']-inv['SF'],inv['CD'])
  eq(num(ts[3][i+1][4]),100*inv['CD']/(inv['SI']+inv['EN']))
  eq(num(ts[5][i+1][3]),100*inv['DQS']/inv['DD'])
  obs=[r for r in detail if r['period']==key and r['category']==c]
  for col,val in [(1,sum(float(r['actual']) for r in obs)),(2,sum(float(r['prediction']) for r in obs)),(3,mean(abs(float(r['actual'])-float(r['prediction']))/float(r['actual']) for r in obs)*100)]:
   eq(num(ts[10][i+1][col]),val);eq(num(ts[25][i+1][col]),val)
  eq(num(ts[23][i+1][3]),num(ts[23][i+1][1])/num(ts[23][i+1][2]))
 rows=ts[7][1:]+ts[8][1:]
 for row,day in zip(rows,q['days']):
  actual=sum(float(r['actual']) for r in detail if r['period']==key and r['date']==day)
  eq(num(row[2]),actual);eq(num(row[1]),actual/261.4)
 products=[r for t in ts[12:22] for r in t[1:] if r[0].isdigit()]
 for row in products: eq(num(row[2]),num(row[3]));eq(num(row[2]),audit['products'][key][row[1]])
 eq(sum(num(r[2]) for r in products),q['global']['DR'])

# Reuse only the typography and table helper from the existing guide.
base=(R/'tmp/scrum_plan/metodologia.py').read_text(encoding='utf-8')
exec(base.split("d.add_heading('Metodología Scrum")[0])
d.add_heading('Explicación de los números del pretest y postest',0)
p('Guía de lectura del Anexo 2 · Minimarket Taully\nPreparada el 4 de octubre de 2026')
p('Esta guía explica las ocho fichas de cada documento: el significado de las columnas, la operación que produce cada resultado y cómo leer los totales. Incluye el detalle diario de estacionalidad y un listado conjunto de los productos de ambas ventanas.')
h('Qué se comprobó')
p('Se contrastaron las cifras de los dos Word recibidos con el detalle de cálculo conservado del proyecto. Las operaciones revisadas coinciden al redondear a dos decimales. Esta comprobación explica la consistencia aritmética; la autenticidad de los registros de la tienda debe respaldarse con sus documentos fuente.')
table(['Dato','Pretest','Postest'],[
['Período de ventas','18/07/2026 al 16/08/2026','17/08/2026 al 15/09/2026'],['Días incluidos','30','30'],['Corte del pronóstico','17/07/2026','16/08/2026'],['Pronóstico utilizado','PMS-7 reconstruido','Modelo del programa'],['Ventas acumuladas','7 916 unidades','7 851 unidades'],['Observaciones para el MAPE','30 por categoría; 150 globales','30 por categoría; 150 globales']],[5.2,5.7,5.7])
h('Cómo leer las cifras')
p('1 433 significa mil cuatrocientas treinta y tres unidades: el espacio separa los miles. 74,61 % es un porcentaje. 1,11 es un índice, sin símbolo de porcentaje. Las cantidades son unidades registradas de productos, no soles ni número de clientes.')
p('En las fichas por categoría: 1 = Abarrotes; 2 = Bebidas; 3 = Golosinas; 4 = Helados; 5 = Limpieza. En estacionalidad diaria el número identifica un día; en volumen de demanda identifica un producto.')
p('Los nombres pretest y postest describen las dos ventanas retrospectivas. Las fechas de estas ventanas no confirman por sí mismas cuándo comenzó a utilizarse el sistema en la tienda.')

page('1 Cantidad demandada por inventario')
p('Ficha 1. CD = (SI + EN) − SF. SI es el saldo inicial, EN las entradas y SF el saldo final. La resta obtiene la salida por balance del inventario. Su identificación con ventas requiere conciliar también ajustes, devoluciones y mermas.')
p('En el cálculo del programa, SI suma el primer saldo disponible de cada producto dentro de la ventana y SF suma el último. EN suma los movimientos de entrada de todas las filas de la ventana. No se suman los saldos iniciales de todos los días, porque eso contaría varias veces las mismas existencias.')
for idx,key in enumerate(keys):
 h(labels[idx])
 table(['Categoría','SI + EN − SF','CD en unidades'],[[cats[i].title() if i<5 else 'TOTAL',f'{r[1]} + {r[2]} − {r[3]}',r[4]] for i,r in enumerate(inputs[idx]['tables'][1][1:])],[3.4,9.4,3.8])
p('Ejemplo del postest: en Abarrotes, 372 unidades iniciales más 949 entradas dan 1 321 disponibles; al restar 290 finales quedan 1 031 unidades de salida por balance. El 949 viene de sumar las entradas del archivo de inventario, no de sumar las ventas.')
p('Fuente de estas cifras: inventario demostrativo del proyecto. Hay productos sin saldo en las fechas exactas de inicio o fin; por ello los primeros y últimos registros disponibles requieren revisión con el kardex de la tienda.')

page('2 Índice de salida de inventario')
p('Ficha 2. ISI = CD ÷ (SI + EN). La fórmula produce una proporción; para mostrarla en porcentaje se multiplica por 100. Describe qué parte del stock disponible salió según el balance.')
for idx in range(2):
 h(labels[idx])
 table(['Categoría','Operación para mostrar porcentaje','ISI'],[[cats[i].title() if i<5 else 'TOTAL',f'{r[1]} ÷ ({r[2]} + {r[3]}) × 100',r[4]] for i,r in enumerate(inputs[idx]['tables'][3][1:])],[3.4,9.4,3.8])
p('Ejemplo del postest: 1 031 ÷ 1 321 = 0,780469…; expresado en porcentaje es 78,05 %. Se interpreta como aproximadamente 78,05 unidades de salida por cada 100 unidades disponibles en ese balance.')
p('El total usa la suma de cantidades: 4 540 ÷ (2 127 + 3 958) × 100 = 74,61 %. No se obtiene promediando los cinco porcentajes. Estas operaciones usan el mismo inventario demostrativo de la ficha 1.')

page('3 Tasa de quiebre de stock')
p('Ficha 3. TQS = (DQS ÷ DD) × 100. DQS cuenta fechas distintas en las que algún producto de la categoría tiene Dias_Sin_Stock mayor que cero. DD es el número de días calendario evaluados, incluidos el primero y el último.')
for idx in range(2):
 h(labels[idx])
 table(['Categoría','DQS ÷ DD × 100','TQS'],[[cats[i].title() if i<5 else 'TOTAL',f'{r[1]} ÷ {r[2]} × 100',r[3]] for i,r in enumerate(inputs[idx]['tables'][5][1:])],[3.4,9.4,3.8])
p('Bebidas del postest: 2 indica dos fechas con quiebre registradas, 30 es la duración de la ventana y 6,67 % es 2 ÷ 30 × 100. Si dos productos tienen quiebre el mismo día en Bebidas, esa fecha cuenta una sola vez para la categoría.')
p('El total postest suma 1 + 2 + 7 + 3 + 0 = 13 categoría-días con quiebre. El denominador es 5 categorías × 30 días = 150 categoría-días. No significa que el negocio estuvo 13 días distintos sin stock ni que se evaluaron 150 días calendario.')
p('Un cero indica que el archivo no registra fechas con quiebre para esa categoría. No demuestra ausencia real de quiebres si los registros están incompletos. Fuente: el mismo inventario demostrativo.')

page('4 Qué fechas producen los quiebres')
p('Estas fechas proceden del registro de cálculo conservado del inventario. Permiten entender los conteos DQS de la ficha 3. Su comprobación operativa requiere contrastarlas con el kardex de la tienda.')
for idx,key in enumerate(keys):
 h(labels[idx])
 rows=[]
 for c in cats:
  v=audit['periods'][key]['inventory'][c]
  rows.append([c.title(),str(v['DQS']),', '.join(date.fromisoformat(x).strftime('%d/%m') for x in v['stockout_dates']) or 'Sin fechas registradas'])
 table(['Categoría','DQS','Fechas registradas de 2026'],rows,[3.4,1.6,11.6])
h('Por qué no coincide el balance con las ventas')
table(['Período','Ventas de las fichas','Salida por balance','Diferencia'],[['Pretest','7 916','4 447','3 469'],['Postest','7 851','4 540','3 311']],[3.4,4.4,4.4,4.4])
p('La diferencia se obtiene restando salida por balance a ventas: 7 916 − 4 447 = 3 469 y 7 851 − 4 540 = 3 311. La fórmula de CD puede estar bien aplicada y, a la vez, los movimientos no representar un inventario conciliado. No se debe interpretar esa diferencia automáticamente como pérdida o demanda no atendida.')

page('5 Referencia para los índices de demanda')
p('Las fichas 4 y 7 comparan las ventas con junio de 2026. Junio tiene 30 días. Se conserva la interpretación usada en los Word: VPM es el promedio diario de junio y VPR es el equivalente de 30 días de ese promedio. Esta definición operativa debe confirmarse con el docente.')
table(['Categoría','Ventas de junio','Promedio diario','VPR para 30 días'],[[c.title(),fmt(audit['periods']['pre']['by_category'][c]['VPR']),fmt(audit['periods']['pre']['by_category'][c]['VPR']/30,4),fmt(audit['periods']['pre']['by_category'][c]['VPR'])] for c in cats]+[['TOTAL','7 842','261,4000','7 842']],[4.0,3.6,4.5,4.5])
p('El 261,40 se obtiene de 7 842 ÷ 30. Se repite en las 30 filas de estacionalidad diaria porque todos los días se comparan con la misma referencia. Se divide unidades por día entre unidades por día; el índice no tiene unidad.')
h('Cómo interpretar IE')
p('Un índice mayor que 1 indica ventas superiores a la referencia; menor que 1, inferiores. Un valor cercano a 1 indica ventas similares. Un índice de 1,10 representa aproximadamente 10 % por encima de la referencia, no 1,10 % de ventas.')
p('El total del patrón de demanda postest es 7 851 ÷ 7 842 = 1,001147… y se muestra 1,00 por redondeo. Los valores originales no son exactamente iguales: la diferencia es de 9 unidades. Estos índices describen la comparación; por sí solos no prueban un patrón estacional recurrente.')

for idx,key in enumerate(keys):
 page(f'6 Estacionalidad diaria del {labels[idx].lower()}')
 p('Ficha 4. IE = VD ÷ VPM. VD suma las unidades vendidas de las cinco categorías en esa fecha. VPM = 261,40 unidades diarias. Cada fila siguiente explica la división de la ficha original.')
 rows=inputs[idx]['tables'][7][1:]+inputs[idx]['tables'][8][1:]
 table(['N°','Fecha','Operación VD ÷ VPM','IE'],[[r[0],date.fromisoformat(day).strftime('%d/%m/%Y'),f'{r[2]} ÷ {r[3]}',r[1]] for r,day in zip(rows,audit['periods'][key]['days'])],[1.0,3.5,9.1,3.0])
 for row in d.tables[-1].rows:
  for cell in row.cells:
   for pp in cell.paragraphs:
    pp.paragraph_format.space_before=Pt(3)
    pp.paragraph_format.space_after=Pt(3)
 p(f'La suma de las 30 ventas diarias es {fmt(audit["periods"][key]["global"]["DR"])} unidades. Los índices se redondean después de dividir. N° identifica el día de la ventana, no la categoría.')

page('7 Error y precisión del pronóstico')
p('Las fichas 5 y 8 repiten los mismos pares de datos. DR es la suma de ventas observadas de los 30 días y DP la suma de los 30 pronósticos. DP puede tener decimales porque es una estimación. MAPE expresa el error porcentual absoluto medio; un valor menor indica menor error en esa evaluación.')
p('La fórmula se aplica así: para cada día se resta DP a DR, se divide entre DR, se toma el valor absoluto y se multiplica por 100. Se suman esos 30 errores porcentuales y se divide entre 30. En el global se usan 150 pares: cinco categorías por 30 días. No se divide la diferencia de los totales mensuales entre DR.')
for idx in range(2):
 h(labels[idx]+(' con PMS-7' if idx==0 else ' con el modelo'))
 table(['Categoría','DR acumulada','DP acumulada','MAPE diario medio'],[[cats[i].title() if i<5 else 'GLOBAL']+r[1:] for i,r in enumerate(inputs[idx]['tables'][10][1:])],[3.6,4.0,4.3,4.7])
p('En pretest, PMS-7 usa las siete observaciones anteriores para calcular el promedio y añade cada estimación a la serie para pronosticar el día siguiente. Es un método de referencia reconstruido; su uso histórico por la tienda no está confirmado. El postest usa las predicciones del modelo registradas para el corte 16/08/2026.')

page('8 Cómo se obtiene cada MAPE')
p('La tabla muestra la suma de errores porcentuales diarios, calculada desde el detalle conservado, y su división entre el número de pares. Se muestran cuatro decimales en las sumas para explicar el resultado; el cálculo utiliza la precisión completa.')
rows=[]
for idx,key in enumerate(keys):
 for c in cats+['GLOBAL']:
  obs=[r for r in detail if r['period']==key and (c=='GLOBAL' or r['category']==c)]
  errors=[abs(float(r['actual'])-float(r['prediction']))/float(r['actual'])*100 for r in obs]
  rows.append([labels[idx],c.title(),f'{fmt(sum(errors),4)} ÷ {len(obs)}',fmt(mean(errors),2)+' %'])
table(['Período','Categoría','Suma de errores ÷ pares','MAPE'],rows,[2.4,3.2,7.2,3.8])
h('Ejemplos con una observación real del detalle')
p('Pretest, Abarrotes del 18/07/2026: DR = 56 y DP = 54,285714…; el error es 1,714285… unidades. Al dividir entre 56 y multiplicar por 100 resulta 3,061224… %. Este es solo un día de los 30 usados para obtener 21,20 %.')
p('Postest, Abarrotes del 17/08/2026: DR = 48 y DP = 45,79. La diferencia absoluta es 2,21; 2,21 ÷ 48 × 100 = 4,604166… %. El promedio de los 30 errores de Abarrotes es 21,19 %.')
h('Por qué el global es 26,21 %')
p('Se promedian los 150 errores individuales. Como cada categoría tiene 30 pares válidos, equivale a promediar los cinco MAPE sin redondear. Un promedio de errores no es un porcentaje de unidades totales mal pronosticadas. La evaluación revisada no contiene DR igual a cero.')
p('En pretest, sumar las DP ya redondeadas de las cinco filas produce 7 767,15, mientras que sumar primero las predicciones completas y redondear al final produce 7 767,14. La diferencia de 0,01 es de presentación.')

page('9 Volumen de demanda por producto')
p('Ficha 6. CD es la suma de unidades vendidas de cada producto en los reportes de la ventana. Cantidad Vendida y Cantidad Demandada repiten el mismo número porque el instrumento define CD mediante las ventas realizadas. No se suman ambas columnas entre sí.')
table(['Ejemplo de producto','Pretest','Postest'],[['AGUA CIELO SIN GAS 1 L','68 unidades','67 unidades'],['AGUA CIELO SIN GAS 2.5 L','45 unidades','46 unidades'],['AGUA SAN MATEO SIN GAS 600 ML','105 unidades','103 unidades']],[10.0,3.3,3.3])
p('El 68 de AGUA CIELO SIN GAS 1 L significa que, al reunir sus cantidades de los reportes entre el 18 de julio y el 16 de agosto, se acumulan 68 unidades. El 67 corresponde a la suma independiente del 17 de agosto al 15 de septiembre. No son stock inicial ni pronóstico.')
p('El pretest lista 195 productos y el postest 197. El número de fila puede cambiar entre documentos: compare por el nombre completo del producto. Por ejemplo, un envase de 1 L y uno de 2,5 L son registros distintos; los litros del nombre no son las unidades vendidas.')
table(['Categoría','Ventas pretest','Ventas postest'],[[c.title(),fmt(audit['periods']['pre']['by_category'][c]['DR']),fmt(audit['periods']['post']['by_category'][c]['DR'])] for c in cats]+[['TOTAL','7 916','7 851']],[7.4,4.6,4.6])
p('Al sumar todos los productos se obtienen los totales de ventas. Estos coinciden con DR y VP de las fichas de pronóstico y patrón de demanda. Las ventas registradas miden la demanda atendida; no contienen automáticamente pedidos que no se pudieron vender por falta de stock.')
p('El anexo de esta guía conserva cada producto y sus números de fila para localizarlo en ambos Word. Cada cantidad allí presentada es una suma de unidades del período correspondiente.')

page('10 Patrón de demanda por categoría')
p('Ficha 7. IE = VP ÷ VPR. VP es la venta acumulada de 30 días. VPR es la referencia equivalente de 30 días de junio. Ambos valores se expresan en unidades; el resultado es un índice.')
for idx in range(2):
 h(labels[idx])
 table(['Categoría','VP ÷ VPR','IE'],[[cats[i].title() if i<5 else 'TOTAL',f'{r[1]} ÷ {r[2]}',r[3]] for i,r in enumerate(inputs[idx]['tables'][23][1:])],[3.4,9.4,3.8])
p('Abarrotes del postest: 1 653 ÷ 1 491 = 1,108651…; se muestra 1,11, aproximadamente 10,87 % sobre la referencia. Golosinas: 2 129 ÷ 2 381 = 0,894162…; se muestra 0,89, aproximadamente 10,58 % por debajo.')
p('El total divide la suma de VP entre la suma de VPR. No se suman los índices. Las referencias de junio explican la columna VPR; no son predicciones del modelo.')

page('11 Qué resultados se pueden comparar')
table(['Indicador global','Pretest','Postest','Lectura'],[['Ventas','7 916','7 851','65 unidades menos en la segunda ventana.'],['MAPE de la ficha','23,76 %','26,21 %','El error mostrado aumenta aproximadamente 2,45 puntos porcentuales.'],['ISI','73,81 %','74,61 %','Cambia el balance del archivo demostrativo; requiere kardex conciliado.'],['TQS','15,33 %','8,67 %','Disminuye el conteo del archivo; requiere verificar registros reales.']],[3.0,2.5,2.5,8.6])
p('Los MAPE de las dos fichas utilizan períodos y métodos diferentes. Por eso no bastan para afirmar que el sistema mejoró o empeoró causalmente el negocio. En la misma ventana postest, el registro conservado muestra PMS-7 de 29,14 % frente al modelo de 26,21 %; esa comparación controla el período, aunque sigue siendo retrospectiva.')
h('Dónde comprobar los números')
p('Para el pretest use corte 17/07/2026 y horizonte de 30 días, y lea las columnas PMS-7. Para el postest use corte 16/08/2026 y 30 días, y lea las columnas del modelo. El detalle diario permite sumar DR y DP y calcular los errores de cada par.')
p('Para volumen, sume cantidades por nombre de producto en los reportes del período. Para IE diario, sume primero las ventas de las cinco categorías del día y divida entre 261,40. Para el patrón de demanda, divida los acumulados de cada categoría entre sus referencias de junio.')
p('La exportación revisada no incluía toda la ficha VP–VPR–IE: la presencia de DR permite verificar VP, pero VPR e IE requieren el cálculo y la fuente de referencia. Para inventario, filtre exactamente la ventana y revise saldos, entradas y fechas con quiebre en el kardex.')
h('Documentos y alcance de esta revisión')
p('Documentos leídos: Pretest_Anexo2_30dias (1).docx y Postest_Anexo2_30dias.docx, recibidos del usuario. Apoyo de cálculo: calculos_30dias_verificados.json y detalle_30dias_verificado.csv, conservados en outputs/anexo2_30dias_hasta_15septiembre del proyecto.')
p('Se verificaron operaciones y coincidencia con el detalle de cálculo conservado. En esta guía no se certifica la procedencia empresarial de los reportes ni del inventario. Las fichas de inventario siguen vinculadas al archivo demostrativo identificado en los documentos.')

prod=[]
look=[]
for idx in range(2):
 look.append({r[1]:(r[0],r[2]) for t in inputs[idx]['tables'][12:22] for r in t[1:] if r[0].isdigit()})
names=sorted(set(look[0])|set(look[1]))
for name in names:prod.append([name,look[0].get(name,('—','—'))[0],look[0].get(name,('—','—'))[1],look[1].get(name,('—','—'))[0],look[1].get(name,('—','—'))[1]])
for offset in range(0,len(prod),20):
 page('Anexo Lectura de cada producto'+(' continuación' if offset else ''))
 p('Cada cantidad es la suma de unidades vendidas del producto en la ventana. Pretest: 18/07 al 16/08. Postest: 17/08 al 15/09 de 2026. N° es su fila en el Word original. El guion indica que el producto no aparece en esa ficha; no acredita por sí solo una venta cero.')
 table(['Producto','N° pre','Unidades pre','N° post','Unidades post'],prod[offset:offset+20],[9.0,1.5,2.3,1.5,2.3])
 if offset+20>=len(prod):p('Totales de las fichas: 195 productos y 7 916 unidades en pretest; 197 productos y 7 851 unidades en postest. Para reproducir una cantidad, filtre el nombre del producto y sume sus registros del período.')
footer=d.sections[0].footer.paragraphs[0]
footer.alignment=2
footer.add_run('Guía de explicación · Página ')
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
out=R/'outputs/Explicacion_numeros_Pretest_Postest_Anexo2.docx'
d.save(out)
print('Documento:',out,'Comprobaciones:',checks,'Productos unidos:',len(names))
