from pathlib import Path
import json,csv
from docx import Document
from docx.shared import Cm,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
R=Path('C:/taully_demand_forecast');O=R/'outputs/anexo2_30dias_hasta_15septiembre'
j=json.loads((O/'calculos_30dias_verificados.json').read_text(encoding='utf-8'));P=j['periods']
d=Document();s=d.sections[0];s.page_width=Cm(21);s.page_height=Cm(29.7);s.left_margin=s.right_margin=Cm(2);s.top_margin=s.bottom_margin=Cm(1.9)
for n in ['Normal','Title','Heading 1','Heading 2']:
 st=d.styles[n];st.font.name='Arial';st.font.color.rgb=RGBColor(0,0,0)
for el in d.styles.element.xpath('.//w:pBdr'):el.getparent().remove(el)
d.styles['Normal'].font.size=Pt(10.5);d.styles['Normal'].paragraph_format.space_after=Pt(7);d.styles['Normal'].paragraph_format.line_spacing=1.05
d.styles['Title'].font.size=Pt(21);d.styles['Heading 1'].font.size=Pt(16);d.styles['Heading 2'].font.size=Pt(12)
def p(s):d.add_paragraph(s)
def h(s):d.add_heading(s,2)
def page(s):d.add_page_break();d.add_heading(s,1)
def f(n):return f'{n:,.2f}'.replace(',',' ').replace('.',',')
def table(headers,rows,widths):
 t=d.add_table(rows=1,cols=len(headers));t.autofit=False
 for col,w in zip(t.columns,widths):col.width=Cm(w)
 for i,vals in enumerate([headers]+rows):
  cells=t.rows[0].cells if i==0 else t.add_row().cells
  for cell,txt,w in zip(cells,vals,widths):
   cell.width=Cm(w);cell.text=str(txt);pr=cell._tc.get_or_add_tcPr()
   sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'17365D' if i==0 else ('F0F4F8' if i%2 else 'FFFFFF'));pr.append(sh)
   b=OxmlElement('w:tcBorders')
   for side in ['top','left','bottom','right']:
    e=OxmlElement('w:'+side);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');b.append(e)
   pr.append(b)
   for pp in cell.paragraphs:
    pp.paragraph_format.space_before=Pt(4);pp.paragraph_format.space_after=Pt(4);pp.paragraph_format.line_spacing=1
    for r in pp.runs:r.font.size=Pt(9.5);r.bold=i==0;r.font.color.rgb=RGBColor.from_string('FFFFFF' if i==0 else '000000')
  t.rows[i]._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
 p('')
d.add_heading('Guía para verificar el pretest y el postest de 30 días',0)
p('Anexo 2 de la tesis del Minimarket Taully Comas. Fecha final del postest: 15 de septiembre de 2026.')
p('Esta guía permite al docente reproducir los resultados, comprobar la fuente de cada ficha y reconocer sus límites. Las fórmulas de las fichas se conservan. Las cifras de inventario proceden de una demostración y requieren sustitución por el kardex real.')
table(['Configuración','Pretest retrospectivo','Postest retrospectivo'],[
 ['Corte de prueba','17/07/2026','16/08/2026'],['Horizonte','30 días','30 días'],['Período','18/07 al 16/08/2026','17/08 al 15/09/2026'],['Pronóstico que se usa','PMS-7 reconstruido','Modelo predictivo'],['Unidades vendidas','7 916','7 851'],['Pronóstico acumulado','7 767,14','7 684,29'],['MAPE de la ficha','23,76 %','26,21 %'],['Observaciones','150 categoría-días','150 categoría-días']],[5.3,5.85,5.85])
h('Por qué se usan estas fechas')
p('Se toman los últimos 30 días disponibles hasta el 15/09 y los 30 días anteriores, sin superposición. No se eligió una fecha posterior ni se modificaron valores para mejorar el error. El horizonte se definió antes del recálculo. Se conserva la prueba anterior de siete días, que arrojó otros resultados.')
p('Falta confirmar cuándo se implementó el sistema y qué método se utilizaba antes. Por ello estos documentos son una comparación retrospectiva reproducible, no una acreditación de que ocurrió una intervención en esas fechas. El docente debe validar esa correspondencia y la duración del estudio.')
h('Qué se comprobó')
p(str(len(j['checks']))+' comprobaciones numéricas conformes: reportes contra historial, ventas diarias, PMS-7, MAPE, inventario y exportaciones del programa. Se evaluaron 150 pares diarios por período. El MAPE no contiene observaciones con demanda real cero. Los cálculos mantienen precisión completa y los Word muestran dos decimales.')

page('Cómo reproducirlo en el programa')
p('1. Abra http://127.0.0.1:5000/ con el servidor del programa en ejecución. Compruebe que el historial termina el 15/09/2026 y contiene 840 registros en total: 168 fechas por cinco categorías. Esos registros abarcan todo el historial, no solamente el postest.')
p('2. Para el postest, seleccione Horizonte 30 días y Corte de prueba 16/08/2026. Pulse Ejecutar Postest. No use Generar pronóstico para sustituir esta evaluación: se necesita comparar con ventas ya registradas.')
p('3. En Resultado del Postest debe aparecer 17/08/2026 al 15/09/2026, 30 días, 150 observaciones y cinco categorías. En la tarjeta Modelo predictivo, MAPE debe ser 26,21 %. La tarjeta PMS-7 del mismo período muestra 29,14 %.')
p('4. Pulse Descargar Postest. En Métricas categoría, MODELO MAPE es la columna E. En Detalle diario, C es demanda real, D es el pronóstico del modelo y E es PMS-7. Hay 150 filas de observaciones, de la fila 2 a la 151.')
p('5. Para reproducir el pretest, cambie el corte a 17/07/2026, mantenga 30 días y vuelva a ejecutar. Debe aparecer 18/07 al 16/08/2026. Use la tarjeta PMS-7, cuyo MAPE es 23,76 %, y la columna PMS-7 MAPE, I, de Métricas categoría. La tarjeta del modelo corresponde a otra evaluación y no es el pretest usado en el Word.')
p('6. Descargue esa ejecución. En Detalle diario, use =SUMA(E2:E151) para el pronóstico del pretest y =SUMA(D2:D151) para el postest. Las ventas reales son =SUMA(C2:C151). Por categoría use =SUMAR.SI(B2:B151;"ABARROTES";C2:C151); para pronósticos cambie C por E o D. SUMA incluye las filas ocultas por filtros.')
h('Cifras de control por categoría')
table(['Categoría','Pre DR','Pre MAPE','Post DR','Post MAPE'],[[c,int(P['pre']['by_category'][c]['DR']),f(P['pre']['by_category'][c]['MAPE'])+' %',int(P['post']['by_category'][c]['DR']),f(P['post']['by_category'][c]['MAPE'])+' %'] for c in j['categories']],[4.5,2.5,3.5,2.5,4])
p('El panel presenta WAPE en la tabla por categoría. No copie ese porcentaje en una columna MAPE. El Excel permite ver ambos encabezados. Las fichas del Anexo 2 mantienen MAPE y no añaden MAE, RMSE ni WAPE.')

page('Cómo comprobar cada instrumento')
table(['Ficha','Fuente y operación de verificación'],[
 ['1 Cantidad por inventario','Hoja Inventario del Excel del período. Compruebe SI, EN, SF y la resta que produce CD.'],
 ['2 ISI','Divida CD entre SI + EN. El Word presenta el cociente como porcentaje.'],
 ['3 TQS','Cuente fechas distintas con quiebre por categoría. Divida DQS entre 30 y multiplique por 100. El global usa 150 categoría-días.'],
 ['4 IE diario','Sume la demanda real de las cinco categorías por fecha. Divida esa venta diaria entre 261,4, promedio diario de junio. No aparece calculado en el panel.'],
 ['5 Error de pronóstico','Use cada par diario DR y DP para el MAPE y luego promedie los errores absolutos porcentuales.'],
 ['6 Volumen de demanda','Sume las cantidades vendidas por producto en los 30 reportes. El total debe coincidir con la suma de DR del Excel.'],
 ['7 Patrón de demanda','Divida las ventas del período entre la referencia equivalente de 30 días de junio para la misma categoría.'],
 ['8 Precisión del pronóstico','Reproduce el MAPE de la ficha 5 con las mismas observaciones; no es una muestra distinta.']],[5,12])
h('Ejemplo de MAPE en Excel')
p('En una copia del Excel del postest, cree una columna auxiliar en Detalle diario. En H2 escriba =ABS((C2-D2)/C2)*100 y copie hasta H151. =PROMEDIO(H2:H151) debe dar aproximadamente 26,206831, que se presenta como 26,21 %. Para el pretest sustituya D2 por E2; el promedio debe dar aproximadamente 23,757021. Los valores de la columna auxiliar ya están expresados en porcentaje: no aplique otra multiplicación por 100.')
p('Para abarrotes use =PROMEDIO.SI(B2:B151;"ABARROTES";H2:H151). Cambie el nombre para otra categoría. PROMEDIO sobre todo el rango incluye también filas ocultas por un filtro. No calcule MAPE con los totales acumulados. El Excel redondea métricas y PMS-7 a seis decimales; esa precisión conserva los resultados del Word a dos decimales.')
h('Detalle por producto')
p('Los Word conservan Producto en la ficha de volumen: 195 productos en pretest y 197 en postest. Esos números son productos distintos con ventas en cada ventana. No equivalen a las cinco categorías del modelo. Para contrastar un producto, búsquelo en cada reporte del período y sume sus unidades; para la categoría, utilice el catálogo maestro.')

page('Referencia mensual e inventario')
h('La referencia de estacionalidad')
p('Se utiliza junio de 2026 para ambas ventanas porque es un mes completo anterior a los dos cortes. Agosto no sería una referencia completamente previa al pretest. Esta selección es un supuesto de la comparación y debe confirmarse con el docente antes de cerrar la metodología.')
table(['Categoría','Ventas de junio','VPM diario','VPR de 30 días'],[[c,f(x*30),f(x),f(x*30)] for c,x in j['reference_daily'].items()]+[['GLOBAL','7 842,00','261,40','7 842,00']],[5.3,3.9,3.9,3.9])
p('La ficha conserva IE = VD / VPM, pero interpreta VPM como promedio diario del mes de referencia. La tabla registra 30 días del negocio: VD suma las cinco categorías por fecha. No se promedian sus porcentajes para crear el índice global. Para IE = VP / VPR, ambas cantidades abarcan 30 días. Los cocientes describen la relación con junio y no acreditan por sí solos estacionalidad recurrente.')
h('Cifras demostrativas de inventario')
table(['Dato','Pretest','Postest'],[[k,f(P['pre']['inventory_total'][key]),f(P['post']['inventory_total'][key])] for k,key in [('Stock inicial','SI'),('Entradas','EN'),('Stock final','SF'),('CD por balance','CD'),('ISI porcentaje','ISI'),('DQS categoría-días','DQS'),('TQS porcentaje','TQS'),('Ventas menos CD','sales_difference')]],[7,5,5])
p('El archivo activo de inventario se identifica como demostrativo y no apto para evidencia de tesis. Las diferencias con las ventas son 3 469 unidades en pretest y 3 311 en postest. Existen saldos incompletos en los límites de las ventanas y discontinuidades entre registros. No se corrigieron datos para forzar la igualdad. Se necesita kardex real conciliado antes de interpretar cambios de ISI o TQS como efectos en el negocio.')

page('Interpretación y evidencia para el docente')
h('Qué resultado permite afirmar esta comparación')
p('En la ventana del postest de 30 días, el modelo obtiene MAPE de 26,21 % y PMS-7 obtiene 29,14 % sobre las mismas ventas: el modelo reduce el error en 2,93 puntos porcentuales frente a esa referencia. Es una comparación retrospectiva de pronósticos. No demuestra por sí sola causalidad, significancia estadística ni mejora operativa del inventario.')
p('El pretest de la ventana anterior obtiene 23,76 %. Compararlo directamente con el 26,21 % del postest mezcla semanas distintas y condiciones de venta distintas. No debe describirse esa diferencia como mejora. Las 150 observaciones representan cinco categorías durante 30 días, no 150 días independientes.')
p('La prueba previa del 09 al 15/09/2026 arrojó 33,89 % para el modelo y 31,54 % para PMS-7. Se conserva ese resultado desfavorable y se informa la ampliación a 30 días; no se reemplaza silenciosamente. La duración final debe justificarse por el protocolo y el objetivo del estudio.')
h('Qué falta confirmar')
p('Fecha de implementación y método previo real; procedencia y autenticidad de los reportes del negocio; kardex real; autorización para usar los datos; duración y unidad de análisis aprobadas; interpretación de VPM y VPR. También debe concordar el criterio de selección del modelo: el programa usa WAPE, mientras el texto de la tesis menciona selección por MAPE.')
h('Archivos que permiten reproducir los resultados')
p('Carpeta: C:/taully_demand_forecast/outputs/anexo2_30dias_hasta_15septiembre. Los Excel Exportacion_programa_pre_30dias.xlsx y Exportacion_programa_post_30dias.xlsx son descargas del servicio del programa, verificadas contra sus observaciones guardadas. El historial y los reportes fuente permanecen en data.')
for label in ['pre','post']:p(('Pretest PMS-7: ' if label=='pre' else 'Postest modelo: ')+P[label]['run_id'])
p('Las fechas incluidas en esos identificadores registran la ejecución del cálculo, no la fecha de implementación en la tienda. La ejecución es reproducible mientras se conserven los archivos fuente y la misma versión del código.')
p('calculos_30dias_verificados.json conserva las comprobaciones y fuentes; detalle_30dias_verificado.csv contiene los 300 pares diarios de ambas ventanas. Los Word son resultados estáticos de esta revisión: si cambia el historial, el inventario o el código, deberán recalcularse y versionarse.')
d.save(O/'Guia_verificacion_30dias_para_docente_REVISADA.docx')
print(O/'Guia_verificacion_30dias_para_docente_REVISADA.docx')
