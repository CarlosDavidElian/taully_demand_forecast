from pathlib import Path
import json,csv
from docx import Document
from docx.shared import Cm,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
R=Path('C:/taully_demand_forecast'); O=R/'outputs/verificacion_pretest_postest_2026-09-30'
j=json.loads((O/'calculos_verificados.json').read_text(encoding='utf-8'))
detail=list(csv.DictReader((O/'detalle_diario_verificado.csv').open(encoding='utf-8-sig')))
d=Document(); sec=d.sections[0]; sec.page_width=Cm(21);sec.page_height=Cm(29.7)
sec.top_margin=sec.bottom_margin=Cm(1.8);sec.left_margin=sec.right_margin=Cm(2)
for name in ['Normal','Title','Heading 1','Heading 2']:
 s=d.styles[name];s.font.name='Calibri';s.font.color.rgb=RGBColor(0,0,0)
for el in d.styles.element.xpath('.//w:pBdr'):
 el.getparent().remove(el)
d.styles['Normal'].font.size=Pt(10.5);d.styles['Normal'].paragraph_format.space_after=Pt(6)
d.styles['Normal'].paragraph_format.line_spacing=1.06
d.styles['Title'].font.size=Pt(23);d.styles['Heading 1'].font.size=Pt(17);d.styles['Heading 2'].font.size=Pt(12)
p=sec.footer.paragraphs[0];p.alignment=WD_ALIGN_PARAGRAPH.RIGHT
r=p.add_run('Verificación del 30 de septiembre de 2026  |  ');r.font.size=Pt(8)
fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');p._p.append(fld)
def p(t,style=None):return d.add_paragraph(t,style)
def h(t):d.add_heading(t,2)
def page(t):d.add_page_break();d.add_heading(t,1)
def f(x):return f'{x:,.2f}'.replace(',',' ').replace('.',',')
def table(head,rows,widths=None):
 t=d.add_table(rows=1,cols=len(head));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
 widths=widths or [17/len(head)]*len(head)
 for c,w in zip(t.columns,widths):c.width=Cm(w)
 for i,row in enumerate([head]+rows):
  cells=t.rows[0].cells if i==0 else t.add_row().cells
  for c,txt,w in zip(cells,row,widths):
   c.width=Cm(w);c.text=str(txt);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
   pr=c._tc.get_or_add_tcPr(); shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'17365D' if i==0 else ('F0F4F8' if i%2 else 'FFFFFF'));pr.append(shade)
   borders=OxmlElement('w:tcBorders')
   for edge in ['top','left','bottom','right']:
    el=OxmlElement('w:'+edge);el.set(qn('w:val'),'single');el.set(qn('w:sz'),'4');el.set(qn('w:color'),'D9D9D9');borders.append(el)
   pr.append(borders)
   for para in c.paragraphs:
    para.paragraph_format.space_before=Pt(4);para.paragraph_format.space_after=Pt(4);para.paragraph_format.line_spacing=1
    for run in para.runs:
     run.font.size=Pt(9.5);run.bold=i==0;run.font.color.rgb=RGBColor.from_string('FFFFFF' if i==0 else '000000')
  trpr=t.rows[i]._tr.get_or_add_trPr();trpr.append(OxmlElement('w:cantSplit'))
  if i==0:trpr.append(OxmlElement('w:tblHeader'))
 p('')
 return t
d.add_heading('Verificación del pretest y del postest',0)
p('Minimarket Taully Comas\nRevisión de cálculos y correspondencia con el programa\n30 de septiembre de 2026')
h('Resultado de la revisión')
p('Las 18 tablas de datos de Pretest_corregido.docx y Postest_corregido.docx coinciden con el recálculo independiente de los archivos revisados. Las ocho fórmulas de los instrumentos se conservan. Sin embargo, esa coincidencia no permite presentar todavía ambos documentos como evidencia definitiva de una intervención real: el inventario es demostrativo y falta sustentar la ejecución del pretest y del postest en el negocio.')
table(['Aspecto','Resultado'],[
 ['Cálculos de los dos Word','78 comprobaciones automáticas sin discrepancias'],
 ['Exportación del postest','40 comprobaciones de categorías y registros diarios conformes; redondeo de métricas a seis decimales'],
 ['Pruebas del programa','30 pruebas de evaluación, inventario, pronóstico y web aprobadas'],
 ['Mejora del pronóstico','No demostrada en la semana evaluada'],
 ['Inventario como evidencia real','No apto mientras se utilice el archivo demostrativo']],[5,12])
h('Períodos y cifras centrales')
table(['Dato','Pretest','Postest'],[
 ['Período evaluado','02 al 08/09/2026','09 al 15/09/2026'],
 ['Última fecha previa al pronóstico','01/09/2026','08/09/2026'],
 ['Ventas registradas en unidades','1 784','1 671'],
 ['Demanda pronosticada en unidades','2 016,97','1 931,38'],
 ['MAPE global','27,69 %','33,89 %'],
 ['Observaciones','35 categoría-días','35 categoría-días']],[7,5,5])
p('Son siete fechas por período y cinco categorías. Las 35 observaciones no equivalen a 35 días independientes. El pretest usa PMS-7 reconstruido; el postest procede de la evaluación retrospectiva del programa.')

page('Resultados de ventas y error de pronóstico')
p('La demanda real de estas fichas corresponde a unidades vendidas. No mide por sí sola pedidos no atendidos o demanda perdida por falta de stock. La demanda pronosticada es la suma de las siete estimaciones diarias.')
for period,title in [('pre','Pretest del 02 al 08 de septiembre'),('post','Postest del 09 al 15 de septiembre')]:
 h(title); v=j['periods'][period]
 rows=[[c,int(s['actual']),f(s['prediction']),f(s['mape'])+' %'] for c,s in v['by_category'].items()]
 rows.append(['GLOBAL',int(v['global']['actual']),f(v['global']['prediction']),f(v['global']['mape'])+' %'])
 table(['Categoría','DR unidades','DP unidades','MAPE'],rows,[5.3,3.6,4.2,3.9])
h('Cómo interpretar la comparación')
p('El MAPE pasa de 27,69 % a 33,89 %, pero corresponden a semanas distintas. Esta diferencia no aísla el efecto del sistema. Para comparar los pronósticos sobre las mismas ventas, el programa también evalúa PMS-7 del 09 al 15 de septiembre: obtiene 31,54 %, frente a 33,89 % del modelo, una diferencia desfavorable de 2,35 puntos porcentuales.')
p('El PMS-7 de 31,54 % que aparece junto al modelo en el programa no es el pretest de 27,69 %. Es una referencia calculada sobre la semana del postest. Con esta prueba no corresponde afirmar que el modelo mejoró la precisión ni que se confirmó una mejora estadísticamente significativa.')
p('Las dos fichas de error de cada Word utilizan los mismos pares de demanda real y pronosticada. Su repetición no crea una segunda muestra.')

page('Cómo comprobar las operaciones de las fichas')
h('Volumen de demanda y MAPE')
p('Para CD por ventas, filtre las siete fechas y la categoría en los reportes; sume Cantidad Vendida en unidades. Para MAPE, aplique la fórmula de la ficha a cada par diario DR y DP, sume los errores porcentuales absolutos y divida entre siete. Para el global se utilizan los 35 pares, sin redondear los resultados intermedios. No se calcula MAPE dividiendo los totales semanales.')
rows=[x for x in detail if x['period']=='post' and x['category']=='ABARROTES']
table(['Fecha de 2026','DR','DP','Error absoluto porcentual'],[[x['date'][8:10]+'/09',int(float(x['actual'])),f(float(x['prediction'])),f(float(x['absolute_percentage_error']))+' %'] for x in rows],[4,3,4,6])
p('Ejemplo del 09/09: la venta fue 31 y el pronóstico 47,75. La diferencia absoluta es 16,75; al dividir entre 31 y expresar en porcentaje se obtiene 54,03 %. El promedio de los siete errores diarios de abarrotes es 39,51 %. No hay demanda real cero en los 70 registros evaluados; si la hubiera, el MAPE requeriría un tratamiento explícito, no sustituirla silenciosamente.')
h('Índices de estacionalidad de ambas fichas')
p('La ficha diaria mantiene IE = VD / VPM. El cálculo de los Word interpreta VPM como el promedio diario de agosto: ventas del mes divididas entre sus 31 días. La ficha de patrón mantiene IE = VP / VPR; para un VP de siete días, VPR es siete veces ese promedio diario. Esta interpretación debe quedar expresamente validada por el docente, porque el instrumento denomina VPM “Venta Promedio Mensual”.')
p('Ejemplo de abarrotes: agosto suma 1 748 unidades; el promedio diario es 56,38709677. El 09/09, IE diario es 31 dividido entre ese promedio: 0,55. La referencia de siete días es 394,7096774; el IE semanal del postest es 330 dividido entre esa referencia: 0,84. Para las cinco categorías juntas, el IE semanal es 0,93 en pretest y 0,87 en postest.')
p('Estos cocientes describen las ventas respecto a la referencia elegida. Una sola semana no acredita por sí misma un patrón estacional recurrente. Los índices se comprobaron fuera del programa; no aparecen calculados en el panel ni en su exportación actual.')

page('Inventario y quiebres de stock')
p('Las operaciones CD = (SI + EN) − SF, ISI = CD / (SI + EN) y TQS = (DQS / DD) × 100 coinciden con las fichas. ISI se presenta como porcentaje. La validez de sus resultados depende de un kardex completo y auténtico.')
table(['Magnitud','Pretest','Postest'],[
 ['Stock inicial SI','1 988','1 965'],['Entradas EN','778','786'],['Stock final SF','1 355','1 440'],
 ['CD del balance de inventario','1 411','1 311'],['ISI global','51,01 %','47,66 %'],
 ['DQS sumados entre categorías','2','3'],['DD global','35 categoría-días','35 categoría-días'],['TQS global','5,71 %','8,57 %'],
 ['Ventas registradas','1 784','1 671'],['Ventas menos CD del inventario','373','360']],[8,4.5,4.5])
h('Verificación de los denominadores')
p('Para una categoría se evalúan siete días: un día con quiebre produce 14,29 % y dos producen 28,57 %. En el total hay cinco categorías por siete días: pretest 2/35 y postest 3/35. Este total es una proporción de categoría-días; no representa directamente días del negocio sin stock. Cada fecha con quiebre dentro de una categoría se cuenta una vez.')
h('Problemas de la fuente de inventario')
p('El archivo activo inventario_producto_Taully_2026-04-01_a_2026-09-15.xlsx está identificado por el programa como demostración y no apto para evidencia. El balance de inventario difiere de las ventas en 373 unidades para pretest y 360 para postest. No corresponde forzar la igualdad ni atribuir la diferencia a una causa sin revisar movimientos y ajustes.')
p('Además, faltan registros de los límites del período para numerosos productos: pretest, 107 de 167 sin fila del primer día y 107 sin fila del último; postest, 111 de 169 sin fila del primer día y 101 sin fila del último. El cálculo toma la primera y la última fila disponibles por producto, que pueden no coincidir con el inicio y cierre reales. Se detectaron 31 y 23 discontinuidades entre saldos sucesivos, respectivamente.')
p('Se necesita el kardex real con saldos de apertura y cierre y movimientos conciliados. Hasta entonces, CD de inventario, ISI y TQS quedan como resultados demostrativos. Tampoco debe afirmarse reducción de quiebres: las cifras demostrativas pasan de 5,71 % a 8,57 %.')

page('Dónde comprobar el postest en el programa')
p('El postest del pronóstico sí se obtiene del programa. El pretest documenta la situación previa y debe respaldarse con reportes y el método previo acreditado. El botón Postest realiza una evaluación histórica; su nombre no acredita que el sistema se implementó realmente en el negocio antes del período.')
for s in [
 '1. Abra http://127.0.0.1:5000/ con el programa en ejecución.',
 '2. En Corte de prueba seleccione 08/09/2026 y en Horizonte seleccione 7 días.',
 '3. Pulse Ejecutar Postest y espere a que termine la evaluación.',
 '4. En Resultado del Postest compruebe el período 09 al 15/09/2026, cinco categorías, 35 observaciones y MAPE del modelo 33,89 %.',
 '5. Pulse Descargar Postest. Abra Métricas categoría para el MAPE por categoría y Detalle diario para las ventas reales y los pronósticos.',
 '6. Para inventario utilice Indicadores de inventario del período y la hoja Inventario del Excel. La vista general del inventario abarca abril a septiembre y muestra otros acumulados.']:p(s)
table(['Ficha','Ubicación de la comprobación'],[
 ['1 CD de inventario','Panel del período y hoja Inventario; SI, EN y SF en el Excel'],
 ['2 ISI','Panel del período y hoja Inventario'],
 ['3 TQS','Panel del período y hoja Inventario'],
 ['4 IE diario','Cálculo externo en Word con ventas diarias y referencia de agosto'],
 ['5 Error MAPE','MAPE global en panel; MAPE por categoría en Métricas categoría'],
 ['6 Volumen de demanda','Demanda real por categoría y suma de Detalle diario'],
 ['7 IE semanal','Cálculo externo en Word con ventas semanales y referencia de agosto'],
 ['8 Precisión MAPE','Mismos datos y ubicaciones que la ficha 5']],[5,12])
p('Atención: la tabla del panel compara WAPE por categoría. Ese porcentaje no debe copiarse en la columna MAPE del Word. El Excel exportado conserva ambos indicadores y permite identificarlos por su encabezado.')
p('MAE y RMSE aparecen como métricas complementarias del programa y se mencionan en el cuerpo de la tesis. No se añadieron a las fichas corregidas: los instrumentos de error proporcionados usan MAPE. Es necesario armonizar el texto metodológico con las fichas finalmente aprobadas.')

page('Pendientes para sustentar los resultados')
h('Diseño y ejecución del estudio')
p('Documentar la fecha real de implementación, el período autorizado de observación y el método utilizado antes de implementar el sistema. El PMS-7 reconstruido es una referencia técnica; no acredita que fuera el procedimiento previo del minimarket. Sin esos antecedentes, la evaluación demuestra funcionamiento retrospectivo, no el efecto de una intervención real.')
p('Justificar con el docente el horizonte de siete días y el análisis estadístico previsto. Siete días pueden servir para una prueba operativa semanal, pero no bastan por sí solos para sostener significancia, estacionalidad o causalidad. Cualquier ampliación debe definirse por razones metodológicas y mantener también los resultados desfavorables.')
h('Coherencia de la tesis con la implementación')
p('La tesis plantea abril a diciembre de 2026 y menciona estimaciones de 320 productos y 10 categorías. Los archivos revisados cubren del 01/04 al 15/09/2026, con 203 productos y cinco categorías. Debe distinguirse lo planificado de lo efectivamente analizado. Las 11 748 filas de ventas no equivalen necesariamente a 11 748 transacciones o comprobantes.')
p('El texto menciona ARIMA, Random Forest, CatBoost y XGBoost y selección por MAPE. El programa auditado compara promedio estacional y Random Forest y selecciona por WAPE; en esta ejecución utiliza promedio estacional en cuatro categorías y Random Forest en golosinas. La validación interna usa un corte cronológico y la prueba externa pronostica siete días de forma recursiva. Se requiere describir y justificar el procedimiento que realmente se utilice.')
h('Sustento documental y trazabilidad')
p('Conservar los reportes originales del negocio, el kardex real, el respaldo de autorización para usar los datos y los documentos que correspondan al protocolo institucional. La coincidencia numérica entre archivos no certifica su autenticidad ni sustituye la revisión del docente o del comité.')
p('Se contrastaron 168 reportes diarios, 11 748 filas y 43 105 unidades contra 840 registros del historial, sin diferencias de cantidades ni productos sin correspondencia en el catálogo. El recálculo fue independiente del servicio de indicadores. La exportación cotejada contiene cinco categorías y 35 pares diarios; las métricas exportadas coinciden dentro de su redondeo a seis decimales.')
p('Archivos revisados: Pretest_corregido.docx y Postest_corregido.docx, en outputs/postest_visualizado. Fuentes de cálculo: historial_demanda.csv, catalogo_maestro.xlsx, reportes diarios e inventario activo. El manifiesto de archivos y sus huellas SHA-256 queda en calculos_verificados.json; el detalle de las 70 observaciones queda en detalle_diario_verificado.csv, en la carpeta de esta revisión.')
p('Ejecución exportada y verificada: '+j['live_run']['run_id']+'.')
h('Referencia metodológica para la comprobación')
p('Hyndman y Athanasopoulos, Forecasting Principles and Practice, tercera edición: evaluación fuera de muestra y limitaciones del MAPE cuando la demanda real es cero (https://otexts.com/fpp3/accuracy.html); validación temporal con información anterior al corte (https://otexts.com/fpp3/tscv.html).')
path=O/'Verificacion_integral_Pretest_Postest.docx';d.save(path);print(path)
