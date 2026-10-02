from pathlib import Path
import json,csv
from docx import Document
from docx.shared import Cm,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
R=Path('C:/taully_demand_forecast');O=R/'outputs/verificacion_pretest_postest'
a=json.loads((O/'verificacion_calculos.json').read_text(encoding='utf-8'));pre=a['periods']['Pretest'];post=a['periods']['Postest'];cats=list(pre['bycat'])
d=Document();s=d.sections[0];s.page_width=Cm(21);s.page_height=Cm(29.7);s.top_margin=s.bottom_margin=Cm(1.8);s.left_margin=s.right_margin=Cm(1.8)
for name,size in [('Normal',10.5),('Title',21),('Heading 1',15),('Heading 2',12)]:
 st=d.styles[name];st.font.name='Arial';st.font.size=Pt(size);st.font.color.rgb=RGBColor(0,0,0);st.paragraph_format.space_after=Pt(7);st.paragraph_format.line_spacing=1.06
def p(t,style=None):return d.add_paragraph(t,style)
def h(t):p(t,'Heading 1')
def page(t):d.add_page_break();h(t)
def f(v):return f'{v:.2f}'.replace('.',',')
def table(head,rows,widths):
 t=d.add_table(rows=1,cols=len(head));t.autofit=False;t.alignment=WD_TABLE_ALIGNMENT.CENTER
 for c,v in zip(t.rows[0].cells,head):c.text=str(v)
 for row in rows:
  for c,v in zip(t.add_row().cells,row):c.text=str(v)
 for col,w in zip(t.columns,widths):col.width=Cm(w)
 borders=OxmlElement('w:tblBorders')
 for edge in ['top','bottom','left','right','insideH','insideV']:
  e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
 t._tbl.tblPr.append(borders)
 for ri,row in enumerate(t.rows):
  row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
  if ri==0:row._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
  for ci,(c,w) in enumerate(zip(row.cells,widths)):
   c.width=Cm(w);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER;pr=c._tc.get_or_add_tcPr()
   sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'193854' if ri==0 else ('F1F4F6' if ri%2==0 else 'FFFFFF'));pr.append(sh)
   mar=OxmlElement('w:tcMar')
   for edge in ['top','bottom','left','right']:
    e=OxmlElement('w:'+edge);e.set(qn('w:w'),'80');e.set(qn('w:type'),'dxa');mar.append(e)
   pr.append(mar)
   for pp in c.paragraphs:
    pp.paragraph_format.space_after=Pt(0);pp.paragraph_format.line_spacing=1
    for run in pp.runs:run.font.size=Pt(9.5);run.bold=ri==0;run.font.color.rgb=RGBColor.from_string('FFFFFF' if ri==0 else '000000')
 p('')
 return t

p('Verificación del pretest y postest','Title')
p('Minimarket Taully • Revisión técnica del 29 de septiembre de 2026')
p('Documentos revisados: Pretest_corregido.docx y Postest_corregido.docx. Destinatarios: investigador, docente de tesis y comité de ética.')
h('Resultado de la revisión')
p('Los valores de las fichas son reproducibles con los archivos disponibles y el postest calculado por el programa. No están sustentados todavía como resultados definitivos de una intervención real: el inventario es demostrativo, sus saldos no concilian con las ventas y no se ha confirmado la fecha de aplicación en la tienda ni el método previo usado en la práctica.')
table(['Comprobación','Resultado'],[
('Cálculos en Word, programa y Excel',f"{a['checks']} comparaciones numéricas sin diferencias fuera del redondeo indicado."),
('Reportes frente al historial','168 reportes; 11 748 filas de venta; 43 105 unidades. Cero diferencias al consolidar por fecha y categoría con el catálogo actual.'),
('Historial','840 registros categoría-día; cinco categorías; del 01/04 al 15/09/2026. Sin claves fecha-categoría duplicadas.'),
('Repetición del postest','La ejecución del 29/09/2026 reprodujo las 35 predicciones del resultado guardado del 27/09. Diferencia máxima: 0,00 unidades.'),
('Alcance de la conformidad','Conformidad aritmética y reproducción técnica. La autenticidad empresarial, la validez del diseño y la aprobación ética requieren sustento adicional.')],[5,12.4])
p('La coincidencia entre archivos comprueba consistencia, no acredita por sí sola que sean registros originales del negocio. Las filas de venta no equivalen al número de tickets o transacciones.')
p('Esta revisión conserva las fórmulas de los instrumentos y no cambia los datos ni el código del programa. No constituye aprobación del docente, juicio de expertos ni dictamen del comité.')

page('Pronóstico y comparación de períodos')
table(['Evaluación','Fechas','Corte','Ventas','Pronóstico','MAPE'],[
('Pretest PMS-7','02–08/09','01/09','1784','2016,97','27,69 %'),
('Postest modelo','09–15/09','08/09','1671','1931,38','33,89 %'),
('PMS-7 de comparación','09–15/09','08/09','1671','1801,98','31,54 %')],[3.7,2.5,2,2,3,2.2])
p('Todas las fechas corresponden a 2026. El PMS-7 mostrado junto al modelo en el programa se calcula para la misma semana del postest. No es el pretest de la semana anterior.')
table(['Categoría','DR pre','DP pre','MAPE pre','DR post','DP post','MAPE post'],[[c,int(pre['bycat'][c]['dr']),f(pre['bycat'][c]['dp']),f(pre['bycat'][c]['mape'])+' %',int(post['bycat'][c]['dr']),f(post['bycat'][c]['dp']),f(post['bycat'][c]['mape'])+' %'] for c in cats],[3.6,1.7,2.3,2.6,1.7,2.6,2.9])
p('Fórmula del instrumento: MAPE = (100/n) × Σ |(DR − DP) / DR|. Se utilizan siete parejas diarias por categoría y 35 parejas para el global. No hay DR igual a cero en estas ventanas. Los totales semanales de DR y DP no se usan como una sola observación para el MAPE.')
p('El MAPE del modelo en la misma semana es 2,35 puntos porcentuales mayor que el de PMS-7. Por tanto, esta ejecución no respalda una mejora global de precisión. Tampoco permite afirmar que el modelo sea inferior en cualquier otro período. La diferencia entre 27,69 % y 33,89 % mezcla semanas y métodos distintos; no permite atribuir un efecto al programa.')
p('El PMS-7 del pretest se reconstruyó usando las siete últimas observaciones hasta el 01/09 y agregando cada estimación al siguiente paso. La tesis contempla PMS como referencia, pero no se aportó evidencia de que PMS-7 fuera el método usado previamente por la tienda.')
p('En esta ejecución el programa selecciona promedio estacional para ABARROTES, BEBIDAS, HELADOS y LIMPIEZA; Random Forest para GOLOSINAS. Esta selección debe describirse de forma explícita en el informe de tesis.')

page('Inventario y estacionalidad')
table(['Indicador agregado','Pretest','Postest'],[
('SI','1988','1965'),('EN','778','786'),('SF','1355','1440'),('CD por balance','1411','1311'),('CD por ventas','1784','1671'),('Ventas menos balance','373','360'),('ISI en porcentaje','51,01 %','47,66 %'),('DQS / categoría-días','2 / 35','3 / 35'),('TQS','5,71 %','8,57 %'),('IE semanal frente a agosto','0,93','0,87')],[8.4,4.5,4.5])
p('CD por balance = (SI + EN) − SF. CD por ventas = Σ ventas realizadas. Son los dos cálculos presentes en los instrumentos; sus resultados no deben confundirse. La diferencia requiere saldos y movimientos verificables. En presencia de quiebres, las ventas observadas tampoco miden toda la demanda no atendida.')
p('El ISI conserva CD / (SI + EN), presentado como porcentaje. El total se calcula con las sumas de CD, SI y EN. TQS conserva (DQS / DD) × 100. El total usa 35 categoría-días; no representa 35 días calendario de la tienda ni exige que toda una categoría se haya agotado.')
p('Los saldos del programa son la primera y última fila disponible de cada producto dentro de la semana. En pretest, 107 de 167 productos no tienen fila al inicio y 107 al final; en postest, faltan esas filas en 111 y 101 de 169 productos. Se observan 31 y 23 diferencias entre saldos de registros sucesivos, respectivamente. Con días faltantes, no equivalen automáticamente a pérdidas, pero impiden afirmar que se cuenta con un kardex semanal completo.')
p('Estacionalidad: se comprobaron los 35 registros diarios de cada Word y los índices semanales. VPM se interpreta como promedio diario de agosto; VPR es su equivalente de siete días. Las operaciones coinciden. El docente debe confirmar esta definición operacional, porque “venta promedio mensual” puede interpretarse de otra manera. No se demuestra una estacionalidad recurrente con un único cociente frente a agosto.')

page('Qué se muestra en el programa')
p('El postest del 09 al 15/09/2026 se ejecutó y se visualizó en http://127.0.0.1:5000/. El panel confirma corte 08/09, horizonte siete días y 35 observaciones. El Excel descargado corresponde a esa ejecución y coincide en detalle diario, MAPE por categoría e indicadores de inventario.')
table(['Contenido del Word','Dónde comprobarlo'],[
('MAPE global 33,89 %','Resultado del postest → tarjeta Modelo predictivo → MAPE.'),
('MAPE y ventas por categoría','Excel → Métricas categoría. La tabla visible del panel muestra WAPE por categoría, no MAPE.'),
('Demanda real y pronosticada diaria','Excel → Detalle diario. La suma del modelo es 1931,38 y la suma real es 1671.'),
('CD, ISI, DQS y TQS','Indicadores de inventario del período. SI, EN y SF se verifican también en la hoja Inventario del Excel.'),
('IE diario y semanal','Se calcula en las fichas desde el historial. El panel y esta exportación no muestran actualmente un indicador IE.'),
('Pretest 02–08/09','Word y cálculo independiente sobre reportes e historial. No existe una ficha específica de pretest en este panel.')],[5.7,11.7])
p('Para reproducirlo: elegir horizonte 7 días; ingresar 08/09/2026 en Corte de prueba; pulsar Ejecutar Postest; comprobar 09–15/09; pulsar Descargar Postest. No tomar el resumen de todo el inventario de abril a septiembre como resultado de esta semana.')
p('El programa sigue mostrando métricas adicionales y utiliza WAPE para su comparación y selección interna. Las fichas solicitadas conservan MAPE. La tesis indica seleccionar por MAPE: esa diferencia entre metodología y programa debe resolverse y documentarse antes de la evaluación final; ocultar columnas no cambia el criterio del modelo.')
d.add_picture(str(O/'Resultado_postest_programa.png'),width=Cm(17.1))

page('Observaciones para la revisión académica')
p('1. Naturaleza del estudio. Las fichas actuales son una evaluación histórica. Para presentarlas como antes y después de una intervención, falta documentar la fecha real de implementación, el procedimiento anterior y la manera en que se utilizaron los pronósticos. La respuesta a esta consulta sigue pendiente al cerrar la revisión.')
p('2. Origen del inventario. El propio programa identifica el archivo como demostración y “no apto para evidencia”. Para resultados de negocio deben sustituirse los saldos y días de quiebre por registros validados y conservar los resultados de demostración como tales. No procede ajustar cifras para obtener una mejora.')
p('3. Alcance declarado. La tesis original propone abril–diciembre de 2026, aproximadamente 8500 transacciones, 320 productos y diez categorías. Lo disponible llega al 15/09, con 203 productos de catálogo y cinco categorías. Debe distinguirse lo planificado de lo recolectado; septiembre no permite certificar resultados de diciembre.')
p('4. Métodos y métricas. Los anexos de error de pronóstico contienen MAPE. Sin embargo, la hipótesis y el análisis estadístico de la tesis también nombran MAE y RMSE. Se respeta la instrucción de mantener solo MAPE en las fichas; corresponde armonizar el texto metodológico con el docente. Los modelos propuestos incluyen ARIMA, Random Forest, CatBoost y XGBoost; esta ruta del programa compara Random Forest y promedio estacional.')
p('5. Validación. El panel ejecuta un corte histórico externo y una partición interna cronológica 80/20. Eso no documenta por sí solo múltiples cortes de validación hacia adelante. El Random Forest seleccionado se conserva sin reajustarlo con todo el tramo previo al corte. La selección interna se evalúa con rezagos observados y el pronóstico externo es recursivo de siete días. Estas decisiones deben estar justificadas; si se cambian, hay que generar y conservar una nueva evaluación.')
p('6. Interpretación. Siete días pueden formar una ventana de evaluación, pero no justifican por sí solos el tamaño de muestra ni una mejora estadísticamente significativa. No se encontró un contraste inferencial para estas fichas. Los 35 registros comparten fechas y categorías; no debe suponerse independencia sin examinarla.')
p('7. Integridad y acceso a datos. Para la revisión ética conviene presentar autorización de uso de la información comercial, procedencia verificable, responsable de custodia y medidas de confidencialidad. Esos documentos no se acreditaron en esta revisión. Los requisitos específicos los establece el comité de la institución; esta comprobación no los sustituye.')
p('Conclusión para el docente: las operaciones están verificadas; permanecen pendientes la fuente real de inventario, la definición del antes y después y la coherencia entre protocolo, fichas y programa. Los resultados actuales no prueban la mejora planteada.')

page('Cómo reproducir y conservar la verificación')
p('Ejemplos con las fórmulas originales','Heading 2')
table(['Cálculo','Sustitución y resultado'],[
('CD pretest ABARROTES','(372 + 189) − 316 = 245 unidades.'),
('CD postest ABARROTES','(383 + 175) − 276 = 282 unidades.'),
('ISI postest ABARROTES','282 / (383 + 175) = 0,505376… = 50,54 %.'),
('TQS postest GOLOSINAS','(2 / 7) × 100 = 28,57 %.'),
('Volumen postest','330 + 611 + 496 + 146 + 88 = 1671 unidades.'),
('IE diario postest ABARROTES 09/09','31 / (1748 / 31) = 0,55, redondeado.'),
('IE semanal postest ABARROTES','330 / [(1748 / 31) × 7] = 0,84, redondeado.'),
('MAPE postest ABARROTES','Promediar los siete errores porcentuales absolutos diarios da 39,51 %. El del 09/09 es |(31 − 47,75) / 31| × 100 = 54,03 %; un solo día no es el promedio semanal.')],[6,11.4])
p('Archivos conservados','Heading 2')
p('En outputs/verificacion_pretest_postest: Postest_exportado_programa.xlsx; detalle_diario.csv con las 70 parejas diarias y sus errores; comprobaciones.csv con 715 verificaciones; verificacion_calculos.json; fuentes_sha256.json con la identificación de los documentos y fuentes; y la captura del programa. Las huellas permiten detectar cambios posteriores, no certificar autenticidad.')
p('Ejecución comprobada: '+a['run_id']+'. Fecha local: 29/09/2026, 20:25:23, Lima. El identificador usa UTC y por eso contiene la fecha 30/09/2026. Los Word conservan como fuente el resultado del 27/09, cuyas predicciones se reprodujeron exactamente.')
p('Fuentes de la revisión','Heading 2')
p('Pretest_corregido.docx; Postest_corregido.docx; LIMA-NORTE_PI_GURRERRO.docx; 168 reportes diarios; catalogo_maestro.xlsx; historial_demanda.csv; inventario demostrativo y código de cálculo y exportación. El archivo Informe de tesis (1).docx no se usó como sustituto de la tesis original indicada por el investigador.')
p('Referencia metodológica: Hyndman y Athanasopoulos, Forecasting: Principles and Practice, 3.ª edición, secciones 5.8 y 5.10. Fundamentan evaluar con observaciones no utilizadas en el ajuste y respetar el orden temporal. Consultadas el 29/09/2026.')
p('https://otexts.com/fpp3/accuracy.html\nhttps://otexts.com/fpp3/tscv.html')
d.core_properties.title='Verificación del pretest y postest'
d.save(O/'Verificacion_pretest_postest_para_revision.docx')
print(O/'Verificacion_pretest_postest_para_revision.docx')
