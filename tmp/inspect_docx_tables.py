from pathlib import Path
from docx import Document


path = Path(r"C:\taully_demand_forecast\outputs\01a0b651-7620-73f0-8f6a-8b68b14f7a17\Anexo_2_Instrumentos_Completos_Pretest_Postest.docx")
document = Document(path)
print(f"paragraphs={len(document.paragraphs)} tables={len(document.tables)}")
for index, table in enumerate(document.tables):
    print(f"--- table {index} ({len(table.rows)}x{len(table.columns)}) ---")
    for row in table.rows:
        print(" | ".join(cell.text.replace("\n", " / ") for cell in row.cells))
