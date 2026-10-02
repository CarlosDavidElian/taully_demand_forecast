from pathlib import Path
from copy import deepcopy
import csv, json
from statistics import mean
from datetime import date, timedelta
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT

root=Path('C:/taully_demand_forecast')
src=Document('C:/taully_demand_forecast/outputs/postest_visualizado/Pretest_corregido.docx')
metadata=[deepcopy(src.tables[i]._tbl) for i in [0,2,4,6,9,11,13,15]]
doc=Document('C:/taully_demand_forecast/outputs/postest_visualizado/Pretest_corregido.docx')
for el in list(doc._element.body):
    if el.tag!=qn('w:sectPr'): doc._element.body.remove(el)
for sec in doc.sections:
    sec.page_width=Cm(21);sec.page_height=Cm(29.7)
    sec.top_margin=sec.bottom_margin=Cm(1.7)
    sec.left_margin=sec.right_margin=Cm(2)
for name in ['Normal','Title','Heading 1','Heading 2']:
    s=doc.styles[name];s.font.name='Arial';s.font.color.rgb=RGBColor(0,0,0)
    s.font.size=Pt(10.5 if name=='Normal' else 13)
    s.paragraph_format.space_after=Pt(6)
    s.paragraph_format.line_spacing=1

audit=json.loads((root/'tmp/auditoria_20260929/resultados.json').read_text(encoding='utf-8'))['periods']['post']
cats=list(audit['baseline_by_category'])
hist=list(csv.DictReader((root/'data/historial_demanda.csv').open(encoding='utf-8-sig')))
lookup={(r['date'][:10],r['category']):float(r['quantity']) for r in hist}
days=[(date(2026,9,9)+timedelta(days=i)).isoformat() for i in range(7)]
aug={c:sum(float(r['quantity']) for r in hist if r['category']==c and r['date'].startswith('2026-08'))/31 for c in cats}
pre={c:sum(lookup[(d,c)] for d in days) for c in cats}
assert sum(pre.values())==1671

saved=json.loads((root/'data/posttests/postest_20260927t191247z_20260908_7d_70578bbb.json').read_text(encoding='utf-8'))['result']
assert saved['cutoff_date']=='2026-09-08' and saved['horizon_days']==7
observations=saved['observations']
assert len(observations)==35
assert len({(o['date'],o['category']) for o in observations})==35
assert all(o['date'] in days and o['actual_demand']==lookup[o['date'],o['category']] and o['actual_demand']>0 for o in observations)
def model_score(rows):
    return {'actual':sum(r['actual_demand'] for r in rows),'prediction':sum(r['model_prediction'] for r in rows),'mape':mean(abs(r['actual_demand']-r['model_prediction'])/r['actual_demand'] for r in rows)*100}
audit['baseline_by_category']={c:model_score([r for r in observations if r['category']==c]) for c in cats}
audit['baseline']=model_score(observations)
assert round(audit['baseline']['mape'],2)==33.89
assert round(audit['baseline']['prediction'],2)==1931.38

def fmt(x,n=2):return f'{x:.{n}f}'.replace('.',',')
def integer(x):return str(int(x))
def p(text,style=None):
    pp=doc.add_paragraph(text,style)
    return pp
def layout(t,widths,header=True,size=10):
    t.autofit=False;t.alignment=WD_TABLE_ALIGNMENT.CENTER
    pr=t._tbl.tblPr
    borders=pr.find(qn('w:tblBorders'))
    if borders is not None:pr.remove(borders)
    borders=OxmlElement('w:tblBorders')
    for edge in ['top','left','bottom','right','insideH','insideV']:
        e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
    pr.append(borders)
    for col,w in zip(t.columns,widths): col.width=Cm(w)
    for ri,row in enumerate(t.rows):
        trpr=row._tr.get_or_add_trPr()
        for tag in ['trHeight']:
            for old in list(trpr.findall(qn('w:'+tag))):trpr.remove(old)
        cant=OxmlElement('w:cantSplit');trpr.append(cant)
        if ri==0 and header:trpr.append(OxmlElement('w:tblHeader'))
        for ci,(cell,w) in enumerate(zip(row.cells,widths)):
            cell.width=Cm(w);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcp=cell._tc.get_or_add_tcPr()
            for tag in ['tcBorders','shd','tcMar']:
                for old in list(tcp.findall(qn('w:'+tag))):tcp.remove(old)
            shd=OxmlElement('w:shd');shd.set(qn('w:fill'),'EEEEEE' if ri==0 and header else 'FFFFFF');tcp.append(shd)
            mar=OxmlElement('w:tcMar')
            for edge,val in [('top','65'),('bottom','65'),('left','85'),('right','85')]:
                ee=OxmlElement('w:'+edge);ee.set(qn('w:w'),val);ee.set(qn('w:type'),'dxa');mar.append(ee)
            tcp.append(mar)
            for pp in cell.paragraphs:
                pp.paragraph_format.space_before=Pt(0);pp.paragraph_format.space_after=Pt(0)
                pp.paragraph_format.line_spacing=1
                pp.alignment=WD_ALIGN_PARAGRAPH.LEFT if not header or ci==1 else WD_ALIGN_PARAGRAPH.CENTER
                for r in pp.runs:
                    r.font.name='Arial';r.font.size=Pt(size);r.font.color.rgb=RGBColor(0,0,0)
                    r.font.highlight_color=None;r.bold=(ri==0 if header else ci==0)
def table(headers,rows,widths,size=10):
    t=doc.add_table(rows=1,cols=len(headers))
    for c,s in zip(t.rows[0].cells,headers):c.text=s
    for row in rows:
        for c,s in zip(t.add_row().cells,row):c.text=str(s)
    layout(t,widths,True,size)
    return t
def start(i,title):
    if i:doc.add_page_break()
    p('Postest '+title,'Heading 1')
    el=deepcopy(metadata[i]);doc._element.body.insert(len(doc._element.body)-1,el)
    t=doc.tables[-1]
    for row in t.rows:
        if row.cells[0].text.strip().startswith('Momento'):
            row.cells[1].text='Postest. Del 09 al 15 de septiembre de 2026. Siete días.'
    layout(t,[4.4,12.6],False)
    p('Registro de resultados','Heading 2')
def note(text):
    pp=p(text)
    pp.paragraph_format.space_before=Pt(8)
    for r in pp.runs:r.font.size=Pt(9.5)
inv=audit['inventory'];total=audit['inventory_total']
start(0,'Cantidad demandada por balance de inventario')
table(['N°','Categoría','SI','EN','SF','CD'],[[i+1,c]+[integer(inv[c][k]) for k in ['si','en','sf','cd']] for i,c in enumerate(cats)]+[['','TOTAL']+[integer(total[k]) for k in ['si','en','sf','cd']]],[.8,5,2.8,2.8,2.8,2.8])
note('Fuente: inventario_producto_Taully_2026-04-01_a_2026-09-15.xlsx, archivo demostrativo. Se conservan sus cifras: el balance suma 1311 unidades y las ventas registradas suman 1671. La diferencia de 360 unidades está pendiente de conciliación con un kardex validado; no se modifica para forzar igualdad.')
start(1,'Índice de salida de inventario')
table(['N°','Categoría','CD','SI','EN','ISI'],[[i+1,c]+[integer(inv[c][k]) for k in ['cd','si','en']]+[fmt(inv[c]['isi'])+' %'] for i,c in enumerate(cats)]+[['','TOTAL']+[integer(total[k]) for k in ['cd','si','en']]+[fmt(total['isi'])+' %']],[.8,5,2.8,2.8,2.8,2.8])
note('Se conserva la fórmula del instrumento. El cociente se presenta en formato porcentual: por ejemplo, 0,5054 equivale a 50,54 %. Fuente: el mismo inventario demostrativo; los resultados requieren conciliación.')
start(2,'Tasa de quiebre de stock')
table(['N°','Categoría','DQS','DD','TQS'],[[i+1,c,inv[c]['dqs'],'7',fmt(inv[c]['tqs'])+' %'] for i,c in enumerate(cats)]+[['','TOTAL','3','35 categoría-días','8,57 %']],[.8,5.2,2.5,5,2.5])
note('El total comprende cinco categorías durante siete días: 35 categoría-días. DQS cuenta fechas distintas con quiebre dentro de cada categoría. Fuente: inventario demostrativo; estos registros no acreditan por sí solos quiebres reales de la tienda.')
start(3,'Estacionalidad')
note('VD corresponde a las unidades vendidas en cada fecha. VPM se interpreta como el promedio diario de ventas del mes de referencia, agosto de 2026 (31 días), para comparar magnitudes de igual duración. Se conserva IE = VD / VPM. Fuente: historial_demanda.csv.')
def daily_rows(ds):
    return [[date.fromisoformat(d).strftime('%d/%m'),c,integer(lookup[d,c]),fmt(aug[c]),fmt(lookup[d,c]/aug[c])] for d in ds for c in cats]
table(['Fecha','Categoría','VD','VPM','IE'],daily_rows(days[:3]),[2,5.5,3,3.5,3],9.5)
doc.add_page_break()
p('Postest Estacionalidad continuación','Heading 1')
p('Período del 09 al 15/09/2026. Referencia: promedio diario de agosto de 2026.')
table(['Fecha','Categoría','VD','VPM','IE'],daily_rows(days[3:]),[2,5.5,3,3.5,3],10)
note('Los índices se calculan con el promedio de referencia sin redondear. La tabla muestra dos decimales. Este cociente describe la relación con agosto; no demuestra por sí solo estacionalidad recurrente.')
def forecast_rows():
    return [[i+1,c,integer(pre[c]),fmt(audit['baseline_by_category'][c]['prediction']),fmt(audit['baseline_by_category'][c]['mape'])+' %'] for i,c in enumerate(cats)]+[['','GLOBAL','1671',fmt(audit['baseline']['prediction']),fmt(audit['baseline']['mape'])+' %']]
def forecast_note():
    note('DR y DP son totales semanales. El MAPE se calcula con las parejas diarias, no con estos totales: n = 7 por categoría y n = 35 para el global. No hay demanda real igual a cero en estas 35 observaciones. Fuente: ejecución guardada postest_20260927t191247z_20260908_7d_70578bbb.json e historial_demanda.csv. El modelo pronostica los siete días posteriores al corte del 08/09/2026. Esta evaluación es histórica y no demuestra por sí sola una mejora frente al método previo.')
start(4,'Error de pronóstico')
table(['N°','Categoría','DR semanal','DP semanal','MAPE'],forecast_rows(),[.8,5.2,3.5,4,3.5])
forecast_note()
start(5,'Volumen de demanda')
table(['N°','Categoría','Cantidad vendida','Cantidad demandada CD'],[[i+1,c,integer(pre[c]),integer(pre[c])] for i,c in enumerate(cats)]+[['','TOTAL','1671','1671']],[.8,5.2,5,6])
note('Fuente: ventas registradas en historial_demanda.csv del 09 al 15/09/2026. CD conserva la fórmula de esta dimensión: suma de ventas realizadas. Su total de 1671 unidades difiere de las 1311 unidades calculadas por balance de inventario; la diferencia pendiente de conciliación es de 360 unidades.')
start(6,'Patrón de demanda')
table(['N°','Categoría','VP','VPR','IE'],[[i+1,c,integer(pre[c]),fmt(aug[c]*7),fmt(pre[c]/(aug[c]*7))] for i,c in enumerate(cats)]+[['','TOTAL','1671',fmt(sum(aug.values())*7),fmt(1671/(sum(aug.values())*7))]],[.8,5.2,3.5,4,3.5])
note('VP representa las ventas acumuladas de los siete días. VPR representa las ventas equivalentes de siete días según el promedio diario de agosto de 2026. Ambos valores tienen la misma duración de referencia. Se conserva IE = VP / VPR. Fuente: historial_demanda.csv. El cálculo utiliza valores sin redondear; la presentación utiliza dos decimales.')
start(7,'Precisión del pronóstico')
table(['N°','Categoría','DR semanal','DP semanal','MAPE'],forecast_rows(),[.8,5.2,3.5,4,3.5])
forecast_note()
note('Esta ficha reproduce el mismo indicador de la ficha Error de pronóstico; no constituye una muestra adicional. Los datos corresponden a una evaluación histórica y requieren validación antes de presentarse como resultados de una intervención real.')

out=root/'outputs/postest_visualizado/Postest_corregido.docx'
doc.save(out)
check=Document(out)
txt=' '.join(p.text for p in check.paragraphs)+' '.join(c.text for t in check.tables for r in t.rows for c in r.cells)
assert not any(s in txt for s in ['RMSE','MAE','WAPE'])
assert len(check.tables)==17
assert txt.count('33,89 %')==2
assert 'Pretest' not in txt
assert '27,69' not in txt
print(out)
