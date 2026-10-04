from pathlib import Path
from copy import deepcopy
from io import BytesIO
import zipfile,re
from docx import Document
from docx.shared import Pt,Cm,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
R=Path('C:/taully_demand_forecast')
src=Path(r'C:/Users/Acer/Downloads/Formatos de Plan de Prácticas 2/Formatos de Plan de Prácticas/Metodología SCRUM/PPP1-C1-PRA03-PLAN_PRACTICAS-[APELLIDOS NOMBRES].docx')
d=Document(src)
for section in d.sections:
 section.header_distance=Cm(0.9)
def setp(p,text):
 if p.runs:
  p.runs[0].text=text
  for r in p.runs[1:]:r.text=''
 else:p.add_run(text)
def cell(c,text):
 setp(c.paragraphs[0],text)
 for p in c.paragraphs[1:]:p._element.getparent().remove(p._element)
def fields(t,vals):
 for i,v in enumerate(vals):cell(d.tables[t].cell(i*2+1,0),v)
fields(0,['Guerrero Mora, Carlos David Elian','Pendiente de completar por el practicante','Pendiente de completar por el practicante','Pendiente de confirmar correo institucional'])
fields(1,['TAULLY MARKET E.I.R.L.','20610066748','Comas, Lima. Dirección exacta pendiente de confirmar','Pendiente de acreditar representante legal','Pendiente de confirmar','Pendiente de confirmar'])
fields(2,['Apoyo tecnológico a ventas y planificación de inventario (área propuesta)','Pendiente de designación por la empresa','Pendiente de confirmar','Pendiente de confirmar','Pendiente de confirmar'])
fields(3,['Sistema de predicción de demanda diaria por categoría de producto para Minimarket Taully, Comas','Scrum para el desarrollo; estructura de gestión PMI del formato institucional','Ventas, compras y planificación de inventario'])
setp(d.paragraphs[1],'METODOLOGÍA SCRUM · 3 MESES / 12 SEMANAS\nPROYECTO MINIMARKET TAULLY')
setp(d.paragraphs[15],'PLAN PROPUESTO: Las X indican semanas previstas de trabajo, no actividades ya realizadas. Se organizan seis Sprints de dos semanas para validar y completar el sistema existente. El inicio y fin del plan se definirán con la empresa; las fechas del historial de ventas no son fechas de prácticas.')
setp(d.paragraphs[16],'Las fases de gestión del formato institucional se conservan. Scrum organiza el desarrollo mediante Product Goal, Product Backlog, Sprint Goal, incrementos y eventos de inspección y adaptación. El anexo precisa las entregas, los criterios de aceptación y la evidencia disponible.')
setp(d.paragraphs[18],'Duración propuesta: 12 semanas. La empresa y el docente deberán aprobar las fechas, los responsables y el alcance antes de la ejecución. Datos de razón social y RUC tomados del documento previo proporcionado; pendiente corroboración administrativa. Contacto informado: Jaimes Mays Miguel Angel; no se presume su cargo ni representación legal.')
cell(d.tables[5].cell(1,0),'Pendiente de designación')
cell(d.tables[5].cell(1,1),'Pendiente de confirmación')
cell(d.tables[5].cell(1,2),'Guerrero Mora, Carlos David Elian')
activities={
1:('Identificar necesidades de predicción de ventas y reposición en Taully',[1]),
2:('Identificar usuarios de ventas, compras y responsable de validación',[1]),
3:('Definir valor esperado del pronóstico diario por categoría',[1]),
4:('Acordar objetivos y resultados verificables del sistema',[1]),
5:('Delimitar alcance, fuentes autorizadas y entregables',[1,2]),
6:('Revisar reportes, catálogo, historial y requisitos del panel',[1,2]),
7:('Definir Product Goal para pronóstico y evaluación trazables',[1]),
8:('Detallar historias de usuario y criterios de aceptación',[1,2]),
9:('Priorizar y estimar el Product Backlog con el responsable',[1,2]),
10:('Planificar seis Sprints con entregas y dependencias',[1,2]),
11:('Confirmar roles Scrum, disponibilidad y comunicación',[1]),
12:('Definir calidad, riesgos, fuentes y Definition of Done',[1,2]),
13:('Sprint Planning y selección de trabajo según Sprint Goal',[1,3,5,7,9,11]),
14:('Sprint 1: validar alcance, reportes y catálogo maestro',[1,2]),
15:('Sprint 2: comprobar historial y calidad de datos',[3,4]),
16:('Sprint 3: validar entrenamiento por categoría y métricas',[5,6]),
17:('Sprint 4: comprobar pronósticos y exportaciones',[7,8]),
18:('Sprint 5: reproducir postest y completar tablas del Anexo 2',[9,10]),
19:('Sprint 6: validar kardex real, documentación y entrega',[11,12]),
20:('Integrar cambios, pruebas y documentación de cada incremento',list(range(1,13))),
21:('Daily Scrum y seguimiento del objetivo del Sprint',list(range(1,13))),
22:('Actualizar tareas y evidencia en Sprint Backlog',list(range(1,13))),
23:('Gestionar impedimentos, disponibilidad de datos y riesgos',list(range(1,13))),
24:('Verificar criterios de aceptación y Definition of Done',list(range(1,13))),
25:('Sprint Review con demostración del incremento',[2,4,6,8,10,12]),
26:('Registrar observaciones y validación de interesados',[2,4,6,8,10,12]),
27:('Refinar Product Backlog y prioridades del siguiente Sprint',[2,4,6,8,10,12]),
28:('Revisar tareas terminadas, defectos y métricas de pronóstico',list(range(1,13))),
29:('Sprint Retrospective y acuerdos de mejora',[2,4,6,8,10,12]),
30:('Aplicar acciones de mejora del Sprint anterior',list(range(3,13))),
31:('Validar producto y registrar pendientes críticos',[11,12]),
32:('Consolidar manual, cálculos, pruebas y evidencias',[11,12]),
33:('Transferir solución y solicitar aceptación formal',[12]),
34:('Documentar lecciones aprendidas y riesgos abiertos',[12]),
35:('Presentar resultados y obtener conformidad del plan',[12])}
for row in d.tables[4].rows:
 m=re.match(r'^(\d+)\.',row.cells[0].text)
 if m:
  n=int(m.group(1));txt,weeks=activities[n];cell(row.cells[0],f'{n}. {txt}')
  for w in range(1,13):cell(row.cells[w],'X' if w in weeks else '')

def para(text,bold=False):
 p=d.add_paragraph();p.paragraph_format.space_after=Pt(6);p.paragraph_format.line_spacing=1.1
 r=p.add_run(text);r.font.name='Arial';r.font.size=Pt(10);r.bold=bold
 return p
def page(title):
 d.add_page_break();p=para(title,True);p.runs[0].font.size=Pt(14)
def tbl(headers,rows,widths):
 t=d.add_table(rows=1,cols=len(headers));t.autofit=False
 for col,w in zip(t.columns,widths):col.width=Cm(w)
 for i,vals in enumerate([headers]+rows):
  rr=t.rows[0] if i==0 else t.add_row()
  for c,v,w in zip(rr.cells,vals,widths):
   c.width=Cm(w);c.text=v
   pr=c._tc.get_or_add_tcPr();sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'A00020' if i==0 else 'FFFFFF');pr.append(sh)
   borders=OxmlElement('w:tcBorders')
   for side in ['top','bottom','left','right']:
    el=OxmlElement('w:'+side);el.set(qn('w:val'),'single');el.set(qn('w:sz'),'4');el.set(qn('w:color'),'808080');borders.append(el)
   pr.append(borders)
   for p in c.paragraphs:
    p.paragraph_format.space_after=Pt(4);p.paragraph_format.space_before=Pt(4)
    for r in p.runs:r.font.name='Arial';r.font.size=Pt(9);r.bold=i==0;r.font.color.rgb=RGBColor.from_string('FFFFFF' if i==0 else '000000')
  rr._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
 para('')
page('Anexo 1 Aplicación de Scrum al proyecto Taully')
para('Objetivo general',True)
para('Desarrollar, validar y documentar un sistema que estime la demanda diaria por categoría del Minimarket Taully a partir de reportes de ventas y permita contrastar sus resultados con datos históricos y exportaciones reproducibles.')
para('Product Goal',True)
para('Disponer de un sistema trazable de carga, clasificación, pronóstico y evaluación que el responsable del negocio pueda revisar con sus fuentes y utilizar como apoyo a la planificación, sin confundir estimaciones con ventas realizadas.')
para('Alcance y situación inicial',True)
para('El programa existente incluye interfaz Flask, catálogo maestro, consolidación de ventas, entrenamiento por categoría, pronóstico y evaluación retrospectiva con exportación Excel. El plan organiza su validación y las mejoras pendientes; no acredita que los seis Sprints ya se hayan ejecutado.')
para('Quedan fuera la compra automática, la sustitución del POS y la afirmación de mejoras operativas no demostradas. La distribución del pronóstico a productos no equivale a un modelo entrenado directamente para cada SKU.')
tbl(['Responsabilidad','Asignación propuesta'],[
['Product Owner','Representante del negocio por designar. Prioriza necesidades y valida el valor de las entregas.'],
['Scrum Master','Persona por designar. Facilita el trabajo Scrum y gestiona impedimentos.'],
['Desarrollo','Guerrero Mora, Carlos David Elian, como practicante propuesto; equipo y disponibilidad por confirmar.'],
['Revisión académica','Docente o coordinador por confirmar. Revisa coherencia metodológica y evidencias.']],[5,12.5])
para('Se propone Sprint Planning al inicio, Daily Scrum de seguimiento, Sprint Review y Sprint Retrospective al cierre de cada Sprint. Se conservarán fecha, participantes, acuerdos y evidencias reales. No se consignan reuniones, firmas ni aprobaciones que no estén documentadas.')
para('Definition of Done',True)
para('Una historia se considera terminada cuando cumple sus criterios, sus cálculos se contrastan con la fuente, pasan las pruebas pertinentes, se identifican los datos demostrativos, se actualiza la documentación y el incremento queda disponible para la revisión. La aceptación se registra cuando ocurra.')
page('Anexo 2 Product Backlog propuesto')
para('Prioridades y Sprints son propuestos y deberán acordarse con el responsable. Cada historia describe una capacidad del producto y una comprobación observable.')
tbl(['ID y Sprint','Historia y prioridad','Criterio de aceptación'],[
['HU01 · S1','Como encargado, quiero cargar ventas y catálogo para clasificar productos. Alta.','Procesar reportes autorizados; comprobar categorías y productos sin correspondencia. Función existente a validar.'],
['HU02 · S2','Como revisor, quiero un historial consistente para rastrear las ventas. Alta.','Contrastar fechas y totales contra reportes; revisar duplicados y registros sin categoría.'],
['HU03 · S3','Como analista, quiero comparar modelos por categoría. Alta.','Verificar separación cronológica, variables y selección por WAPE; reportar exclusiones y métricas.'],
['HU04 · S4','Como encargado, quiero pronósticos y exportaciones para planificar. Alta.','Comprobar horizonte, fecha base y totales por categoría; identificar asignación estimada a productos.'],
['HU05 · S5','Como docente, quiero reproducir el postest. Alta.','Ejecutar corte 16/08/2026 y 30 días con la base revisada; contrastar 150 pares, DR 7851 y MAPE 26,21 %.'],
['HU06 · S5','Como revisor, quiero visualizar las tablas del Anexo 2. Alta.','Implementar y exportar VP, VPR, IE y detalle diario con fuente y referencia explícitas. Actualmente pendientes en el panel.'],
['HU07 · S6','Como encargado, quiero indicadores con kardex real. Alta.','Conciliar saldos, entradas y salidas; recalcular ISI y TQS y distinguirlos de la demostración actual.'],
['HU08 · S6','Como usuario, quiero manual y entrega verificable. Media.','Entregar guía con capturas, evidencia de pruebas pertinentes y pendientes; solicitar conformidad real.']],[2.1,7.2,8.2])
para('La estimación del esfuerzo se acordará durante la planificación. No se atribuyen puntos, velocidad histórica ni horas efectivamente trabajadas sin registro. Las seis ventanas de dos semanas son la propuesta del formato.')
page('Anexo 3 Incrementos y control de evidencia')
tbl(['Sprint','Objetivo e incremento esperado'],[
['1 · Semanas 1 y 2','Alcance y datos: fuentes autorizadas, catálogo revisado, Product Backlog y criterios acordados.'],
['2 · Semanas 3 y 4','Historial trazable: conciliación de ventas, fechas y categorías con registro de incidencias.'],
['3 · Semanas 5 y 6','Entrenamiento verificable: validación cronológica y documentación del método elegido por categoría.'],
['4 · Semanas 7 y 8','Pronóstico utilizable: revisión de horizontes, exportaciones y distribución estimada a productos.'],
['5 · Semanas 9 y 10','Postest reproducible y tablas de demanda integradas: comparación de pantalla, Excel y fichas.'],
['6 · Semanas 11 y 12','Validación con kardex real y entrega: manual, evidencias, limitaciones y conformidad solicitada.']],[4,13.5])
para('Evidencia técnica disponible',True)
para('La base revisada contiene 840 registros diarios por categoría, correspondientes a 168 fechas y cinco categorías; el catálogo contiene 203 productos. Son cifras de la versión examinada y deberán recalcularse si cambian las fuentes.')
para('El postest del 17/08 al 15/09/2026 utiliza 150 observaciones. El modelo obtiene MAPE 26,21 % y PMS-7 29,14 % en la misma ventana. Esta comparación retrospectiva no demuestra causalidad ni una reducción real del inventario. El pretest reconstruido anterior tiene MAPE 23,76 % y requiere confirmar el método previo y la fecha real de intervención.')
para('El código de entrenamiento aplica requisitos de historia por categoría y compara candidatos usando WAPE. Las fichas académicas utilizan MAPE; ambos criterios deben explicarse de forma coherente. No se trasladan automáticamente al documento resultados de pruebas de versiones anteriores.')
para('Riesgos y acciones previstas',True)
para('Inventario demostrativo: obtener kardex conciliado antes de usar ISI y TQS como evidencia. Referencia de junio: validar la definición de VPR y VPM con el docente. Tablas incompletas en pantalla: implementar HU06 y verificar sus exportaciones. Disponibilidad de responsables: confirmar roles y fechas antes de iniciar.')
para('Registro por Sprint',True)
para('Conservar Sprint Goal, historias seleccionadas, tareas, responsables, impedimentos, pruebas, evidencia de la demostración, observaciones de Review y acción de retrospectiva. Las X del cronograma solo representan planificación. La aceptación y las firmas se completarán al realizar la revisión formal.')
page('Anexo 4 Evidencia visual y fuentes de verificación')
para('Captura real del programa obtenida durante la verificación previa. La interfaz acredita funciones visibles y un resultado de ejecución; no acredita fechas de Sprints ni conformidad de la empresa.')
cap=R/'tmp/guia_capturas/resultado.png'
if cap.exists():d.add_picture(str(cap),width=Cm(17))
para('Figura 1. Evaluación retrospectiva del 17 de agosto al 15 de septiembre de 2026. La tarjeta del modelo muestra MAPE 26,21 %. El panel identifica el inventario como demostrativo.')
para('Fuentes utilizadas',True)
para('Formato institucional PPP1-C1-PRA03; plan previo de evidencia técnica entregado por el practicante; informe LIMA-NORTE_PI_GURRERRO.docx; código de entrenamiento y evaluación en application/services; exportación postest_pronostico_2026-08-17_a_2026-09-15.xlsx; cálculos y guía conservados en outputs/anexo2_30dias_hasta_15septiembre.')
para('Los datos administrativos proceden del documento entregado por el practicante. No se realizó una validación registral del RUC. El contacto informado no se asigna como representante legal o jefe inmediato sin confirmación.')
out=R/'outputs/Plan_Practicas_SCRUM_Taully_Formato_Institucional.docx'
buffer=BytesIO();d.save(buffer);buffer.seek(0)
# Preserve original opaque package parts, replacing body and image relationships only.
with zipfile.ZipFile(src) as original,zipfile.ZipFile(buffer) as built,zipfile.ZipFile(out,'w',zipfile.ZIP_DEFLATED) as dest:
 editable={'word/document.xml','word/_rels/document.xml.rels','[Content_Types].xml'}
 for n in original.namelist():dest.writestr(n,built.read(n) if n in editable else original.read(n))
 for n in built.namelist():
  if n not in original.namelist():dest.writestr(n,built.read(n))
print(out)
