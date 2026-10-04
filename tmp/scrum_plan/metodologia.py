from pathlib import Path
from docx import Document
from docx.shared import Cm,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
R=Path('C:/taully_demand_forecast')
d=Document();s=d.sections[0]
s.page_width=Cm(21);s.page_height=Cm(29.7)
s.left_margin=s.right_margin=Cm(2.2);s.top_margin=Cm(2);s.bottom_margin=Cm(1.8)
for name,size in [('Normal',11),('Title',23),('Heading 1',16),('Heading 2',12)]:
 st=d.styles[name];st.font.name='Arial';st.font.size=Pt(size);st.font.color.rgb=RGBColor(0,0,0)
d.styles['Normal'].paragraph_format.space_after=Pt(8)
d.styles['Normal'].paragraph_format.line_spacing=1.1
for el in d.styles.element.xpath('.//w:pBdr'):el.getparent().remove(el)
def p(t):d.add_paragraph(t)
def h(t):d.add_heading(t,2)
def page(t):d.add_page_break();d.add_heading(t,1)
def table(headers,rows,widths):
 t=d.add_table(rows=1,cols=len(headers));t.autofit=False
 for col,w in zip(t.columns,widths):col.width=Cm(w)
 for i,vals in enumerate([headers]+rows):
  rr=t.rows[0] if i==0 else t.add_row()
  for c,text,w in zip(rr.cells,vals,widths):
   c.width=Cm(w);c.text=text;c.vertical_alignment=1
   pr=c._tc.get_or_add_tcPr();sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'17365D' if i==0 else ('F0F3F6' if i%2 else 'FFFFFF'));pr.append(sh)
   borders=OxmlElement('w:tcBorders')
   for side in ['top','bottom','left','right']:
    e=OxmlElement('w:'+side);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
   pr.append(borders)
   for pp in c.paragraphs:
    pp.paragraph_format.space_before=Pt(5);pp.paragraph_format.space_after=Pt(5);pp.paragraph_format.line_spacing=1
    for r in pp.runs:r.font.size=Pt(9.5);r.bold=i==0;r.font.color.rgb=RGBColor.from_string('FFFFFF' if i==0 else '000000')
  rr._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
 t.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
 p('')
d.add_heading('Metodología Scrum aplicada al sistema de pronóstico de demanda',0)
p('Minimarket Taully · Comas, Lima\nGuerrero Mora, Carlos David Elian\nVersión de trabajo · 3 de octubre de 2026')
h('Propósito')
p('Organizar el desarrollo, la validación y la mejora del sistema de pronóstico de demanda diaria por categoría de producto. La aplicación de Scrum se plantea mediante entregas pequeñas y verificables, revisadas con el responsable del negocio y vinculadas a los objetivos académicos del proyecto.')
h('Contexto del proyecto')
p('El sistema utiliza reportes de ventas y un catálogo de productos para consolidar la demanda diaria, entrenar modelos por categoría, producir pronósticos y evaluar su error frente a ventas conocidas. Su finalidad es apoyar la planificación de compras e inventario del Minimarket Taully.')
p('La aplicación existente cuenta con interfaz web, procesamiento de reportes, entrenamiento, pronóstico, evaluación retrospectiva y exportaciones. Estas funciones constituyen el punto de partida de esta metodología. La organización en Sprints que se describe es una propuesta de trabajo y no un registro de reuniones o ciclos ya ejecutados.')
h('Objetivo del producto')
p('Disponer de un sistema que permita al usuario rastrear cada resultado desde su fuente, conocer el período y método utilizados, comparar predicciones con ventas reales y obtener reportes consistentes para la revisión operativa y académica.')
h('Alcance')
p('Comprende datos de ventas, catálogo, historial, modelos, panel, evaluación y documentación. La incorporación del kardex real y la visualización completa de los indicadores del Anexo 2 se consideran trabajo pendiente. La compra automática, la sustitución del POS y una mejora causal del negocio no forman parte de los resultados acreditados.')

page('1 Organización del trabajo')
h('Responsabilidades propuestas')
table(['Responsabilidad','Aplicación en Taully'],[
['Product Owner','Representante del negocio por designar. Ordena el Product Backlog y precisa qué información necesita para decidir sobre compras e inventario.'],
['Scrum Master','Persona por designar. Facilita el uso de Scrum, ayuda a resolver impedimentos y promueve la revisión del trabajo.'],
['Developers','Guerrero Mora, Carlos David Elian, como responsable técnico propuesto, y los participantes que se confirmen. Construyen, prueban y documentan los incrementos.'],
['Interesados','Jefatura, usuarios de ventas e inventario y docente revisor. Aportan necesidades, revisan resultados y observan su coherencia académica.']],[4.1,12.5])
p('Los nombres de responsables, su disponibilidad y las fechas se confirmarán antes de iniciar. La existencia del código no demuestra la conformación de un Scrum Team ni la ejecución de eventos Scrum. Si no se cuenta con esas responsabilidades y prácticas, deberá describirse como una adaptación del trabajo inspirada en Scrum.')
h('Artefactos de trabajo')
p('El Product Backlog reúne y ordena las necesidades del sistema. Cada elemento incluye propósito, prioridad y criterios de aceptación. Su orientación común es el objetivo del producto establecido en la primera página.')
p('El Sprint Backlog contiene el Sprint Goal, los elementos seleccionados y el plan de tareas del equipo. Se actualiza durante el Sprint con avances e impedimentos. El incremento es el resultado integrado y utilizable que cumple la Definition of Done; una captura de pantalla por sí sola no demuestra que esté terminado.')
h('Duración propuesta')
p('Se propone un horizonte de 12 semanas, distribuido en seis Sprints de dos semanas. Las fechas reales dependerán del acuerdo con la empresa. Las fechas del historial de ventas y del postest son períodos de datos y no se usarán como fechas de ejecución de los Sprints.')

page('2 Product Backlog del sistema')
p('La siguiente priorización parte de las capacidades existentes y de las brechas identificadas. Las estimaciones de esfuerzo se acordarán con quienes realizarán el trabajo, sin atribuir horas ni puntos históricos no registrados.')
table(['ID','Necesidad y prioridad','Criterio de aceptación'],[
['HU01','Como encargado, quiero cargar reportes y clasificar sus productos. Alta.','Una carga autorizada identifica fechas, unidades, categorías y productos sin correspondencia.'],
['HU02','Como revisor, quiero rastrear el historial hasta los reportes. Alta.','Los totales por fecha y categoría concuerdan con las fuentes; duplicados e inconsistencias se documentan.'],
['HU03','Como analista, quiero validar el modelo por categoría. Alta.','Se revisan separación cronológica, categorías excluidas, método elegido y métricas de validación.'],
['HU04','Como usuario, quiero consultar y exportar pronósticos. Alta.','El archivo coincide con el horizonte y las categorías mostradas; las asignaciones por producto se identifican como estimadas.'],
['HU05','Como docente, quiero reproducir el postest. Alta.','Los pares diarios, totales y MAPE coinciden entre evaluación y exportación al conservar datos, corte y horizonte.'],
['HU06','Como revisor, quiero visualizar los índices de las fichas. Alta.','El panel y el Excel incluyen VP, VPR e IE y el índice diario, con fórmula, fuente y período de referencia. Pendiente.'],
['HU07','Como encargado, quiero indicadores de inventario reales. Alta.','El kardex se concilia y permite recalcular CD, ISI y TQS. Los registros demostrativos no se presentan como reales.'],
['HU08','Como usuario, quiero una guía y una entrega comprobable. Media.','El manual permite repetir la revisión; se conservan pruebas, fuentes, limitaciones y decisiones de aceptación.']],[1.3,6.5,8.8])
p('HU01 a HU05 cuentan con funciones existentes que deben validarse en cada entrega. HU06 requiere completar la visualización y exportación. HU07 depende de disponer de kardex real; la función de cálculo existente no elimina esa dependencia.')

page('3 Plan de Sprints e incrementos')
table(['Sprint','Objetivo del Sprint','Incremento previsto'],[
['1 · Semanas 1 y 2','Acordar alcance y asegurar fuentes utilizables.','HU01 validada, catálogo revisado y backlog ordenado con criterios de aceptación.'],
['2 · Semanas 3 y 4','Garantizar trazabilidad del historial.','HU02 validada; conciliación por fecha y categoría y registro de incidencias de datos.'],
['3 · Semanas 5 y 6','Explicar y comprobar el entrenamiento.','HU03 validada; evidencia de separación temporal, selección por categoría y métricas.'],
['4 · Semanas 7 y 8','Entregar pronósticos y exportaciones coherentes.','HU04 validada; revisión del horizonte, fechas, totales y asignación por producto.'],
['5 · Semanas 9 y 10','Facilitar la comprobación del postest y sus fichas.','HU05 validada y HU06 implementada y comprobada contra fuentes y fórmulas.'],
['6 · Semanas 11 y 12','Validar inventario y preparar transferencia.','HU07 con kardex real disponible y HU08; manual, evidencia, pendientes y revisión de entrega.']],[3.1,5.8,7.7])
h('Dependencias que condicionan el plan')
p('La empresa debe facilitar reportes autorizados y definir un responsable para revisar los incrementos. La validación operativa de inventario requiere saldos y movimientos conciliados. La duración del estudio, la referencia de junio y el método previo deben concordar con la metodología académica.')
p('Si el kardex real no está disponible al llegar al Sprint 6, HU07 permanece pendiente. Se documenta el impedimento y se revisa el alcance con el Product Owner; no se sustituye su validación con cifras demostrativas.')
h('Cómo se convierte el plan en trabajo diario')
p('En cada Sprint se descompone la historia en tareas concretas: preparar fuentes, implementar o revisar el cálculo, probar casos, contrastar exportaciones y actualizar la guía. El equipo selecciona una cantidad de trabajo compatible con su capacidad y ajusta las tareas durante el Sprint sin perder su objetivo.')

page('4 Eventos y condiciones de calidad')
h('Planificación y seguimiento')
p('En Sprint Planning se acuerda el objetivo del Sprint, se seleccionan los elementos del backlog y se define el plan para completarlos. El Daily Scrum permite inspeccionar el avance hacia ese objetivo y ajustar el trabajo de los Developers. Los impedimentos se registran y se atienden con los responsables correspondientes.')
p('En Sprint Review se inspecciona el incremento con los interesados mediante una demostración reproducible y se adapta el backlog según lo aprendido. En Sprint Retrospective se revisa cómo trabajó el equipo y se acuerda una mejora concreta para el siguiente ciclo. El refinamiento del backlog se realiza conforme aparecen nuevas necesidades o información.')
h('Definition of Done para este proyecto')
p('Una historia estará terminada cuando cumpla sus criterios de aceptación, esté integrada en el sistema, pase las pruebas pertinentes y tenga resultados trazables. Debe identificar fuente, período, unidades y fórmula; manejar datos incompletos; conservar las advertencias sobre datos demostrativos; y contar con documentación actualizada para su revisión.')
p('Para un cálculo, se contrastará un conjunto de casos con operaciones independientes. Para una exportación, se comprobará que los encabezados y valores correspondan al resultado ejecutado. Para una interfaz, se verificará que un usuario pueda encontrar el indicador sin confundirlo con otro.')
h('Métricas de seguimiento')
p('Se registrarán elementos terminados respecto de los seleccionados, defectos abiertos y resueltos, impedimentos pendientes y resultados de pruebas. No se consignará una velocidad histórica sin registros reales. Las métricas del pronóstico se analizarán por separado: MAPE y WAPE no son equivalentes ni deben intercambiarse en las fichas.')
h('Evidencias a conservar')
p('Por Sprint se conservarán fecha, participantes, Sprint Goal, backlog seleccionado, cambios, pruebas, demostración, observaciones y acción de mejora. Las actas y firmas se incorporarán cuando ocurran. La aceptación empresarial o académica no se inferirá de un resultado técnico satisfactorio.')

page('5 Validación aplicada y estado actual')
h('Comprobación del pronóstico')
p('En la base revisada, el postest abarca del 17/08/2026 al 15/09/2026, con corte 16/08/2026 y horizonte de 30 días. Se comparan 150 pares diarios de cinco categorías. La demanda real total es 7851 unidades y el pronóstico acumulado del modelo es 7684,29.')
p('El MAPE del modelo es 26,21 % y el de PMS-7, sobre la misma ventana, es 29,14 %. El MAPE se calcula a partir de los errores de los pares diarios, no dividiendo la diferencia de los totales. Estos resultados son una evaluación retrospectiva, no una prueba suficiente de mejora causal del negocio.')
h('Brechas identificadas para los próximos incrementos')
p('El panel y la exportación muestran ventas y métricas de pronóstico, pero no la ficha completa VP–VPR–IE. La referencia de junio y esos índices se calcularon aparte para los Word. HU06 debe incorporarlos de manera trazable; la metodología no presenta esa mejora como implementada.')
p('El inventario activo es demostrativo y no concilia completamente con las ventas. HU07 requiere sustituirlo por kardex real antes de interpretar ISI o TQS como resultados del negocio. También debe confirmarse la fecha real de implementación y el método previo para sostener una comparación pretest–postest.')
h('Condición de cierre')
p('La entrega reúne el sistema, sus fuentes autorizadas, exportaciones, pruebas y guía de uso. Se revisa qué criterios se cumplieron y qué queda pendiente. Los responsables registran su decisión y las observaciones; no se incorporan marcas de aceptación ni firmas anticipadas.')
h('Base documental')
p('Se utilizaron el informe LIMA-NORTE_PI_GURRERRO.docx, el documento previo de evidencia técnica entregado por el practicante, el código de entrenamiento y evaluación del proyecto taully_demand_forecast, y la exportación postest_pronostico_2026-08-17_a_2026-09-15.xlsx. Los resultados deben recalcularse si cambian datos o código.')
p('Este documento adopta una estructura propia de metodología del proyecto. No constituye un acta de ejecución ni una certificación de cumplimiento de Scrum; describe su aplicación propuesta sobre el sistema y las evidencias disponibles.')
page('6 Requerimientos funcionales del programa')
p('Los requerimientos describen el comportamiento esperado del sistema. La condición existente indica una función observada en la versión revisada; su aceptación formal se registra en el acta después de la prueba correspondiente.')
table(['Código','Requerimiento','Comprobación y condición'],[
['RF01','Incorporar reportes de ventas y consolidar cantidades por fecha y categoría.','Comparar una carga autorizada con su reporte fuente. Función existente; revisar registros rechazados o sin categoría.'],
['RF02','Administrar el catálogo maestro y asociar productos con categorías.','Verificar el catálogo activo y los productos sin correspondencia. Función existente.'],
['RF03','Mostrar cobertura, categorías y resumen del historial.','Conciliar fechas y totales del panel con el historial procesado. Función existente.'],
['RF04','Entrenar y evaluar modelos por categoría con datos anteriores al corte.','Revisar selección del método, métricas y exclusiones. Función existente.'],
['RF05','Generar pronósticos para el horizonte elegido y exportarlos.','Comprobar fechas y totales entre pantalla y archivo. Identificar asignaciones estimadas por producto. Función existente.'],
['RF06','Evaluar retrospectivamente el modelo frente a ventas reales y PMS-7.','Reproducir corte, horizonte, pares diarios y MAPE en el Excel. Función existente.'],
['RF07','Calcular CD, ISI y TQS con el inventario del período.','Función existente con datos demostrativos. Requiere kardex real conciliado para aceptación operativa.'],
['RF08','Mostrar y exportar VP, VPR, IE y el índice diario de las fichas.','Comprobar referencias, unidades y fórmulas con el Anexo 2. Pendiente de implementación en el panel y la exportación.'],
['RF09','Facilitar la reproducción de resultados mediante guía y evidencia.','Un revisor sigue la guía e identifica período, fuente, método y resultado. Documentación disponible; validación por el usuario pendiente.']],[1.7,6.7,8.2])

page('7 Requerimientos de calidad y operación')
table(['Código','Requerimiento','Criterio de revisión'],[
['RNF01','Trazabilidad de resultados.','Identificar fuente, período, corte, horizonte y ejecución. Conservar las exportaciones utilizadas.'],
['RNF02','Consistencia numérica.','Pantalla y Excel deben coincidir para el mismo indicador y ejecución. El redondeo de presentación no altera el cálculo base.'],
['RNF03','Calidad de datos.','Detectar registros incompletos, duplicados o sin categoría y documentar su tratamiento. No forzar cifras para hacerlas coincidir.'],
['RNF04','Claridad de uso.','Distinguir ventas, pronósticos, MAPE y WAPE; indicar las unidades y el período. Probar la guía con el responsable del negocio.'],
['RNF05','Protección y conservación de datos.','Utilizar fuentes autorizadas y controlar quién accede al equipo y a los archivos. Verificar respaldo y recuperación antes de la entrega; no presumir controles no comprobados.'],
['RNF06','Mantenimiento y reproducibilidad.','Conservar versión del código, dependencias, fuentes y pruebas pertinentes. Recalcular los resultados cuando cambien datos o código.']],[1.8,5.4,9.4])
h('Condiciones para ejecutar el programa')
p('Se requiere el proyecto instalado con su entorno Python y dependencias, el servidor local en ejecución, un navegador y permisos sobre las carpetas de datos y resultados. La dirección http://127.0.0.1:5000/ funciona en el equipo que ejecuta el programa; no implica acceso remoto desde otra computadora.')
p('Las fuentes necesarias son reportes de ventas, catálogo maestro e historial consolidado. Para una revisión oficial del inventario se requiere además un kardex real conciliado. Las pruebas deben realizarse sobre una versión identificada y conservar los archivos antes de introducir cambios.')
p('No se fija un tiempo máximo de respuesta ni un mínimo de memoria sin mediciones. Durante la revisión se registrarán el equipo utilizado, la duración de las operaciones y las incidencias para acordar los requisitos de rendimiento.')

page('8 Acta de revisión y aceptación del sistema')
p('Proyecto: Sistema de pronóstico de demanda del Minimarket Taully. Complete esta acta durante la revisión real. La firma acredita únicamente la decisión y el alcance consignados; no convierte los datos demostrativos en evidencia del negocio.')
table(['Dato de la revisión','Registro'],[
['Número de acta','________________________________________'],
['Fecha y lugar','________________________________________'],
['Versión del programa o entrega','________________________________________'],
['Sprint o etapa revisada','________________________________________'],
['Responsable que revisa y cargo','________________________________________'],
['Archivos y período utilizados','________________________________________']],[6,10.6])
p('En Estado escriba: Conforme, Con observaciones, No conforme o No evaluado. Ningún resultado se encuentra aprobado de antemano.')
table(['Elemento revisado','Estado','Evidencia u observación'],[
['RF01 a RF03 · Datos, catálogo e historial','',''],
['RF04 y RF05 · Modelo y pronóstico','',''],
['RF06 · Postest y exportación','',''],
['RF07 · Inventario y kardex real','',''],
['RF08 · Tablas VP, VPR e IE','',''],
['RF09 · Guía y reproducción por el usuario','',''],
['RNF01 a RNF06 · Calidad y operación','','']],[7.4,3.4,5.8])
h('Decisión del responsable')
p('[  ] Aceptado en el alcance revisado\n[  ] Aceptado con observaciones\n[  ] No aceptado\n[  ] Revisión pendiente')
p('Alcance aceptado y exclusiones: __________________________________\n________________________________________________________________')

page('9 Observaciones compromisos y firmas')
p('Continuación del acta número: __________________  Fecha: __________________')
h('Observaciones y acciones acordadas')
table(['Observación o pendiente','Acción y responsable','Fecha compromiso'],[
['\n\n','\n\n','\n\n'],['\n\n','\n\n','\n\n'],['\n\n','\n\n','\n\n']],[6.8,6.2,3.6])
p('Registrar expresamente cualquier requisito no evaluado o incumplido. Si RF08 permanece pendiente o no existe kardex real para RF07, consignarlo en el alcance y los compromisos; no marcar esas funciones como aceptadas sin verificarlas.')
h('Conformidad de quienes participan')
table(['Responsable de la empresa','Responsable técnico'],[
['Firma:\n\n\n____________________________','Firma:\n\n\n____________________________'],
['Nombre: ____________________\nCargo: _____________________\nFecha: _____________________\nSello, si corresponde:','Guerrero Mora, Carlos David Elian\nCargo: _____________________\nFecha: _____________________']],[8.3,8.3])
p('Docente o revisor académico, si corresponde\n\nFirma: ______________________________\nNombre: ____________________________\nFecha: ______________________________')
p('Las firmas se incorporan después de consignar la decisión, las evidencias y los pendientes. La revisión técnica del sistema no sustituye la evaluación metodológica de la tesis ni una autorización para usar los datos.')
out=R/'outputs/Metodologia_Scrum_Taully_Con_Requerimientos_y_Acta.docx';d.save(out);print(out)
