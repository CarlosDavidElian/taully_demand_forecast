from docx import Document
src='outputs/Metodologia_Scrum_Taully_Adaptada_al_Ejemplo.docx'
d=Document(src)
replacements={
'Se documenta el impedimento y se revisa el alcance con el Product Owner; no se sustituye su validación con cifras demostrativas.':'Se documenta la disponibilidad de las fuentes y se acuerda el alcance de la revisión con el Product Owner.',
'conservar las advertencias sobre datos demostrativos':'identificar el origen y el tipo de los datos',
'El inventario activo es demostrativo y no concilia completamente con las ventas. HU07 requiere sustituirlo por kardex real antes de interpretar ISI o TQS como resultados del negocio.':'Fuente del inventario: archivo demostrativo del proyecto. Su conciliación con las ventas está pendiente. Para validar ISI y TQS del negocio, HU07 contempla utilizar el kardex real conciliado.',
'No constituye un acta de ejecución ni una certificación de cumplimiento de Scrum; describe su aplicación propuesta sobre el sistema y las evidencias disponibles.':'Describe la aplicación propuesta de Scrum sobre el sistema y las evidencias disponibles. La ejecución se documentará mediante los registros y actas correspondientes.',
'La firma acredita únicamente la decisión y el alcance consignados; no convierte los datos demostrativos en evidencia del negocio.':'La firma registra la decisión, el alcance revisado y las fuentes utilizadas.',
'La revisión técnica del sistema no sustituye la evaluación metodológica de la tesis ni una autorización para usar los datos.':'La evaluación metodológica de la tesis y la autorización de uso de los datos se documentan por separado.',
'Los registros demostrativos no se presentan como reales.':'Se identifica el origen de los registros utilizados.',
}
ps=list(d.paragraphs)+[p for t in d.tables for r in t.rows for c in r.cells for p in c.paragraphs]
count=0
for p in ps:
    for a,b in replacements.items():
        if a in p.text:
            for run in p.runs:
                if a in run.text:
                    run.text=run.text.replace(a,b);count+=1;break
            else: raise RuntimeError('Texto dividido entre fragmentos')
out='outputs/Metodologia_Scrum_Taully_Redaccion_Revisada.docx'
d.save(out)
print(out, 'Cambios:', count)
