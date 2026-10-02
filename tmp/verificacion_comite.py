from pathlib import Path
import csv,json,re,hashlib,unicodedata
from collections import defaultdict
from statistics import mean
from datetime import date,timedelta
from openpyxl import load_workbook
from docx import Document

R=Path('C:/taully_demand_forecast'); O=R/'outputs/verificacion_pretest_postest';O.mkdir(exist_ok=True)
def sha(p):return hashlib.sha256(p.read_bytes()).hexdigest()
def norm(s):return ' '.join(unicodedata.normalize('NFKD',str(s)).encode('ascii','ignore').decode().upper().split())
def read_x(p,sheet=None):
    w=load_workbook(p,read_only=True,data_only=True);rows=list(w[sheet].values if sheet else w.active.values);w.close();return rows
catalog=read_x(R/'data/catalogo_maestro.xlsx','PRODUCTOS'); cm={norm(r[0]):r[2] for r in catalog[1:] if r[0]}
sales=defaultdict(float);report_rows=0;unmatched=[];source_cat_diff=[];manifest=[]
files=sorted((R/'data').glob('reporte_ventas_2026-*.xlsx'))
for f in files:
    rows=read_x(f);m=re.search(r'(\d{2})/(\d{2})/(\d{4})',str(rows[0][0]));ds=date(int(m[3]),int(m[2]),int(m[1])).isoformat()
    idx=next(i for i,r in enumerate(rows) if r and norm(r[0])=='PRODUCTO' and norm(r[1])=='CANTIDAD VENDIDA')
    for row in rows[idx+1:]:
        if not row or not row[0] or norm(row[0]).startswith('TOTAL'):continue
        qty=float(row[1]);key=norm(row[0]);report_rows+=1
        if key not in cm:unmatched.append([f.name,key]);continue
        if row[2] and norm(row[2])!=norm(cm[key]):source_cat_diff.append([f.name,key,row[2],cm[key]])
        sales[ds,cm[key]]+=qty
    manifest.append({'file':str(f.relative_to(R)),'sha256':sha(f)})
h=list(csv.DictReader((R/'data/historial_demanda.csv').open(encoding='utf-8-sig')))
hist={(r['date'][:10],r['category']):float(r['quantity']) for r in h};cats=sorted({c for d,c in hist})
diffs=[(k,hist.get(k,0),sales.get(k,0)) for k in set(hist)|set(sales) if hist.get(k,0)!=sales.get(k,0)]
savedfile=R/'data/posttests/postest_20260927t191247z_20260908_7d_70578bbb.json'
latest=max((R/'data/posttests').glob('*.json'),key=lambda p:p.stat().st_mtime)
old=json.loads(savedfile.read_text(encoding='utf-8'))['result'];record=json.loads(latest.read_text(encoding='utf-8'));live=record['result']
assert live['cutoff_date']=='2026-09-08' and live['horizon_days']==7
replay=max(abs(a['model_prediction']-b['model_prediction']) for a,b in zip(old['observations'],live['observations']))
assert [(r['date'],r['category']) for r in old['observations']]==[(r['date'],r['category']) for r in live['observations']]
invpath=R/'data/inventario_producto_Taully_2026-04-01_a_2026-09-15.xlsx';ir=read_x(invpath);inventory=[dict(zip(ir[0],r)) for r in ir[1:] if r[0]]
for r in inventory:r['ds']=r['Fecha'].date().isoformat()
aug={c:sum(v for (d,k),v in hist.items() if d.startswith('2026-08') and k==c)/31 for c in cats}
checks=[];detail=[];periods={}
def check(label,expected,actual,tol=.00501):
    ok=abs(float(expected)-float(actual))<tol;checks.append(dict(label=label,expected=expected,actual=actual,ok=ok))
def num(s):return float(re.search(r'-?\d+(?:[.,]\d+)?',s.replace(' ','')).group().replace(',','.'))
def stats(rows):return dict(dr=sum(x['dr'] for x in rows),dp=sum(x['dp'] for x in rows),mape=mean(abs(x['dr']-x['dp'])/x['dr'] for x in rows)*100)
for label,start in [('Pretest',date(2026,9,2)),('Postest',date(2026,9,9))]:
    days=[(start+timedelta(days=i)).isoformat() for i in range(7)];cut=(start-timedelta(days=1)).isoformat()
    rr=[];invent={}
    for c in cats:
        past=[v for (d,k),v in sorted(hist.items()) if k==c and d<=cut][-7:]
        for d in days:
            base=mean(past[-7:]);past.append(base)
            dp=base if label=='Pretest' else next(x['model_prediction'] for x in live['observations'] if x['date']==d and x['category']==c)
            dr=hist[d,c];assert dr>0
            rr.append(dict(period=label,date=d,category=c,dr=dr,dp=dp,pms7=base,ape=abs(dr-dp)/dr*100))
        records=[r for r in inventory if r['Categoría']==c and r['ds'] in days];groups=defaultdict(list)
        for row in records:groups[row['Producto']].append(row)
        si=en=sf=missing_first=missing_last=continuity=0
        for prod,rows in groups.items():
            rows.sort(key=lambda x:x['ds']);si+=rows[0]['Stock_Inicial'];sf+=rows[-1]['Stock_Final'];en+=sum(r['Entradas'] for r in rows)
            missing_first+=rows[0]['ds']!=days[0];missing_last+=rows[-1]['ds']!=days[-1]
            continuity+=sum(a['Stock_Final']!=b['Stock_Inicial'] for a,b in zip(rows,rows[1:]))
        cd=si+en-sf;dqs=len({r['ds'] for r in records if r['Dias_Sin_Stock']>0})
        invent[c]=dict(si=si,en=en,sf=sf,cd=cd,isi=cd/(si+en)*100,dqs=dqs,dd=7,tqs=dqs/7*100,products=len(groups),missing_first=missing_first,missing_last=missing_last,continuity=continuity,sales=sum(r['Cantidad_Vendida'] for r in records))
    total={k:sum(v[k] for v in invent.values()) for k in ['si','en','sf','cd','dqs','dd','sales','products','missing_first','missing_last','continuity']}
    total['isi']=total['cd']/(total['si']+total['en'])*100;total['tqs']=total['dqs']/total['dd']*100
    overall=stats(rr);bycat={c:stats([r for r in rr if r['category']==c]) for c in cats}
    docpath=R/f'outputs/postest_visualizado/{label}_corregido.docx';doc=Document(docpath)
    # Check every reported numeric result, including global rows and both repeated MAPE instruments.
    for ti,keys in [(1,['si','en','sf','cd']),(3,['cd','si','en','isi']),(5,['dqs','dd','tqs'])]:
        for row in doc.tables[ti].rows[1:]:
            c=row.cells[1].text;expected=total if c=='TOTAL' else invent[c]
            for col,key in enumerate(keys,2):check(f'{label} T{ti} {c} {key}',expected[key],num(row.cells[col].text))
    for ti in [7,8]:
        for row in doc.tables[ti].rows[1:]:
            ds='2026-'+row.cells[0].text[3:5]+'-'+row.cells[0].text[:2];c=row.cells[1].text
            for col,v in enumerate([hist[ds,c],aug[c],hist[ds,c]/aug[c]],2):check(f'{label} T{ti} {ds} {c} C{col}',v,num(row.cells[col].text))
    for ti in [10,16]:
        for row in doc.tables[ti].rows[1:]:
            c=row.cells[1].text;st=overall if c=='GLOBAL' else bycat[c]
            for col,key in enumerate(['dr','dp','mape'],2):check(f'{label} T{ti} {c} {key}',st[key],num(row.cells[col].text))
    for row in doc.tables[12].rows[1:]:
        c=row.cells[1].text;dr=overall['dr'] if c=='TOTAL' else bycat[c]['dr']
        for col in [2,3]:check(f'{label} volume {c} C{col}',dr,num(row.cells[col].text))
    for row in doc.tables[14].rows[1:]:
        c=row.cells[1].text;dr=overall['dr'] if c=='TOTAL' else bycat[c]['dr'];ref=sum(aug.values())*7 if c=='TOTAL' else aug[c]*7
        for col,v in enumerate([dr,ref,dr/ref],2):check(f'{label} pattern {c} C{col}',v,num(row.cells[col].text))
    detail.extend(rr);periods[label]=dict(start=days[0],end=days[-1],cutoff=cut,forecast=overall,bycat=bycat,inventory=invent,total_inventory=total,seasonal_global=overall['dr']/(sum(aug.values())*7))
    manifest.append({'file':str(docpath.relative_to(R)),'sha256':sha(docpath)})
for c,values in live['model']['by_category'].items():
    for k,v in [('actual_total','dr'),('prediction_total','dp'),('mape','mape')]:check('program '+c+' '+k,periods['Postest']['bycat'][c][v],values[k],.0000011)
for r in live['observations']:check('program actual '+r['date']+' '+r['category'],hist[r['date'],r['category']],r['actual_demand'])
postrows=[r for r in detail if r['period']=='Postest'];base=mean(abs(r['dr']-r['pms7'])/r['dr'] for r in postrows)*100
check('PMS7 same week MAPE',base,live['baseline']['metrics']['mape'],.0000011)
for r in live['inventory']['metrics']:
    exp=periods['Postest']['inventory'][r['category']]
    for app,key in [('SI','si'),('EN','en'),('SF','sf'),('CD','cd'),('ISI','isi'),('DQS','dqs'),('DD','dd'),('TQS','tqs')]:check('program inventory '+r['category']+' '+key,exp[key],r[app],.0000011)
export_path=O/'Postest_exportado_programa.xlsx'
if export_path.exists():
    rows=read_x(export_path,'Detalle diario'); assert len(rows)==36
    for row in rows[1:]:
        obs=next(x for x in live['observations'] if (x['date'],x['category'])==(row[0],row[1]))
        for col,key in enumerate(['actual_demand','model_prediction','pms_7_prediction'],2):check('Excel diario '+row[0]+' '+row[1]+' '+key,obs[key],row[col],.0000011)
    for row in read_x(export_path,'Métricas categoría')[1:]:
        check('Excel MAPE '+row[0],periods['Postest']['bycat'][row[0]]['mape'],row[4],.0000011)
    for row in read_x(export_path,'Inventario')[6:11]:
        exp=periods['Postest']['inventory'][row[0]]
        for col,key in [(1,'si'),(2,'en'),(3,'sf'),(4,'cd'),(5,'isi'),(6,'dqs'),(7,'dd'),(10,'tqs')]:check('Excel inventario '+row[0]+' '+key,exp[key],row[col],.0000011)
    manifest.append({'file':str(export_path.relative_to(R)),'sha256':sha(export_path)})
for f in [savedfile,latest,R/'data/historial_demanda.csv',invpath,R/'data/catalogo_maestro.xlsx',Path('C:/Users/Acer/OneDrive/Tesis Sr/LIMA-NORTE_PI_GURRERRO.docx')]:manifest.append({'file':str(f),'sha256':sha(f)})
summary=dict(reports=len(files),report_rows=report_rows,report_units=sum(sales.values()),history_rows=len(h),history_duplicates=len(h)-len(hist),report_history_diffs=diffs,unmatched=unmatched,category_disagreements=source_cat_diff,periods=periods,checks=len(checks),failed=[c for c in checks if not c['ok']],replay_max_diff=replay,run_id=record['run_id'],created_at=record['created_at'],same_week_pms_mape=base,source=live['inventory']['inventory'],methods=live['training']['validation_by_category'])
(O/'verificacion_calculos.json').write_text(json.dumps(summary,ensure_ascii=False,indent=2),encoding='utf-8')
(O/'fuentes_sha256.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2),encoding='utf-8')
for name,rows in [('detalle_diario.csv',detail),('comprobaciones.csv',checks)]:
    with (O/name).open('w',newline='',encoding='utf-8-sig') as f:w=csv.DictWriter(f,fieldnames=list(rows[0]));w.writeheader();w.writerows(rows)
print(json.dumps({k:v for k,v in summary.items() if k not in ['periods','source','methods']},ensure_ascii=True,indent=2))
