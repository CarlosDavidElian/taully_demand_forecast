from pathlib import Path
from collections import defaultdict, Counter
from statistics import mean
from datetime import date, timedelta, datetime
import csv, json, hashlib, re, unicodedata
from openpyxl import load_workbook
from docx import Document

R=Path('C:/taully_demand_forecast')
O=R/'outputs/verificacion_pretest_postest_2026-09-30'
O.mkdir(exist_ok=True)
def sha(p): return hashlib.sha256(p.read_bytes()).hexdigest()
def norm(s): return ' '.join(unicodedata.normalize('NFKD',str(s)).encode('ascii','ignore').decode().upper().split())
def read_sheet(p,sheet=None):
    w=load_workbook(p,read_only=True,data_only=True)
    rows=list((w[sheet] if sheet else w.active).values);w.close();return rows
def f(v):return f'{v:.2f}'.replace('.',',')
def integer(v):return str(int(v))
checks=[]; mismatches=[]
def check(name,yes,detail=None):
    checks.append({'check':name,'passed':bool(yes),'detail':detail})
    if not yes:mismatches.append(checks[-1])

history=list(csv.DictReader((R/'data/historial_demanda.csv').open(encoding='utf-8-sig')))
H={(r['date'][:10],r['category']):float(r['quantity']) for r in history}
cats=sorted({c for d,c in H})
check('Historial sin claves duplicadas',len(H)==len(history))
catalog=read_sheet(R/'data/catalogo_maestro.xlsx','PRODUCTOS')
names={norm(r[0]):r[2] for r in catalog[1:] if r[0]}
sales=defaultdict(float); sale_rows=0; missing=[];label_diffs=[]; manifest=[]
for path in sorted((R/'data').glob('reporte_ventas_2026-*.xlsx')):
    rows=read_sheet(path); dates=[]
    for row in rows[:8]:
        for v in row:
            m=re.search(r'REPORTE DE VENTAS\s*-\s*(\d{2}/\d{2}/\d{4})',str(v))
            if m:dates.append(datetime.strptime(m[1],'%d/%m/%Y').date().isoformat())
    assert len(set(dates))==1,path
    day=dates[0]; header=next(i for i,row in enumerate(rows) if row and row[0]=='PRODUCTO')
    for row in rows[header+1:]:
        if not row or row[0] is None or norm(row[0]) in ['TOTAL','TOTALES']:continue
        if len(row)<2 or not isinstance(row[1],(int,float)):continue
        name=norm(row[0]);c=names.get(name)
        if c is None:missing.append([path.name,row[0]]);continue
        if len(row)>2 and row[2]!=c:label_diffs.append([path.name,row[0],row[2],c])
        sales[(day,c)]+=row[1];sale_rows+=1
    manifest.append({'file':str(path.relative_to(R)),'sha256':sha(path)})
diff=[{'key':k,'history':H.get(k),'reports':sales.get(k)} for k in set(H)|set(sales) if H.get(k,0)!=sales.get(k,0)]
check('Reportes coinciden con historial',not diff,{'differences':diff,'unmatched':missing,'category_labels':label_diffs})
invpath=R/'data/inventario_producto_Taully_2026-04-01_a_2026-09-15.xlsx'
rawinv=read_sheet(invpath);inv=[dict(zip(rawinv[0],r)) for r in rawinv[1:] if r[0]]
for r in inv:r['date']=r['Fecha'].date().isoformat()
check('Inventario sin duplicados',len({(r['date'],r['Categoría'],r['Producto']) for r in inv})==len(inv))
savedpath=R/'data/posttests/postest_20260927t191247z_20260908_7d_70578bbb.json'
saved=json.loads(savedpath.read_text(encoding='utf-8'))['result']
latestpath=max((R/'data/posttests').glob('*.json'),key=lambda p:p.stat().st_mtime)
latestrec=json.loads(latestpath.read_text(encoding='utf-8'));live=latestrec['result']
check('Ejecucion visible corresponde al periodo del postest',live['cutoff_date']=='2026-09-08' and live['horizon_days']==7)
check('Predicciones repetidas coinciden con ejecucion original',live['observations']==saved['observations'])
aug={c:sum(q for (d,k),q in H.items() if k==c and d.startswith('2026-08'))/31 for c in cats}
periods={}; detail=[]; allnumeric=0
for label,start,end,filename in [('pre','2026-09-02','2026-09-08','Pretest_corregido.docx'),('post','2026-09-09','2026-09-15','Postest_corregido.docx')]:
    days=[(date.fromisoformat(start)+timedelta(days=i)).isoformat() for i in range(7)]
    observed={c:[H[d,c] for d in days] for c in cats};pred={};base={};im={}
    for c in cats:
        past=[q for (d,k),q in sorted(H.items()) if k==c and d<start][-7:]
        v=[]
        for d in days: v.append(mean(past[-7:]));past.append(v[-1])
        base[c]=v
        pred[c]=v if label=='pre' else [next(r['model_prediction'] for r in live['observations'] if r['date']==d and r['category']==c) for d in days]
        groups=defaultdict(list)
        window=[r for r in inv if start<=r['date']<=end and r['Categoría']==c]
        for r in window:groups[r['Producto']].append(r)
        si=en=sf=ms=me=gaps=rowdiff=0
        for rows in groups.values():
            rows.sort(key=lambda r:r['date']);si+=rows[0]['Stock_Inicial'];sf+=rows[-1]['Stock_Final'];en+=sum(r['Entradas'] for r in rows)
            ms+=rows[0]['date']!=start;me+=rows[-1]['date']!=end
            gaps+=sum(a['Stock_Final']!=b['Stock_Inicial'] for a,b in zip(rows,rows[1:]))
            rowdiff+=sum(r['Stock_Inicial']+r['Entradas']-r['Stock_Final']!=r['Cantidad_Vendida'] for r in rows)
        dqs=sorted({r['date'] for r in window if r['Dias_Sin_Stock']>0})
        im[c]={'si':si,'en':en,'sf':sf,'cd':si+en-sf,'isi':100*(si+en-sf)/(si+en),'dqs':len(dqs),'tqs':len(dqs)/7*100,'stockout_dates':dqs,'products':len(groups),'missing_start':ms,'missing_end':me,'continuity_gaps':gaps,'row_differences':rowdiff,'sold':sum(r['Cantidad_Vendida'] for r in window)}
    scores={c:{'actual':sum(observed[c]),'prediction':sum(pred[c]),'mape':mean(abs(a-p)/a for a,p in zip(observed[c],pred[c]))*100} for c in cats}
    baseline_mape=mean(abs(a-p)/a for c in cats for a,p in zip(observed[c],base[c]))*100
    global_score={'actual':sum(v['actual'] for v in scores.values()),'prediction':sum(v['prediction'] for v in scores.values()),'mape':mean(v['mape'] for v in scores.values())}
    totals={k:sum(v[k] for v in im.values()) for k in ['si','en','sf','cd','dqs','sold','products','missing_start','missing_end','continuity_gaps','row_differences']}
    totals.update(isi=totals['cd']/(totals['si']+totals['en'])*100,tqs=totals['dqs']/35*100)
    totals['difference']=totals['sold']-totals['cd']
    for c in cats:
        for d,a,p,b in zip(days,observed[c],pred[c],base[c]):detail.append({'period':label,'date':d,'category':c,'actual':a,'prediction':p,'absolute_percentage_error':abs(a-p)/a*100,'pms_same_week':b})
    path=R/'outputs/postest_visualizado'/filename;doc=Document(path)
    def compare(tidx,rows):
        global allnumeric
        actual=[[c.text.strip() for c in r.cells] for r in doc.tables[tidx].rows[1:]]
        expected=[[str(v) for v in r] for r in rows]
        check(f'{label} tabla {tidx} completa',actual==expected,{'actual':actual,'expected':expected} if actual!=expected else None)
        allnumeric+=sum(bool(re.fullmatch(r'[0-9]+(?:,[0-9]+)?(?: %)?',v)) for row in expected for v in row)
    compare(1,[[i+1,c]+[integer(im[c][k]) for k in ['si','en','sf','cd']] for i,c in enumerate(cats)]+[['','TOTAL']+[integer(totals[k]) for k in ['si','en','sf','cd']]])
    compare(3,[[i+1,c]+[integer(im[c][k]) for k in ['cd','si','en']]+[f(im[c]['isi'])+' %'] for i,c in enumerate(cats)]+[['','TOTAL']+[integer(totals[k]) for k in ['cd','si','en']]+[f(totals['isi'])+' %']])
    compare(5,[[i+1,c,im[c]['dqs'],'7',f(im[c]['tqs'])+' %'] for i,c in enumerate(cats)]+[['','TOTAL',totals['dqs'],'35 categoría-días',f(totals['tqs'])+' %']])
    def daily(ds):return [[date.fromisoformat(d).strftime('%d/%m'),c,integer(H[d,c]),f(aug[c]),f(H[d,c]/aug[c])] for d in ds for c in cats]
    compare(7,daily(days[:3]));compare(8,daily(days[3:]))
    forecasts=[[i+1,c,integer(scores[c]['actual']),f(scores[c]['prediction']),f(scores[c]['mape'])+' %'] for i,c in enumerate(cats)]+[['','GLOBAL',integer(global_score['actual']),f(global_score['prediction']),f(global_score['mape'])+' %']]
    compare(10,forecasts);compare(16,forecasts)
    compare(12,[[i+1,c,integer(scores[c]['actual']),integer(scores[c]['actual'])] for i,c in enumerate(cats)]+[['','TOTAL',integer(global_score['actual']),integer(global_score['actual'])]])
    compare(14,[[i+1,c,integer(scores[c]['actual']),f(aug[c]*7),f(scores[c]['actual']/(aug[c]*7))] for i,c in enumerate(cats)]+[['','TOTAL',integer(global_score['actual']),f(sum(aug.values())*7),f(global_score['actual']/(sum(aug.values())*7))]])
    # Compare the original eight instrument formulas including native Word equations.
    thesis=Document('C:/Users/Acer/OneDrive/Tesis Sr/LIMA-NORTE_PI_GURRERRO.docx')
    def formulas(d):
        out=[]
        for t in d.tables:
            for row in t.rows:
                if norm(row.cells[0].text).startswith('FORMULA'):
                    out.append(''.join(row.cells[1]._tc.xpath('.//w:t/text()|.//m:t/text()')).replace(' ','').replace('\n',''))
        return out
    check(label+' formulas originales conservadas',formulas(doc)==formulas(thesis)[:8],{'document':formulas(doc),'thesis':formulas(thesis)[:8]})
    for tidx in [0,2,4,6,9,11,13,15]:
        moment=doc.tables[tidx].rows[-1].cells[1].text
        check(label+f' fecha ficha {tidx}',('02 al 08' if label=='pre' else '09 al 15') in moment and '2026' in moment)
    periods[label]={'dates':days,'by_category':scores,'global':global_score,'pms_same_week_mape':baseline_mape,'inventory':im,'inventory_total':totals,'word_sha256':sha(path)}
    manifest.append({'file':str(path.relative_to(R)),'sha256':sha(path)})
check('Postest MAPE servicio igual al recalculo',abs(live['model']['metrics']['mape']-periods['post']['global']['mape'])<0.000001)
check('Pretest sin ceros y postest sin ceros',all(r['actual']>0 for r in detail))
for c in cats:
    if not live['inventory']['available']:break
    m=next(r for r in live['inventory']['metrics'] if r['category']==c)
    for short,long in [('si','SI'),('en','EN'),('sf','SF'),('cd','CD'),('isi','ISI'),('dqs','DQS'),('tqs','TQS')]:
        check('Postest inventario servicio '+c+' '+short,abs(m[long]-periods['post']['inventory'][c][short])<1e-5)
for path in [R/'data/historial_demanda.csv',R/'data/catalogo_maestro.xlsx',invpath,savedpath,latestpath,R/'application/services/posttest_evaluation_service.py',R/'application/services/inventory_service.py',R/'interfaces/web.py']:
    manifest.append({'file':str(path.relative_to(R)),'sha256':sha(path)})
result={'review_date':'2026-09-30','reports':{'files':len(list((R/'data').glob('reporte_ventas_2026-*.xlsx'))),'sales_rows':sale_rows,'units':sum(sales.values()),'history_rows':len(H),'catalog_products':len(names),'unmatched_products':missing,'differences':diff},'periods':periods,'august_daily_reference':aug,'numeric_cells_checked':allnumeric,'checks':checks,'failures':mismatches,'live_run':{'path':str(latestpath),'run_id':latestrec['run_id'],'created_at':latestrec['created_at'],'model_methods':live['training']['validation_by_category'],'inventory_source':live['inventory'].get('inventory')},'manifest':manifest}
(O/'calculos_verificados.json').write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf-8')
with (O/'detalle_diario_verificado.csv').open('w',encoding='utf-8-sig',newline='') as file:
    writer=csv.DictWriter(file,fieldnames=list(detail[0]));writer.writeheader();writer.writerows(detail)
print(json.dumps({k:result[k] for k in ['reports','numeric_cells_checked','failures','live_run']},ensure_ascii=True,indent=2))
print('CHECKS',len(checks),'FAILURES',len(mismatches))
