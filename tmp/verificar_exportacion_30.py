import json, csv, urllib.request, hashlib
from pathlib import Path
from openpyxl import load_workbook
R=Path('C:/taully_demand_forecast'); O=R/'outputs/verificacion_pretest_postest_2026-09-30'
j=json.loads((O/'calculos_verificados.json').read_text(encoding='utf-8'))
req=urllib.request.Request('http://127.0.0.1:5000/api/posttest/export',data=json.dumps({'run_id':j['live_run']['run_id']}).encode(),headers={'Content-Type':'application/json'},method='POST')
raw=urllib.request.urlopen(req,timeout=30).read(); p=O/'Postest_exportado_verificado.xlsx'; p.write_bytes(raw)
w=load_workbook(p,data_only=True); checks=[]
for row in list(w['Métricas categoría'].values)[1:]:
    if row[0] not in j['periods']['post']['by_category']: continue
    x=j['periods']['post']['by_category'][row[0]]
    checks.append(abs(row[1]-x['actual'])<1e-8 and abs(row[4]-x['mape'])<=0.0000005)
details=list(csv.DictReader((O/'detalle_diario_verificado.csv').open(encoding='utf-8-sig')))
for row in list(w['Detalle diario'].values)[1:]:
    x=next(x for x in details if x['period']=='post' and x['date']==row[0] and x['category']==row[1])
    checks.append(abs(row[2]-float(x['actual']))<1e-8 and abs(row[3]-float(x['prediction']))<1e-8)
assert len(checks)==40 and all(checks)
result={'run_id':j['live_run']['run_id'],'sheets':w.sheetnames,'checks':len(checks),'failures':0,'sha256':hashlib.sha256(raw).hexdigest()}
(O/'exportacion_verificada.json').write_text(json.dumps(result,indent=2,ensure_ascii=False),encoding='utf-8')
print(json.dumps(result,ensure_ascii=True))
print('INVENTORY TOTALS', {k:v['inventory_total'] for k,v in j['periods'].items()})
