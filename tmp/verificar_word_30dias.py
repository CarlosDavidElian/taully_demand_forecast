from pathlib import Path
import json,zipfile,hashlib
from docx import Document
R=Path('C:/taully_demand_forecast');O=R/'outputs/anexo2_30dias_hasta_15septiembre'
S=Path('C:/Users/Acer/Downloads/LIMA-NORTE_PI_GURRERRO_removed_removed.docx');src=Document(S)
expected=json.loads((O/'tablas_esperadas.json').read_text(encoding='utf-8'))
checks=[]
def check(name,yes):
 assert yes,name
 checks.append(name)
def norm(s):return ' '.join(s.split())
for label,title in [('pre','Pretest'),('post','Postest')]:
 fp=O/(title+'_Anexo2_30dias.docx');d=Document(fp);groups=[]
 for t in d.tables:
  if t.cell(0,0).text.startswith('INDICADOR'):
   groups.append({'metadata':t,'tables':[]})
  else:groups[-1]['tables'].append(t)
 check(title+' ocho instrumentos',len(groups)==8)
 for idx,g in enumerate(groups):
  original=src.tables[idx*2];meta=g['metadata']
  for rn in range(5):
   check(title+f' metadato {idx} fila {rn}',[norm(c.text) for c in meta.rows[rn].cells]==[norm(c.text) for c in original.rows[rn].cells])
  hdr=[norm(c.text) for c in src.tables[idx*2+1].rows[0].cells]
  actual=[]
  for t in g['tables']:
   check(title+f' encabezado {idx}',[norm(c.text) for c in t.rows[0].cells]==hdr)
   actual.extend([[c.text for c in r.cells] for r in t.rows[1:]])
  wanted=[[str(x) for x in row] for row in expected[label][idx]['rows']]
  check(title+f' todos resultados ficha {idx+1}',actual==wanted)
 check(title+' ocho secciones',len(d.sections)==len(src.sections)==8)
 for a,b in zip(d.sections,src.sections):
  check(title+' geometría sección',all(getattr(a,k)==getattr(b,k) for k in ['page_width','page_height','left_margin','right_margin','top_margin','bottom_margin']))
 with zipfile.ZipFile(S) as z,zipfile.ZipFile(fp) as q:
  check(title+' paquete preservado',all(z.read(n)==q.read(n) for n in z.namelist() if n!='word/document.xml'))
manifest=json.loads((R/'tmp/anexo30_reference_manifest.json').read_text())
check('Original intacto',hashlib.sha256(S.read_bytes()).hexdigest()==manifest['sha256'])
(O/'verificacion_final_word.json').write_text(json.dumps({'checks_passed':len(checks),'checks':checks},ensure_ascii=False,indent=2),encoding='utf-8')
print('COMPROBACIONES WORD',len(checks),'OK')
