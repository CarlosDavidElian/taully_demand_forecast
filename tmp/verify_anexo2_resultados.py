from pathlib import Path

from docx import Document


path = Path(r"C:\taully_demand_forecast\outputs\01a0b651-7620-73f0-8f6a-8b68b14f7a17\Anexo_2_Instrumentos_Resultados_Tecnicos_Pretest_Postest_2026-09-02_a_2026-09-15.docx")
document = Document(path)

assert len(document.tables) == 25, f"Se esperaban 25 tablas y hay {len(document.tables)}."
assert "con resultados" in document.paragraphs[0].text.lower()
assert "02 al 08" in document.paragraphs[2].text
assert "09 al 15" in document.paragraphs[2].text

data_tables = (3, 5, 7, 9, 11, 13, 15, 17, 19, 21)
for table_index in data_tables:
    table = document.tables[table_index]
    for row_index, row in enumerate(table.rows):
        for column_index, cell in enumerate(row.cells):
            if table_index in (3, 5, 7, 9, 11, 13, 15, 17, 19, 21) and row_index == 6 and column_index == 0:
                continue
            value = cell.text.strip()
            assert value, f"Celda vacía en tabla {table_index}, fila {row_index}, columna {column_index}."
            assert "___" not in value, f"Marcador sin completar en tabla {table_index}, fila {row_index}, columna {column_index}."

for table_index in (2, 4, 6, 8, 10, 12, 14, 16, 18, 20):
    table = document.tables[table_index]
    for row in table.rows[1:]:
        assert row.cells[1].text.strip(), f"Metadato vacío en tabla {table_index}."
        assert "___" not in row.cells[1].text, f"Marcador sin completar en tabla {table_index}."

checks = {
    (3, 1, 5): "245",
    (3, 6, 5): "1,411",
    (7, 2, 4): "14.29%",
    (17, 3, 4): "28.57%",
    (11, 6, 4): "27.69%",
    (21, 6, 4): "33.89%",
}
for (table_index, row_index, column_index), expected in checks.items():
    actual = document.tables[table_index].cell(row_index, column_index).text.strip()
    assert actual == expected, f"Tabla {table_index}, fila {row_index}, columna {column_index}: {actual!r} != {expected!r}"

print("VERIFICACION_ESTRUCTURAL_OK")
print("Páginas esperadas: 12; tablas: 25; fichas con resultados: 10.")
print("O1 error total:", [cell.text for cell in document.tables[11].rows[6].cells])
print("O2 error total:", [cell.text for cell in document.tables[21].rows[6].cells])
